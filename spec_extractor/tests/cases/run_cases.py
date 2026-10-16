#!/usr/bin/env python3
"""
run_cases.py - run spec_extractor.py on every test-case document and check its output, case by case.

    python tests/cases/run_cases.py                   (all cases, a few minutes)
    python tests/cases/run_cases.py --only TAB        (only the tab-list cases; --only TAB-07 for one case)
    python tests/cases/run_cases.py path/to/other/spec_extractor.py    (test another version of the tool)
    python tests/cases/run_cases.py --only MY         (only your own scans, kept in tests/cases/my_cases)
    python tests/cases/run_cases.py --record tests/cases/my_cases/page12.pdf
                                                      (turn your own scan into a case - see tests/README.md)

The documents are built first when they are missing or out of date (build_cases.py). Each document is run
through the real tool, offline (--no-combined --no-cache --no-word), and the 'Extracted Items' sheet of its
workbook is compared with the case:
  * the number of Requirement rows is exactly the expected number,
  * every row's Doc Ref is exactly the expected one, in the expected order,
  * every row's Content is similar to the expected text (difflib ratio >= 0.9, whitespace normalised), both the
    whole Content AND the part after the lead-in, so a list item that lost its value cannot pass on the
    strength of a long shared lead-in,
  * the checks written in brackets on a row (review note, no list-marker note, Section, Page(s), Title,
    OCR Corrections, Quality Notes),
  * the number of Table items, the text each table must contain, and the caption labels when the case says so,
  * the headings of the Document Structure sheet (level, number and title) and the documents of the Applicable
    Documents sheet, when the case gives them.

Result of each case: PASS, FAIL, or - for a case marked as a known limitation - KNOWN (still wrong, as
recorded) or KNOWN-PASSES (the limitation seems to be fixed: check it and remove the mark).

Your own scans (tests/cases/my_cases) are run into a folder of their own under tests/cases/my_cases/_results (one
for each outdir, named like it) instead of <outdir>, with their reports: no copy of their pages or text is left
outside my_cases, which git ignores.

The results are kept in <outdir>/case_results.json; the next run lists every case whose result changed since
then - the quickest way to see what an edit improved or broke. The full expected-vs-actual listing is written
to <outdir>/case_report.txt and a one-line-per-case table to <outdir>/case_results.xlsx.

Options: --only TEXT  only these cases: an area ('TAB'), a case ('TAB-07') or words of a title ('wrapped
                      label'); several with commas: --only TAB,LST-03
         --jobs N     run the tool as N parallel processes (default: half the processor cores)
         --reuse      do not run the tool again when <outdir> already holds the workbooks (re-check only)
         -v           print the expected/actual rows of failing cases
The tool runs with the Python of the project's .venv (the one setup_windows.bat makes); set SPEC_PYTHON
to the path of another Python to run it with that one instead - to test another Python version.
Exit status: 0 when every selected case passes (known limitations do not count), 1 otherwise, 2 on errors.
"""
from __future__ import annotations

import os
import sys

_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))      # the spec_extractor folder
VENV_PY = os.path.join(_ROOT, ".venv", "Scripts", "python.exe") if os.name == "nt" else os.path.join(_ROOT, ".venv", "bin", "python")
try:
    import openpyxl  # noqa: F401
except ImportError:                                   # started with a Python without openpyxl: use the tool's venv
    if os.path.exists(VENV_PY) and os.path.realpath(sys.executable) != os.path.realpath(VENV_PY):
        os.execv(VENV_PY, [VENV_PY] + sys.argv)
    raise

import argparse
import contextlib
import difflib
import json
import re
import shutil
import subprocess
import tempfile
import time
from collections import Counter
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import casekit  # noqa: E402
import build_cases  # noqa: E402

# Everything the tool writes for your own scans - its workbooks, the page images and text it keeps in _cache, and the
# case reports, which quote the requirement text - stays here, inside my_cases (which git ignores), and never in
# tests/cases/_results: deleting the my_cases folder removes every copy.
OWN_RESULTS = casekit.LOCAL_FOLDER / "_results"

THRESHOLD = 0.9


def norm(s) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip()


def ratio(a: str, b: str) -> float:
    return difflib.SequenceMatcher(None, a, b, autojunk=False).ratio()


def tail_after_stem(actual: str, expected: str, stem_len: int) -> str:
    """The part of `actual` that corresponds to expected[stem_len:] (the list item after the lead-in)."""
    if stem_len <= 0:
        return actual
    pos = 0
    for a, b, size in difflib.SequenceMatcher(None, expected, actual, autojunk=False).get_matching_blocks():
        if a > stem_len or size == 0:
            break
        pos = b + min(stem_len - a, size)
    return actual[pos:].strip()


def run_tool(spec: Path, documents: list, outdir: Path, jobs: int) -> None:
    py = os.environ.get("SPEC_PYTHON") or (VENV_PY if os.path.exists(VENV_PY) else sys.executable)
    jobs = max(1, min(jobs, len(documents)))
    groups = [documents[k::jobs] for k in range(jobs)]
    procs = []
    for k, group in enumerate(groups):
        target = outdir / f"run{k + 1}"
        target.mkdir(parents=True, exist_ok=True)
        cmd = [py, str(spec), *map(str, group), "-o", str(target), "--no-combined", "--no-cache", "--no-word"]
        log = open(target / "tool_output.txt", "w", encoding="utf-8")
        procs.append((subprocess.Popen(cmd, stdout=log, stderr=subprocess.STDOUT), log, target))
    for p, log, target in procs:
        rc = p.wait()
        log.close()
        if rc != 0:
            print(f"WARNING: spec_extractor exited with status {rc} (see {target / 'tool_output.txt'})", file=sys.stderr)


def workbooks_of(outdir: Path, document: Path) -> list:
    """The tool's workbooks for this document under outdir (the name matched as it is: a scan may be called
    'SRS [Rev B].pdf', which a search pattern would read as a set of letters)."""
    name = f"{document.stem}_extracted.xlsx"
    return sorted(p for p in outdir.glob("run*/*_extracted.xlsx") if p.name == name)    # (run_tool's folders only)


def find_workbook(outdir: Path, document: Path):
    hits = workbooks_of(outdir, document)
    return hits[0] if hits else None


@contextlib.contextmanager
def _temporary_files_in(folder: Path):
    """While a workbook is saved, keep its temporary files (openpyxl writes each sheet to one) in a folder of their
    own inside `folder`, and delete it afterwards - not in the system's temp folder: the reports of your own scans
    quote their text. (The same as spec_extractor.py's helper of this name.)"""
    private = Path(tempfile.mkdtemp(prefix="_saving_", dir=folder))
    before, tempfile.tempdir = tempfile.tempdir, str(private)
    try:
        yield
    finally:
        tempfile.tempdir = before
        shutil.rmtree(private, ignore_errors=True)


def _sheet(wb, name: str) -> list:
    """The rows of a sheet as dicts by column name ([] when the workbook has no such sheet)."""
    if name not in wb.sheetnames:
        return []
    rows = list(wb[name].iter_rows(values_only=True))
    if not rows:
        return []
    col = {n: k for k, n in enumerate(rows[0]) if n}
    return [{n: (r[k] if k < len(r) else None) for n, k in col.items()} for r in rows[1:] if any(v is not None for v in r)]


def read_items(xlsx: Path) -> list:
    """The Extracted Items rows (dicts), with the workbook's structure and documents on the first one's list."""
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    try:
        items = _sheet(wb, "Extracted Items")
        structure, documents = _sheet(wb, "Document Structure"), _sheet(wb, "Applicable Documents")
        for r in items:          # a table's row links to the sheet holding it: its text is read from there
            sheet = r.get("Table Sheet")
            if sheet and sheet in wb.sheetnames:
                cells = [" | ".join(norm(v) for v in row if v is not None) for row in
                         wb[sheet].iter_rows(min_row=5, values_only=True)]
                r["Content"] = " || ".join(c for c in cells if c)
    finally:
        wb.close()
    get = lambda r, name: norm(r.get(name))
    out = [{"type": get(r, "Type"), "doc_ref": get(r, "Doc Ref"), "content": get(r, "Content"),
            "reasons": str(r.get("Review Reasons") or ""), "section": get(r, "Section"), "pages": get(r, "Page(s)"),
            "item": get(r, "Item #"), "title": get(r, "Title"), "corrections": str(r.get("OCR Corrections") or ""),
            "quality": str(r.get("Quality Notes") or "")} for r in items]
    headings = [f"{r.get('Level')} | {norm(r.get('Heading / Text'))}" for r in structure if r.get("Is Heading") == "Yes"]
    docs = [" | ".join(norm(r.get(k)) for k in ("Document", "Title", "Referenced In (paragraphs)")) for r in documents
            if r.get("Document")]
    return [{"type": "_structure", "headings": headings, "documents": docs}] + out


def compare_row(e: casekit.Row, a: dict) -> list:
    e_ref, e_text = norm(e.doc_ref), norm(e.content)
    stem_len = len(norm(e.stem)) + (1 if e.stem else 0)
    full = ratio(a["content"], e_text)
    tail = ratio(tail_after_stem(a["content"], e_text, stem_len), e_text[stem_len:].strip())
    bad = []
    if a["doc_ref"] != e_ref:
        bad.append(f"Doc Ref {a['doc_ref']!r} != {e_ref!r}")
    if full < THRESHOLD or tail < THRESHOLD:
        bad.append(f"text similarity {full:.2f} (after the lead-in {tail:.2f})")
    reasons = a["reasons"].lower()
    if e.note and e.note.lower() not in reasons:
        bad.append(f"review note {e.note!r} missing")
    if e.no_marker_note and ("list marker" in reasons or "the list goes from" in reasons):
        bad.append("unexpected list-marker note")
    if e.section is not None and a["section"] != norm(e.section):
        bad.append(f"Section {a['section']!r} != {e.section!r}")
    if e.pages is not None and a["pages"] != norm(e.pages):
        bad.append(f"Page(s) {a['pages']!r} != {e.pages!r}")
    if e.title is not None and a["title"] != norm(e.title):
        bad.append(f"Title {a['title']!r} != {e.title!r}")
    if e.corrected and e.corrected.lower() not in a["corrections"].lower():
        bad.append(f"OCR Corrections {a['corrections']!r} lack {e.corrected!r}")
    if e.quality and e.quality.lower() not in a["quality"].lower():
        bad.append(f"Quality Notes {a['quality']!r} lack {e.quality!r}")
    return bad


def align(exp: list, got: list) -> list:
    """Pair expected and actual rows by Doc Ref order -> [(expected or None, actual or None)]."""
    a, b = [norm(e.doc_ref) for e in exp], [r["doc_ref"] for r in got]
    pairs = []
    for _op, i1, i2, j1, j2 in difflib.SequenceMatcher(None, a, b, autojunk=False).get_opcodes():
        for k in range(max(i2 - i1, j2 - j1)):
            pairs.append((exp[i1 + k] if i1 + k < i2 else None, got[j1 + k] if j1 + k < j2 else None))
    return pairs


def evaluate(case: casekit.Case, rows: list) -> tuple:
    """-> (passed, one-line reason, detail lines)"""
    exp = case.expect
    structure = next((r for r in rows if r["type"] == "_structure"), {"headings": [], "documents": []})
    reqs = [r for r in rows if r["type"] == "Requirement"]
    tables = [r for r in rows if r["type"] == "Table"]
    captions = [r for r in rows if r["type"] in ("Table Caption", "Figure Caption")]
    problems, detail = [], []
    if len(reqs) != len(exp):
        want, got = Counter(norm(e.doc_ref) for e in exp), Counter(r["doc_ref"] for r in reqs)
        missing, extra = list((want - got).elements()), list((got - want).elements())
        msg = f"{len(reqs)} rows, expected {len(exp)}"
        if missing:
            msg += "; missing " + ", ".join(repr(m) for m in missing)
        if extra:
            msg += "; unexpected " + ", ".join(repr(x) for x in extra)
        problems.append(msg)
    else:
        for k, (e, a) in enumerate(zip(exp, reqs), 1):
            bad = compare_row(e, a)
            if bad:
                problems.append(f"row {k}: " + ", ".join(bad))
    # the listing aligns rows by Doc Ref, so one lost row does not shift all the others
    for e, a in align(exp, reqs):
        bad = compare_row(e, a) if e and a else []
        mark = "!!" if bad or not (e and a) else "  "
        checks = f"   ({e.label()})" if e and e.label() else ""
        detail.append(f"   {mark} expected  " + (f"[{norm(e.doc_ref)}] {norm(e.content)}{checks}" if e else "(none)"))
        detail.append(f"   {mark} actual    " + (f"[{a['doc_ref']}] {a['content']}" if a else "(none)")
                      + (f"   <- {'; '.join(bad)}" if bad else "")
                      + (f"   {{review: {norm(a['reasons'])}}}" if a and a["reasons"].strip() else ""))
    if len(tables) != case.tables:
        problems.append(f"{len(tables)} Table item(s), expected {case.tables}")
    for want in case.table_has:
        if not any(norm(want) in t["content"] for t in tables):
            problems.append(f"no table contains {want!r}")
    for t in tables:
        detail.append(f"   {'  ' if len(tables) == case.tables else '!!'} TABLE {t['doc_ref']}: {t['content'][:160]}")
    if case.captions is not None:
        got_caps = sorted(c["doc_ref"] for c in captions)
        if got_caps != sorted(case.captions):
            problems.append(f"captions {got_caps}, expected {sorted(case.captions)}")
    for c in captions:
        detail.append(f"      CAPTION {c['doc_ref']}: {c['content'][:120]}")
    if case.outline is not None:
        want = [norm(h) for h in case.outline]
        if structure["headings"] != want:
            problems.append(f"headings {structure['headings']}, expected {want}")
        for h in structure["headings"]:
            detail.append(f"   {'  ' if h in want else '!!'} HEADING {h}")
    if case.documents is not None:
        want = [norm(x) for x in case.documents]
        got = structure["documents"]

        def same(a: str, b: str) -> bool:
            (ai, at, ap), (bi, bt, bp) = (a.split(" | ") + ["", ""])[:3], (b.split(" | ") + ["", ""])[:3]
            return ai == bi and ap == bp and ratio(at, bt) >= THRESHOLD
        if len(got) != len(want) or not all(same(a, b) for a, b in zip(got, want)):
            problems.append(f"documents {got}, expected {want}")
        for x in got:
            detail.append(f"      DOCUMENT {x}")
    return (not problems), "; ".join(problems[:3]) + (" ..." if len(problems) > 3 else ""), detail


def record(spec: Path, scan: Path) -> int:
    """Run the tool on your own scan and write <scan>.expected.txt beside it: what the tool produces NOW, one line
    per expectation, for you to correct to what it SHOULD produce. An existing file is never overwritten. The scan
    must be in my_cases (git ignores that folder, and only scans there are run as cases); the tool's output for it
    is deleted once the file is written."""
    if not scan.is_file() or scan.suffix.lower() not in casekit.SCAN_TYPES:
        print(f"no scan here: {scan} (a PDF, TIFF, PNG or JPG file)")
        return 2
    if scan.parent != casekit.LOCAL_FOLDER.resolve():
        print(f"{scan.name} is not in {casekit.LOCAL_FOLDER}: put it there first, not in a folder inside it (git "
              f"ignores that folder; a file written anywhere else could be committed, and only scans directly in it "
              f"are run as cases)")
        return 2
    expected = scan.with_name(scan.name + ".expected.txt")
    if expected.exists():
        print(f"{expected.name} already exists - correct it by hand, or delete it to record again")
        return 2
    target = OWN_RESULTS / "_record"
    shutil.rmtree(target, ignore_errors=True)
    target.mkdir(parents=True, exist_ok=True)
    try:
        run_tool(spec, [scan], target, 1)
        xlsx = find_workbook(target, scan)
        if xlsx is None:
            print("the tool wrote no workbook - see", target / "run1" / "tool_output.txt")
            return 2
        rows = read_items(xlsx)
    finally:
        if (target / "run1" / f"{scan.stem}_extracted.xlsx").exists():
            shutil.rmtree(target, ignore_errors=True)     # (kept when the tool failed: its tool_output.txt says why)
    structure = rows[0]
    lines = [f"# What spec_extractor produced for {scan.name} on {time.strftime('%Y-%m-%d')}.",
             "# CORRECT THIS FILE to what the tool SHOULD produce, then run:  python tests/cases/run_cases.py --only MY",
             "# One expectation per line ('#' starts a comment). Kinds:",
             "#   requirement: <Doc Ref> | <Content>      [title: ...] [note: ...] [corrected: ...] may follow",
             "#   tables: <number of tables>    table has: <text a table holds>    caption: <TABLE I>",
             "#   heading: <level> | <number and title>   document: <number> | <title> | <paragraphs>", ""]
    for r in rows[1:]:
        if r["type"] == "Requirement":
            checks = f"  [title: {r['title']}]" if r["title"] else ""
            lines.append(f"requirement: {r['doc_ref']} | {r['content']}{checks}")
    tables = [r for r in rows[1:] if r["type"] == "Table"]
    lines += ["", f"tables: {len(tables)}"]
    lines += [f"caption: {r['doc_ref']}" for r in rows[1:] if r["type"] in ("Table Caption", "Figure Caption") and r["doc_ref"]]
    lines += [""] + [f"heading: {h}" for h in structure["headings"]]
    lines += [""] + [f"document: {d}" for d in structure["documents"]]
    expected.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"wrote {expected}\nOpen it, correct it to what the tool should produce, and run:  "
          f"python tests/cases/run_cases.py --only MY")
    return 0


def status_of(case: casekit.Case, ok: bool) -> str:
    if case.known:
        return "KNOWN-PASSES" if ok else "KNOWN"
    return "PASS" if ok else "FAIL"


def write_workbook(path: Path, results: list) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Case results"
    ws.append(["Case", "Area", "Title", "Result", "Problem", "Known limitation", "Document"])
    for c, st, why in results:
        ws.append([c.id, c.area, c.title, st, why, c.known, c.filename])
    for col, width in zip("ABCDEFG", (9, 24, 60, 13, 70, 50, 60)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    with _temporary_files_in(path.parent):
        wb.save(path)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Run the test-case library against spec_extractor.py")
    ap.add_argument("spec", nargs="?", default=str(HERE.parent.parent / "spec_extractor.py"),
                    help="path of the spec_extractor.py to test (default: the one in this project)")
    ap.add_argument("outdir", nargs="?", default=str(HERE / "_results"),
                    help="folder for the tool's output and the reports (default: tests/cases/_results)")
    ap.add_argument("--only", default="", help="only these cases: an area (TAB), a case (TAB-07) or words of a title; "
                                                "several with commas")
    ap.add_argument("--jobs", type=int, default=max(1, (os.cpu_count() or 2) // 2), help="parallel tool processes")
    ap.add_argument("--reuse", action="store_true", help="re-check existing workbooks in outdir without re-running")
    ap.add_argument("-v", "--verbose", action="store_true", help="print expected/actual rows of failing cases")
    ap.add_argument("--record", metavar="SCAN", help="turn your own scan (in tests/cases/my_cases) into a case: write "
                                                      "<SCAN>.expected.txt with what the tool produces now")
    args = ap.parse_args(argv)
    if args.record:
        return record(Path(args.spec).resolve(), Path(args.record).resolve())

    spec, outdir = Path(args.spec).resolve(), Path(args.outdir).resolve()
    if not spec.is_file():
        ap.error(f"no such file: {spec}")
    try:
        cases = casekit.load_cases() + casekit.load_local_cases()
    except ValueError as exc:                         # a case that cannot be read, or two scans with one name
        print(exc)
        return 2
    wanted = [w.strip() for w in args.only.split(",") if w.strip()]

    def chosen(c) -> bool:
        for w in wanted:
            if c.id.lower() == w.lower():
                return True                                         # a case's ID, also your own (MY-PAGE12)
            if re.fullmatch(r"[A-Za-z]{2,4}(-\d{2}[A-Za-z]?)?", w):      # an area or a case ID
                if c.id.lower().startswith(w.lower() + "-"):
                    return True
            elif w.lower() in c.title.lower():
                return True
        return not wanted

    selected = [c for c in cases if chosen(c)]
    if not selected:
        print(f"no case matches {args.only!r}")
        return 2
    docs = build_cases.build_documents(selected)
    counts = Counter()
    library, own = [c for c in selected if not c.scan], [c for c in selected if c.scan]
    if library:
        counts += check(spec, library, docs, outdir, args)
    if own:
        if library:
            print()
        # (a folder of their own for each outdir: a run of another copy of the tool into another outdir has its own)
        own_dir = OWN_RESULTS / outdir.name
        print(f"your own scans (their results stay in {own_dir}):")
        counts += check(spec, own, docs, own_dir, args)
    return 0 if counts["FAIL"] == 0 else 1


def check(spec: Path, selected: list, docs: dict, outdir: Path, args) -> Counter:
    """Run the tool on these cases' documents into outdir, compare its output with each case, print the result of
    each and write the reports (case_report.txt, case_results.json / .xlsx) there. -> the count of each result."""
    outdir.mkdir(parents=True, exist_ok=True)
    t0 = time.time()
    results_file = outdir / "case_results.json"
    previous, previous_spec = {}, None
    if results_file.exists():
        try:
            saved = json.loads(results_file.read_text(encoding="utf-8"))
            previous, previous_spec = saved.get("cases", {}), saved.get("spec")
        except (ValueError, OSError):
            previous = {}
    # --reuse re-checks only the workbooks this copy of the tool made (each case's result records which made it)
    made_here = {c.id for c in selected if previous.get(c.id, {}).get("spec", previous_spec) == str(spec)}
    if args.reuse and len(made_here) < len(selected):
        print(f"(some workbooks in {outdir} were made by another copy of the tool, or by an unknown one: running it "
              f"again on those instead of --reuse)", flush=True)
    todo = [docs[c.id] for c in selected
            if not (args.reuse and c.id in made_here and find_workbook(outdir, docs[c.id]))]
    if todo:
        print(f"running {spec.name} on {len(todo)} case document(s) -> {outdir}", flush=True)
        for d in todo:                                # never compare against a stale workbook
            for old in workbooks_of(outdir, d):
                old.unlink()
        run_tool(spec, todo, outdir, args.jobs)
    report, results, counts = [f"spec_extractor under test: {spec}", ""], [], Counter()
    for c in selected:
        xlsx = find_workbook(outdir, docs[c.id])
        if xlsx is None:
            ok, why, detail = False, "no Excel output for this document (did the tool fail?)", []
        else:
            ok, why, detail = evaluate(c, read_items(xlsx))
        st = status_of(c, ok)
        counts[st] += 1
        results.append((c, st, "" if ok else why))
        line = f"{st:<12s} {c.id:<8s} {c.title}" + ("" if ok else f"  - {why}")
        print(line, flush=True)
        report += [line, f"   document: {docs[c.id].name}", f"   situation: {c.situation}"]
        if c.known:
            report.append(f"   known limitation: {c.known}")
        report += detail + [""]
        if args.verbose and not ok:
            print("\n".join(detail))
    changes = [(c.id, previous[c.id]["status"], st) for c, st, _ in results
               if c.id in previous and previous[c.id]["status"] != st]
    summary = (f"TOTAL: {counts['PASS']} passed, {counts['FAIL']} failed, {counts['KNOWN']} known limitations"
               + (f", {counts['KNOWN-PASSES']} known limitation(s) now passing" if counts["KNOWN-PASSES"] else "")
               + f"  ({time.time() - t0:.0f} s)")
    print(summary)
    report.append(summary)
    if changes:
        print("Changed since the last run:")
        report.append("Changed since the last run:")
        for cid, before, after in changes:
            print(f"   {cid}: {before} -> {after}")
            report.append(f"   {cid}: {before} -> {after}")
    own = any(c.scan for c in selected)           # (a folder holds the results of your own scans or of the library)
    kept = {k: v for k, v in previous.items() if k.startswith("MY-") == own}
    kept.update({c.id: {"status": st, "problem": why, "title": c.title, "spec": str(spec)} for c, st, why in results})
    results_file.write_text(json.dumps({"spec": str(spec), "time": time.strftime("%Y-%m-%d %H:%M"), "cases": kept},
                                       indent=1), encoding="utf-8")
    (outdir / "case_report.txt").write_text("\n".join(report) + "\n", encoding="utf-8")
    write_workbook(outdir / "case_results.xlsx", results)
    print(f"details: {outdir / 'case_report.txt'}  (table: {outdir / 'case_results.xlsx'})")
    return counts


if __name__ == "__main__":
    sys.exit(main())
