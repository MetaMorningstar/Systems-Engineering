#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
===============================================================================
 SPEC EXTRACTOR  -  requirements, tables and captions from scanned documents
===============================================================================

WHAT IT DOES
  1. Reads scanned PDF and TIFF files (PNG / JPG / BMP also work), page by page.
  2. Cleans every page image: fixes sideways or upside-down pages, straightens
     skew, flattens stains/shadows/yellowed paper, normalizes resolution.
  3. Reads the text with a modern deep-learning OCR model (PaddleOCR PP-OCRv6,
     run through RapidOCR + ONNX Runtime). It runs on an ordinary CPU and needs
     NO internet connection - the model files ship inside the pip package, and
     ONNX Runtime's usage telemetry is switched off (see section 1).
  4. Rebuilds the page layout: headings, paragraphs, lists, tables (cell by
     cell), figures, captions, page headers/footers, in reading order.
  5. Finds Systems Engineering requirements written as paragraphs ("shall"
     statements), list items under a stem ("The system shall: a. ... b. ..."),
     and rows of requirement tables. Also collects every table and caption.
  6. Writes, for each input document:
       <name>_extracted.xlsx     one row per requirement, caption and table
                                 (each table also gets its own worksheet)
       <name>_verification.docx  the document rebuilt page by page next to
                                 the scanned image, with words worth checking
                                 highlighted, to check nothing was missed
     plus ALL_DOCUMENTS_<date>.xlsx when more than one document is processed.

QUICK START
  1. Install a 64-bit Python 3.11 or 3.12 from python.org (3.10, 3.13 and 3.14
     also work - see "Python versions" in README.md). Tick "Add python.exe to
     PATH" if the installer offers it.
  2. Windows: double-click setup_windows.bat. Other systems: open a terminal in
     this folder and run
         python -m pip install -r requirements.txt -c constraints.txt
     (constraints.txt holds the exact package versions the tool was tested with)
  3. Run one of:
         python spec_extractor.py "C:\\Scans\\MySpec.pdf"      (one file)
         python spec_extractor.py "C:\\Scans"                  (every file in a folder)
         python spec_extractor.py                             (opens a file picker)
     Results appear in an "extraction_output" folder next to the input.

USEFUL OPTIONS  (run  python spec_extractor.py --help  for the full list)
  -o / --output FOLDER      put results somewhere else
  --pages 1-5,9             only process some pages (handy for a first test)
  --keywords "shall;must"   words that mark a requirement, separated by commas or
                            semicolons (default: shall, must, is/are required to,
                            should, will)
  --no-will                 do not treat "will" statements as requirements
  --keep-list-stems         one row per stem+list block instead of per item
  --review-threshold 0.70   flag numbers read below this OCR confidence (default 0.70)
  --engine docling          optional: IBM Docling AI layout + table models
                            (needs extra install + one-time model download)
  --no-word                 skip the Word verification report
  --debug-images            also save the cleaned page images

HOW TO CHECK THE RESULTS (recommended routine)
  * Open <name>_verification.docx. Each page shows the scan, then the text the
    tool read. Yellow highlight = a word worth checking against the scan (it
    failed a spelling / punctuation / look-alike check, or is a number the OCR
    was unsure of); red = very unsure. Blue tags like [R-0012] show which
    Excel row a sentence became.
  * In <name>_extracted.xlsx filter "Needs Review" = Yes first. The
    "Review Reasons" column says why (suspect word or number, keyword that had
    to be repaired, text that continued across a page, list stem without items).
  * "Page Quality" sheet: pages graded Poor deserve a closer look; a page graded
    FAILED could not be read at all (the other pages still were).

DESIGN NOTES (for maintainers)
  * Heavy lifting is done by proven, permissively licensed libraries:
      RapidOCR (Apache-2.0) running PaddleOCR PP-OCRv6 models (Apache-2.0)
      img2table (MIT)       - OpenCV ruled / borderless table detection
      OpenCV (Apache-2.0), pypdfium2 (Apache-2.0/BSD), Pillow (MIT-CMU)
      openpyxl (MIT), python-docx (MIT); optional Docling (MIT)
    AGPL/GPL and revenue-restricted model licenses (e.g. PyMuPDF, MinerU,
    Marker/Surya weights) were deliberately avoided.
  * Generative "vision-language" OCR models were NOT used for transcription:
    they read degraded text well but can silently invent or "correct" words
    and numbers, which is unacceptable for requirements. PP-OCRv6 gives a
    confidence score for every word, so uncertain text can be flagged.
  * Requirement detection is rule-based (deterministic and auditable):
    modal verbs per MIL-STD-961 / INCOSE usage, paragraph numbering, list
    stems, requirement-table column detection. All rules are in section 9.
  * OCR results are cached per file, so re-running with different requirement
    settings is fast. Delete the "_cache" folder (or use --no-cache) to force
    a fresh OCR. Editing the OCR/layout code (sections 3-6, 8) invalidates the
    cache automatically.
  * Tests: tests/test_rules.py (seconds, layout + requirement rules) and
    tests/run_regression.py (minutes, whole tool on synthetic aged scans with
    exact ground truth). Run both after any change.
  * Verified configurations only: both test levels passed on 64-bit Python
    3.10 to 3.14 with the exact package versions in constraints.txt. Every run
    compares its own Python and package versions with that tested set and
    reports the result in the log, the "Run Info" sheet and the Word report
    (section 1b). The page-orientation classifier is only installed where its
    maker supports it (Python 3.12 and older); 3.13 and 3.14 use the slower,
    equally accurate OCR-based orientation check.

FILE LAYOUT OF THIS SCRIPT
   1. Imports and dependency check          7. Document-level clean-up
   2. Settings (defaults you can change)    8. Optional Docling engine
   3. Data model                            9. Requirement / caption / table extraction
   4. Loading pages from PDF / TIFF        10. Excel export
   5. Image clean-up (preprocessing)       11. Word verification report
   6. OCR + page layout analysis           12. Batch runner, caching, command line
"""
from __future__ import annotations

# =============================================================================
# 1. IMPORTS AND DEPENDENCY CHECK
# =============================================================================
import argparse
import copy
import dataclasses
import datetime as _dt
import functools
import hashlib
import io
import json
import logging
import math
import os
import posixpath
import re
import statistics
import sys
import sysconfig
import time
import traceback
import unicodedata
from collections import Counter, defaultdict, namedtuple
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, Iterator, Optional

__version__ = "1.3.0"


@functools.lru_cache(maxsize=1)
def build_id() -> str:
    """A short identifier of the code that is actually running, printed beside the version.

    It is a fingerprint of this file's own bytes, so it changes whenever ANY line of the tool
    changes - a merge, a hand edit, a half-applied patch - without the version number having to be
    remembered. Two people can compare build ids to be sure they are running the same thing. When
    the folder is a git checkout the commit is shown as well, which is what identifies a merge;
    a folder handed out without git still gets the fingerprint. Neither needs the network."""
    try:
        fingerprint = hashlib.sha1(Path(__file__).read_bytes()).hexdigest()[:7]
    except OSError:
        return "unknown"
    commit = ""
    try:                                   # read .git directly: no git command, no subprocess
        head = (Path(__file__).resolve().parent.parent / ".git" / "HEAD").read_text(encoding="utf-8").strip()
        if head.startswith("ref:"):
            ref = Path(__file__).resolve().parent.parent / ".git" / head.split(" ", 1)[1].strip()
            commit = ref.read_text(encoding="utf-8").strip()[:7] if ref.exists() else ""
        else:
            commit = head[:7]
    except (OSError, IndexError):
        commit = ""
    return f"{fingerprint} commit {commit}" if commit else fingerprint


def version_line() -> str:
    """'spec_extractor 1.3.0 (build a1b2c3d)' - the one string every report shows."""
    return f"spec_extractor {__version__} (build {build_id()})"

# The Python versions this release passed both test levels on (64-bit, standard build).
# Change this only after running tests/test_rules.py and tests/run_regression.py on the new version
# (and update setup_windows.bat, constraints.txt and README.md to match).
SUPPORTED_PYTHON = ((3, 10), (3, 14))          # (oldest, newest) as (major, minor)


def python_support_problem(version: Optional[tuple] = None, is_64bit: Optional[bool] = None,
                           free_threaded: Optional[bool] = None) -> str:
    """Why this Python is not one of the tested ones, as one sentence ("" when it is tested).
    The arguments exist for the tests; normally the running Python is examined."""
    version = tuple(sys.version_info[:3]) if version is None else tuple(version)
    is_64bit = sys.maxsize > 2 ** 32 if is_64bit is None else is_64bit
    if free_threaded is None:            # the optional build without the GIL, e.g. python3.14t
        free_threaded = bool(sysconfig.get_config_var("Py_GIL_DISABLED"))
    name = "Python " + ".".join(str(part) for part in version)
    lowest, highest = SUPPORTED_PYTHON
    tested = f"a 64-bit Python {lowest[0]}.{lowest[1]} to {highest[0]}.{highest[1]} in its standard build"
    if free_threaded:
        return f"{name} is the free-threaded build, but Spec Extractor was tested on {tested}"
    if not is_64bit:
        return f"{name} is a 32-bit build, but Spec Extractor was tested on {tested}"
    if not lowest <= version[:2] <= highest:
        return f"{name} is not a tested version - Spec Extractor was tested on {tested}"
    return ""


# Keep all traffic on this computer. ONNX Runtime (which runs the OCR models) otherwise sends usage
# telemetry to Microsoft - measured: 4 connection attempts per run, no document content - and the
# Hugging Face libraries used by the optional Docling engine report downloads. These switches must be
# set before those libraries load. (An administrator can still override them in the environment.)
os.environ.setdefault("ORT_DISABLE_TELEMETRY", "1")
os.environ.setdefault("HF_HUB_DISABLE_TELEMETRY", "1")
os.environ.setdefault("DO_NOT_TRACK", "1")

_MISSING = []
try:
    import numpy as np
except ImportError:  # pragma: no cover - only hit on a broken install
    _MISSING.append("numpy")
try:
    import cv2
except ImportError:  # pragma: no cover
    _MISSING.append("opencv-contrib-python (installed automatically with img2table)")
try:
    from PIL import Image
except ImportError:  # pragma: no cover
    _MISSING.append("pillow")
try:
    import pypdfium2 as pdfium
    import pypdfium2.raw as pdfium_c
except ImportError:  # pragma: no cover
    _MISSING.append("pypdfium2")
try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table as XlTable, TableStyleInfo
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.worksheet.hyperlink import Hyperlink
except ImportError:  # pragma: no cover
    _MISSING.append("openpyxl")
try:
    import docx  # python-docx
    from docx.enum.section import WD_ORIENT
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_LINE_SPACING
    from docx.shared import Emu, Inches, Pt, RGBColor
except ImportError:  # pragma: no cover
    _MISSING.append("python-docx")

if _MISSING:  # pragma: no cover - only reached on an incomplete installation
    _problem = python_support_problem()
    sys.stderr.write(
        "\nMissing Python packages: " + ", ".join(_MISSING) + "\n"
        # an unsupported Python is the usual reason the packages could not be installed: say so plainly
        + (f"{_problem}.\nThe OCR and table packages cannot be installed on it, or were not tested with it.\n"
           "Install the 64-bit Python 3.12 from python.org (it can sit next to other Python versions),\n"
           "then set up again as described below.\n" if _problem else "")
        + "Windows:        double-click setup_windows.bat in the spec_extractor folder\n"
          "Other systems:  python -m pip install -r requirements.txt -c constraints.txt\n\n")
    sys.exit(2)

# Scanned engineering drawings can be very large; these are trusted local files.
Image.MAX_IMAGE_PIXELS = 400_000_000

LOG = logging.getLogger("spec_extractor")


# ---- 1b. Tested configuration: which Python and package versions this run uses ----------------------------
# constraints.txt (next to this script) pins every package to the version that passed the tests. Each run
# compares the installed versions of the packages that decide what the tool reads and writes with that file
# and states the result in the log, the Excel "Run Info" sheet and the Word report. A difference does not
# stop the run, but the documented accuracy was not measured for it, so it is reported prominently.
CONSTRAINTS_FILE = Path(__file__).resolve().with_name("constraints.txt")
# How the Word report looks (fonts, sizes, spacing, page layout, your own Word template): see section 11
DEFAULT_WORD_STYLE_FILE = Path(__file__).resolve().with_name("word_style") / "report_style.ini"
RESULT_PACKAGES = ("rapidocr", "onnxruntime", "opencv-python", "opencv-contrib-python", "img2table", "numpy",
                   "pillow", "pypdfium2", "pyclipper", "shapely", "rapid-orientation", "pyspellchecker",
                   "openpyxl", "python-docx")
DOCLING_NOTE = "the optional Docling engine was used - its accuracy has not been measured"
# one condition of a constraints.txt line, e.g.   python_version < "3.11"   or   sys_platform == "win32"
_CONDITION = re.compile(r"""^(python_version|sys_platform)\s*(==|!=|<=|>=|<|>)\s*["']([^"']+)["']$""")


def canonical_package_name(name: str) -> str:
    """Compare package names the way pip does: 'Rapid_Orientation' is the same as 'rapid-orientation'."""
    return re.sub(r"[-_.]+", "-", name.strip()).lower()


def _conditions_hold(conditions: str, python: tuple, platform: str) -> bool:
    """True when every condition of a constraints.txt line (joined with 'and') holds for this computer."""
    for condition in conditions.split(" and "):
        m = _CONDITION.match(condition.strip())
        if not m:
            raise ValueError(f"constraints.txt: unsupported condition {condition.strip()!r} "
                             "(use python_version or sys_platform comparisons joined with 'and')")
        key, op, value = m.groups()
        if key == "python_version":        # compared as numbers, so "3.10" comes after "3.9"
            left, right = tuple(python[:2]), tuple(int(part) for part in value.split("."))
        else:
            left, right = platform, value
        holds = {"==": left == right, "!=": left != right, "<": left < right,
                 "<=": left <= right, ">": left > right, ">=": left >= right}[op]
        if not holds:
            return False
    return True


def tested_versions(path: Optional[Path] = None, python: Optional[tuple] = None,
                    platform: Optional[str] = None) -> dict:
    """{package: version} that constraints.txt pins for this Python and operating system.
    Returns {} when the file is missing (e.g. only spec_extractor.py was copied somewhere else)."""
    path = CONSTRAINTS_FILE if path is None else Path(path)
    python = tuple(sys.version_info[:2]) if python is None else tuple(python)
    platform = sys.platform if platform is None else platform
    try:
        lines = path.read_text(encoding="utf-8").splitlines()
    except OSError:
        return {}
    pinned = {}
    for line in lines:
        entry = line.split("#", 1)[0].strip()             # drop comments and blank lines
        if not entry:
            continue
        requirement, _, conditions = entry.partition(";")
        name, is_pin, version = requirement.partition("==")
        if not (is_pin and name.strip() and version.strip()):
            raise ValueError(f"constraints.txt: expected 'name==version', got {line.strip()!r}")
        if conditions.strip() and not _conditions_hold(conditions, python, platform):
            continue                                       # this line is for another Python / system
        pinned[canonical_package_name(name)] = version.strip()
    return pinned


def installed_versions(names: Iterable = RESULT_PACKAGES) -> dict:
    """{package: installed version, or None when it is not installed}."""
    from importlib import metadata
    found = {}
    for name in names:
        try:
            found[canonical_package_name(name)] = metadata.version(name)
        except metadata.PackageNotFoundError:
            found[canonical_package_name(name)] = None
    return found


def configuration_differences(installed: dict, tested: dict, python_problem: str = "") -> list:
    """Every way this computer differs from the tested configuration, as readable sentences.
    An empty list means: a tested Python and exactly the tested version of every checked package."""
    differences = [python_problem] if python_problem else []
    if not tested:
        differences.append(f"{CONSTRAINTS_FILE.name} was not found next to spec_extractor.py, "
                           "so the package versions could not be checked")
        return differences
    for name, have in installed.items():
        want = tested.get(name)
        if want is None and have is not None:        # e.g. a package forced onto a Python it does not support
            differences.append(f"{name} {have} is installed, but it is not part of the tested set for this Python")
        elif want is not None and have is None:
            differences.append(f"{name} is not installed (tested with {want})")
        elif want is not None and have != want:
            differences.append(f"{name} {have} is installed (tested with {want})")
    return differences


@functools.lru_cache(maxsize=1)
def _environment_status() -> tuple:
    """(software versions text, differences) for this computer - worked out once per run."""
    python = "Python " + ".".join(map(str, sys.version_info[:3])) + (" 64-bit" if sys.maxsize > 2 ** 32 else " 32-bit")
    try:
        installed = installed_versions()
        software = "; ".join([python] + [f"{n} {v}" if v else f"{n} (not installed)" for n, v in installed.items()])
        differences = configuration_differences(installed, tested_versions(), python_support_problem())
    except Exception as exc:                        # the check itself must never stop a run
        software, differences = python, [f"the package versions could not be checked ({exc})"]
    return software, tuple(differences)


def configuration_status(docling_used: bool = False) -> tuple:
    """-> (software versions text, [differences from the tested configuration]) for the log and reports."""
    software, differences = _environment_status()
    return software, list(differences) + ([DOCLING_NOTE] if docling_used else [])


# =============================================================================
# 2. SETTINGS  -  every tunable value lives here, with an explanation.
#    Command-line options override these defaults (see section 12).
# =============================================================================
@dataclass
class Settings:
    # ---- Input / output -----------------------------------------------------
    recursive: bool = False             # also look inside sub-folders
    output_dir: Optional[Path] = None   # default: <input folder>/extraction_output
    page_selection: str = ""            # e.g. "1-3,7"; empty = all pages
    use_cache: bool = True              # reuse OCR results for unchanged files

    # ---- Page images ----------------------------------------------------------
    # OCR works best around 200-220 dpi for PP-OCR (tested on degraded scans:
    # 300 dpi was no more accurate and took twice as long). Pages are resized
    # toward this resolution; very low-resolution scans are gently upscaled.
    target_dpi: int = 220
    fallback_dpi: int = 300             # assumed when a file has no resolution info
    max_page_megapixels: float = 60.0   # huge drawings are scaled down to this
    auto_rotate: bool = True            # fix pages scanned sideways / upside-down
    deskew: bool = True                 # straighten slightly rotated scans
    max_skew_degrees: float = 8.0       # largest skew the deskew step searches for
    clean_background: bool = True       # flatten stains, shadows, yellowed paper
    denoise: bool = True                # light median filter against speckle
    remove_color_marks: bool = False    # fade colored stamps/highlighter (keeps black text)

    # ---- OCR --------------------------------------------------------------------
    engine: str = "standard"            # "standard" (pip-only) or "docling" (optional)
    ocr_model_size: str = "small"       # PP-OCRv6 "small" ships with RapidOCR (no download)
    min_text_confidence: float = 0.30   # text below this is treated as noise and dropped
    # Tested on degraded scans: most real misreads still score > 0.95, so confidence alone is a
    # weak signal. Words are ALSO flagged by independent checks (dictionary, accent marks,
    # punctuation inside words, numbers with letters) - see WordChecker in section 7.
    review_confidence: float = 0.70     # words below this OCR confidence are flagged "Needs Review"
    very_low_confidence: float = 0.50   # shown in red in the Word report
    english_text: bool = True           # documents are English: remove accent marks OCR adds (telémetry)
    spell_check_flags: bool = True      # flag words missing from an English dictionary (never auto-corrects)
    # your program's own words that the dictionary check should accept, e.g. {"xr7", "radome"}
    extra_dictionary_words: set = field(default_factory=set)
    ocr_threads: int = 0                # 0 = let ONNX Runtime decide

    # ---- Tables -------------------------------------------------------------------
    detect_tables: bool = True
    borderless_tables: bool = True      # also look for tables without ruling lines
    merge_continued_tables: bool = True # join a table that continues on the next page
    # Many specifications rule a table's columns but print no line between the rows; verification
    # cross-reference matrices nearly always do. A row only exists between two horizontal lines, so
    # such a table is otherwise missed completely. With this on, the separators that were never
    # printed are worked out from where the text sits - but only where the printed lines clearly do
    # not describe the table, so a table that draws all of its lines keeps exactly the grid it had.
    infer_missing_rules: bool = True
    # Column headers printed sideways (turned a quarter turn, because the words are far wider than
    # the columns they head) are normal in verification matrices, and the recogniser reads upright
    # text only: it returns nothing for such a band, or short fragments ("INS OON" for INSPECTION).
    # Either way the band is cut out, turned upright and read again.
    sideways_headers: bool = True

    # ---- Requirements -------------------------------------------------------------
    # keyword -> requirement class. Matching ignores case. Multi-word phrases are fine.
    keywords: dict = field(default_factory=lambda: {
        "shall": "Mandatory",
        "must": "Mandatory",
        "is required to": "Mandatory",
        "are required to": "Mandatory",
        "should": "Goal / recommended",
        "will": "Statement of intent",
    })
    split_list_items: bool = True       # one row per list item ("stem + item")
    fix_ocr_keywords: bool = True       # repair OCR-garbled keywords (e.g. "sha11" -> "shall")
    quality_checks: bool = True         # INCOSE-style notes (TBD, vague terms, and/or ...)
    # a table column with one of these headers holds requirements (every row becomes a requirement)
    requirement_table_headers: tuple = ("requirement", "requirements", "req text", "req. text",
                                        "shall statement", "requirement statement")
    id_column_headers: tuple = ("req id", "req. id", "requirement id", "req no", "req. no", "req. no.",
                                "req #", "id", "no.", "no", "number", "paragraph", "para", "para.",
                                "ref", "ref.", "reference", "item", "req")

    # ---- Outputs ----------------------------------------------------------------------
    make_excel: bool = True
    make_word_report: bool = True
    word_page_images: bool = True       # put the scanned page image in the Word report
    word_image_dpi: int = 110           # resolution of those images (file size vs. clarity)
    # How the Word report LOOKS - fonts, sizes, spacing, page layout and your own Word template - is set in
    # word_style/report_style.ini (explained inside that file), not here. None = the tool's own look, no file.
    word_style_file: Optional[Path] = DEFAULT_WORD_STYLE_FILE
    word_look: Optional["ReportLook"] = field(default=None, repr=False, compare=False)  # (set at start-up)
    combined_workbook: bool = True      # ALL_DOCUMENTS_<date>.xlsx when >1 document
    save_debug_images: bool = False

    # ---- Optional Docling engine -----------------------------------------------------
    docling_models_dir: Optional[Path] = None  # offline model folder (docling-tools models download)
    docling_device: str = "auto"               # auto / cpu / cuda / mps

    def ocr_fingerprint(self) -> str:
        """Settings that change OCR/layout results (used for the cache key)."""
        keys = ["target_dpi", "fallback_dpi", "max_page_megapixels", "auto_rotate", "deskew",
                "max_skew_degrees", "clean_background", "denoise", "remove_color_marks", "engine",
                "ocr_model_size", "min_text_confidence", "detect_tables", "borderless_tables",
                "infer_missing_rules", "sideways_headers"]
        values = {k: getattr(self, k) for k in keys}
        values["keywords"] = sorted(self.keywords)      # the layout rules protect keyword sentences
        blob = json.dumps(values, sort_keys=True, default=str)
        return hashlib.sha1((blob + CACHE_FORMAT + __version__ + _cached_code_hash()).encode()).hexdigest()[:10]


CACHE_FORMAT = "spec-extractor-cache-v1"


def _cached_code_hash() -> str:
    """Fingerprint of the code whose results are cached (data model, page loading, image clean-up, OCR,
    page layout, Docling mapping: sections 3-6 and 8). Editing that code makes old caches unused
    automatically; changes to the requirement rules and exports (sections 7, 9-12) reuse the OCR cache."""
    try:
        src = Path(__file__).read_text(encoding="utf-8")

        def banner(title):                # position of a section title line (not this function's text)
            return re.search(r"^# " + re.escape(title), src, re.M).start()

        cached = src[banner("3. DATA MODEL"):banner("7. DOCUMENT-LEVEL CLEAN-UP")]
        cached += src[banner("8. OPTIONAL ENGINE: DOCLING"):banner("9. REQUIREMENT / CAPTION / TABLE EXTRACTION")]
        return hashlib.sha1(cached.encode("utf-8")).hexdigest()[:8]
    except (OSError, AttributeError):
        return ""
SUPPORTED_EXTENSIONS = {".pdf", ".tif", ".tiff", ".png", ".jpg", ".jpeg", ".bmp"}


# =============================================================================
# 3. DATA MODEL
#    Everything the engines produce is stored in these small classes, so the
#    requirement extraction and the exports do not care which engine ran.
#    Boxes are (x0, y0, x1, y1) in pixels of the cleaned page image, origin top-left.
# =============================================================================
Box = tuple  # (x0, y0, x1, y1)


def box_union(boxes: Iterable[Box]) -> Box:
    boxes = list(boxes)
    return (min(b[0] for b in boxes), min(b[1] for b in boxes),
            max(b[2] for b in boxes), max(b[3] for b in boxes))


def box_area(b: Box) -> float:
    return max(0, b[2] - b[0]) * max(0, b[3] - b[1])


def box_intersection(a: Box, b: Box) -> float:
    return max(0, min(a[2], b[2]) - max(a[0], b[0])) * max(0, min(a[3], b[3]) - max(a[1], b[1]))


def box_center_inside(inner: Box, outer: Box, pad: float = 0) -> bool:
    cx, cy = (inner[0] + inner[2]) / 2, (inner[1] + inner[3]) / 2
    return outer[0] - pad <= cx <= outer[2] + pad and outer[1] - pad <= cy <= outer[3] + pad


@dataclass
class Word:
    text: str
    conf: float
    box: Box


@dataclass
class Line:
    """One OCR text line (a horizontal run of words)."""
    text: str
    conf: float
    box: Box
    words: list = field(default_factory=list)   # list[Word]

    @property
    def height(self) -> float:
        return self.box[3] - self.box[1]

    @property
    def min_word_conf(self) -> float:
        return min((w.conf for w in self.words), default=self.conf)


@dataclass
class TableData:
    """A table as a rectangular grid. Merged cells repeat their text in every
    grid position they cover and are listed in `merges` (r0, c0, r1, c1)."""
    grid: list                                  # list[list[str]]
    conf: list                                  # list[list[float]] min word confidence per cell (1.0 if empty)
    merges: list = field(default_factory=list)  # list[(r0, c0, r1, c1)]
    header_rows: int = 1
    cell_boxes: list = field(default_factory=list)  # list[list[Box]] (may be empty)
    continued_pages: list = field(default_factory=list)  # extra pages merged into this table
    cell_words: list = field(default_factory=list)  # list[list[list[[word, confidence]]]] (may be empty)

    @property
    def n_rows(self) -> int:
        return len(self.grid)

    @property
    def n_cols(self) -> int:
        return max((len(r) for r in self.grid), default=0)

    def as_text(self, sep: str = " | ") -> str:
        """Plain-text rendering used in the one-row-per-table Excel column."""
        out = []
        for i, row in enumerate(self.grid):
            out.append(sep.join((c or "").replace("\n", " ").strip() for c in row))
            if i == self.header_rows - 1 and self.n_rows > 1:
                out.append(sep.join("---" for _ in row))
        return "\n".join(out)


@dataclass
class Block:
    """A layout element in reading order.
    kind: heading | paragraph | list_item | caption | table | figure | page_header | page_footer"""
    kind: str
    page: int
    box: Box
    lines: list = field(default_factory=list)       # list[Line]
    text: str = ""
    table: Optional[TableData] = None
    meta: dict = field(default_factory=dict)       # marker, caption kind/number, image file, ...
    uid: int = 0                                   # unique number within the document
    pages: list = field(default_factory=list)      # all pages this block spans (after merging)

    @property
    def conf(self) -> float:
        confs = [ln.conf for ln in self.lines]
        if self.table is not None:
            confs += [c for row in self.table.conf for c in row]
        return float(statistics.fmean(confs)) if confs else 1.0

    @property
    def min_conf(self) -> float:
        confs = [ln.min_word_conf for ln in self.lines]
        if self.table is not None:
            confs += [c for row in self.table.conf for c in row]
        return float(min(confs)) if confs else 1.0

    def page_label(self) -> str:
        pages = sorted(set(self.pages or [self.page]))
        return str(pages[0]) if len(pages) == 1 else f"{pages[0]}-{pages[-1]}"


@dataclass
class PageInfo:
    number: int
    width: int
    height: int
    source_dpi: float
    dpi: float                    # resolution of the cleaned image the OCR saw
    rotation: int = 0             # degrees clockwise that were applied to make the page upright
    skew: float = 0.0             # degrees of skew correction applied
    n_lines: int = 0
    mean_conf: float = 1.0
    low_conf_lines: int = 0
    image_file: str = ""          # downscaled page image for the Word report (in cache folder)
    notes: list = field(default_factory=list)


@dataclass
class DocResult:
    source: str
    engine: str
    pages: list = field(default_factory=list)    # list[PageInfo]
    blocks: list = field(default_factory=list)   # list[Block] in reading order
    warnings: list = field(default_factory=list)
    seconds: float = 0.0
    cache_dir: str = ""

    # ---- JSON round trip (for the OCR cache) --------------------------------------
    def to_json(self) -> dict:
        return dataclasses.asdict(self)

    @staticmethod
    def from_json(d: dict) -> "DocResult":
        def line(ld):
            return Line(ld["text"], ld["conf"], tuple(ld["box"]), [Word(w["text"], w["conf"], tuple(w["box"])) for w in ld["words"]])

        blocks = []
        for b in d["blocks"]:
            t = b.get("table")
            table = None
            if t:
                table = TableData(grid=t["grid"], conf=t["conf"], merges=[tuple(m) for m in t["merges"]],
                                  header_rows=t["header_rows"], cell_boxes=t.get("cell_boxes", []),
                                  continued_pages=t.get("continued_pages", []), cell_words=t.get("cell_words", []))
            blocks.append(Block(kind=b["kind"], page=b["page"], box=tuple(b["box"]),
                                lines=[line(ld) for ld in b["lines"]], text=b["text"], table=table,
                                meta=b["meta"], uid=b["uid"], pages=b.get("pages", [])))
        pages = [PageInfo(**p) for p in d["pages"]]
        return DocResult(source=d["source"], engine=d["engine"], pages=pages, blocks=blocks,
                         warnings=d.get("warnings", []), seconds=d.get("seconds", 0.0), cache_dir=d.get("cache_dir", ""))


@dataclass
class ExtractedItem:
    """One row of the final Excel sheet."""
    item_type: str            # Requirement | Table | Table Caption | Figure Caption
    source_file: str
    pages: str
    section: str = ""
    doc_ref: str = ""         # paragraph number, requirement ID, or table/figure number
    title: str = ""           # inline paragraph title, e.g. "Detection range"
    text: str = ""
    keyword: str = ""
    req_class: str = ""
    found_in: str = ""        # Paragraph | List item | Table row | Table cell | ...
    context: str = ""         # list stem, table caption, or surrounding sentence(s)
    attributes: str = ""      # other table columns, e.g. "Method: T"
    verification: str = ""    # verification method(s): from the requirement table or the VCRM
    references: str = ""      # standards / tables / figures / sections the requirement cites
    ocr_conf: Optional[float] = None
    review_reasons: list = field(default_factory=list)
    quality_notes: list = field(default_factory=list)
    table: Optional[TableData] = None
    block_uids: list = field(default_factory=list)
    item_id: str = ""
    table_sheet: str = ""

    @property
    def needs_review(self) -> bool:
        return bool(self.review_reasons)


# =============================================================================
# 4. LOADING PAGES FROM PDF / TIFF / IMAGE FILES
# =============================================================================
@dataclass
class RawPage:
    number: int            # 1-based page number inside the file
    image: "np.ndarray"    # uint8, grayscale (H, W) or RGB (H, W, 3), square pixels
    dpi: float             # resolution of `image`
    dpi_known: bool        # False when the file had no trustworthy resolution
    notes: list = field(default_factory=list)
    error: str = ""        # set when this page could not be read (the other pages still are)


def parse_page_selection(spec: str, n_pages: int) -> list:
    """Turn '1-3,7' into [1, 2, 3, 7] (1-based, limited to the pages that exist). Accepts en/em dashes
    and reversed ranges ('5-3'); raises ValueError with a readable message for anything else."""
    if not spec or not spec.strip():
        return list(range(1, n_pages + 1))
    pages = set()
    for part in re.sub(r"[\u2012\u2013\u2014\u2212]", "-", spec).replace(" ", "").split(","):
        if not part:
            continue
        m = re.fullmatch(r"(\d*)-(\d*)|(\d+)", part)
        if not m or part == "-":
            raise ValueError(f"page selection '{part}' is not understood - use for example 1-5,9")
        if m.group(3):
            pages.add(int(m.group(3)))
        else:
            first = int(m.group(1)) if m.group(1) else 1
            last = int(m.group(2)) if m.group(2) else n_pages
            if m.group(1) and m.group(2):
                first, last = min(first, last), max(first, last)     # "5-3" means pages 3 to 5
            pages.update(range(first, min(last, n_pages) + 1))
    return [p for p in sorted(pages) if 1 <= p <= n_pages]


def count_pages(path: Path) -> int:
    if path.suffix.lower() == ".pdf":
        pdf = pdfium.PdfDocument(str(path))
        try:
            return len(pdf)
        finally:
            pdf.close()
    with Image.open(path) as im:
        return int(getattr(im, "n_frames", 1))


def iter_raw_pages(path: Path, settings: Settings) -> Iterator[RawPage]:
    """Yield the selected pages of a document one at a time (keeps memory low)."""
    if path.suffix.lower() == ".pdf":
        yield from _iter_pdf_pages(path, settings)
    else:
        yield from _iter_image_pages(path, settings)


def _pdf_scan_dpi(page) -> Optional[float]:
    """Resolution of the scanned picture that fills a PDF page, or None when the
    page is not a scan (born-digital text, vector drawing, ...)."""
    try:
        page_w, page_h = page.get_size()
        best = None
        for obj in page.get_objects(filter=(pdfium_c.FPDF_PAGEOBJ_IMAGE,), max_depth=3):
            left, bottom, right, top = obj.get_bounds()
            if (right - left) * (top - bottom) < 0.4 * page_w * page_h:
                continue  # a logo or signature, not the page scan
            px_w, px_h = obj.get_px_size()
            # the image may be placed rotated; use the larger of the two estimates
            dpi = max(px_w, px_h) / max(max(right - left, top - bottom) / 72.0, 1e-6)
            best = max(best or 0.0, dpi)
        return best
    except Exception:  # never let an odd PDF object stop the run
        return None


def _iter_pdf_pages(path: Path, settings: Settings) -> Iterator[RawPage]:
    pdf = pdfium.PdfDocument(str(path))
    try:
        for number in parse_page_selection(settings.page_selection, len(pdf)):
            try:
                page = pdf[number - 1]
            except Exception as exc:
                yield RawPage(number, None, 0.0, False, [], error=f"page could not be opened ({exc})")
                continue
            try:
                notes = []
                page_w, page_h = page.get_size()                       # PDF points (1/72 inch)
                scan_dpi = _pdf_scan_dpi(page)
                dpi_known = scan_dpi is not None and 70 <= scan_dpi <= 1600
                # Render at the scan's own resolution (no needless resampling), within 150-400 dpi.
                render_dpi = float(np.clip(scan_dpi if dpi_known else settings.fallback_dpi, 150, 400))
                megapixels = (page_w / 72 * render_dpi) * (page_h / 72 * render_dpi) / 1e6
                if megapixels > settings.max_page_megapixels:
                    render_dpi *= math.sqrt(settings.max_page_megapixels / megapixels)
                    notes.append(f"very large page rendered at {render_dpi:.0f} dpi")
                try:
                    if page.get_textpage().count_chars() > 50:
                        notes.append("PDF already has a text layer (ignored - the page image is OCR'd)")
                except Exception:
                    pass
                bitmap = page.render(scale=render_dpi / 72.0, grayscale=not settings.remove_color_marks)
                pil = bitmap.to_pil()
                img = np.array(pil.convert("RGB" if settings.remove_color_marks else "L"))
            except Exception as exc:
                page.close()
                yield RawPage(number, None, 0.0, False, [], error=f"page could not be rendered ({exc})")
                continue
            page.close()
            yield RawPage(number, img, render_dpi, dpi_known, notes)
    finally:
        pdf.close()


def _pil_to_uint8(frame: "Image.Image", keep_color: bool) -> "np.ndarray":
    """Convert any Pillow image mode (1-bit fax, 16-bit, palette, CMYK, alpha) to uint8."""
    mode = frame.mode
    if mode in ("I;16", "I;16B", "I;16L", "I", "F"):
        arr = np.asarray(frame, dtype=np.float32)
        lo, hi = np.percentile(arr, (0.5, 99.5))
        return np.clip((arr - lo) * 255.0 / max(hi - lo, 1.0), 0, 255).astype(np.uint8)
    if mode in ("RGBA", "LA") or (mode == "P" and "transparency" in frame.info):
        rgba = frame.convert("RGBA")
        white = Image.new("RGBA", rgba.size, (255, 255, 255, 255))
        frame = Image.alpha_composite(white, rgba)
    return np.array(frame.convert("RGB" if keep_color else "L"))


def _iter_image_pages(path: Path, settings: Settings) -> Iterator[RawPage]:
    with Image.open(path) as im:
        n_frames = int(getattr(im, "n_frames", 1))
        for number in parse_page_selection(settings.page_selection, n_frames):
            try:
                raw = _read_image_frame(im, number, settings)
            except Exception as exc:
                raw = RawPage(number, None, 0.0, False, [], error=f"page could not be decoded ({exc})")
            yield raw


def _read_image_frame(im: "Image.Image", number: int, settings: Settings) -> RawPage:
    """Read one frame of a (possibly multi-page) image file, with its scan resolution."""
    im.seek(number - 1)
    notes = []
    # Pillow reports TIFF resolution as IFDRational objects -> convert to float.
    raw_dpi = im.info.get("dpi")
    try:
        x_dpi, y_dpi = (float(v) for v in raw_dpi) if raw_dpi else (0.0, 0.0)
    except (TypeError, ValueError):
        x_dpi, y_dpi = 0.0, 0.0
    if im.info.get("resolution_unit") == 3 or im.info.get("unit") == "cm":  # dots per cm
        x_dpi, y_dpi = x_dpi * 2.54, y_dpi * 2.54
    img = _pil_to_uint8(im, keep_color=settings.remove_color_marks)
    # Fax "normal" mode stores 204 x 98 dpi: stretch to square pixels first.
    if x_dpi > 0 and y_dpi > 0 and abs(x_dpi - y_dpi) / max(x_dpi, y_dpi) > 0.05:
        new_h = int(round(img.shape[0] * x_dpi / y_dpi))
        img = cv2.resize(img, (img.shape[1], new_h), interpolation=cv2.INTER_CUBIC)
        notes.append(f"non-square pixels ({x_dpi:.0f}x{y_dpi:.0f} dpi) resampled")
    dpi = max(x_dpi, y_dpi)
    # 72/96 dpi are software defaults, not real scan resolutions.
    dpi_known = 100 <= dpi <= 1600 and round(dpi) not in (72, 96)
    megapixels = img.shape[0] * img.shape[1] / 1e6
    if megapixels > settings.max_page_megapixels:
        f = math.sqrt(settings.max_page_megapixels / megapixels)
        img = cv2.resize(img, None, fx=f, fy=f, interpolation=cv2.INTER_AREA)
        dpi *= f
        notes.append("very large image scaled down")
    return RawPage(number, img, dpi if dpi_known else float(settings.fallback_dpi), dpi_known, notes)


# =============================================================================
# 5. IMAGE CLEAN-UP (PREPROCESSING)
#    Tested on degraded scans (photocopy, office scan, fax): deskew + background
#    flattening + a 3x3 median filter gave the best OCR word accuracy. Hard
#    black/white binarization was worse for the deep-learning OCR, so the OCR
#    sees a cleaned grayscale image.
# =============================================================================
@dataclass
class PreparedPage:
    ocr_image: "np.ndarray"    # cleaned grayscale image used for OCR and table detection
    view_image: "np.ndarray"   # same geometry, not flattened (figures, Word report)
    dpi: float
    rotation: int = 0          # degrees clockwise applied to make the page upright
    skew: float = 0.0          # degrees of rotation applied to remove skew
    notes: list = field(default_factory=list)


def _odd(n: float) -> int:
    n = max(1, int(round(n)))
    return n if n % 2 else n + 1


def to_grayscale(img: "np.ndarray", remove_color_marks: bool) -> "np.ndarray":
    if img.ndim == 2:
        return img
    if remove_color_marks:
        # Colored ink (red stamps, blue pen, highlighter) is bright in at least one color
        # channel while black print is dark in all of them: the brightest channel keeps
        # the black text and fades the colored marks.
        return img.max(axis=2).astype(np.uint8)
    return cv2.cvtColor(img, cv2.COLOR_RGB2GRAY)


def estimate_dpi_from_text(gray: "np.ndarray") -> Optional[float]:
    """Guess the scan resolution from the typical height of printed characters
    (used only when the file has no usable dpi information)."""
    scale = min(1.0, 1600.0 / max(gray.shape))
    small = cv2.resize(gray, None, fx=scale, fy=scale, interpolation=cv2.INTER_AREA) if scale < 1 else gray
    _, bw = cv2.threshold(small, 0, 255, cv2.THRESH_BINARY_INV + cv2.THRESH_OTSU)
    n, _, stats, _ = cv2.connectedComponentsWithStats(bw, connectivity=8)
    heights = [h for x, y, w, h, a in stats[1:] if 3 <= h <= 120 and w <= 3 * h and a >= 6]
    if len(heights) < 50:
        return None
    median_h = statistics.median(heights) / scale
    # The floor is low on purpose: a page exported as a screen-sized image really can be near
    # 60 dpi, and saying so is what lets it be enlarged enough to read (see prepare_page).
    return float(np.clip(median_h / 0.085, 40, 800))   # ~0.085 in = glyph height at 10-12 pt


def rotate_bound(gray: "np.ndarray", angle: float, border: int) -> "np.ndarray":
    """Rotate by `angle` degrees (counter-clockwise) without cutting off the corners."""
    h, w = gray.shape[:2]
    m = cv2.getRotationMatrix2D((w / 2, h / 2), angle, 1.0)
    cos, sin = abs(m[0, 0]), abs(m[0, 1])
    new_w, new_h = int(h * sin + w * cos), int(h * cos + w * sin)
    m[0, 2] += new_w / 2 - w / 2
    m[1, 2] += new_h / 2 - h / 2
    return cv2.warpAffine(gray, m, (new_w, new_h), flags=cv2.INTER_CUBIC,
                          borderMode=cv2.BORDER_CONSTANT, borderValue=int(border))


def _binary_thumbnail(gray: "np.ndarray", max_side: int = 1200) -> "np.ndarray":
    scale = max_side / max(gray.shape)
    small = cv2.resize(gray, None, fx=scale, fy=scale, interpolation=cv2.INTER_AREA) if scale < 1 else gray
    _, bw = cv2.threshold(small, 0, 255, cv2.THRESH_BINARY_INV + cv2.THRESH_OTSU)
    # ignore dark scanner/photocopier borders, they carry no text-line information
    m = max(2, int(0.03 * min(bw.shape)))
    bw[:m, :], bw[-m:, :], bw[:, :m], bw[:, -m:] = 0, 0, 0, 0
    return bw


def _profile_energy(bw: "np.ndarray", angle: float, axis: int = 1) -> float:
    """How 'striped' the image is along one axis after rotating by `angle`.
    Text lines produce sharp alternations between ink rows and blank rows."""
    if angle:
        m = cv2.getRotationMatrix2D((bw.shape[1] / 2, bw.shape[0] / 2), angle, 1.0)
        bw = cv2.warpAffine(bw, m, (bw.shape[1], bw.shape[0]), flags=cv2.INTER_NEAREST)
    profile = bw.sum(axis=axis, dtype=np.float64)
    return float(np.sum(np.diff(profile) ** 2))


def estimate_skew(gray: "np.ndarray", max_angle: float) -> float:
    """Projection-profile deskew: try angles coarse-to-fine and keep the one that makes
    text lines most horizontal. Accurate to ~0.05 degree; takes ~0.1 s per page."""
    bw = _binary_thumbnail(gray)
    if bw.sum() == 0:
        return 0.0
    coarse = np.arange(-max_angle, max_angle + 1e-6, 0.5)
    best = max(coarse, key=lambda a: _profile_energy(bw, a))
    fine = np.arange(best - 0.5, best + 0.5 + 1e-6, 0.05)
    best = max(fine, key=lambda a: _profile_energy(bw, a))
    return round(float(best), 2)


class OrientationDetector:
    """Decide whether a page must be turned 0/90/180/270 degrees clockwise.
    Uses the small RapidOrientation classifier when installed (Python <= 3.12);
    otherwise compares OCR confidence of the page read in both directions."""

    def __init__(self):
        self._model = None
        try:
            from rapid_orientation import RapidOrientation  # optional dependency
            _quiet_third_party_loggers()
            self._model = RapidOrientation()
        except Exception:
            self._model = None

    def detect(self, gray: "np.ndarray", ocr: "RapidOcrEngine") -> int:
        if self._model is not None:
            try:
                scale = min(1.0, 1024 / max(gray.shape))
                small = cv2.resize(gray, None, fx=scale, fy=scale, interpolation=cv2.INTER_AREA)
                label, _elapsed = self._model(cv2.cvtColor(small, cv2.COLOR_GRAY2RGB))[:2]
                # The classifier reports how far the page content is rotated CLOCKWISE
                # (verified: content turned 90 deg clockwise -> label "90"), so the
                # correction is the remaining angle to a full turn.
                rotation = (360 - int(label)) % 360
                if rotation or self._sparse(gray):
                    # Rotated pages are rare and the classifier can confuse opposite directions on
                    # sparse pages (title pages, a page with a few lines at its top and foot): confirm
                    # against the opposite turn with OCR - on a sparse page also when it says upright.
                    rotation = self._ocr_vote(gray, ocr, (rotation, (rotation + 180) % 360))
                return rotation
            except Exception:
                pass
        # Fallback without the classifier: find the reading direction, then let OCR decide.
        return self._ocr_vote(gray, ocr, (90, 270) if self._text_runs_vertically(gray, ocr) else (0, 180))

    @staticmethod
    def _sparse(gray: "np.ndarray") -> bool:
        """A page with little text - a title page, the last lines of a section: under a fifth of its rows of
        pixels hold any ink (a page of text has a third or more)."""
        scale = min(1.0, 600 / max(gray.shape))
        small = cv2.resize(gray, None, fx=scale, fy=scale, interpolation=cv2.INTER_AREA)
        return float((small < 128).any(axis=1).mean()) < 0.2

    @staticmethod
    def _text_runs_vertically(gray: "np.ndarray", ocr: "RapidOcrEngine") -> bool:
        """The OCR text detector draws wide boxes around horizontal text lines and tall boxes
        around lines of a page lying on its side."""
        scale = min(1.0, 1200 / max(gray.shape))
        small = cv2.resize(gray, None, fx=scale, fy=scale, interpolation=cv2.INTER_AREA)
        sizes = ocr.detect_boxes(small)
        wide = sum(1 for w, h in sizes if w > 1.8 * h)
        tall = sum(1 for w, h in sizes if h > 1.8 * w)
        return tall > 1.5 * max(wide, 1)

    @staticmethod
    def _ocr_vote(gray: "np.ndarray", ocr: "RapidOcrEngine", candidates: tuple) -> int:
        """Read a band of the page in each candidate direction; upright text reads confidently. The band is the
        part of the page that holds the most text (a page whose text sits at its top and foot has none in the
        middle), and when neither direction reads better the page is left upright - turned pages are rare."""
        scale = min(1.0, 1400 / max(gray.shape))
        small = cv2.resize(gray, None, fx=scale, fy=scale, interpolation=cv2.INTER_AREA)
        best, best_score = candidates[0], -1.0
        for rot in candidates:
            img = rotate_right_angle(small, rot)
            h = img.shape[0]
            win = max(1, int(0.45 * h))
            ink = np.concatenate([[0], np.cumsum((img < 128).sum(axis=1))])
            top = int(np.argmax(ink[win:] - ink[:-win])) if h > win else 0
            band = img[top:top + win, :]
            lines = ocr.read(band, min_conf=0.3)
            score = sum(ln.conf * len(ln.text) for ln in lines if ln.conf >= 0.6)
            if score > best_score or (score == best_score and rot == 0):
                best, best_score = rot, score
        return best


def rotate_right_angle(img: "np.ndarray", clockwise_degrees: int) -> "np.ndarray":
    code = {90: cv2.ROTATE_90_CLOCKWISE, 180: cv2.ROTATE_180, 270: cv2.ROTATE_90_COUNTERCLOCKWISE}.get(clockwise_degrees % 360)
    return img if code is None else cv2.rotate(img, code)


def remove_punch_holes(gray: "np.ndarray", dpi: float) -> tuple:
    """Paint over binder punch holes (round dark blobs near the page edges). OCR otherwise
    reads them as a large letter 'O' that gets glued to the text line next to it."""
    h, w = gray.shape[:2]
    _, bw = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY_INV + cv2.THRESH_OTSU)
    # ignore the outermost strip, where photocopiers leave dark borders touching the holes
    border = max(2, int(0.02 * min(h, w)))
    bw[:border, :], bw[-border:, :], bw[:, :border], bw[:, -border:] = 0, 0, 0, 0
    contours, _ = cv2.findContours(bw, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
    rects = [cv2.boundingRect(c) for c in contours]
    paper = int(np.percentile(gray, 90))
    edge_zone = 0.9 * dpi
    removed = 0
    for c in contours:
        x, y, cw, ch = cv2.boundingRect(c)
        diameter = (cw + ch) / 2.0
        if not (0.15 * dpi <= diameter <= 0.5 * dpi) or abs(cw - ch) > 0.25 * max(cw, ch):
            continue
        cx, cy = x + cw / 2.0, y + ch / 2.0
        if min(cx, w - cx, cy, h - cy) > edge_zone:
            continue                                          # not near an edge: could be a real symbol
        perimeter = cv2.arcLength(c, True)
        circularity = 4 * math.pi * cv2.contourArea(c) / max(perimeter * perimeter, 1.0)
        if circularity < 0.7:
            continue
        # isolated: no character-sized ink on the same line within 0.3 inch (a letter 'O' inside
        # a word has neighbours; specks and the tall dark photocopier edge are ignored)
        neighbour = False
        for ox, oy, ow, oh in rects:
            if (ox, oy, ow, oh) == (x, y, cw, ch) or not (0.25 * diameter <= oh <= 1.5 * diameter) or ow > 3 * diameter:
                continue
            if abs((oy + oh / 2.0) - cy) < diameter and min(abs(ox - (x + cw)), abs(x - (ox + ow))) < 0.3 * dpi:
                neighbour = True
                break
        if neighbour:
            continue
        radius = int(diameter / 2 + 0.03 * dpi)
        ring_mask = np.zeros(gray.shape, np.uint8)
        cv2.circle(ring_mask, (int(cx), int(cy)), radius + int(0.08 * dpi), 255, -1)
        cv2.circle(ring_mask, (int(cx), int(cy)), radius, 0, -1)
        local = gray[(ring_mask > 0) & (bw == 0)]
        fill = int(np.median(local)) if local.size > 20 else paper
        cv2.circle(gray, (int(cx), int(cy)), radius, fill, -1)
        removed += 1
    return gray, removed


def flatten_background(gray: "np.ndarray", dpi: float) -> "np.ndarray":
    """Remove stains, shadows, yellowing and uneven lighting: estimate the paper
    brightness around every pixel (ink removed by a max-filter) and divide it out."""
    k = _odd(max(7, dpi * 0.035))                  # wider than a pen stroke (~0.035 in)
    background = cv2.dilate(gray, np.ones((k, k), np.uint8))
    background = cv2.medianBlur(background, _odd(max(21, dpi * 0.1)))
    flat = cv2.divide(gray, background, scale=255)
    # stretch contrast so the darkest ink is close to black again
    lo = float(np.percentile(flat, 0.5))
    if lo < 245:
        flat = np.clip((flat.astype(np.float32) - lo) * (255.0 / (255.0 - lo)), 0, 255).astype(np.uint8)
    return flat


def prepare_page(raw: RawPage, settings: Settings, orientation: OrientationDetector,
                 ocr: "RapidOcrEngine") -> PreparedPage:
    notes = list(raw.notes)
    gray = to_grayscale(raw.image, settings.remove_color_marks)
    dpi = raw.dpi
    if not raw.dpi_known:
        guess = estimate_dpi_from_text(gray)
        if guess:
            dpi = guess
            notes.append(f"resolution unknown - estimated {dpi:.0f} dpi from text size")

    # 1) resolution normalization (the OCR model is most accurate/fast around target_dpi).
    # Enlarging is allowed to go further than shrinking: a page saved as a screen-sized image can
    # be a third of the resolution the recogniser wants, and doubling alone does not get it there.
    scale = float(np.clip(settings.target_dpi / max(dpi, 1.0), 0.35, 4.0))
    megapixels = gray.shape[0] * gray.shape[1] * scale * scale / 1e6
    if megapixels > settings.max_page_megapixels:      # never enlarge past the memory limit
        scale *= math.sqrt(settings.max_page_megapixels / megapixels)
    if abs(scale - 1.0) > 0.12:
        interp = cv2.INTER_AREA if scale < 1 else cv2.INTER_CUBIC
        gray = cv2.resize(gray, None, fx=scale, fy=scale, interpolation=interp)
        dpi *= scale

    paper = int(np.percentile(gray, 90))
    # 2) page orientation (sideways / upside-down scans)
    rotation = 0
    if settings.auto_rotate:
        rotation = orientation.detect(gray, ocr)
        if rotation:
            gray = rotate_right_angle(gray, rotation)
            notes.append(f"page turned {rotation} degrees clockwise")

    # 3) deskew
    skew = 0.0
    if settings.deskew:
        skew = estimate_skew(gray, settings.max_skew_degrees)
        if abs(skew) >= 0.1:
            gray = rotate_bound(gray, skew, paper)
        else:
            skew = 0.0

    # 4) binder punch holes
    gray, holes = remove_punch_holes(gray, dpi)
    if holes:
        notes.append(f"{holes} punch hole(s) removed")

    view = gray
    # 5) background flattening + light denoise
    ocr_img = flatten_background(gray, dpi) if settings.clean_background else gray.copy()
    if settings.denoise:
        ocr_img = cv2.medianBlur(ocr_img, 3)
    return PreparedPage(ocr_img, view, dpi, rotation, skew, notes)


# =============================================================================
# 6. OCR AND PAGE LAYOUT ANALYSIS  (the "standard" engine)
# =============================================================================

# ---- 6a. Text patterns shared by layout analysis and requirement extraction ----
# Paragraph / heading numbers: "3", "3.2", "3.2.1.4", "A.1.2" (optionally followed by a period)
RE_PARA_NUM = re.compile(r"^(?P<num>(?:\d{1,2}(?:\.(?:\d{1,3}|[lIi|](?=[.\s'`·]))){0,7}|[A-Z](?:\.\d{1,3}){1,7}))(?P<dot>\.?)(?:\s+|['`´’·,]\s*(?=[A-Z]))(?P<rest>\S.*)$", re.S)
# A paragraph number that starts a NEW paragraph part-way through a block: it follows the end of a
# sentence and a capitalised word follows it.  A single-level number ("1.") is a list marker, so at
# least one sub-level is required; the space after the number may be lost ("3.2.2.2.Input power.").
RE_INNER_PARA_NUM = re.compile(r"(?<=[.:;])\s+(?P<num>\d{1,2}(?:\.[\dlIi|]{1,3}){1,7})\.?\s*(?=[A-Z])")
# List markers: a.  a)  (a)  (1)  1)  iv)  (iv)  bullets.  "1." is handled separately.
RE_LIST_MARKER = re.compile(
    r"^(?P<marker>\([a-z]\)|[a-z][.)]|\(\d{1,2}\)|\d{1,2}\)|\((?:i|ii|iii|iv|v|vi|vii|viii|ix|x|xi|xii)\)"
    r"|(?:i|ii|iii|iv|v|vi|vii|viii|ix|x|xi|xii)[.)]|[•●▪■‣⁃∙·◦]"
    r"|[–—*\-o](?=\s+[A-Z0-9]))"          # dash / 'o' bullets only before a capital or digit
    # The space after a marker may be lost ("c.Built-in test mode."), so a capital letter hard
    # against the marker also starts the text. A marker still has to be at the start of the line.
    r"(?:\s+|(?=[A-Z]))(?P<rest>\S.*)$", re.S)
RE_BARE_MARKER = re.compile(r"^\s*(?:\(?[a-z]\)|[a-z]\.|\(?\d{1,2}[.)]|\((?:i|ii|iii|iv|v|vi|vii|viii|ix|x)\)|[•·◦*\-o])\s*$")
RE_NUMBERED_ITEM = re.compile(r"^(?P<marker>\d{1,2}\.)(?:\s+|(?=[A-Z][a-z]))(?P<rest>\S.*)$", re.S)   # "1.Input" too
# Something shaped like a list marker, right or wrong: "c." "C." "(b)" "3)" "iv." - one to four
# characters closed by "." or ")" - used when the OCR has misread the marker itself.
RE_MARKER_SHAPE = re.compile(r"^(?P<marker>\(?(?:[A-Za-z]|\d{1,2}|[ivxIVX]{2,4})[.)])[.,]?(?:\s+|(?=[A-Z]))(?P<rest>\S.*)$", re.S)
_ROMANS = ("i", "ii", "iii", "iv", "v", "vi", "vii", "viii", "ix", "x", "xi", "xii", "xiii", "xiv", "xv",
           "xvi", "xvii", "xviii", "xix", "xx")


def _marker_parts(marker: str):
    """'b.' -> ('', 'b', '.'), '(iv)' -> ('(', 'iv', ')'), '12)' -> ('', '12', ')'); None for a bullet."""
    m = re.fullmatch(r"(\(?)([A-Za-z]{1,5}|\d{1,3})([.)]?)", (marker or "").strip())
    return (m.group(1), m.group(2), m.group(3)) if m and (m.group(1) or m.group(3)) else None


def _next_marker(marker: str, roman: bool = False) -> str:
    """The marker that follows `marker` in the same list: 'b.' -> 'c.', '(2)' -> '(3)', 'iv)' -> 'v)'.
    Letters and numbers count up one at a time, so a list's markers can be checked against each other."""
    parts = _marker_parts(marker)
    if not parts:
        return ""
    opening, core, closing = parts
    if core.isdigit():
        following = str(int(core) + 1)
    elif roman and core.lower() in _ROMANS[:-1]:
        following = _ROMANS[_ROMANS.index(core.lower()) + 1]
        following = following.upper() if core.isupper() else following
    elif len(core) == 1 and core.isalpha() and core.lower() != "z":
        following = chr(ord(core) + 1)
    else:
        return ""
    return opening + following + closing


def _previous_marker(marker: str, roman: bool = False) -> str:
    """The marker before `marker`: 'b.' -> 'a.', '(3)' -> '(2)'; '' for the first one."""
    parts = _marker_parts(marker)
    if not parts:
        return ""
    opening, core, closing = parts
    if core.isdigit():
        before = str(int(core) - 1) if int(core) > 1 else ""
    elif roman and core.lower() in _ROMANS[1:]:
        before = _ROMANS[_ROMANS.index(core.lower()) - 1]
        before = before.upper() if core.isupper() else before
    elif len(core) == 1 and core.isalpha() and core.lower() != "a":
        before = chr(ord(core) - 1)
    else:
        return ""
    return opening + before + closing if before else ""
# Captions: "TABLE II. Title", "Figure 3-1: Title", "FIG. 4 Title", "Exhibit A - Title"
RE_CAPTION = re.compile(
    r"^(?P<kind>TABLE|Table|FIGURE|Figure|FIG\.?|Fig\.?|EXHIBIT|Exhibit|CHART|Chart|DRAWING|Drawing|PLATE|Plate)"
    # a speck read as a period or comma between the word and its number ("TABLE. IV.")
    r"[.,]?\s*(?P<num>(?:[A-Z]-)?\d{1,3}(?:[.\-]\d{1,3})*[A-Za-z]?|[IVXLCl1|]{1,8}(?=[\s.:\-–—]))"
    r"(?P<delim>\s*[.:\-–—]+)?\s*(?P<title>.*)$", re.S)
RE_CONTINUED = re.compile(r"\(?\s*(?:continued|cont'?d\.?|concluded)\s*\)?\s*\.?$", re.I)
# Explicit requirement identifiers: [SRS-101], SRS-103:, PRF-201, SYS_PWR_012 ...
_ID_STOP = (r"(?:MIL|DOD|DI|NPR|NPD|NASA|ISO|IEC|IEEE|SAE|ASTM|ANSI|RTCA|STANAG|FED|JSSG|ECSS|AES|SHA|RS|DO|AS|MS|"
            r"NAS|ARINC|TIA|EIA|UL|NFPA|CFR|FAR|DFARS|IPC|ASME|AWG|UTC|GMT|CAGE|NSN|DTL|PRF-\d{5,}|TB[DRSC])")
RE_REQ_ID = r"(?!" + _ID_STOP + r"[-_ ])[A-Z][A-Z0-9]{0,9}(?:[-_.][A-Z0-9]{1,9}){0,3}[-_.]?\d{1,6}[a-z]?"
RE_LEADING_ID = re.compile(r"^\s*[\[(]?(?P<id>" + RE_REQ_ID + r")[\])]?\s*(?:[:.\-–—]\s*|\s+)(?P<rest>\S.*)$", re.S)
RE_TRAILING_ID = re.compile(r"\s*[\[(](?P<id>" + RE_REQ_ID + r")[\])]\s*\.?$")
RE_ID_CELL = re.compile(r"^\s*[\[(]?(?:" + RE_REQ_ID + r"|(?:\d{1,2}(?:\.\d{1,3}){1,7}\.?)|[A-Z]{1,4}-?\d{1,4}[a-z]?)[\])]?\s*$")
RE_PAGE_NUMBER = re.compile(r"^\s*(?:page\s*)?[-–—]?\s*(?:\d{1,4}|[ivxlcdm]{1,7})\s*[-–—]?\s*(?:of\s*\d{1,4})?\s*$", re.I)
RE_TERMINAL = re.compile(r"[.:;!?][\"')\]]*\s*$")
OCR_KEYWORD_REPAIRS = [
    # (correct word, pattern of OCR-garbled spellings). Whole tokens only; real English
    # words such as "shell", "small", "wall", "most" do NOT match.
    ("shall", re.compile(r"(?<![A-Za-z0-9])[sS5$][hbn][aA@][l1I|!]{2}(?![A-Za-z0-9])")),
    ("should", re.compile(r"(?<![A-Za-z0-9])[sS5][hb][o0]u[l1I|!][dc](?![A-Za-z0-9])")),
    ("must", re.compile(r"(?<![A-Za-z0-9])(?:m|rn|nn)[uv][s5]t(?![A-Za-z0-9])")),
    ("will", re.compile(r"(?<![A-Za-z0-9])(?:[wW]|vv)[i1l|!][l1I|!]{2}(?![A-Za-z0-9])")),
]
RE_UNIT = re.compile(r"(?:[kMGmu\u00b5]?(?:A|V|W|Hz|VA|VAC|VDC|Vdc|Vac|s|ms|g|m|Pa|N|J|F|H|dB|dBm|dBi|bps|B|b)|"
                     r"mm|cm|km|nmi|kg|lb|lbs|oz|in|ft|psi|psig|psia|rpm|\u00b0[CFK]?|%|ohms?|\u03a9|kts|knots|AWG|min|hr|hrs)")

LIGATURES = {"ﬀ": "ff", "ﬁ": "fi", "ﬂ": "fl", "ﬃ": "ffi", "ﬄ": "ffl",
             "‘": "'", "’": "'", "“": '"', "”": '"', " ": " ", "‐": "-", "‑": "-"}


def clean_para_number(num: str) -> str:
    """'3.2.i.4' -> '3.2.1.4' (a letter l/I/i/| inside a paragraph number is a misread 1)."""
    return re.sub(r"(?<=\.)[lIi|](?=\.|$)", "1", num)


def _number_parts(number: str):
    """'3.2.4' -> (3, 2, 4); 'A.1' -> (-1, 1) (an appendix letter, counted as its own scheme); None when
    the text is not a paragraph number."""
    parts = clean_para_number((number or "").strip().rstrip(".")).split(".")
    out = []
    for k, part in enumerate(parts):
        if part.isdigit():
            out.append(int(part))
        elif k == 0 and len(part) == 1 and part.isalpha() and part.isupper():
            out.append(-(ord(part) - ord("A") + 1))           # appendix A, B, ... as -1, -2, ...
        else:
            return None
    return tuple(out) if out else None


def plausible_next_number(previous: str, candidate: str, max_skip: int = 3) -> bool:
    """Is `candidate` a believable next paragraph number after `previous`?

    Paragraph numbers count up. After 3.2.4 may come 3.2.4.1 (its first sub-paragraph), 3.2.5 (the next
    one) or 3.3 / 4 (the next one a level up), and a few numbers may be skipped because the scan missed a
    heading. A number that fits none of these - '1.5' or '20' after 3.2.4 - is a value that happens to
    start a line ("1.5 V peak-to-peak", "20 Hz to 2000 Hz"), not a section. Deeper levels of a new number
    start low (3.3.1, not 3.3.7). With nothing to compare against, any number is accepted."""
    p, c = _number_parts(previous), _number_parts(candidate)
    if p is None or c is None:
        return True
    if (p[0] < 0) != (c[0] < 0):
        return True                                   # the body moves into an appendix, or back: a new scheme

    def starts_low(levels) -> bool:
        return all(0 <= x <= 2 for x in levels)

    if len(c) > len(p) and c[:len(p)] == p:
        return starts_low(c[len(p):])                  # a sub-paragraph: 3.2.4 -> 3.2.4.1
    for k in range(min(len(p), len(c))):
        if c[k] != p[k]:
            step = (p[k] - c[k]) if c[k] < 0 else (c[k] - p[k])       # appendix letters count -1, -2, ...
            return 0 < step <= max_skip and starts_low(c[k + 1:])
    return c == p                                     # the same number again (a paragraph under its own heading)


def normalize_ocr_text(text: str) -> str:
    """Tidy raw OCR output without changing its meaning (no NFKC: that would turn m² into m2)."""
    text = unicodedata.normalize("NFC", text or "")
    for bad, good in LIGATURES.items():
        text = text.replace(bad, good)
    text = "".join(ch for ch in text if ch == "\n" or unicodedata.category(ch)[0] != "C")
    return re.sub(r"[ \t]+", " ", text).strip()


_FILE_EXTENSIONS = {"xml", "txt", "csv", "pdf", "doc", "docx", "xls", "xlsx", "exe", "dll", "cfg", "ini", "log", "dat",
                    "bin", "htm", "html", "com", "org", "net", "gov", "mil", "edu", "json", "yaml", "yml", "py"}


def strip_accents(text: str) -> str:
    """'telémetry' -> 'telemetry'. Only Latin letters lose accents; symbols such as °, ±, µ, ² stay."""
    out = []
    for ch in text:
        if ord(ch) > 127 and unicodedata.category(ch).startswith("L") and "LATIN" in unicodedata.name(ch, ""):
            base = unicodedata.normalize("NFD", ch)[0]
            out.append(base if base.isascii() else ch)
        else:
            out.append(ch)
    return "".join(out)


def remove_ocr_specks(text: str) -> str:
    """Remove dots/apostrophes that are paper specks read as punctuation on photocopies and
    faxes ("The.power", "radar' shall", "finished'in", ".interface"). Numbers, decimals,
    abbreviations (e.g.), possessives (operator's) and file names (config.xml) are left alone."""
    # speck read as a middle dot next to a word ("be·as", ".·The"); "N·m" (units) is kept
    text = re.sub(r"(?:(?<=[a-z]{2})|(?<=[.,;:]))·(?=[A-Za-z])", " ", text)
    # dot glued between two ordinary lowercase words: "steady-state.voltages" -> "steady-state voltages"
    text = re.sub(r"(?<![\w.])([A-Za-z][a-z]{2,})\.([a-z]{2,})(?!\w|\.\w)",
                  lambda m: m.group(0) if m.group(2) in _FILE_EXTENSIONS or m.group(1).lower() == "www"
                  else f"{m.group(1)} {m.group(2)}", text)
    # the space after a full stop lost between two words: "inspection.Each" -> "inspection. Each".
    # Only a lower-case word followed by a capitalised one, so MIL-STD-461G, file names and
    # initials such as "U.S.Navy" are left alone.
    text = re.sub(r"(?<![\w.])([a-z]{3,})\.([A-Z][a-z]{2,})", r"\1. \2", text)
    # the space after a paragraph number lost, or read as a period: "4.2Extended verification." and
    # "3.2.2.2.Input power." both become "4.2 Extended ..." / "3.2.2.2. Input ...", so the heading is
    # seen as its own paragraph.  Only numbers that already have a sub-level ("4.2", not "1.") are
    # touched, which leaves list markers ("1.Built-in") and part numbers (MIL-STD-1553B) alone; a
    # period that is really a full stop is kept, so "in 3.2.1.The unit" -> "in 3.2.1. The unit".
    text = re.sub(r"(?<![\w.])(\d{1,2}(?:\.\d{1,3}){1,7}\.?)(?=[A-Z][a-z])", r"\1 ", text)
    # dot stuck to the front of a word: "The .interface" -> "The interface"
    text = re.sub(r"(?:(?<=\s)|^)\.(?=[a-z]{3,}\b)", "", text)
    # apostrophe speck after a word (not a plural possessive): "radar' shall" -> "radar shall"
    text = re.sub(r"(?<=[A-Za-rt-z])['’`](?=\s|[.,;:]|$)", "", text)
    # colon/dot/quote specks between a word and a number or unit: "of:1" "450:kg" "16.AWG" 'Aircraft"Electric'
    text = re.sub(r"(?<=[a-z]):(?=\d)|(?<=\d):(?=[a-z]{2})|(?<=\d)\.(?=[A-Z]{2,}\b)|(?<=[a-z])\"(?=[A-Z])", " ", text)
    # ",." speck after a word: "(I),.analysis" -> "(I), analysis"  (".," is kept: "i.e., etc.,")
    text = re.sub(r"(?<=[A-Za-z)\]]),\.(?=\s*[a-z])", ", ", text)
    # doubled period after a word ("II..") and a speck apostrophe in front of a lowercase word ("'mode")
    text = re.sub(r"(?<=[A-Za-z)\]])\.\.(?!\.)", ".", text)
    text = re.sub(r"(?:(?<=\s)|^)['`](?=[a-z]{2,}\b)(?![a-z]*['`])", "", text)
    # apostrophe speck between words: "finished'in" -> "finished in" (contractions are kept)
    text = re.sub(r"(?<=[A-Za-z0-9][a-z0-9])['’`](?!(?:s|t|ll|re|ve|d|m)\b)(?=[a-z]{2,})", " ", text)
    return re.sub(r"[ \t]+", " ", text).strip()


def word_count(text: str) -> int:
    return len(text.split())


def _quiet_third_party_loggers():
    for name in ("RapidOCR", "rapidocr", "rapid_orientation", "img2table", "onnxruntime", "docling"):
        logging.getLogger(name).setLevel(logging.ERROR)


# ---- 6b. OCR engine -------------------------------------------------------------
class RapidOcrEngine:
    """PaddleOCR PP-OCRv6 text detection + recognition through RapidOCR / ONNX Runtime.
    The PP-OCRv6 'small' models are bundled inside the rapidocr pip package."""

    def __init__(self, settings: Settings):
        try:
            from rapidocr import ModelType, RapidOCR
        except ImportError as exc:  # pragma: no cover
            raise SystemExit("RapidOCR is not installed. Run: python -m pip install -r requirements.txt") from exc
        params = {
            "Global.log_level": "error",
            "Global.use_cls": False,          # page orientation is already fixed; faster and safer
            "Global.max_side_len": 4096,      # never shrink a normal page before OCR (hurts accuracy)
            "Det.limit_type": "max",
            "Det.limit_side_len": 2600,       # cap detection resolution for very large drawings
        }
        size = (settings.ocr_model_size or "small").lower()
        if size in ("tiny", "medium"):        # these sizes download once from ModelScope
            model_type = ModelType.TINY if size == "tiny" else ModelType.MEDIUM
            params.update({"Det.model_type": model_type, "Rec.model_type": model_type})
        if settings.ocr_threads > 0:
            params["EngineConfig.onnxruntime.intra_op_num_threads"] = int(settings.ocr_threads)
        _quiet_third_party_loggers()
        self.engine = RapidOCR(params=params)
        _quiet_third_party_loggers()
        self.settings = settings

    def detect_boxes(self, gray: "np.ndarray") -> list:
        """Text-line detection only (no recognition): [(width, height)] of each line box."""
        rgb = cv2.cvtColor(gray, cv2.COLOR_GRAY2RGB) if gray.ndim == 2 else gray
        res = self.engine(rgb, use_det=True, use_cls=False, use_rec=False, box_thresh=0.5, unclip_ratio=1.6)
        boxes = getattr(res, "boxes", None)
        if boxes is None:
            return []
        return [(float(np.ptp(np.asarray(b)[:, 0])), float(np.ptp(np.asarray(b)[:, 1]))) for b in boxes]

    def read(self, gray: "np.ndarray", min_conf: Optional[float] = None) -> list:
        """OCR one image -> list of Line (each with word boxes and confidences)."""
        rgb = cv2.cvtColor(gray, cv2.COLOR_GRAY2RGB) if gray.ndim == 2 else gray
        min_conf = self.settings.min_text_confidence if min_conf is None else min_conf
        # NOTE: RapidOCR resets text_score/box_thresh/unclip_ratio on every call, so pass all three.
        res = self.engine(rgb, use_det=True, use_cls=False, use_rec=True, text_score=min_conf,
                          box_thresh=0.5, unclip_ratio=1.6, return_word_box=True)
        if res is None or res.boxes is None or res.txts is None:
            return []
        word_results = getattr(res, "word_results", None) or ()
        lines = []
        for i, (quad, raw_text, score) in enumerate(zip(res.boxes, res.txts, res.scores)):
            text = normalize_ocr_text(raw_text)
            if not text:
                continue
            q = np.asarray(quad, dtype=float)
            box = (int(q[:, 0].min()), int(q[:, 1].min()), int(math.ceil(q[:, 0].max())), int(math.ceil(q[:, 1].max())))
            words = []
            if i < len(word_results) and word_results[i]:
                for entry in word_results[i]:
                    try:
                        w_text, w_score, w_quad = normalize_ocr_text(entry[0]), float(entry[1]), np.asarray(entry[2], dtype=float)
                    except (TypeError, ValueError, IndexError):
                        words = []
                        break
                    if w_text:
                        words.append(Word(w_text, w_score, (int(w_quad[:, 0].min()), int(w_quad[:, 1].min()),
                                                            int(math.ceil(w_quad[:, 0].max())), int(math.ceil(w_quad[:, 1].max())))))
            # word boxes must spell the same text as the line, otherwise estimate them
            if not words or "".join(w.text for w in words).replace(" ", "") != text.replace(" ", ""):
                words = _estimate_word_boxes(text, box, float(score))
            lines.append(Line(text, float(score), box, words))
        return lines


def _estimate_word_boxes(text: str, box: Box, conf: float) -> list:
    """Split a line box into word boxes proportionally to character counts."""
    x0, y0, x1, y1 = box
    n = max(len(text), 1)
    words = []
    for m in re.finditer(r"\S+", text):
        wx0 = x0 + (x1 - x0) * m.start() / n
        wx1 = x0 + (x1 - x0) * m.end() / n
        words.append(Word(m.group(), conf, (int(wx0), y0, int(math.ceil(wx1)), y1)))
    return words


# ---- 6c. Tables (img2table finds the grid; our OCR words fill the cells) --------------
def detect_tables(gray: "np.ndarray", lines: list, dpi: float, settings: Settings, ocr=None) -> list:
    """Return [(TableData, table_box, set_of_line_indexes_inside, caption_lines)] for one page.
    img2table finds the table grid (ruling lines, or aligned text for borderless tables);
    the words we already OCR'd are then placed into the cells."""
    try:
        from img2table.document import Image as I2TImage
        from img2table.ocr._types import OCRData, OCRInstance
    except Exception as exc:  # pragma: no cover
        LOG.warning("Table detection unavailable (img2table could not be imported: %s)", exc)
        return []

    records = []
    for li, ln in enumerate(lines):
        for wi, wd in enumerate(ln.words):
            records.append({"id": f"w{li}_{wi}", "parent": f"l{li}", "value": wd.text,
                            "confidence": max(1, int(round(wd.conf * 100))),
                            "x1": int(wd.box[0]), "y1": int(wd.box[1]), "x2": int(wd.box[2]), "y2": int(wd.box[3])})

    class _PrecomputedOCR(OCRInstance):
        """Hands img2table the words we already read, so each page is OCR'd only once."""

        def of(self, document):
            return OCRData(records={0: records}) if records else None

    ok, png = cv2.imencode(".png", gray)

    def run_img2table(borderless: bool, infer: bool = False) -> list:
        try:
            return I2TImage(src=png.tobytes(), detect_rotation=False).extract_tables(
                ocr=_PrecomputedOCR() if records else None, implicit_rows=infer, implicit_columns=infer,
                borderless_tables=borderless, min_confidence=1)
        except Exception as exc:
            LOG.warning("Table detection failed on a page (%s) - continuing without tables", exc)
            return []

    # Pass 1: ruled tables, using only the ruling lines that are really printed (most reliable).
    ruled = [(et, "ruled") for et in run_img2table(False)]
    # Pass 2: the same page with the separators that were never printed filled in from the text.
    # This is what finds a table whose columns are ruled but whose rows are not - the usual layout
    # of a verification matrix, which pass 1 does not see at all, because a row needs two
    # horizontal lines before it exists. _better_row_split then decides, area by area, which
    # reading to believe, and returns each area as the readings to try in order, best first.
    candidates = (_better_row_split(ruled, run_img2table(False, infer=True))
                  if settings.infer_missing_rules else [[alt] for alt in ruled])
    # Pass 3: borderless candidates, which must pass stricter checks because aligned lists and
    # indented text can look like tables.
    if settings.borderless_tables:
        ruled_boxes = [_et_box(alts[0][0]) for alts in candidates]
        for et in run_img2table(True):
            eb = _et_box(et)
            if not any(box_intersection(eb, rb) > 0.3 * min(box_area(eb), box_area(rb)) for rb in ruled_boxes):
                candidates.append([(et, "borderless")])

    page_area = gray.shape[0] * gray.shape[1]
    results = []
    claimed = set()

    def build(et, found_by: str):
        """Turn one reading of one area into a table, or return None if it is not a table at all."""
        rows = list(et.content.values())
        if len(rows) < 2 or max((len(r) for r in rows), default=0) < 2:
            return None
        n_cols = max(len(r) for r in rows)
        cell_boxes = [[(int(c.bbox.x1), int(c.bbox.y1), int(c.bbox.x2), int(c.bbox.y2)) for c in r]
                      + [None] * (n_cols - len(r)) for r in rows]
        if found_by == "resplit":
            cell_boxes = _resplit_rows(cell_boxes, lines)
            if cell_boxes is None:
                return None
        if found_by in ("inferred", "resplit"):
            # rows worked out from the text put every LINE on its own row; put the wrapped ones back
            cell_boxes = _merge_continuation_rows(cell_boxes, lines)
        tbox = (int(et.bbox.x1), int(et.bbox.y1), int(et.bbox.x2), int(et.bbox.y2))
        pad = 0.02 * dpi
        inside = {i for i, ln in enumerate(lines) if box_center_inside(ln.box, tbox, pad) and i not in claimed}
        # a caption close under the last row ("TABLE I. Interface characteristics.") can be read into that row's
        # cells ("100 Mohm at 500 VDC TABLE I. ..."): it is given back before the cells are filled
        under = _caption_under_table(inside, lines)
        if under is not None:
            inside.discard(under)
            tbox = (tbox[0], min(lines[i].box[1] for i in inside), tbox[2], max(lines[i].box[3] for i in inside))
        inside_order = sorted(inside)
        grid, conf, merges, n_unique, n_filled, words_in_cells, words_grid, row_lines = _fill_cells(
            cell_boxes, [lines[i] for i in inside_order])

        # A caption that img2table swallowed as the first/last row is given back. It may fill one merged
        # cell or be split over a few cells ("TABLE II." | "Interface signals"), so the row's joined text
        # is tested. A "list of tables" (many rows starting "Table n") is left alone.
        caption_lines = [under] if under is not None else []
        for edge in (0, -1):
            if len(grid) < 3:
                break
            distinct = list(dict.fromkeys(c for c in grid[edge] if c))   # a merged cell repeats its text
            joined = " ".join(distinct)
            other_caption_rows = sum(1 for row in grid if row is not grid[edge] and row and RE_CAPTION.match(row[0] or ""))
            if not (distinct and len(distinct) <= 3 and word_count(joined) <= 40
                    and RE_CAPTION.match(joined) and other_caption_rows == 0):
                continue
            row_index = edge % len(grid)
            # every line whose words were placed in that row (a caption can wrap over several lines)
            elsewhere = set().union(*(row_lines[r] for r in range(len(row_lines)) if r != row_index))
            given_back = sorted(row_lines[row_index] - elsewhere)
            if not given_back:
                continue                                # its lines also fill other rows: leave the table as it is
            caption_lines.extend(inside_order[li] for li in given_back)
            del grid[edge], conf[edge], cell_boxes[edge], words_grid[edge], row_lines[edge]
            merges = _merges_without_row(merges, row_index)
        if caption_lines:
            inside -= set(caption_lines)
            # the table now starts/ends at its first/last real row (so the caption sits next to it)
            remaining = [cb for row in cell_boxes for cb in row if cb]
            if remaining:
                tbox = (tbox[0], min(cb[1] for cb in remaining), tbox[2], max(cb[3] for cb in remaining))

        # A verification matrix is mostly empty on purpose (one tick per row), so it is judged on
        # its shape rather than on how many of its cells carry text.
        if n_filled < max(2, 0.25 * n_unique) and not _is_mark_matrix(grid):
            return None                                 # mostly empty grid: a drawing or form, not a table
        if box_area(tbox) > 0.6 * page_area and words_in_cells / max(n_filled, 1) > 25:
            return None                                 # a page border drawn around ordinary paragraphs
        if found_by == "borderless" and not _plausible_borderless_table(
                grid, n_unique, n_filled, captioned=bool(caption_lines) or _caption_next_to(tbox, lines)):
            return None
        if found_by == "borderless" and _tab_list_in_context(grid, tbox, lines, dpi):
            return None                                 # a list on tab stops under its lead-in
        if found_by == "borderless" and _cut_out_of_a_record(tbox, lines):
            return None                                 # half a table: leave it to a reading that
            # sees the header and the first record as well (detect_aligned_tables, below)
        if found_by in ("inferred", "resplit") and _reads_like_prose(grid):
            return None                                 # paragraphs inside a page border, not a table
        # Headers printed sideways read as nothing, or as fragments; turn that band upright and
        # read it again. The band is put back together first, in case the row inference split it.
        if settings.sideways_headers and ocr is not None:
            merges = _collapse_header_band(grid, conf, cell_boxes, words_grid, merges)
            _add_sideways_headers(gray, grid, cell_boxes, ocr)
        table = TableData(grid=grid, conf=conf, merges=merges, cell_boxes=cell_boxes, cell_words=words_grid,
                          header_rows=_header_row_count(grid, merges))
        return table, tbox, inside, caption_lines

    for alternatives in candidates:
        # try the best reading of this area first and fall back to the other one, so that a reading
        # rejected by the checks above never costs us the table itself
        for et, found_by in alternatives:
            made = build(et, found_by)
            if made is not None:
                claimed |= made[2]
                results.append(made)
                break
    if settings.borderless_tables:
        results += detect_aligned_tables(lines, dpi, claimed)
    return results


def _caption_under_table(inside: set, lines: list):
    """The index of the line that is the table's caption printed UNDER it ("TABLE I. Interface characteristics."):
    the last line of the area, the only one there that reads as a caption, below at least two rows of the table.
    None when there is no such line - or several caption-like lines (a list of tables)."""
    caps = [i for i in inside if RE_CAPTION.match(lines[i].text.strip()) and _is_caption_text(lines[i].text)]
    if len(caps) != 1 or len(inside) < 3:
        return None
    box = lines[caps[0]].box
    h = max(1, box[3] - box[1])
    others = [lines[j].box for j in inside if j != caps[0]]
    if any(b[3] > box[1] + 0.5 * h for b in others):
        return None                                   # a line level with it or below it: not under the table
    return caps[0]


def _et_box(et) -> Box:
    """The page box of one table found by img2table."""
    return (et.bbox.x1, et.bbox.y1, et.bbox.x2, et.bbox.y2)


def _rows_are_under_split(et) -> bool:
    """True when the rows of a detected table have each swallowed several rows of text - the sign
    of row separators that were never printed, or that were too faint to be seen.

    A cell whose text merely WRAPS onto a second line also holds more than one line, so counting
    lines proves nothing on its own. img2table joins the lines that share a cell with newlines, and
    two patterns in those give a missing separator away:

    * MOST columns of one grid row hold several lines at once, and at least one holds three or
      more. When a cell is only wrapping, just the wide column does.
    * A verification matrix is mostly ticks, so only its LEFT-HAND column shows the crowding: one
      column holds more lines than the whole table has rows, while two or more of the others hold
      nothing but marks. The row count is what keeps an ordinary table with one deeply wrapped
      cell out of this: there, the table has more rows than any one cell has lines."""
    n_rows = len(et.content)
    for cells in et.content.values():
        filled = [c for c in cells if (c.value or "").strip()]
        if len(filled) < 2:
            continue
        depths = [1 + (c.value or "").count("\n") for c in filled]
        crowded = sum(1 for d in depths if d >= 2)
        if crowded >= max(2, 0.6 * len(filled)) and max(depths) >= 3:
            return True
        def only_marks(cell):
            return all(RE_MARK_CELL.match(v.strip()) for v in (cell.value or "").split("\n") if v.strip())

        marks_only = sum(1 for c in filled if only_marks(c))
        deepest = filled[depths.index(max(depths))]
        if max(depths) > n_rows and marks_only >= 2 and not only_marks(deepest):
            return True
    return False


def _better_row_split(printed: list, inferred: list) -> list:
    """Choose, area by area, between a table as its printed ruling lines describe it and the same
    table with the missing separators worked out from the text.

    Each area comes back as a list of readings to TRY IN ORDER, best first, so that a reading later
    judged not to be a table at all falls back to the other one instead of losing the table
    altogether. The inferred reading goes first only where the printed one cannot be right:
      * nothing at all was found in that area - the columns are ruled but the rows are not, so
        there are no two horizontal lines to make a row from; or
      * the printed rows have each swallowed several rows of text (see _rows_are_under_split).
    A table that prints all of its ruling lines therefore keeps exactly the grid it already had."""
    kept = [[alt] for alt in printed]
    for cand in inferred:
        cb = _et_box(cand)
        overlaps = [i for i, (et, _) in enumerate(printed)
                    if box_intersection(cb, _et_box(et)) > 0.5 * min(box_area(cb), box_area(_et_box(et)))]
        # The inference can split a table and still leave several records sharing a row - it gives
        # up early on a badly printed page. Where its own reading is STILL under-split, cutting its
        # rows from the text (_resplit_rows) is tried first, with the inference behind it.
        readings = ([(cand, "resplit")] if _rows_are_under_split(cand) else []) + [(cand, "inferred")]
        if not overlaps:
            kept.append(readings)                    # nothing was found here without the inferred rows
            continue
        for i in overlaps:
            et = printed[i][0]
            if len(cand.content) > len(et.content) and _rows_are_under_split(et):
                kept[i] = readings + [printed[i]]    # prefer them, but keep the printed fallback
    # Where the printed rules still do not describe the rows and the inference offered nothing
    # better, cut the rows ourselves from where the text sits (_resplit_rows).
    for i, alts in enumerate(kept):
        if len(alts) == 1 and alts[0][1] == "ruled" and _rows_are_under_split(alts[0][0]):
            kept[i] = [(alts[0][0], "resplit"), alts[0]]
    return kept


def _cut_out_of_a_record(tbox: Box, lines: list) -> bool:
    """True when a candidate's top edge sits INSIDE a record instead of above the table.

    A borderless table is found from where the text lines up, and that can lock on to the lower
    half of one - taking the wrapped lines for rows and missing the header and the first record
    above them. The giveaway is that the text immediately above almost touches it: a table starts
    after a gap, never a fraction of a line below the line before it."""
    inside = [ln for ln in lines if box_center_inside(ln.box, tbox)]
    above = [ln for ln in lines
             if ln.box[3] <= tbox[1] + 2 and min(ln.box[2], tbox[2]) > max(ln.box[0], tbox[0])]
    if not inside or not above:
        return False
    median_h = statistics.median(ln.box[3] - ln.box[1] for ln in inside)
    return tbox[1] - max(ln.box[3] for ln in above) < 0.5 * median_h


def _gap_threshold(gaps: list, median_h: float) -> Optional[float]:
    """A table's own dividing line between "the next wrapped line" and "the next record".

    Both spacings come from the table's leading, so both are constant through it while differing
    from table to table - which is why this is measured per table instead of being a fixed number.

    The split is taken at the LOWEST clear step in the sorted gaps, not the biggest one. The gaps
    INSIDE a record are all alike, but the gaps BETWEEN records are not: a tall header band, or a
    legend, sits closer to the first record than the records sit to one another. Taking the biggest
    step then draws the line between two kinds of record gap and swallows a whole record into the
    header. None when no step is clear, meaning every line is spaced alike and so is a record of
    its own."""
    if len(gaps) < 3:
        return None
    ordered = sorted(gaps)
    least = 0.25 * max(median_h, 1.0)
    for i in range(len(ordered) - 1):
        if ordered[i + 1] - ordered[i] >= least:
            return (ordered[i] + ordered[i + 1]) / 2
    return None


def _looks_like_row_key(text: str) -> bool:
    """A left-hand cell that NAMES a record: a paragraph number (3.2.1.4) or an identifier
    (SRS-101, IFC-010). A cell holding one of these cannot be a line that something wrapped onto."""
    t = (text or "").strip()
    return bool(RE_MATRIX_PARA.match(t) or (RE_ID_CELL.match(t) and any(ch.isdigit() for ch in t)))


def _gap_starts(row_ink: list, median_h: float) -> set:
    """Rows that begin a record judged on spacing alone (see _gap_threshold). Where every line is
    spaced alike there is nothing to group, so every line is a record of its own."""
    gaps = [row_ink[i][0] - row_ink[i - 1][1] for i in range(1, len(row_ink))]
    cut = _gap_threshold(gaps, median_h)
    if cut is None:
        return set(range(len(row_ink)))
    return {0} | {i for i, gap in enumerate(gaps, start=1) if gap > cut}


def _record_breaks(row_ink: list, row_columns: list, median_h: float, row_keys: list = None) -> list:
    """Which rows of text BEGIN a record of a table, given each row's ink (y0, y1), the set of
    columns it fills, and (where known) the text of its left-hand cell. A record is one row of the
    table plus the lines its cells wrap onto.

    A row whose left-hand cell holds a new key - a paragraph number or an identifier - always
    begins a record, whatever the spacing says. That is what stops a tall header band, whose ink
    reaches down close to the first record, from swallowing it.

    Otherwise three things are tried, strongest first:
      1. The key column. Where the left-hand column carries a value on some rows but not others -
         an ID, a paragraph number, a signal name - those rows are the records and the rest are
         wrapped lines. This is what the table itself says, so it is believed first.
      2. The leading. Lines wrapped inside a record sit closer together than records do, so one
         threshold per table separates them (_gap_threshold).
      3. Neither: every line is a record, which is what a matrix of one-line entries looks like."""
    n = len(row_ink)
    if n < 2:
        return list(range(n))
    by_spacing = _gap_starts(row_ink, median_h)
    named = {i for i, key in enumerate(row_keys or []) if _looks_like_row_key(key)}
    if 2 <= len(named) < n:
        # The left-hand column NAMES the records, which settles the data rows outright. What sits
        # above the first named row is the table's header and legend bands - they have no key of
        # their own, so their spacing is what separates them from one another.
        first = min(named)
        return sorted({0} | named | {i for i in by_spacing if i < first})
    keyed = [i for i, cols in enumerate(row_columns) if 0 in cols]
    if 2 <= len(keyed) < n:
        return sorted({0} | set(keyed))
    return sorted(by_spacing)


def _merge_continuation_rows(cell_boxes: list, lines: list) -> list:
    """Join the grid rows that are really the wrapped lines of one record.

    Row separators that were never printed have to be worked out from the text, and both img2table
    and _resplit_rows put every LINE of text on a row of its own. That is right for a matrix whose
    entries are one line each, and wrong for a requirement table, where a cell wraps over two or
    three lines and all of them belong to the record that started it."""
    if len(cell_boxes) < 2:
        return cell_boxes
    row_ink, row_columns, row_keys = [], [], []
    for row in cell_boxes:
        inside, columns = [], set()
        for c, cb in enumerate(row):
            here = [ln for ln in lines if cb is not None and box_center_inside(ln.box, cb)]
            if here:
                columns.add(c)
                inside += here
        present = [cb for cb in row if cb is not None]
        fallback = box_union(present) if present else (0, 0, 0, 0)
        row_ink.append((min((ln.box[1] for ln in inside), default=fallback[1]),
                        max((ln.box[3] for ln in inside), default=fallback[3])))
        row_columns.append(columns)
        first = [ln for ln in lines if row and row[0] is not None and box_center_inside(ln.box, row[0])]
        row_keys.append(_join_words_in_reading_order([w for ln in first for w in ln.words]))
    heights = [ln.box[3] - ln.box[1] for ln in lines]
    starts = _record_breaks(row_ink, row_columns, statistics.median(heights) if heights else 0.0, row_keys)
    if len(starts) == len(cell_boxes):
        return cell_boxes
    merged = []
    for first, stop in zip(starts, starts[1:] + [len(cell_boxes)]):
        block = cell_boxes[first:stop]
        merged.append([box_union([row[c] for row in block if c < len(row) and row[c] is not None])
                       if any(c < len(row) and row[c] is not None for row in block) else None
                       for c in range(max(len(row) for row in block))])
    return merged


def _resplit_rows(cell_boxes: list, lines: list) -> Optional[list]:
    """New cell boxes for a table whose printed ruling lines do not describe its rows.

    This is the last resort, for when img2table's own row inference gives up as well - it does that
    on badly degraded scans, where the ink is too broken for it to measure a row pitch. Only the
    horizontal cuts are new: a grid row that has swallowed several rows of text is cut half-way
    down the white gap between them, and the columns are left exactly as the printed rules give
    them. Returns None when there was nothing to cut."""
    out, changed, cut_to = [], False, None
    for row in cell_boxes:
        present = [cb for cb in row if cb is not None]
        if not present:
            out.append(row)
            continue
        band = box_union(present)
        # A cell merged down the table gives every row it covers the same tall box, so the same
        # band arrives again and again. Cutting it each time would repeat the whole table, so an
        # area already cut is passed over.
        if cut_to is not None and band[3] <= cut_to + 1:
            continue
        text_rows = _rows_of([ln for ln in lines if box_center_inside(ln.box, band)])
        if len(text_rows) < 2:
            out.append(row)
            cut_to = band[3] if cut_to is None else max(cut_to, band[3])
            continue
        cut_to = band[3] if cut_to is None else max(cut_to, band[3])
        cuts = [band[1]]
        for above, below in zip(text_rows, text_rows[1:]):
            cuts.append(int((max(ln.box[3] for ln in above) + min(ln.box[1] for ln in below)) / 2))
        cuts.append(band[3])
        out += [[None if cb is None else (cb[0], y0, cb[2], y1) for cb in row]
                for y0, y1 in zip(cuts, cuts[1:])]
        changed = True
    return out if changed else None


def _is_merged_band(row: list) -> bool:
    """A row that is one piece of text merged across the table - a title, a banner, or the LEGEND
    a quality conformance matrix prints above its columns. A merged cell repeats its text in every
    position it covers, which is what this recognises."""
    cells = [(c or "").strip() for c in row if (c or "").strip()]
    return len(cells) > 1 and len(set(cells)) == 1


def _header_row_count(grid: list, merges: list) -> int:
    """How many rows at the top of a table are header rows (nearly always one).

    A second row counts when the first is a band across the whole width and the row underneath
    reads like headers: a "VERIFICATION METHOD" banner over the I / A / D / T letters, or a LEGEND
    explaining what the numbered and lettered columns mean. A legend is ordinary prose and reads
    nothing like a header, so it is recognised by its shape - one text merged across the table -
    which is what used to leave such a matrix with no header row recognised at all."""
    if not grid:
        return 0
    banner = any(r0 == 0 and r1 == 0 and c1 > c0 for r0, c0, r1, c1 in merges or [])
    if (banner or _is_merged_band(grid[0])) and len(grid) > 2 and _looks_like_header(grid[1]):
        return 2
    return 1 if _looks_like_header(grid[0]) else 0


# ---- 6c-2. Sideways (rotated) column headers -------------------------------------------
# Verification cross-reference matrices almost always print their method headers sideways
# ("INSPECTION", "ANALYSIS", ...), because the words are far wider than the columns they head.
# The recogniser reads upright text only, so such a band comes back either empty or as short
# fragments ("INS OON" for INSPECTION), and a matrix whose columns have no usable names cannot fill
# in the Verification column. The pass below cuts out the header band, turns it upright, reads it
# again and puts each header back on the column whose ink it was printed over.
def _unrotate_box(box: Box, rotation: int, shape: tuple) -> Box:
    """Map a box found in a turned-upright crop back onto the crop as it sits on the page.

    `rotation` is the quarter turn that was applied (90 = clockwise, 270 = anticlockwise) and
    `shape` is the (height, width) of the crop BEFORE it was turned."""
    h, w = int(shape[0]), int(shape[1])
    x0, y0, x1, y1 = box
    if rotation % 360 == 90:        # cv2.ROTATE_90_CLOCKWISE sends (x, y) to (h - y, x)
        return (y0, h - x1, y1, h - x0)
    if rotation % 360 == 270:       # cv2.ROTATE_90_COUNTERCLOCKWISE sends (x, y) to (y, w - x)
        return (w - y1, x0, w - y0, x1)
    return box


def _word_like_score(lines: list) -> float:
    """How much a reading looks like real words: confident, and mostly letters. Used to decide
    which way round a sideways header band was printed."""
    return sum(sum(1 for ch in ln.text if ch.isalpha()) * ln.conf
               for ln in lines if sum(1 for ch in ln.text if ch.isalpha()) >= 2)


def read_sideways_headers(gray: "np.ndarray", header_box: Box, column_bands: list, ocr,
                          rotation: Optional[int] = None) -> dict:
    """Read column headers that are printed sideways. Returns {column index: header text}.

    The header band is cut out and turned a quarter turn each way; the turn that produces the more
    word-like text wins. Every line read is mapped back onto the page, so a header lands on the
    column whose ink it was printed over."""
    x0, y0 = max(int(header_box[0]), 0), max(int(header_box[1]), 0)
    x1, y1 = min(int(header_box[2]), gray.shape[1]), min(int(header_box[3]), gray.shape[0])
    if x1 - x0 < 8 or y1 - y0 < 8:
        return {}
    strip = gray[y0:y1, x0:x1]
    best, best_score = {}, 0.0
    for rot in ((90, 270) if rotation is None else (rotation,)):
        try:
            found = ocr.read(rotate_right_angle(strip, rot))
        except Exception as exc:            # a failed extra read must never lose the table
            LOG.debug("Sideways header reading failed (%s)", exc)
            continue
        score = _word_like_score(found)
        if not found or score <= best_score:
            continue
        headers = {}
        for ln in found:
            bx0, _y0, bx1, _y1 = _unrotate_box(ln.box, rot, strip.shape)
            center = x0 + (bx0 + bx1) / 2
            col = next((c for c, (cx0, cx1) in enumerate(column_bands) if cx0 <= center <= cx1), None)
            if col is not None and ln.text.strip():
                headers[col] = (headers[col] + " " + ln.text.strip()).strip() if col in headers else ln.text.strip()
        if headers:
            best, best_score = headers, score
    return best


def _has_ink(gray: "np.ndarray", box: Box, min_fraction: float = 0.004) -> bool:
    """True when a region of the page holds a real amount of print, not just paper."""
    x0, y0 = max(int(box[0]), 0), max(int(box[1]), 0)
    patch = gray[y0:int(box[3]), x0:int(box[2])]
    return patch.size > 0 and float((patch < 160).mean()) >= min_fraction


def _looks_like_words(text: str) -> bool:
    """A cell holding an unbroken run of four or more letters reads like a real word. Sideways
    print read the normal way up comes back as short fragments ("INS OON" for INSPECTION)."""
    return bool(re.search(r"[A-Za-z]{4,}", text or ""))


def _longest_letter_run(text: str) -> int:
    return max((len(m) for m in re.findall(r"[A-Za-z]+", text or "")), default=0)


def _sideways_header_band(grid: list, cell_boxes: list) -> Optional[Box]:
    """The box of a header row that was probably printed sideways, or None if it was not.

    Either sign is enough. The header cells came back EMPTY - the recogniser found no text at all
    in that band. Or the header row is far TALLER than the data rows, which is what a quarter-turned
    word needs, and what was read there does not look like words: upright reading of sideways print
    returns short fragments rather than nothing."""
    boxes = cell_boxes[0]
    present = [cb for cb in boxes if cb is not None]
    if not present:
        return None
    band = box_union(present)
    if band[3] - band[1] < 8:
        return None
    header = grid[0]
    empty = sum(1 for c, cb in enumerate(boxes)
                if cb is not None and not (header[c] if c < len(header) else "").strip())
    if empty >= 2:
        return band
    body_heights = [cb[3] - cb[1] for row in cell_boxes[1:] for cb in row if cb is not None]
    if not body_heights or (band[3] - band[1]) < 2.0 * statistics.median(body_heights):
        return None
    garbled = sum(1 for c in range(len(boxes))
                  if not _looks_like_words(header[c] if c < len(header) else ""))
    return band if garbled >= 2 else None


def _row_height(cell_boxes_row: list) -> float:
    """The height of one grid row, taken from the tallest cell box it has."""
    present = [cb for cb in cell_boxes_row if cb is not None]
    return max((cb[3] - cb[1] for cb in present), default=0.0)


def _collapse_header_band(grid: list, conf: list, cell_boxes: list, words_grid: list, merges: list) -> list:
    """Put a sideways header band back together when the row inference cut it into slices.

    A quarter-turned header needs a tall band, and the fragments that upright reading finds in it
    ("DEM" above "ON") can look like two rows of text, so the band comes back split. A leading row
    holding almost nothing is one of those slices and never a row of data, so it is merged into the
    row below it. Only a band that is much taller than the body rows is treated this way.
    Returns the (possibly shortened) list of merged-cell ranges."""
    while len(grid) > 2 and sum(1 for c in grid[0] if (c or "").strip()) < 2:
        body = [_row_height(cell_boxes[r]) for r in range(2, len(cell_boxes))]
        # The row BELOW the sliver must itself be far taller than the body rows: that is the tall
        # band a sideways word needs. Rows of equal height are an ordinary table and are left alone.
        if not body or _row_height(cell_boxes[1]) < 1.8 * statistics.median(body):
            break
        for c in range(len(grid[1])):
            upper = (grid[0][c] if c < len(grid[0]) else "") or ""
            grid[1][c] = (upper + " " + (grid[1][c] or "")).strip()
            if c < len(conf[0]):
                conf[1][c] = min(conf[0][c], conf[1][c])
            if c < len(words_grid[0]):
                words_grid[1][c] = words_grid[0][c] + words_grid[1][c]
            above = cell_boxes[0][c] if c < len(cell_boxes[0]) else None
            here = cell_boxes[1][c] if c < len(cell_boxes[1]) else None
            if above is not None and here is not None:
                cell_boxes[1][c] = box_union([above, here])    # the crop must cover the whole band
            elif above is not None:
                cell_boxes[1][c] = above
        del grid[0], conf[0], cell_boxes[0], words_grid[0]
        merges = _merges_without_row(merges, 0)
    return merges


def _add_sideways_headers(gray: "np.ndarray", grid: list, cell_boxes: list, ocr) -> list:
    """Put back a table's column headers that were printed sideways. Returns the columns changed.

    A header is only written where it beats what upright reading produced, measured by the longest
    unbroken run of letters: "INSPECTION" beats the "INS OON" that reading it the wrong way up
    gives, while a header that was read properly in the first place is never overwritten."""
    if not grid or not cell_boxes or not cell_boxes[0]:
        return []
    boxes = cell_boxes[0]
    band = _sideways_header_band(grid, cell_boxes)
    if band is None or not _has_ink(gray, band):
        return []
    bands = [(cb[0], cb[2]) if cb is not None else (0, 0) for cb in boxes]
    filled = []
    for col, text in read_sideways_headers(gray, band, bands, ocr).items():
        if col < len(grid[0]) and _longest_letter_run(text) > _longest_letter_run(grid[0][col]):
            grid[0][col] = text
            filled.append(col)
    if filled:
        LOG.debug("Read %d sideways column header(s)", len(filled))
    return filled


# Cells of a verification matrix that mean "this method applies": X, a tick, a dot, a bullet.
RE_MARK_CELL = re.compile(r"^[xX×✓✔✕✖•●■√*+]$")
# A paragraph number as it appears in the left-hand column of a verification matrix ("3.2.1.4").
RE_MATRIX_PARA = re.compile(r"^(?:\d{1,2}|[A-Z])(?:\.[\dlIi|]{1,3}){1,7}\.?$")


def _mark_columns(data_rows: list) -> list:
    """Columns of these DATA rows (headers excluded) holding nothing but a "this applies" mark."""
    out = []
    for c in range(max((len(r) for r in data_rows), default=0)):
        values = [(r[c] or "").strip() for r in data_rows if c < len(r) and (r[c] or "").strip()]
        if values and all(RE_MARK_CELL.match(v) for v in values):
            out.append(c)
    return out


def _method_column_run(data_rows: list, n_cols: int) -> list:
    """The block of columns at the right-hand end of a verification matrix that hold only "this
    applies" marks, or nothing at all - that is, its method columns.

    A method column is often empty on a given page simply because no requirement on that page uses
    that method, so an empty column inside the block still counts; this is what lets the block be
    lined up with the legend even then. Returns [] unless at least two columns really carry a mark,
    or if a column of text sits among them (then it is not a block of method columns)."""
    def values(c):
        return [(r[c] or "").strip() for r in data_rows if c < len(r) and (r[c] or "").strip()]

    def marks_only(c):
        vals = values(c)
        return bool(vals) and all(RE_MARK_CELL.match(v) for v in vals)

    marked = [c for c in range(n_cols) if marks_only(c)]
    if len(marked) < 2:
        return []
    # The block starts at the first marked column and runs right over any empty columns after the
    # last one - a method can be unused on a given page - stopping at the first column holding
    # text. It does NOT have to reach the right-hand edge: a quality conformance matrix puts an
    # AQL and an inspection method to the right of its marks. It is not widened to the LEFT, where
    # an empty label column would be swallowed as though it were a method.
    lo, hi = min(marked), max(marked)
    while hi + 1 < n_cols and not values(hi + 1):
        hi += 1
    if any(values(c) and not marks_only(c) for c in range(lo, hi + 1)):
        return []                       # a column of text sits among them: not a block of methods
    return list(range(lo, hi + 1))


def _is_mark_matrix(grid: list) -> bool:
    """True for a verification cross-reference matrix: a column of paragraph numbers beside two or
    more columns that hold nothing but tick marks. Such a table is mostly empty on purpose, so the
    "is this really a table?" checks must not judge it on how many of its cells are filled."""
    for skip in (1, 2):                 # one header row, or a legend band above it
        data = grid[skip:]
        if len(data) < 2 or len(_mark_columns(data)) < 2:
            continue
        for c in range(max((len(r) for r in grid), default=0)):
            values = [(r[c] or "").strip() for r in data if c < len(r) and (r[c] or "").strip()]
            if len(values) >= max(2, 0.6 * len(data)) and all(RE_MATRIX_PARA.match(v) for v in values):
                return True
    return False


def column_headers(grid: list, header_rows: int) -> list:
    """One header per column, read from however many header rows the table has.

    With a banner merged across the top ("VERIFICATION METHOD" above I / A / D / T) the name of a
    column is the LOWEST header line that carries text, so the banner never hides the real name."""
    n_cols = max((len(r) for r in grid), default=0)
    out = []
    for c in range(n_cols):
        above = [(grid[r][c] or "").strip() for r in range(min(header_rows, len(grid))) if c < len(grid[r])]
        out.append(next((x for x in reversed(above) if x), ""))
    return out


def _names_requirement_column(header_row: list) -> bool:
    """True when a row of headers announces a column of requirement text ("Requirement",
    "Requirement Text", "Shall Statement"), so cells holding whole sentences are expected."""
    return any(re.search(r"\b(?:requirement|shall statement)\b", (c or "").lower()) for c in header_row or [])


def _reads_like_prose(grid: list) -> bool:
    """True when the "cells" are really running text - a paragraph or a numbered heading caught
    inside a border - rather than a table. Monospaced typewriter text lines up very well, so text
    that merely looks tidy must not be taken for a table."""
    prose_rows = 0
    for row in grid:
        cells = [c for c in row if c]
        for cell in cells:
            m = RE_PARA_NUM.match(cell)
            # A heading names something, so what follows the number has to be a word. Without that,
            # a verification matrix cell holding a column number beside its tick ("1 X") reads as
            # the heading "1. X" and the whole table is dismissed as running text.
            heading_word = bool(m and re.search(r"[A-Za-z]{3,}", m.group("rest")))
            if m and "." in m.group("num") + m.group("dot") and heading_word:
                return True                            # "4.1 Verification ..." : a numbered paragraph
            if heading_word and _is_heading_text(cell) and word_count(cell) >= 2:
                return True                            # "4. VERIFICATION"
        # one sentence running across the cells: "manufactured and" | "finished" | "in accordance with".
        # A lone tick in a verification matrix column ("Detection range" | "x") is not a word, so it
        # never counts as a sentence carrying on into the next cell.
        if any(re.search(r"[a-z,]$", a) and re.match(r"[a-z]", b) and word_count(a) + word_count(b) >= 3
               and not RE_MARK_CELL.match(a.strip()) and not RE_MARK_CELL.match(b.strip())
               for a, b in zip(cells, cells[1:])):
            prose_rows += 1
    return prose_rows >= max(2, 0.35 * len(grid))


def detect_aligned_tables(lines: list, dpi: float, claimed: set) -> list:
    """Borderless tables recognised from text alignment alone (the idea behind Camelot's
    'stream' mode): consecutive rows that break into the same columns. Works on noisy
    photocopies where image-based detection fails, because it only needs the OCR boxes."""
    free = [i for i in range(len(lines)) if i not in claimed]
    if len(free) < 6:
        return []
    char_h = statistics.median(lines[i].height for i in free)
    # rows of OCR line boxes; separate boxes on one row are candidate cells
    rows = []
    for row in _rows_of([lines[i] for i in free]):
        ids = [next(i for i in free if lines[i] is ln) for ln in row]
        segs = []
        for i in ids:                      # also split a box at very wide word gaps
            words = sorted(lines[i].words, key=lambda w: w.box[0])
            current = [words[0]] if words else []
            for w in words[1:]:
                if w.box[0] - current[-1].box[2] > 2.0 * char_h:
                    segs.append((i, current))
                    current = [w]
                else:
                    current.append(w)
            if current:
                segs.append((i, current))
        box = box_union(lines[i].box for i in ids)
        rows.append({"ids": ids, "segs": segs, "box": box})

    results = []
    k = 0
    while k < len(rows):
        if len(rows[k]["segs"]) < 2:
            k += 1
            continue
        run = [k]
        m = k + 1
        while m < len(rows):
            gap = rows[m]["box"][1] - rows[run[-1]]["box"][3]
            if gap > 2.2 * char_h:
                break
            if len(rows[m]["segs"]) >= 2:
                run.append(m)                 # another row of the table
                m += 1
                continue
            row_text = " ".join(lines[i].text for i in rows[m]["ids"]).strip()
            if RE_CAPTION.match(row_text) and _is_caption_text(row_text):
                break                         # "TABLE I. Interface characteristics." close under it: its caption
            # A line with no cells of its own is a cell still wrapping, but only while it sits
            # tight against the line above the way wrapped lines do; set off by more than that it
            # is the caption of the next table, or the text below this one. However MANY wrapped
            # lines follow is not the question - demanding that the next row have cells again, as
            # this once did, ended the table at the first cell that wrapped over two lines.
            if gap > 0.8 * char_h:
                break
            run.append(m)
            m += 1
        k = max(m, k + 1)          # always move on, whatever the code below does with its own names
        multi = [r for r in run if len(rows[r]["segs"]) >= 2]
        if len(multi) < 3:
            continue
        # Which rows start a record, and which are the lines its cells wrap onto. A wrapped line
        # can itself break into two segments ("completed without a fault." | "reset." - two columns
        # wrapping at once), so the count of segments cannot decide this; _record_breaks uses the
        # key column and the table's own leading instead.
        provisional = _merge_intervals([(min(w.box[0] for w in ws), max(w.box[2] for w in ws))
                                        for r in multi for _, ws in rows[r]["segs"]], 0.6 * char_h)
        if len(provisional) < 2:
            continue
        row_ink = [(rows[r]["box"][1], rows[r]["box"][3]) for r in run]
        row_cols = [{_nearest_band(w, provisional) for _i, ws in rows[r]["segs"] for w in ws} for r in run]
        row_keys = [_join_words_in_reading_order([w for _i, ws in rows[r]["segs"] for w in ws
                                                  if _nearest_band(w, provisional) == 0]) for r in run]
        starts = set(_record_breaks(row_ink, row_cols, char_h, row_keys))
        # The bands are measured from the record rows only: a wrapped line breaks where its text
        # happens to end, and letting that count would invent columns of its own.
        bands = _merge_intervals([(min(w.box[0] for w in ws), max(w.box[2] for w in ws))
                                  for pos, r in enumerate(run) if pos in starts
                                  for _, ws in rows[r]["segs"]], 0.6 * char_h)
        if len(bands) < 2 or sum(1 for pos in starts if len(rows[run[pos]]["segs"]) >= 2) < 3:
            continue
        grid, conf, words_grid, used = [], [], [], set()
        for pos, r in enumerate(run):
            cells = [[] for _ in bands]
            for i, ws in rows[r]["segs"]:
                for w in ws:
                    cells[_nearest_band(w, bands)].append(w)
                used.add(i)
            texts = [_join_words_in_reading_order(c) for c in cells]
            confs = [min((w.conf for w in c), default=1.0) for c in cells]
            if pos not in starts and grid:                 # a wrapped line: add it to its own record
                for c, t in enumerate(texts):
                    if t:
                        grid[-1][c] = (grid[-1][c] + " " + t).strip()
                        conf[-1][c] = min(conf[-1][c], confs[c])
                        words_grid[-1][c] += [[w.text, round(w.conf, 4)] for w in cells[c]]
                continue
            grid.append(texts)
            conf.append(confs)
            words_grid.append([[[w.text, round(w.conf, 4)] for w in c] for c in cells])
        n_cells = len(grid) * len(bands)
        filled = sum(1 for row in grid for c in row if c)
        words_per_cell = sum(word_count(c) for row in grid for c in row) / max(filled, 1)
        # Long cells normally mean running text rather than a table - except where the header row
        # says outright that a column holds requirement sentences, which are expected to be long.
        wordy_limit = 40 if _names_requirement_column(grid[0] if grid else []) else 10
        if ((filled < 0.6 * n_cells and not _is_mark_matrix(grid)) or words_per_cell > wordy_limit
                or not _plausible_borderless_table(grid, n_cells, filled,
                                                   captioned=_caption_next_to(box_union(lines[i].box for i in used), lines))):
            continue
        tbox = box_union(lines[i].box for i in used)
        if _tab_list_in_context(grid, tbox, lines, dpi):
            continue                                       # a list on tab stops under its lead-in
        table = TableData(grid=grid, conf=conf, cell_words=words_grid, header_rows=_header_row_count(grid, []))
        results.append((table, tbox, used, []))
        claimed |= used
    return results


def _fill_cells(cell_boxes: list, table_lines: list) -> tuple:
    """Assign each OCR word to the grid cell it overlaps most; merged cells share one box."""
    unique = {}
    for r, row in enumerate(cell_boxes):
        for c, cb in enumerate(row):
            if cb is not None:
                unique.setdefault(cb, []).append((r, c))
    cells = list(unique.keys())
    cell_words = defaultdict(list)
    cell_lines = defaultdict(set)                         # cell -> indexes of the table_lines its words came from
    words = [wd for ln in table_lines for wd in ln.words]
    word_line = [li for li, ln in enumerate(table_lines) for _ in ln.words]
    if cells and words:
        C = np.asarray(cells, dtype=np.float64)
        centers = np.stack([(C[:, 0] + C[:, 2]) / 2, (C[:, 1] + C[:, 3]) / 2], axis=1)
        for start in range(0, len(words), 400):          # chunks keep memory small on huge tables
            chunk = words[start:start + 400]
            W = np.asarray([w.box for w in chunk], dtype=np.float64)
            ix = np.clip(np.minimum(W[:, None, 2], C[None, :, 2]) - np.maximum(W[:, None, 0], C[None, :, 0]), 0, None)
            iy = np.clip(np.minimum(W[:, None, 3], C[None, :, 3]) - np.maximum(W[:, None, 1], C[None, :, 1]), 0, None)
            area = np.maximum((W[:, 2] - W[:, 0]) * (W[:, 3] - W[:, 1]), 1.0)
            ratio = ix * iy / area[:, None]
            best = ratio.argmax(axis=1)
            wc = np.stack([(W[:, 0] + W[:, 2]) / 2, (W[:, 1] + W[:, 3]) / 2], axis=1)
            nearest = np.abs(wc[:, None, :] - centers[None, :, :]).sum(axis=2).argmin(axis=1)
            for k, wd in enumerate(chunk):
                # a word sitting on a ruling line (little overlap with any cell) goes to the nearest cell
                idx = best[k] if ratio[k, best[k]] >= 0.3 else nearest[k]
                cell_words[cells[int(idx)]].append(wd)
                cell_lines[cells[int(idx)]].add(word_line[start + k])
    n_rows, n_cols = len(cell_boxes), max((len(r) for r in cell_boxes), default=0)
    grid = [["" for _ in range(n_cols)] for _ in range(n_rows)]
    conf = [[1.0 for _ in range(n_cols)] for _ in range(n_rows)]
    words_grid = [[[] for _ in range(n_cols)] for _ in range(n_rows)]
    row_lines = [set() for _ in range(n_rows)]           # grid row -> indexes of the lines placed in it
    merges = []
    for cb, positions in unique.items():
        for r, _c in positions:
            row_lines[r] |= cell_lines.get(cb, set())
        words = cell_words.get(cb, [])
        text = _join_words_in_reading_order(words)
        cmin = min((w.conf for w in words), default=1.0)
        for r, c in positions:
            grid[r][c], conf[r][c] = text, cmin
            words_grid[r][c] = [[w.text, round(w.conf, 4)] for w in words]
        if len(positions) > 1:
            rs, cs = [p[0] for p in positions], [p[1] for p in positions]
            merges.append((min(rs), min(cs), max(rs), max(cs)))
    n_filled = sum(1 for cb in unique if cell_words.get(cb))
    return grid, conf, merges, len(unique), n_filled, sum(len(v) for v in cell_words.values()), words_grid, row_lines


def _pages_text(pages: list) -> str:
    """[1, 2, 3, 7] -> '1-3, 7' (shortened after a few ranges)."""
    ranges = []
    for p in sorted(set(pages)):
        if ranges and p == ranges[-1][1] + 1:
            ranges[-1][1] = p
        else:
            ranges.append([p, p])
    parts = [str(a) if a == b else f"{a}-{b}" for a, b in ranges]
    return ", ".join(parts[:6]) + (" ..." if len(parts) > 6 else "")


def _merges_without_row(merges: list, row: int) -> list:
    """Merged-cell ranges (r0, c0, r1, c1) after grid row `row` was deleted: later rows move up one,
    a range that covered the deleted row shrinks, and a range left with a single cell is dropped."""
    out = []
    for r0, c0, r1, c1 in merges:
        if r0 > row:
            r0, r1 = r0 - 1, r1 - 1
        elif r1 >= row:
            r1 -= 1
        if r1 >= r0 and (r1 > r0 or c1 > c0):
            out.append((r0, c0, r1, c1))
    return out


def _first_column_counts_like_a_list(cells: list, n_cols: int = 2, captioned: bool = False) -> bool:
    """True when most cells of a first column BEGIN with list markers that count up - "a. Acceleration:",
    "b. Vibration:", "c. Shock:" or "1.", "2.", "3." - i.e. a list laid out on tab stops, which lines up
    exactly like a table without ruling lines. A misread marker ("C." for "c.") still counts. So does a
    first column of short labels ("POWER:", "FAULT:"): a bulleted list whose bullets the OCR did not see
    (unless a caption - "TABLE I. ..." - right above or below it says it is a table)."""
    texts = [(c or "").strip() for c in cells if (c or "").strip()]
    # "POWER: | Green, ..." - labels and values, the bullets not read. Only a grid of exactly two columns
    # whose EVERY first cell is a label: a header row ("Parameter") or a third column makes it a table.
    if n_cols == 2 and len(texts) >= 2 and not captioned and all(_is_item_label(t) and word_count(t) <= 4 for t in texts):
        return True
    markers = []
    for text in texts:
        m = RE_LIST_MARKER.match(text) or RE_NUMBERED_ITEM.match(text) or RE_MARKER_SHAPE.match(text)
        if m and _marker_parts(m.group("marker")):
            markers.append(m.group("marker"))
        elif _marker_parts(text):
            markers.append(text)                        # the cell is only a marker: "ii."
    if len(markers) < max(2, 0.5 * len(texts)):
        return False
    roman = any((_marker_parts(mk) or ("", "", ""))[1].lower() in _ROMANS[1:9] for mk in markers)
    steps = sum(1 for a, b in zip(markers, markers[1:]) if _next_marker(a, roman).lower() == b.lower())
    return steps >= max(1, 0.5 * (len(markers) - 1))


def _has_real_extra_column(grid: list) -> bool:
    """Does a grid of three or more columns have a real column after its second one - short cells
    lined up on every row, such as a Method column ("Test", "Test", "Analysis") or a Units column?
    A list item's value that the OCR split at its word gaps also gives extra 'columns', but ragged ones:
    pieces of one phrase ("150 W | maximum.", "as in | MIL-STD-704F.", "50 ms | after loss of | input.")."""
    rows = [r for r in grid if r and any((c or "").strip() for c in r)]
    for c in range(2, max((len(r) for r in rows), default=0)):
        cells = [((r[c - 1] or "").strip(), (r[c] or "").strip()) for r in rows if c < len(r) and (r[c] or "").strip()]
        if len(cells) < max(2, 0.75 * len(rows)) or any(word_count(t) > 3 for _, t in cells):
            continue
        # a piece of a phrase: it ends the sentence, carries on in lower case, or finishes the cell before it
        pieces = sum(1 for before, t in cells if t.endswith(".") or (t[:1].islower() and word_count(t) > 1)
                     or _ENDS_MID_PHRASE.search(before))
        if len({t.lower() for _, t in cells}) <= max(2, 0.6 * len(cells)) or pieces <= 0.2 * len(cells):
            return True
    return False


def _join_wrapped_labels(cells: list) -> list:
    """First cells of a grid, with a label that wraps onto the next row put back together:
    'Storage' + 'temperature:' -> 'Storage temperature:' (a list on tab stops whose labels wrap)."""
    out, k = [], 0
    while k < len(cells):
        t = cells[k]
        if k + 1 < len(cells) and not t.endswith(":") and cells[k + 1][:1].islower() and cells[k + 1].endswith(":"):
            joined = f"{t} {cells[k + 1]}"
            label = joined.split(":", 1)[0]
            if RE_ITEM_LABEL.match(joined) and word_count(label) <= 8 and not _modal_re().search(label):
                out.append(joined)
                k += 2
                continue
        out.append(t)
        k += 1
    return out


def _is_caption_text(text: str) -> bool:
    """'TABLE I. Interface characteristics.' is a caption; 'Table 4 gives the limits ...' is a sentence
    (the same test the layout step uses for caption blocks)."""
    t = (text or "").strip()
    cap = RE_CAPTION.match(t)
    if not cap or word_count(t) > 40 or _modal_re().search(t):
        return False
    title = (cap.group("title") or "").strip()
    return bool(cap.group("delim") or not title or title[:1].isupper() or RE_CONTINUED.search(title))


def _caption_next_to(tbox, lines: list) -> bool:
    """Is the line right above a grid (a blank line may stand between them), or the line straight under
    it, a table caption ("TABLE I. Interface ...")? A caption further down, after a blank line, is more
    often the NEXT table's or figure's than this one's."""
    if not lines:
        return False
    h = statistics.median([ln.box[3] - ln.box[1] for ln in lines])
    centre = lambda ln: (ln.box[1] + ln.box[3]) / 2
    for near in ([ln for ln in lines if tbox[1] - 3.5 * h <= ln.box[3] <= tbox[1] + 0.3 * h and centre(ln) < tbox[1]
                  and ln.box[0] < tbox[2] and ln.box[2] > tbox[0]],
                 [ln for ln in lines if tbox[3] - 0.3 * h <= ln.box[1] <= tbox[3] + 1.2 * h and centre(ln) > tbox[3]
                  and ln.box[0] < tbox[2] and ln.box[2] > tbox[0]]):
        if near:
            y = centre(min(near, key=lambda ln: abs(centre(ln) - (tbox[1] + tbox[3]) / 2)))
            row = sorted((ln for ln in near if abs(centre(ln) - y) <= 0.5 * h), key=lambda ln: ln.box[0])
            if _is_caption_text(" ".join(ln.text for ln in row)):
                return True
    return False


def _tab_list_in_context(grid: list, tbox, lines: list, dpi: float) -> bool:
    """Is a borderless 'table' really a list on tab stops? Judged from what is around it, when its first
    column holds labels or markers ("Forward:", "Power (J1):" - one label may have lost its colon):
    list markers standing just left of its rows (1., 2., 3. that were not taken into the grid), a
    lead-in sentence right above it ("... under the following loads:"), or an item of the same list
    right above it ("Ripple:  1 V peak to peak maximum." - the list began before the grid did).
    A caption above or below it ("TABLE I. ..."), a header row (which is not a label), or a real third
    column ("Test", "Analysis") keeps it a table."""
    cells = [(r[0] or "").strip() for r in grid if r and (r[0] or "").strip()]
    first = _join_wrapped_labels(cells)
    wrapped = set(first) - set(cells)                  # 'Storage temperature:' put back together
    if not first or not lines:
        return False
    if max(len(r) for r in grid) > 2 and _has_real_extra_column(grid):
        return False
    if _caption_next_to(tbox, lines):
        return False                                   # "TABLE I. Interface characteristics." above or under it
    h = statistics.median([ln.box[3] - ln.box[1] for ln in lines])
    centre = lambda ln: (ln.box[1] + ln.box[3]) / 2
    # what is around the grid
    markers_left = [ln for ln in lines if tbox[0] - 1.0 * dpi <= ln.box[2] <= tbox[0] + 0.02 * dpi
                    and tbox[1] - 0.5 * h <= centre(ln) <= tbox[3] + 0.5 * h and _marker_parts(ln.text.strip())]
    above = [ln for ln in lines if tbox[1] - 2.5 * h <= ln.box[3] <= tbox[1] + 0.3 * h and ln.box[0] < tbox[2]
             and centre(ln) < tbox[1]]
    text, row = "", []
    if above:
        y = centre(max(above, key=lambda ln: ln.box[3]))
        row = sorted((ln for ln in above if abs(centre(ln) - y) <= 0.5 * h), key=lambda ln: ln.box[0])
        text = " ".join(ln.text for ln in row).strip()
    if _is_caption_text(text):
        return False
    lead_in = _introduces_list(text)
    # the row above is an item of the same list: a label at the grid's left edge with its value beside it
    item_above = bool(row) and abs(row[0].box[0] - tbox[0]) <= 1.0 * h and _is_item_label(text)
    if not (len(markers_left) >= 2 or lead_in or item_above):
        return False
    # the first cells: labels or markers. One label may have lost its colon ("Power  150 W maximum."), but a
    # first row that is not a label is a header ("Parameter | Value") - unless the list began above the grid
    is_label = lambda t: bool(_is_item_label(t) or RE_LIST_MARKER.match(t) or RE_NUMBERED_ITEM.match(t) or _marker_parts(t)
                              or t in wrapped)
    odd = [k for k, t in enumerate(first) if not is_label(t)]
    if not odd:
        return True
    t = first[odd[0]]
    return (len(odd) == 1 and (odd[0] > 0 or item_above) and len(first) >= (3 if item_above else 4)
            and word_count(t) <= 4 and t[:1].isupper() and not _ends_sentence(t))


def _plausible_borderless_table(grid: list, n_unique: int, n_filled: int, captioned: bool = False) -> bool:
    """Borderless candidates must look like real tables, not lists, headings or prose that happens
    to line up (monospaced typewriter text lines up very well).

    A verification cross-reference matrix is the one shape that breaks the usual "most cells are
    filled" rule, because a tick in one method column per row is exactly what it is meant to look
    like, so it is recognised first and judged on its shape instead."""
    n_rows, n_cols = len(grid), max(len(r) for r in grid)
    if n_rows < 3 or n_cols < 2:
        return False
    mark_matrix = _is_mark_matrix(grid)
    if n_cols > (20 if mark_matrix else 15):
        return False
    if n_filled < 0.5 * n_unique and not mark_matrix:
        return False
    first_col = [r[0] for r in grid if r and r[0]]
    if first_col and sum(1 for c in first_col if RE_BARE_MARKER.match(c)) >= 0.5 * len(first_col):
        return False                                   # "a. | Search mode." -> that is a list
    if _first_column_counts_like_a_list(first_col, n_cols, captioned):
        return False                                   # "a. Acceleration: | 5 g ..." -> a list on tab stops
    if _reads_like_prose(grid):
        return False
    well_filled = sum(1 for c in range(n_cols) if sum(1 for r in grid if c < len(r) and r[c]) >= 0.6 * n_rows)
    return well_filled >= 2 or mark_matrix


def _join_words_in_reading_order(words: list) -> str:
    """Order words by visual rows (top to bottom), then left to right."""
    if not words:
        return ""
    words = sorted(words, key=lambda w: (w.box[1] + w.box[3]) / 2)
    med_h = statistics.median(w.box[3] - w.box[1] for w in words)
    rows, current, current_y = [], [], None
    for w in words:
        cy = (w.box[1] + w.box[3]) / 2
        if current and abs(cy - current_y) > 0.5 * med_h:
            rows.append(current)
            current = []
        current.append(w)
        current_y = statistics.fmean((x.box[1] + x.box[3]) / 2 for x in current)
    rows.append(current)
    parts = []
    for row in rows:
        row_text = " ".join(w.text for w in sorted(row, key=lambda w: w.box[0]))
        if parts and parts[-1].endswith("-") and row_text[:1].islower():
            parts[-1] = parts[-1] + row_text       # word hyphenated across lines inside a cell
        else:
            parts.append(row_text)
    return " ".join(parts).strip()


def _looks_like_header(row: list) -> bool:
    cells = [c for c in row if c and c.strip()]
    if not cells:
        return False
    short = all(word_count(c) <= 6 for c in cells)
    numeric = sum(1 for c in cells if re.fullmatch(r"[\d.,%\-+ ]+", c)) / len(cells)
    # "Shall Statement" is a header label; "The unit shall ..." is a requirement sentence
    return short and numeric < 0.5 and not any(word_count(c) > 3 and _modal_re().search(c) for c in cells)


# ---- 6d. Figures (drawings, photos, diagrams) -----------------------------------------
def detect_figures(prep: PreparedPage, lines: list, table_boxes: list) -> list:
    """Find non-text graphics: ink that remains after masking text lines and tables."""
    dpi = prep.dpi
    view = prep.view_image
    h, w = view.shape
    _, bw = cv2.threshold(prep.ocr_image, 0, 255, cv2.THRESH_BINARY_INV + cv2.THRESH_OTSU)
    bw = cv2.bitwise_or(bw, (view < 140).astype(np.uint8) * 255)     # keep mid-gray photos too
    for ln in lines:                                                   # remove text
        p = int(0.25 * ln.height)
        cv2.rectangle(bw, (ln.box[0] - p, ln.box[1] - p), (ln.box[2] + p, ln.box[3] + p), 0, -1)
    tpad = int(0.04 * dpi)
    for tb in table_boxes:                                             # remove tables (incl. their border lines)
        cv2.rectangle(bw, (tb[0] - tpad, tb[1] - tpad), (tb[2] + tpad, tb[3] + tpad), 0, -1)
    edge = int(0.035 * min(h, w))                                      # scanner edges / punch holes margin
    bw[:edge, :], bw[-edge:, :], bw[:, :edge], bw[:, -edge:] = 0, 0, 0, 0
    k = _odd(dpi * 0.03)
    closed = cv2.morphologyEx(bw, cv2.MORPH_CLOSE, np.ones((k, k), np.uint8))
    n, _, stats, _ = cv2.connectedComponentsWithStats(closed, connectivity=8)
    parts = []
    for x, y, cw, ch, area in stats[1:]:
        if max(cw, ch) < 0.3 * dpi or min(cw, ch) < 0.05 * dpi:
            continue                                                   # specks, rules, underlines
        if cw * ch < 0.12 * dpi * dpi and area > 0.6 * cw * ch:
            continue                                                   # small solid blobs (punch holes)
        parts.append((x, y, x + cw, y + ch))
    # merge parts that are close to each other into whole figures
    gap = 0.25 * dpi
    merged = True
    while merged and len(parts) > 1:
        merged = False
        out = []
        while parts:
            a = parts.pop()
            for i, b in enumerate(out):
                if a[0] - gap <= b[2] and b[0] - gap <= a[2] and a[1] - gap <= b[3] and b[1] - gap <= a[3]:
                    out[i] = box_union([a, b])
                    merged = True
                    break
            else:
                out.append(a)
        parts = out
    figures = []
    for fb in parts:
        fw, fh = fb[2] - fb[0], fb[3] - fb[1]
        if fw < 0.75 * dpi or fh < 0.4 * dpi or fw * fh < 0.01 * h * w:
            continue
        if any(box_intersection(fb, tb) > 0.3 * box_area(fb) for tb in table_boxes):
            continue                                                   # leftovers of a table's frame
        figures.append(tuple(int(v) for v in fb))
    return figures


def _frame_is_caption_cell(fbox: Box, lines_inside: list, gray: "np.ndarray", dpi: float, text: str) -> bool:
    """A box drawn around a table's caption (the caption row of the table frame), not a figure: the text is
    a TABLE caption, the box is only a few text lines tall, and apart from that text and the frame itself
    there is no drawing inside it."""
    cap = RE_CAPTION.match(text)
    if not cap or not cap.group("kind").lower().startswith("tab") or not lines_inside:
        return False
    line_h = statistics.median(ln.height for ln in lines_inside)
    x0, y0, x1, y1 = (int(v) for v in fbox)
    if y1 - y0 > 5 * line_h + 0.15 * dpi:
        return False
    margin = max(3, int(0.04 * dpi))                    # the frame lines themselves
    ink = gray[y0 + margin:y1 - margin, x0 + margin:x1 - margin] < 128
    if ink.size == 0:
        return True
    pad = max(2, int(0.02 * dpi))
    for ln in lines_inside:                             # remove the text
        a0, b0 = max(0, int(ln.box[1]) - y0 - margin - pad), max(0, int(ln.box[0]) - x0 - margin - pad)
        a1, b1 = int(ln.box[3]) - y0 - margin + pad, int(ln.box[2]) - x0 - margin + pad
        ink[a0:max(a0, a1), b0:max(b0, b1)] = False
    return float(ink.mean()) < 0.002


def _figure_is_really_text(fig_box: Box, lines_inside: list) -> bool:
    """Guard against losing text: a 'figure' that is mostly text (e.g. a table whose ruling
    lines were not recognised) or that contains requirement sentences stays as text."""
    if not lines_inside:
        return False
    text_area = sum(box_area(ln.box) for ln in lines_inside)
    long_lines = sum(1 for ln in lines_inside if word_count(ln.text) >= 8)
    has_modal = any(_modal_re().search(ln.text) for ln in lines_inside)
    return text_area > 0.35 * box_area(fig_box) or long_lines >= 3 or has_modal


# ---- 6e. Reading order: recursive XY-cut ----------------------------------------------
def _nearest_band(word, bands: list) -> int:
    """The column band a word sits in, or the nearest one when it sits between two."""
    cx = (word.box[0] + word.box[2]) / 2
    return min(range(len(bands)), key=lambda b: 0 if bands[b][0] <= cx <= bands[b][1]
               else min(abs(cx - bands[b][0]), abs(cx - bands[b][1])))


def _merge_intervals(intervals: list, tolerance: float) -> list:
    out = []
    for a, b in sorted(intervals):
        if out and a <= out[-1][1] + tolerance:
            out[-1][1] = max(out[-1][1], b)
        else:
            out.append([a, b])
    return out


def _is_tab_not_gutter(left: list, right: list) -> bool:
    """True when a vertical band of whitespace is the wide TAB inside the rows of a list, not the
    gutter between two columns of text:

        a.      Acceleration:      5 g in each of the three orthogonal axes
                                   for 1 minute per axis.

    The OCR returns 'a.', 'Acceleration:' and the value as separate pieces, so the band looks like a
    column gutter; cutting there would put every label in one column and every value in another.
    The difference from real columns: the left-hand pieces are list markers or short labels (not lines
    of running text), and nearly each of them sits on the same line as text on the right. Only the
    rows from the first to the last one with text beside them are judged: a lead-in's last line above
    the list, or a NOTE or heading under it, sits in the same band but is not part of the list.

    The right-hand side is checked too, because a column of short list items ("a. MIL-STD-461G") can
    sit beside a column of prose on the same line grid. It is a real column when the right side holds
    numbered paragraphs ("3.1 Power. The unit ...", "3.2 Cooling. ..."), when it carries on the list
    of the left side (a-c on the left, d-f on the right), or when its lines run on as prose - full
    lines that carry on from one line to the next. The values of a tab list start afresh on each
    item's line, and a sub-list inside a value ("(1) Range ...", "(2) Azimuth ...") starts again under
    each item."""
    left_lines = [p for _, p in left if isinstance(p, Line)]
    right_lines = [p for _, p in right if isinstance(p, Line)]
    if len(left_lines) != len(left) or not right_lines:
        return False                                  # a table or figure on one side: a real split
    all_rows = _rows_of(left_lines)
    height = statistics.median([ln.box[3] - ln.box[1] for ln in left_lines + right_lines])
    centre = lambda row: (min(ln.box[1] for ln in row) + max(ln.box[3] for ln in row)) / 2
    right_rows = _rows_of(right_lines)
    right_y = [centre(r) for r in right_rows]
    beside = [any(abs(centre(row) - c) <= 0.5 * height for c in right_y) for row in all_rows]
    if not any(beside):
        return False
    first, last = beside.index(True), len(beside) - 1 - beside[::-1].index(True)
    rows, beside = all_rows[first:last + 1], beside[first:last + 1]
    left_texts = [" ".join(ln.text for ln in row).strip() for row in rows]
    cued = 0
    for text in left_texts:
        m = RE_LIST_MARKER.match(text) or RE_NUMBERED_ITEM.match(text) or RE_MARKER_SHAPE.match(text)
        cued += bool(m or RE_BARE_MARKER.match(text) or text.endswith(":"))
        # a marker and a label ("c. Peak inrush current (cold):") may run to 8 words; other lines of
        # more than 4 words are running text: this is a column
        label = _is_item_label(m.group("rest") if m else text)
        if word_count(text) > (8 if label else 4):
            return False
    n = len(rows)
    # a row ending in ':' with nothing beside it ("b. Random:" over its sub-items, a label's last line
    # "temperature:" under a one-line value) is part of a tab list's shape, not a sign of two columns
    lone_stems = sum(1 for t, b in zip(left_texts, beside) if t.endswith(":") and not b)
    if sum(beside) < max(1, 0.6 * (n - lone_stems)) or cued < 0.5 * n:
        return False
    right_texts = [" ".join(ln.text for ln in r).strip() for r in right_rows]
    left_y = [centre(row) for row in rows]
    all_left_y = [centre(row) for row in all_rows]
    # the left rows where an item STARTS: a marker, or a label after a complete one. A label's second
    # line ("temperature:" under "a. Storage") sits beside the value's second line, which runs on.
    starts_y = [y for k, (t, y) in enumerate(zip(left_texts, left_y))
                if RE_LIST_MARKER.match(t) or RE_NUMBERED_ITEM.match(t) or RE_MARKER_SHAPE.match(t) or _marker_parts(t)
                or (t.endswith(":") and (k == 0 or left_texts[k - 1].endswith(":")))]
    # numbered paragraphs on the right: two in sequence ("3.1 Power. ...", "3.2 Cooling. ..."), one that
    # follows a numbered paragraph on the left ("3.4 Modes. ..." | "3.5 Power. ..."), or one that starts
    # on the same line as something on the left. A titled number alone on a line of its own is a
    # cross-reference a value carries on with ("as specified in" / "4.5.3 Random vibration. Levels ...").
    titled = [(RE_PARA_NUM.match(t).group("num"), y) for t, y in zip(right_texts, right_y) if _is_titled_number(t)]
    left_titled = [RE_PARA_NUM.match(t).group("num") for t in left_texts if _is_titled_number(t)]
    if (any(plausible_next_number(a, b) for (a, _), (b, _) in zip(titled, titled[1:]))
            or any(plausible_next_number(a, b) for a in left_titled for b, _ in titled)
            or any(abs(y - c) <= 0.5 * height for _, y in titled for c in all_left_y)):
        return False
    # a list in a second column: the right-hand markers carry on from the last left-hand one (a-c | d-f),
    # or form one list of their own, item for item beside the left-hand items (a-e | a-e). A sub-list
    # inside the values ("(1) ...", "(2) ...") starts again under each item, beside no item of its own.
    def marker_of(t):
        m = RE_LIST_MARKER.match(t) or RE_NUMBERED_ITEM.match(t) or RE_MARKER_SHAPE.match(t)
        return m.group("marker") if m and _marker_parts(m.group("marker")) else ""
    left_markers = [mk for mk in map(marker_of, left_texts) if mk]
    right_marked = [(mk, y) for mk, y in zip(map(marker_of, right_texts), right_y) if mk]
    if left_markers and right_marked:
        right_markers = [mk for mk, _ in right_marked]
        roman = any(_marker_parts(mk)[1].lower() in _ROMANS[1:9] for mk in left_markers + right_markers)
        in_sequence = all(_next_marker(a, roman).lower() == b.lower() for a, b in zip(right_markers, right_markers[1:]))
        beside_items = sum(1 for _, y in right_marked if any(abs(y - c) <= 0.5 * height for c in starts_y))
        if right_markers[0].lower() == _next_marker(left_markers[-1], roman).lower() or \
                (len(right_markers) >= 3 and in_sequence and beside_items >= 2 / 3 * len(right_markers)):
            return False
    # The left-hand side holds nothing of its own - only markers and labels waiting for their values
    # ("a.", "Voltage:", "Storage" / "temperature:"): no column of text looks like that, so the values
    # beside it are not tested for prose (short values such as "within 22 to 29 VDC," would look like it).
    rests = []
    for t in left_texts:
        m = RE_LIST_MARKER.match(t) or RE_NUMBERED_ITEM.match(t) or RE_MARKER_SHAPE.match(t)
        rests.append("" if RE_BARE_MARKER.match(t) or (not m and _marker_parts(t)) else (m.group("rest") if m else t).strip())
    if all(not r or (r.endswith(":") and RE_ITEM_LABEL.match(r) and word_count(r) <= 8 and not _modal_re().search(r))
           for r in _join_wrapped_labels(rests)):
        return True
    # Prose wraps a line only when the next word does not fit on it. The values of a tab list stop
    # where each value ends: a sentence broken although the next word would have fitted is evidence
    # against running text (a ragged typewriter margin gives some of those too, so it is weighed).
    right_edge = max(ln.box[2] for ln in right_lines)
    right_w = right_edge - min(ln.box[0] for ln in right_lines)
    runs_on = early = shared = 0
    for k in range(1, len(right_rows)):
        if not any(abs(right_y[k] - y) <= 0.5 * height for y in starts_y):
            continue                                  # a value's own continuation line: nothing to compare
        shared += 1
        if _ends_sentence(right_texts[k - 1]) or not right_texts[k][:1].islower():
            continue
        # (a typist's margin wanders by a few characters, so only a clearly short line counts as early)
        room = right_edge - max(ln.box[2] for ln in right_rows[k - 1])
        if _first_word_width(right_rows[k]) + max(height, 0.25 * right_w) <= room:
            early += 1
        else:
            runs_on += 1
    # prose flowing past the list: a column
    return not (shared >= 2 and runs_on >= 0.5 * shared and runs_on > early)


def _first_word_width(row: list) -> float:
    """Width in pixels of the first word on a row of OCR pieces."""
    piece = min(row, key=lambda ln: ln.box[0])
    if getattr(piece, "words", None):
        w = min(piece.words, key=lambda w: w.box[0])
        return w.box[2] - w.box[0]
    text = piece.text.strip()
    return (piece.box[2] - piece.box[0]) * len(text.split()[0]) / max(len(text), 1) if text else 0.0


def _column_cut(items: list, region_w: float, min_vgap: float, tabs: tuple = ()) -> tuple:
    """The widest vertical whitespace band with real text on both sides -> ((gap, left, right) or None,
    [the bands found to be tabs]). The wide tabs inside a list ("a.   Label:   value") are not columns
    (see _is_tab_not_gutter). `tabs` are bands already found to be tabs in a larger region around this
    one: a few rows of the same list, cut off from the rest, must not be judged again on less evidence."""
    x_spans = _merge_intervals([(it[0][0], it[0][2]) for it in items], 0)
    best, found_tabs = None, []
    for (a0, a1), (b0, b1) in zip(x_spans, x_spans[1:]):
        gap = b0 - a1
        if gap < min_vgap or any(min(b0, t1) > max(a1, t0) for t0, t1 in tabs):
            continue
        left = [it for it in items if it[0][2] <= a1]
        right = [it for it in items if it[0][0] >= b0]
        lw = max(it[0][2] for it in left) - min(it[0][0] for it in left)
        rw = max(it[0][2] for it in right) - min(it[0][0] for it in right)
        if lw >= 0.2 * region_w and rw >= 0.2 * region_w and len(left) >= 2 and len(right) >= 2:
            if _is_tab_not_gutter(left, right):
                found_tabs.append((a1, b0))
            elif best is None or gap > best[0]:
                best = (gap, left, right)
    return best, found_tabs


def xy_cut(items: list, min_hgap: float, min_vgap: float, depth: int = 0, tabs: tuple = ()) -> list:
    """items = [(box, payload)]. Returns leaves (lists of items) in reading order.
    Columns are split first (a vertical whitespace band crossing the whole region, with
    real text on both sides); otherwise the region is split at horizontal gaps. `tabs`: whitespace
    bands found to be a list's tab stops in a region around this one (see _column_cut)."""
    if len(items) <= 1 or depth > 60:
        return [sorted(items, key=lambda it: (it[0][1], it[0][0]))] if items else []
    x_min, x_max = min(it[0][0] for it in items), max(it[0][2] for it in items)
    region_w = max(x_max - x_min, 1)
    # --- vertical cut (columns)
    best, found_tabs = _column_cut(items, region_w, min_vgap, tabs)
    tabs = tuple(tabs) + tuple(found_tabs)
    if best:
        return xy_cut(best[1], min_hgap, min_vgap, depth + 1, tabs) + xy_cut(best[2], min_hgap, min_vgap, depth + 1, tabs)
    # --- a line across the whole width (running header/footer, page-wide heading or table) hides the
    #     column gutter: cut the region above and below such lines first, then look for columns again.
    #     (Cutting at every horizontal gap would slice two columns into rows and mix their lines.)
    # (A long line that STARTS well to the right - the value column of a tabbed list - is not page-wide.)
    wide = {id(it) for it in items if it[0][2] - it[0][0] >= 0.6 * region_w and it[0][0] - x_min <= 0.25 * region_w}
    narrow = [it for it in items if id(it) not in wide]
    narrow_cut, narrow_tabs = _column_cut(narrow, region_w, min_vgap, tabs) if wide and len(narrow) >= 4 else (None, [])
    tabs = tabs + tuple(narrow_tabs)          # (the tabs of a list between two page-wide lines hold below, too)
    if narrow_cut:
        bands, current = [], []
        for it in sorted(items, key=lambda it: (it[0][1], it[0][0])):
            if id(it) in wide:
                if current:
                    bands.append(current)
                bands.append([it])
                current = []
            else:
                current.append(it)
        if current:
            bands.append(current)
        if len(bands) > 1:
            return [leaf for band in bands for leaf in xy_cut(band, min_hgap, min_vgap, depth + 1, tabs)]
    # --- horizontal cuts
    y_spans = _merge_intervals([(it[0][1], it[0][3]) for it in items], min_hgap)
    if len(y_spans) > 1:
        leaves = []
        for y0, y1 in y_spans:
            group = [it for it in items if y0 <= (it[0][1] + it[0][3]) / 2 <= y1]
            leaves.extend(xy_cut(group, min_hgap, min_vgap, depth + 1, tabs))
        return leaves
    return [sorted(items, key=lambda it: (it[0][1], it[0][0]))]


# ---- 6f. Grouping lines into headings / paragraphs / list items ---------------------------
_MODAL_RE = None
_MODAL_WORDS = ("shall", "must", "should", "will", "is required to", "are required to")
_SENTENCE_STARTERS = {"The", "A", "An", "This", "These", "Each", "All", "Any", "Every", "It", "When", "While", "If", "Where"}


def configure_keywords(keywords: Iterable) -> None:
    """Layout rules (headings, captions, figure labels) must never swallow a requirement: they use
    the same keywords as the extractor, plus the OCR-garbled spellings the extractor repairs."""
    global _MODAL_WORDS, _MODAL_RE
    _MODAL_WORDS = tuple(sorted({" ".join(str(k).lower().split()) for k in keywords if str(k).strip()}
                                | {"shall", "must", "should", "will"}, key=len, reverse=True))
    _MODAL_RE = None


def _modal_re():
    """Regex matching any requirement keyword (as configured) or an OCR-garbled shall/must/should/will."""
    global _MODAL_RE
    if _MODAL_RE is None:
        words = "|".join(r"\s+".join(re.escape(p) for p in w.split()) for w in _MODAL_WORDS)
        garbled = "|".join(f"(?:{pattern.pattern})" for _, pattern in OCR_KEYWORD_REPAIRS)
        _MODAL_RE = re.compile(r"(?<![A-Za-z])(?:" + words + r")(?![A-Za-z])|" + garbled, re.I)
    return _MODAL_RE
_ENDS_MID_PHRASE = re.compile(r"\b(?:of|in|the|and|or|to|for|with|a|an|by|on|at|from|as|that|which|section|per|than)\s*$", re.I)


def _ends_sentence(text: str) -> bool:
    return bool(RE_TERMINAL.search(text))


# Small letters that have the same shape as their capitals: the scan may read the capital of a title as one of
# these ("3.2.7 output. The power supply ..."), as it reads the list marker 'C.' for 'c.'
_SAME_SHAPE_LETTERS = set("cosvwxz")
_NOT_TITLE_STARTS = {"of", "or", "on", "so", "see", "section", "sub", "subparagraph", "subsection", "clause",
                     "chapter", "sheet", "series", "volume", "version", "shall", "will", "was", "with", "without",
                     "within", "whichever", "over", "outside", "for", "to", "and"}


def _lowercased_title(rest: str) -> bool:
    """'output. The power supply ...' after a paragraph number - a run-in title of one to six words and a full
    stop or colon, then a sentence, whose first letter the scan read in lower case. Only for letters whose
    small and capital forms have the same shape (c, o, s, v, w, x, z), and not for a word that starts a
    phrase ('3.2.4 of MIL-STD-810') or a unit ('3.2.4 s. The ...')."""
    t = re.match(r"(?P<title>[a-z][A-Za-z-]*(?: [A-Za-z][A-Za-z-]*){0,5})[.:]\s+[A-Z(\"]", rest or "")
    if not t or rest[:1] not in _SAME_SHAPE_LETTERS:
        return False
    first = t.group("title").split()[0]
    return (first.lower() not in _NOT_TITLE_STARTS and not RE_UNIT.fullmatch(first)
            and not _modal_re().search(t.group("title")))


def _is_titled_number(text: str) -> bool:
    """'3.1 Power. The unit shall ...' / '3.2.4.1 Mode switching. The radar ...' - a numbered paragraph
    with a run-in title followed by a sentence. A value that starts with a number ('2.5 G sustained
    thereafter.', '3.3 V logic, low level ...') has no such title, so this tells the two apart."""
    m = RE_PARA_NUM.match((text or "").strip())
    if not (m and "." in m.group("num") + m.group("dot")):
        return False
    title = re.match(r"(?P<title>[A-Z][A-Za-z0-9 ,/()&'-]{0,80}?)[.:]\s+[A-Z(\"]", m.group("rest"))
    if not title:
        return False
    words = title.group("title").split()
    first = words[0].rstrip(",")
    # "5.0 V. Ripple shall ...", "1.5 V peak to peak. Measured ...", "2.5 G rms overall. Duration ..." are
    # quantities: a unit (or a lone capital) straight after the number is not the start of a title -
    # unless a capitalised word follows it ("3.2.1 X Band.", "3.4.1 I Channel.", "3.2 W Band.")
    unit_like = RE_UNIT.fullmatch(first) or len(first) == 1
    titled_on = len(words) > 1 and words[1][:1].isupper()
    return not ((unit_like and not titled_on) or (_starts_with_unit(title.group("title")) and not titled_on))


# A sentence that leads into a list. It normally ends in ":", but the colon is often left out, or lost
# on the scan: "... shall meet each of the following.", "... within the following limits.",
# "... as follows.", "... as specified below." (up to four words may follow "the following").
RE_LEAD_IN = re.compile(r"(?:\bthe\s+following(?!\s+(?:working\s+|business\s+|calendar\s+)?(?:day|days|week|weeks|month|"
                        r"months|year|years|morning|afternoon|evening|night|shift|period|cycle)\b)(?:\s+[A-Za-z][A-Za-z-]*){0,4}"
                        r"|\bas\s+follows"
                        r"|\b(?:listed|specified|given|shown|defined|described|stated|detailed)\s+below)\s*[.:;]?\s*$", re.I)


def _introduces_list(text: str) -> bool:
    """True for a sentence that leads into a list: it ends in ":" or with a lead-in phrase."""
    t = (text or "").rstrip()
    return t.endswith(":") or bool(RE_LEAD_IN.search(t))


def _is_heading_text(text: str) -> bool:
    """Short title-like line: '3.2 Characteristics.', '1. SCOPE', 'REQUIREMENTS'."""
    t = text.strip()
    if not t or word_count(t) > 12 or _modal_re().search(t) or RE_LIST_MARKER.match(t):
        return False                               # "(1) SEARCH MODE." is a list item, not a heading
    m = RE_PARA_NUM.match(t)
    if m:
        rest = m.group("rest").strip()
        if RE_CAPTION.match(rest) or _starts_with_unit(rest):
            return False                           # "3.3 V shall ..." / "28 VDC" are quantities
        if re.match(r"[^.:;]{1,40}:\s+\S", rest):
            return False                           # "2. Vibration:  20 Hz ..." - a list item's label and value
        # a heading title has no sentence after it ("Detection range. The radar ..." is a paragraph),
        # does not stop in the middle of a phrase ("... requirement in section") and does not start
        # like a sentence ("3.2.1 The primary display units")
        first = rest.split()[0] if rest.split() else ""
        return (word_count(rest) <= 10 and not re.search(r"[.:;]+['`\u00b7]?\s*[A-Z(]", rest)
                and rest[:1].isupper() and not _ENDS_MID_PHRASE.search(rest) and first not in _SENTENCE_STARTERS)
    m = RE_NUMBERED_ITEM.match(t)
    if m and m.group("rest").isupper():
        return True
    letters = [c for c in t if c.isalpha()]
    return (len(letters) >= 4 and all(c.isupper() for c in letters) and word_count(t) <= 10
            and not RE_PAGE_NUMBER.match(t) and not _ENDS_MID_PHRASE.search(t)
            and not t.endswith(":"))                   # "NOTES:" introduces a list; it is not a section title


_QUANTITY_FOLLOWERS = {"to", "and", "or", "max", "min", "maximum", "minimum", "nominal", "dc", "ac", "rms", "peak",
                       "\u00b1", "+/-", "input", "output", "bus", "power", "supply", "nom"}


def _starts_with_unit(text: str) -> bool:
    """'V shall ...' / 'VDC input' after a number is a quantity ('3.3 V shall'), not a paragraph title.
    A single letter only counts as a unit when a keyword, number or 'to/max/...' follows ('3.1 A radar' is fine)."""
    toks = text.split()
    if not toks:
        return False
    first = toks[0].rstrip(".,;:")
    if not RE_UNIT.fullmatch(first):
        return False
    if len(first) >= 2 and first not in ("in", "min", "hr", "ft", "lb", "oz"):
        return True
    second = toks[1].lower().strip(".,;:") if len(toks) > 1 else ""
    return second in _QUANTITY_FOLLOWERS or bool(_modal_re().fullmatch(second)) or second[:1].isdigit()


def _id_prefix(rid: str) -> str:
    m = re.match(r"[A-Za-z][A-Za-z0-9]*?(?=[-_.]|\d|$)", rid or "")
    return m.group(0).upper() if m else ""


def _starts_with_delimited_id(text: str) -> bool:
    """'[SRS-101] ...', '(SRS-101) ...' or 'SRS-103: ...' - an ID that is clearly marked as one."""
    return bool(re.match(r"^\s*[\[(](?:" + RE_REQ_ID + r")[\])]|^\s*(?:" + RE_REQ_ID + r")\s*:", text))


def _starts_new_item(text: str) -> bool:
    t = text.strip()
    if RE_LIST_MARKER.match(t) or RE_CAPTION.match(t):
        return True
    m = RE_PARA_NUM.match(t)
    if m and "." in m.group("num") and not _starts_with_unit(m.group("rest")) and (
            m.group("rest")[:1].isupper()
            or (m.group("num").count(".") >= 2 and _lowercased_title(m.group("rest")))):   # "3.2.5 weight. The ..."
        return True
    m = RE_NUMBERED_ITEM.match(t)
    if m and m.group("rest")[:1].isupper():
        return True
    return _starts_with_delimited_id(t)


def _rows_of(lines: list) -> list:
    """Group line fragments that sit on the same visual row (left to right)."""
    rows = []
    for ln in sorted(lines, key=lambda l: (l.box[1], l.box[0])):
        if rows:
            last = rows[-1]
            ry0, ry1 = min(l.box[1] for l in last), max(l.box[3] for l in last)
            overlap = min(ry1, ln.box[3]) - max(ry0, ln.box[1])
            if overlap >= 0.5 * min(ln.height, ry1 - ry0):
                last.append(ln)
                continue
        rows.append([ln])
    return [sorted(r, key=lambda l: l.box[0]) for r in rows]


def caption_meta(text: str, default_kind: str) -> dict:
    """Caption kind and label ('TABLE 3-1') for a text already known to be a caption."""
    cap = RE_CAPTION.match(text.strip())
    if not cap:
        return {"caption_kind": default_kind, "caption_label": "", "continued": False}
    kind = "table" if cap.group("kind").lower().startswith("tab") else "figure"
    return {"caption_kind": kind, "caption_label": f"{cap.group('kind').upper().rstrip('.')} {cap.group('num')}",
            "continued": bool(RE_CONTINUED.search(cap.group("title") or ""))}


RE_ITEM_LABEL = re.compile(r"^[A-Z][A-Za-z0-9 /()&,+-]{0,40}:(?:\s|$)")     # "Acceleration:"  "Input voltage:"


_NOT_LABELS = {"NOTE", "NOTES", "CAUTION", "CAUTIONS", "WARNING", "WARNINGS", "RATIONALE", "EXAMPLE", "EXAMPLES",
               "EXCEPTION", "EXCEPTIONS", "REMARK", "REMARKS", "NB", "COMMENT", "COMMENTS"}


def _lowered_label(text: str) -> bool:
    """'current:  The unit shall ...' - an item's label whose capital the scan read as the small letter of the same
    shape (c, o, s, v, w, x, z; see _lowercased_title)."""
    t = (text or "").strip()
    return t[:1] in _SAME_SHAPE_LETTERS and _is_item_label(t[:1].upper() + t[1:])


def _is_item_label(text: str) -> bool:
    """'Acceleration:  5 g ...' / 'Input voltage:' - a list item's label. A notice ('NOTE:', 'WARNING:')
    is not an item of the list above it, and 'SRS-101:' is a requirement ID, not a label."""
    t = (text or "").strip()
    if not RE_ITEM_LABEL.match(t) or _starts_with_delimited_id(t):
        return False
    label = t.split(":", 1)[0].strip()
    # a label names a quantity in a few words; "Search mode, which shall include:" states a requirement
    return label.upper() not in _NOT_LABELS and word_count(label) <= 5 and not _modal_re().search(label)


def _starts_list_item(text: str, row: list) -> bool:
    """Does a block starting with this line (row = its OCR pieces) look like a list item?
    An ordinary marker ("a." "(2)" "iv)" "3."), a capital letter followed by a label ("A.  Flow rate:"),
    or a label on a line laid out with wide tabs ("POWER:        Green, lit when ...") - a bulleted
    list whose bullets the OCR did not see. A capital letter alone ("A. The unit ...") is not enough."""
    if RE_LIST_MARKER.match(text) or RE_NUMBERED_ITEM.match(text):
        return True
    shape = RE_MARKER_SHAPE.match(text)
    if shape and re.fullmatch(r"[A-Z][.)]", shape.group("marker")) and _is_item_label(shape.group("rest")):
        return True
    return bool((_is_item_label(text) or _lowered_label(text)) and len(row) > 1)


def _opens_sub_list(text: str) -> bool:
    """Does a list item's text so far lead into a list of its own ('a. Search mode, which shall
    include:', 'a. Vibration:  as follows:')? Then the lines below it are its sub-items, not more of its
    text. A bare label waiting for its value on the next line ('a. Acceleration:') does not."""
    t = (text or "").strip()
    if not t.endswith(":"):
        return False
    m = RE_LIST_MARKER.match(t) or RE_NUMBERED_ITEM.match(t) or RE_MARKER_SHAPE.match(t)
    rest = (m.group("rest") if m else t).strip()
    bare_label = bool(re.fullmatch(r"[A-Z][A-Za-z0-9 /()&,+-]{0,40}:", rest)) and _is_item_label(rest)
    return not bare_label


def _wrapped_label_rows(rows: list, line_h: float):
    """(value x, number of rows) when a list item's LABEL wraps beside its value on tab stops:

        a.   Storage          -55 C to +85 C for        a.   Operating        -40 C to +71 C.
             temperature:     up to 5 years.                 temperature:

    The label lines come first, one per row from the item's first row, and the label ends with ':'
    on its last line; each line is short. The value may be shorter than the label (right), so a label
    line need not have a value beside it. A first line that already ends with ':' is a whole label;
    lines under it without a ':' to end them are the value wrapping back under the label, or a note,
    and keep their place. None when the label does not wrap."""
    _, value_xs = _item_columns(rows[0])
    if not value_xs:
        return None
    vx = min(value_xs) - 0.3 * line_h
    run = []
    for k, row in enumerate(rows):
        left = [ln for ln in sorted(row, key=lambda l: l.box[0]) if ln.box[0] < vx]
        if not left:
            return None                                  # a value line alone before any ':' - no label here
        if max(ln.box[2] for ln in left) > vx + 0.3 * line_h:
            return None                                  # it runs on into the value column: the value wrapping back
        text = " ".join(ln.text for ln in left).strip()
        if k and (RE_LIST_MARKER.match(text) or RE_NUMBERED_ITEM.match(text) or RE_BARE_MARKER.match(text)):
            return None                                  # a sub-item, not more of the label
        if k and (word_count(text) > 6 or (_ends_sentence(text) and not text.endswith(":"))):
            return None
        run.append(text)
        if text.endswith(":"):
            break
    else:
        return None                                      # no ':' closes the label
    if len(run) < 2:
        return None                                      # the label was complete on its first line
    return float(vx), len(run)


def _item_columns(row: list) -> tuple:
    """(label x, [value x]) - where the text of a list item's first row continues after its marker:
    the word after the marker (the label, or the text) and, after a label's colon or a wide tab,
    the value. Lines that start in one of these columns carry on the item."""
    pieces = sorted(row, key=lambda ln: ln.box[0])
    words = [w for ln in pieces for w in ln.words]
    label_x = words[1].box[0] if len(words) > 1 else None
    values = [ln.box[0] for ln in pieces[1:]]
    for i, w in enumerate(words[:-1]):
        if w.text.endswith(":"):
            values.append(words[i + 1].box[0])
    return label_x, [x for x in values if label_x is None or x > label_x + 1]


def classify_text_block(text: str, n_rows: int) -> tuple:
    """-> (kind, meta) for a block of text lines."""
    t = text.strip()
    meta = {}
    cap = RE_CAPTION.match(t)
    if cap and word_count(t) <= 40 and not _modal_re().search(t):
        kind = "table" if cap.group("kind").lower().startswith("tab") else "figure"
        title = (cap.group("title") or "").strip()
        # "Table 3 lists the ..." is a sentence, not a caption: require a delimiter or a title-like start
        if cap.group("delim") or not title or title[:1].isupper() or RE_CONTINUED.search(title):
            meta.update({"caption_kind": kind, "caption_label": f"{cap.group('kind').upper().rstrip('.')} {cap.group('num')}",
                         "continued": bool(RE_CONTINUED.search(title))})
            return "caption", meta
    if n_rows <= 2 and _is_heading_text(t):
        m = RE_PARA_NUM.match(t)
        if m:
            meta["number"] = clean_para_number(m.group("num"))
            meta["title"] = m.group("rest").strip().rstrip(".")
            meta["level"] = m.group("num").count(".") + 1
        else:
            meta["level"] = 1
            meta["title"] = t.rstrip(".")
        return "heading", meta
    m = RE_LIST_MARKER.match(t) or RE_NUMBERED_ITEM.match(t)
    if m:
        meta["marker"] = m.group("marker")
        return "list_item", meta
    return "paragraph", meta


@dataclass
class _PageStats:
    line_h: float
    line_gap: float
    para_gap: float


def _page_stats(rows: list) -> _PageStats:
    heights = [max(l.box[3] for l in r) - min(l.box[1] for l in r) for r in rows] or [30.0]
    line_h = float(statistics.median(heights))
    gaps = []
    for a, b in zip(rows, rows[1:]):
        ab, bb = box_union(l.box for l in a), box_union(l.box for l in b)
        overlap_x = min(ab[2], bb[2]) - max(ab[0], bb[0])
        gap = bb[1] - ab[3]
        if overlap_x > 0 and 0 <= gap < 2 * line_h:
            gaps.append(gap)
    line_gap = float(np.percentile(gaps, 30)) if gaps else 0.35 * line_h
    if len(gaps) < 3:
        # two gaps cannot tell a double-spaced page from a short one whose lines are all separated by blank lines
        # (the top of a page that carries a list on): single spacing is taken, so a blank line still ends a block
        line_gap = min(line_gap, 0.35 * line_h)
    para_gap = float(np.clip(line_gap + 0.5 * line_h, 0.6 * line_h, 1.5 * line_h))
    return _PageStats(line_h, line_gap, para_gap)


def analyze_page(page_no: int, prep: PreparedPage, lines: list, settings: Settings, ocr=None) -> list:
    """Turn the OCR lines of one page into ordered Blocks."""
    dpi = prep.dpi
    blocks = []
    page_h, page_w = prep.ocr_image.shape[:2]
    # tiny OCR 'lines' in the outer side margins (specks, binder holes, revision bars) are noise
    lines = [ln for ln in lines if not (len(ln.text.strip()) <= 2 and not ln.text.strip().isdigit()
                                        and (ln.box[2] < 0.09 * page_w or ln.box[0] > 0.91 * page_w))]
    free = list(lines)

    # tables first: their words are removed from the running text
    table_captions = []       # (table box, caption lines taken out of that table's grid)
    if settings.detect_tables and lines:
        used = set()
        for table, tbox, inside, caption_lines in detect_tables(prep.ocr_image, lines, dpi, settings, ocr=ocr):
            blocks.append(Block("table", page_no, tbox, lines=[lines[i] for i in sorted(inside)], table=table))
            used |= inside
            if caption_lines:
                table_captions.append((tbox, {id(lines[i]) for i in caption_lines}))
        free = [ln for i, ln in enumerate(lines) if i not in used]

    # figures: graphics that are not text; labels inside a figure belong to the figure
    for fbox in detect_figures(prep, free, [b.box for b in blocks if b.kind == "table"]):
        inside = [ln for ln in free if box_center_inside(ln.box, fbox)]
        if _figure_is_really_text(fbox, inside):
            continue
        framed_caption = " ".join(ln.text for ln in sorted(inside, key=lambda l: (l.box[1], l.box[0]))).strip()
        table_below_or_above = next((tb for tb in blocks if tb.kind == "table"
                                     and abs(fbox[0] - tb.box[0]) < 0.1 * dpi and abs(fbox[2] - tb.box[2]) < 0.1 * dpi
                                     and (abs(tb.box[1] - fbox[3]) < 0.15 * dpi or abs(fbox[1] - tb.box[3]) < 0.15 * dpi)), None)
        if (inside and table_below_or_above is not None
                and _frame_is_caption_cell(fbox, inside, prep.ocr_image, dpi, framed_caption)):
            # the frame around a caption printed as the table's first (or last) row: keep it as text
            table_captions.append((table_below_or_above.box, {id(ln) for ln in inside}))
            continue
        blocks.append(Block("figure", page_no, fbox, lines=inside))
        inside_ids = {id(ln) for ln in inside}
        free = [ln for ln in free if id(ln) not in inside_ids]

    # reading order for text lines + tables + figures together
    items = [(ln.box, ln) for ln in free] + [(b.box, b) for b in blocks]
    line_h = statistics.median([ln.height for ln in free]) if free else 0.12 * dpi
    leaves = xy_cut(items, min_hgap=0.45 * line_h, min_vgap=max(1.2 * line_h, 0.12 * dpi))

    ordered_rows = []   # (leaf_index, row) where row is list[Line]  OR  (leaf_index, Block)
    for li, leaf in enumerate(leaves):
        leaf_lines = [p for _, p in leaf if isinstance(p, Line)]
        leaf_blocks = [p for _, p in leaf if isinstance(p, Block)]
        entries = [(r[0].box[1], "row", r) for r in _rows_of(leaf_lines)] + [(b.box[1], "block", b) for b in leaf_blocks]
        for _, what, obj in sorted(entries, key=lambda e: e[0]):
            ordered_rows.append((li, what, obj))

    stats = _page_stats([obj for _, what, obj in ordered_rows if what == "row"])
    row_boxes = [box_union(ln.box for ln in obj) for _, what, obj in ordered_rows if what == "row"]
    text_x0 = min((b[0] for b in row_boxes), default=0)
    text_w = max(1.0, max((b[2] for b in row_boxes), default=page_w) - text_x0)
    narrow_rows = [(b, None) for b in row_boxes if b[2] - b[0] < 0.6 * text_w]
    has_columns = len(narrow_rows) >= 4 and bool(_column_cut(narrow_rows, text_w, max(1.2 * stats.line_h, 0.12 * dpi))[0])

    def column_right_edge(x0: float, x1: float, in_columns: bool) -> float:
        """Right edge of the text column holding x0..x1: taken from the rows of this page that overlap it
        horizontally, so each column of a two-column page has its own edge. A line that stops well short
        of it, although the next line's first word would have fitted, was ended on purpose (a heading, a
        list item, the end of a paragraph) - it is not a wrapped sentence."""
        overlapping = [b for b in row_boxes if min(b[2], x1) - max(b[0], x0) > 0.5 * max(1.0, min(b[2] - b[0], x1 - x0))]
        if in_columns:                   # a line across both columns (page footer, page-wide title) is not this column
            overlapping = [b for b in overlapping if b[2] - b[0] < 0.6 * text_w] or overlapping
        same = [b[2] for b in overlapping]
        return float(np.percentile(same, 90)) if len(same) >= 3 else float(max(same, default=x1))

    out = []
    current = []          # rows of the block being built
    current_leaf = None
    item_by_position = False   # the block being built starts a list item found by its indent, not its marker
    current_right = None  # right edge of the column of the block being built

    def flush():
        nonlocal current, current_right, item_by_position
        current_right = None
        if not current:
            return
        all_lines = [ln for row in current for ln in row]
        text = " ".join(" ".join(ln.text for ln in row) for row in current)
        kind, meta = classify_text_block(text, len(current))
        if kind == "paragraph" and _starts_list_item(text, current[0]) and not RE_PARA_NUM.match(text):
            shape = RE_MARKER_SHAPE.match(text)
            kind, meta = "list_item", {"marker": shape.group("marker") if shape else ""}
        if item_by_position and kind in ("paragraph", "heading"):     # "C. MIL-STD-810H" is all capitals
            # "C.  Shock: ..." level with "b.  Vibration: ..." - the marker was misread, the item is not
            shape = RE_MARKER_SHAPE.match(text)
            kind, meta = "list_item", {"marker": shape.group("marker") if shape else "", "marker_by_position": True}
        item_by_position = False
        if kind == "list_item" and len(current) > 1:        # (after every way of becoming a list item)
            label_rows = _wrapped_label_rows(current, stats.line_h)
            if label_rows:
                meta["value_x"], meta["label_rows"] = label_rows       # see _block_rows
        out.append(Block(kind, page_no, box_union(ln.box for ln in all_lines), lines=all_lines, text=text, meta=meta))
        current = []

    for leaf_idx, what, obj in ordered_rows:
        if what == "block":
            flush()
            out.append(obj)
            current_leaf = None
            continue
        row = obj
        row_box = box_union(ln.box for ln in row)
        row_text = " ".join(ln.text for ln in row)
        if current:
            prev = current[-1]
            prev_box = box_union(ln.box for ln in prev)
            prev_text = " ".join(ln.text for ln in prev)
            block_box = box_union(ln.box for r in current for ln in r)
            first_box = box_union(ln.box for ln in current[0])
            first_text = " ".join(ln.text for ln in current[0])
            gap = row_box[1] - prev_box[3]
            sentence_ended = _ends_sentence(prev_text)
            if current_right is None:
                # on a page with columns, a block whose lines are all narrow sits in one of the columns
                in_columns = has_columns and max(block_box[2] - block_box[0], row_box[2] - row_box[0]) < 0.6 * text_w
                current_right = column_right_edge(first_box[0], max(first_box[2], prev_box[2]), in_columns)
            first_word = row[0].words[0].box if row[0].words else None
            word_w = (first_word[2] - first_word[0]) if first_word else 0.5 * stats.line_h * len(row_text.split()[0])
            next_word_fits = current_right - prev_box[2] > word_w + 1.0 * stats.line_h
            # The previous line ran to the margin in mid-sentence, so the text simply wrapped. That holds for
            # any line spacing: a double-spaced paragraph on a page with single-spaced notes is not split.
            wrapped = not sentence_ended and not next_word_fits
            # ... but a larger gap before the bottom / after the top 8.5% of the page is the page footer / header
            furniture_edge = (row_box[1] > 0.915 * page_h >= prev_box[1]) or (prev_box[3] < 0.085 * page_h <= row_box[3])
            wrapped = wrapped and not (furniture_edge and gap > stats.para_gap)
            # A larger gap after a wrapped line is ordinary spacing only when the next line continues the
            # sentence: it starts in lower case / with a number, or the block already has this spacing. A last
            # line that lost its final period (faint print) must not swallow "3.2.5 Weight. ..." or "The ...".
            if len(current) >= 2:
                before_box = box_union(ln.box for ln in current[-2])
                same_spacing = gap <= 1.3 * (prev_box[1] - before_box[3]) + 0.1 * stats.line_h
            else:
                same_spacing = False
            # A line that ends in the middle of a phrase ("... twenty targets in track at") cannot end the
            # sentence there: a next line in lower case carries it on, also when the typist broke the line
            # early and the next word would have fitted
            mid_phrase = not sentence_ended and bool(_ENDS_MID_PHRASE.search(prev_text)) and row_text[:1].islower()
            continues = ((wrapped or mid_phrase) and gap < 3 * stats.line_h and not _starts_new_item(row_text)
                         and not RE_LIST_MARKER.match(row_text)
                         and (row_text[:1].islower() or row_text[:1].isdigit() or same_spacing))
            wider_gap = gap > stats.line_gap + 0.25 * stats.line_h and not wrapped
            # A new bullet / list marker at the same indent as the block's first line is a sibling item, even
            # when the previous item has no final punctuation. So is "[SRS-101] ..." or "SRS-103: ...".
            # "(CB12)" at the start of a line needs evidence that the line before really ended: it is often
            # a designator in the middle of a wrapped sentence.
            strong_id = bool(re.match(r"^\s*(?:\[(?:" + RE_REQ_ID + r")\]|(?:" + RE_REQ_ID + r")\s*:)\s*[A-Z\"(\[]", row_text))
            # Inside a list, a line that starts level with the item above begins the NEXT item even
            # when its own marker cannot be recognised - misread ("c." read as "C." or as "3."),
            # run together with the text, or dropped altogether. A marker is one or two characters,
            # which is exactly what OCR loses or mistakes, while a list's indentation is printed the
            # same for every item; so the indentation is given slightly more weight than the marker.
            # The line above must have finished its sentence, which is what keeps a cell still
            # wrapping from being taken for the next item.
            in_list = _starts_list_item(first_text, current[0])
            item_by_indent = (in_list and sentence_ended and not continues
                              and abs(row_box[0] - first_box[0]) <= 0.6 * stats.line_h
                              and (row_text[:1].isupper() or row_text[:1].isdigit() or _lowered_label(row_text)))
            # "C. MIL-STD-810H" level with "b. MIL-STD-704F": the next letter of the list, misread, starts
            # the next item even though items without final punctuation never "end a sentence"
            first_marker = (RE_LIST_MARKER.match(first_text) or RE_NUMBERED_ITEM.match(first_text)) if in_list else None
            row_shape = RE_MARKER_SHAPE.match(row_text)
            next_in_sequence = bool(first_marker and row_shape and _next_marker(first_marker.group("marker")).lower()
                                    == row_shape.group("marker").lower() and abs(row_box[0] - first_box[0]) <= 0.6 * stats.line_h)
            # "C.   Shock:   40 g ..." - a marker, then a label: the row of a list item, also when the scan read the
            # marker as a capital and the block above is not a list item (a value carried on from the page before)
            label_row = bool(row_shape and len(row_shape.group("marker")) <= 3 and _is_item_label(row_shape.group("rest")))
            row_is_item = bool(RE_LIST_MARKER.match(row_text) or strong_id or item_by_indent or next_in_sequence or label_row
                               or (_starts_with_delimited_id(row_text) and (sentence_ended or next_word_fits or wider_gap)))
            sibling_item = row_is_item and row_box[0] <= first_box[0] + 0.6 * stats.line_h
            # a line back left of a list item's first line ends it - also when the item's own marker was
            # lost ("Vibration:  20 Hz ..." then "C.  Shock: ..." back at the markers' column) - unless it
            # simply carries on a wrapped sentence ("A. General: The contractor shall provide all labor," /
            # "and equipment ..." back at the margin: a paragraph with its first line indented)
            outdented = in_list and row_box[0] < first_box[0] - 0.6 * stats.line_h and not continues
            x_overlap = min(row_box[2], block_box[2]) - max(row_box[0], block_box[0])
            narrow = max(1, min(row_box[2] - row_box[0], block_box[2] - block_box[0]))
            prev_h, row_h = prev_box[3] - prev_box[1], row_box[3] - row_box[1]
            # A list item laid out with wide tabs - "a.   Acceleration:   5 g in each ..." - carries on
            # under its label or its value. A line that starts in one of those columns belongs to the
            # item whatever it begins with: "2.5 G peak ...", "20 Hz to 2000 Hz" and "3 hours per axis"
            # are values, not new sections, and the line above a value need not reach the margin.
            # Only a new marker, a line back at the marker, or a jump in type size ends the item.
            # A line that starts with the paragraph number that should come NEXT - "3.5.1 Discrete ..."
            # under "3.5 INTERFACE CONDITIONS" - begins a new paragraph even when the scan turned the
            # line above into something that does not look finished (a speck read as a period).
            row_num, first_num = RE_PARA_NUM.match(row_text), RE_PARA_NUM.match(first_text)
            numbered_next = bool(row_num and first_num and "." in row_num.group("num") and not wrapped
                                 and _starts_new_item(row_text) and not _ENDS_MID_PHRASE.search(prev_text)
                                 and clean_para_number(row_num.group("num")) != clean_para_number(first_num.group("num"))
                                 and plausible_next_number(clean_para_number(first_num.group("num")),
                                                           clean_para_number(row_num.group("num"))))
            label_x, value_xs = _item_columns(current[0]) if in_list else (None, [])
            at_label = label_x is not None and abs(row_box[0] - label_x) <= 0.6 * stats.line_h
            at_value = any(abs(row_box[0] - x) <= 0.6 * stats.line_h for x in value_xs)
            # the item so far leads into a sub-list of its own: the lines below are its sub-items
            opens_sub_list = in_list and _opens_sub_list(" ".join(" ".join(ln.text for ln in r) for r in current))
            # "Shock:  40 g ..." in the label column after a finished item: the next item, its marker lost
            lost_marker = (at_label and sentence_ended and not opens_sub_list and _is_item_label(row_text)
                           and bool(RE_ITEM_LABEL.match(first_text.split(None, 1)[-1])))
            new_requirement = row_text[:1].isupper() and bool(_modal_re().search(row_text))
            hanging = (in_list and not _is_heading_text(first_text) and not lost_marker and not opens_sub_list
                       and (at_label or at_value)
                       and not RE_LIST_MARKER.match(row_text) and not RE_NUMBERED_ITEM.match(row_text)
                       and not _is_titled_number(row_text)          # "3.2.4.1 Mode switching. The ..." is a paragraph
                       and -0.3 * stats.line_h <= gap < 2.5 * stats.line_h
                       # after a blank line, a finished item is only carried on in its VALUE column - and so is an
                       # item without a full stop ("b. Track mode") when the line starts a requirement of its own
                       # ("The operator shall select ...")
                       and not (not at_value and ((gap > stats.para_gap and sentence_ended)
                                                  or ((gap > stats.para_gap or wider_gap) and new_requirement))))
            new_block = (
                gap < -0.3 * stats.line_h                                        # jumped back up (new column)
                or (leaf_idx != current_leaf and x_overlap < 0.3 * narrow)          # different column
                or (gap > stats.para_gap and not continues)                        # paragraph spacing
                or sibling_item or outdented                                       # next list item / end of list
                or (_is_heading_text(prev_text) and len(current) == 1 and not row_text[:1].islower()
                    and (prev_text.rstrip().endswith((".", ":")) or prev_text.isupper() or _starts_new_item(row_text)
                         or _is_heading_text(row_text) or wider_gap or next_word_fits))
                or (_introduces_list(prev_text) and row_box[0] > first_box[0] + 0.6 * stats.line_h      # list starts
                    and (prev_text.rstrip().endswith(":") or len(row) > 1 or RE_LIST_MARKER.match(row_text)
                         or RE_NUMBERED_ITEM.match(row_text) or RE_BARE_MARKER.match(row[0].text)))
                or ((_starts_new_item(row_text) or _is_heading_text(row_text)) and (sentence_ended or wider_gap))
                or (sentence_ended and prev_box[2] < block_box[2] - 0.25 * (block_box[2] - block_box[0])
                    and len(current) > 1)                                           # short last line of a paragraph
                or (max(prev_h, row_h) > 1.45 * max(1, min(prev_h, row_h)))         # font size change
                or lost_marker                                                      # next item, marker lost
                or numbered_next                                                    # 3.5 -> 3.5.1
                # a new sentence after a blank line under a finished item, outside its value column - or a
                # requirement of its own under an item without a full stop ("b. Track mode" / "The operator shall")
                or (in_list and wider_gap and not at_value
                    and ((sentence_ended and row_text[:1].isupper()) or new_requirement))
            )
            if new_block and hanging and max(prev_h, row_h) <= 1.45 * max(1, min(prev_h, row_h)):
                new_block = False
            if new_block:
                flush()
                # only a line that starts with something SHAPED like a marker ("C." "e." "3)"): a numbered
                # paragraph ("3.4 Mass. ...") or a sentence level with the list is not an item
                item_by_position = bool((sibling_item and (item_by_indent or next_in_sequence) and RE_MARKER_SHAPE.match(row_text)
                                         and not RE_LIST_MARKER.match(row_text) and not RE_NUMBERED_ITEM.match(row_text))
                                        or lost_marker)
        current.append(row)
        current_leaf = leaf_idx
    flush()

    # "1. SEARCH MODE." / "1. Note text ..." look like headings on their own. They are list items when
    # they follow a line ending in ":" ("NOTES:", "... the following modes:"), when they sit level with a
    # list item, or when several of them follow each other with consecutive numbers (1., 2., 3.).
    def numbered_heading(blk):
        return RE_NUMBERED_ITEM.match(blk.text or "") if blk.kind == "heading" else None

    k = 0
    while k < len(out):
        if not numbered_heading(out[k]):
            k += 1
            continue
        run = [k]
        while (run[-1] + 1 < len(out) and numbered_heading(out[run[-1] + 1]) and out[run[-1] + 1].page == out[k].page
               and abs(out[run[-1] + 1].box[0] - out[run[-1]].box[0]) <= 0.6 * stats.line_h):
            run.append(run[-1] + 1)
        numbers = [int(numbered_heading(out[x]).group("marker").rstrip(".")) for x in run]
        before = out[k - 1] if k > 0 else None
        after_stem = (before is not None and before.page == out[k].page and _introduces_list(before.text)
                      and (before.kind == "heading" or word_count(before.text) <= 3             # "NOTES:"
                           or out[k].box[0] > before.box[0] + 0.3 * stats.line_h))              # indented under a stem
        # A numbered line sitting level with the list item just above it is that list's next item,
        # not a section heading: "a. ... b. ... 3. ..." is a list whose third marker was misread.
        # "3. REQUIREMENTS" after a list is still a heading, which is what the all-capitals test is for.
        after_item = (before is not None and before.page == out[k].page and before.kind == "list_item"
                      and abs(out[k].box[0] - before.box[0]) <= 0.6 * stats.line_h
                      and not RE_NUMBERED_ITEM.match(out[k].text).group("rest").isupper())
        # ... also when the item above ended in a paragraph of its own, in its text column ("2. Track mode." /
        # "     It shall hold 64 tracks." / "3. Ground mapping mode."): its number is the one just before
        # (or two before - a list of tests may count 1, 3, 5)
        for back in range(k - 1, -1, -1):
            blk = out[back]
            if blk.page != out[k].page or after_item:
                break
            if blk.kind == "paragraph" and blk.box[0] > out[k].box[0] + 0.6 * stats.line_h:
                continue                                # more of the item above
            prev = RE_NUMBERED_ITEM.match(blk.text or "") if blk.kind == "list_item" else None
            after_item = bool(prev and abs(blk.box[0] - out[k].box[0]) <= 0.6 * stats.line_h and back < k - 1
                              and 0 < numbers[0] - int(prev.group("marker").rstrip(".")) <= 2
                              and not RE_NUMBERED_ITEM.match(out[k].text).group("rest").isupper())
            break
        # "1. Note text ...", "2. Note text ..." - but not "3. REQUIREMENTS" / "4. VERIFICATION" (empty section)
        consecutive = (len(run) >= 2 and all(y == x + 1 for x, y in zip(numbers, numbers[1:]))
                       and not all(RE_NUMBERED_ITEM.match(out[x].text).group("rest").isupper() for x in run))
        # "4. Verification" followed by "4.1 General ..." is a section heading with its first sub-paragraph,
        # even straight after a 1., 2., 3. list - the number that follows says so. Only the LAST line of
        # the run can be such a heading, only when the next paragraph is exactly its N.1, and only when
        # it reads as a title (no full stop, or all capitals): "3. Ground mapping mode." followed by
        # "3.2.5 Weight. ..." or "2. MIL-STD-704F" followed by "2.2 Order of precedence." are list items.
        after = out[run[-1] + 1] if run[-1] + 1 < len(out) else None
        next_num = RE_PARA_NUM.match(after.text or "") if after is not None else None
        last_rest = RE_NUMBERED_ITEM.match(out[run[-1]].text).group("rest").strip()
        opens_section = bool(next_num and re.fullmatch(rf"{numbers[-1]}\.0?1", clean_para_number(next_num.group("num")))
                             and word_count(last_rest) <= 8 and not _modal_re().search(last_rest))
        # right under a lead-in ("... the following modes:") consecutive numbers are its list, wherever they sit -
        # after a colon also when they are all capitals ("... shall apply:" / "1. MIL-STD-461G" / "2. MIL-STD-704F")
        lead_before = before is not None and before.page == out[k].page and _introduces_list(before.text)
        from_one = numbers[0] == 1 and all(y == x + 1 for x, y in zip(numbers, numbers[1:]))
        colon_list = lead_before and before.text.rstrip().endswith(":") and from_one
        if after_stem or after_item or (consecutive and lead_before) or colon_list:     # a list
            for x in (run[:-1] if opens_section else run):
                out[x].kind, out[x].meta = "list_item", {"marker": numbered_heading(out[x]).group("marker")}
        elif consecutive and not opens_section:        # "1. Note text ...", "2. ..." (but "5. Packaging." /
            for x in run:                              # "6. Notes." / "6.1 ..." are empty sections)
                out[x].kind, out[x].meta = "list_item", {"marker": numbered_heading(out[x]).group("marker")}
        k = run[-1] + 1

    # a caption that was printed inside a table's frame belongs to that table
    for tbox, caption_line_ids in table_captions:
        for b in out:
            if b.kind == "caption" and any(id(ln) in caption_line_ids for ln in b.lines):
                b.meta["caption_for_box"] = [int(v) for v in tbox]
    return out


# =============================================================================
# 7. DOCUMENT-LEVEL CLEAN-UP
#    Runs after every page is read: page headers/footers, words hyphenated across
#    lines, paragraphs and tables that continue on the next page, caption links.
# =============================================================================
def _similar(a: str, b: str, at_least: float = 0.0) -> float:
    """Similarity 0..1 of two strings. With `at_least`, clearly different strings return 0 early
    (the cheap upper bounds are checked first - the full comparison is slow on long text)."""
    import difflib
    m = difflib.SequenceMatcher(None, a, b, autojunk=False)
    if at_least and (m.real_quick_ratio() < at_least or m.quick_ratio() < at_least):
        return 0.0
    return m.ratio()


# Engineering words missing from the general English dictionary (so they are not flagged as misreads).
# Program-specific terms can be added to Settings.extra_dictionary_words.
ENGINEERING_WORDS = set("""
antitamper backplane backplanes baseband beamformer beamforming beamwidth bitrate bitrates bootloader bootloaders
bootup boresight boresighting checksum checksums circulators config configs connectorization connectorized crypto
cryptos cybersecurity databus datagram datagrams datalink datalinks deconfliction deserializer diplexer downlink
downlinks duplexer duplexers dwg dwgs failover failsafe geofence geofencing geolocation georeference georeferenced
geospatial gimbaled hardcopy hypervisor interoperate interoperates interoperating keyloader lidar lifecycle loopback
mainlobe metadata microamp microamps microservice microservices middleware milliamp milliamps milliwatt milliwatts
monopulse multimode multipath narrowband navaid offboard onboard overcurrent overtemperature overvoltage packetized
passthrough picosecond pinout pinouts pitot postflight powerdown powerup pulsewidth pushbuttons qty readback realtime req
reqs runtime scanline serializer sidelobe sidelobes singlemode softcopy subassemblies subcomponent subcomponents
subtier teleoperated teleoperation testset testsets timestamp timestamped timestamps undervoltage uninterruptible
uplink uplinks verif virtualization waypoint waypoints webserver wideband workaround writeback zeroization zeroize
zeroized
rms vrms grms arms rpm psig psid vdc vac dbm dbw ghz gpm cfm pps bps kbps mbps gbps msec usec nsec hrs typ kva kvar lbf
""".split())


class WordChecker:
    """Spots OCR misreads the OCR engine itself is confident about. Tested on the degraded test
    scans: most real misreads ('telémetry', 'fro', 'governmient', '16.AWG') scored above 0.95,
    so these independent checks are what makes the 'Needs Review' flag useful:
      * accent marks in English text (removed, and flagged)
      * a word that is not in an English dictionary (pyspellchecker, MIT) and occurs only once
      * a rare word one letter away from a common word of the same document ('fro' vs 'from')
      * punctuation inside a word ('Rev'A')
    Nothing is ever auto-corrected except the accent marks."""

    def __init__(self, counts: Counter, settings: "Settings", ids: Counter = None):
        self.counts = counts
        self.ids = ids or Counter()
        self.settings = settings
        self._memo = {}
        self._id_index = defaultdict(set)         # "MIL-STD-810?" -> {"MIL-STD-810H", "MIL-STD-810B"}
        for token in self.ids:
            for i in range(len(token)):
                self._id_index[token[:i] + "\0" + token[i + 1:]].add(token)
        self.spell = None
        if settings.spell_check_flags:
            try:
                from spellchecker import SpellChecker
                self.spell = SpellChecker(language="en", distance=1)
                self.spell.word_frequency.load_words(ENGINEERING_WORDS | settings.extra_dictionary_words)
            except Exception:
                self.spell = None

    def _id_twin(self, token: str) -> str:
        """Identifiers such as MIL-STD-810H / MIL-STD-810B or SPEC-1234-A / SPEC-I234-A that differ in one
        character. A different digit in the running number at the end (SW-001 / SW-002, ICD-1204A /
        ICD-1205A) is ordinary numbering and is not reported: sequential IDs would all be flagged."""
        if self.ids.get(token, 0) >= 2:
            return ""
        last_number = re.search(r"(\d+)[A-Za-z]{0,2}$", token)      # the running number (+ revision letter)
        for i in range(len(token)):
            for other in self._id_index.get(token[:i] + "\0" + token[i + 1:], ()):
                if other == token:
                    continue
                if (last_number and last_number.start(1) <= i < last_number.end(1)
                        and token[i].isdigit() and other[i].isdigit()):
                    continue                     # SW-001 / SW-002, ICD-1204A / ICD-1205A: ordinary numbering
                return other
        return ""

    def _document_twin(self, word: str) -> str:
        """An ordinary English word one edit away from an unknown word that THIS document also uses
        ("systen" printed beside "system"). That pairing is evidence the scan misread a normal word
        rather than that the document has a term of its own: a specification that writes "boresight"
        or "monopulse" does not also write "foresight" or anything one edit from them. Returns "" when
        there is no such word, which leaves genuine project vocabulary unflagged."""
        best, best_freq = "", 0.0
        for cand in self.spell.edit_distance_1(word):
            if len(cand) >= 4 and self.counts.get(cand, 0) and cand in self.spell:
                freq = self.spell.word_usage_frequency(cand)
                if freq > best_freq:
                    best, best_freq = cand, freq
        return best

    def clean(self, raw: str) -> str:
        return remove_ocr_specks(strip_accents(raw) if self.settings.english_text else raw)

    def issue(self, raw: str, clean: str) -> str:
        key = (raw, clean)
        if key not in self._memo:
            self._memo[key] = self._issue(raw, clean)
        return self._memo[key]

    def _issue(self, raw: str, clean: str) -> str:
        core = clean.strip(".,;:()[]{}\"'!?")
        if self.settings.english_text and strip_accents(raw) != raw:
            parts = [x.lower() for x in re.findall(r"[A-Za-z]{2,}", core)]
            if not parts or any(self.spell is not None and x not in self.spell and self.counts.get(x, 0) < 2 for x in parts):
                return f"accent marks removed from '{raw}'"
        if re.fullmatch(r"\d[\d.,]*['\u2019`]", core + ("'" if clean.endswith("'") else "")) or re.search(r"\d['\u2019`]$", clean):
            return "stray mark after a number (or a feet symbol)"
        for part in core.split():
            part = part.strip(".,;:()[]{}\"'!?")
            if re.search(r"\d", part) and re.search(r"[A-Za-z]", part) and len(part) >= 5:
                twin = self._id_twin(part)
                if twin:
                    return f"differs by one character from '{twin}' elsewhere in the document"
        if (re.search(r"(?<=[A-Za-z])[;'\"`](?=[A-Za-z0-9])|(?<=[a-z])[.:](?=[a-z])|(?<=\d)[;'\"`](?=[A-Za-z])", core)
                and not re.fullmatch(r"(?:[A-Za-z]\.){2,}[A-Za-z]?|[A-Za-z]+'(?:s|t|ll|re|ve|d|m)|[\w.-]+\.(?:" + "|".join(_FILE_EXTENSIONS) + r")", core)):
            return "punctuation inside a word"
        if re.search(r"[a-z]{3,}[A-Z]{3,}|[A-Za-z]{3,}[01][A-Za-z]{2,}", core) and not re.fullmatch(r"(?:[A-Z]{2,}\d+[A-Z]*|\w+\.\w+)", core):
            return "letters and digits mixed inside a word (overlapping stamp or misread)"
        if self.spell is None:
            return ""
        for part in re.split(r"[-/]", core):
            if not re.fullmatch(r"[A-Za-z][a-z]{2,}", part):
                continue                       # acronyms, IDs, numbers and short words are not checked
            w = part.lower()
            repeated = self.counts.get(w, 0) >= 2
            if w not in self.spell:
                # A word the document repeats is normally a project term, not a misread. But OCR
                # repeats ITS mistakes too - the same word misread the same way on every page - and
                # those are the errors most easily missed. So a repeated unknown word is still
                # reported when the document ALSO uses the ordinary word one edit away from it.
                twin = self._document_twin(w) if repeated else ""
                if repeated and not twin:
                    continue
                return (f"not in the English dictionary (the document also writes '{twin}')"
                        if twin else "not in the English dictionary")
            if repeated:
                continue                       # repeated in the document: a real term (or a consistent one)
            freq = self.spell.word_usage_frequency(w)
            if freq < 1e-6:
                for cand in self.spell.edit_distance_1(w):
                    if self.counts.get(cand, 0) >= 3 and self.spell.word_usage_frequency(cand) > max(1e-4, 1000 * freq):
                        return f"possible misread of '{cand}'"
        return ""


@dataclass
class TextContext:
    vocab: set
    counts: Counter
    checker: WordChecker


def build_text_context(blocks: list, settings: "Settings") -> TextContext:
    counts, ids = Counter(), Counter()
    for b in blocks:
        for ln in b.lines:
            for w in ln.words:
                plain = remove_ocr_specks(strip_accents(w.text))
                for tok in re.findall(r"[A-Za-z]+(?![A-Za-z-])", plain):
                    counts[tok.lower()] += 1
                for part in plain.split():
                    core = part.strip(".,;:()[]{}\"'!?")
                    if len(core) >= 5 and re.search(r"\d", core) and re.search(r"[A-Za-z]", core):
                        ids[core] += 1
    return TextContext(set(counts), counts, WordChecker(counts, settings, ids))


def document_text_context(doc: "DocResult", settings: "Settings") -> TextContext:
    """The word statistics of a document are computed once and re-used by extraction and exports."""
    ctx = getattr(doc, "_text_ctx", None)
    if ctx is None:
        ctx = build_text_context(doc.blocks, settings)
        doc._text_ctx = ctx
    return ctx


def token_attention(token: tuple, settings: "Settings") -> str:
    """'' (fine), 'check' or 'likely_wrong' for one (text, confidence, issue) token.
    Same rules for the Excel 'Needs Review' flag and the Word highlighting: numbers with low
    OCR confidence, words that failed a WordChecker check, and any word read very poorly."""
    text, conf = token[0], token[1]
    issue = token[2] if len(token) > 2 else ""
    has_digit = bool(re.search(r"\d", text))
    if (has_digit and conf < settings.very_low_confidence) or (conf < 0.35 and re.search(r"[A-Za-z]{2}", text)):
        return "likely_wrong"
    if issue or (has_digit and conf < settings.review_confidence):
        return "check"
    return ""


def word_tokens(words: Iterable, ctx: TextContext) -> list:
    """[(clean text, confidence, issue)] for a sequence of OCR words."""
    out = []
    for w in words:
        clean = ctx.checker.clean(w.text)
        if not clean:
            continue
        if out and re.fullmatch(r"[.,;:]", clean):
            prev = out[-1]
            out[-1] = (prev[0] + clean, min(prev[1], w.conf), prev[2])   # "mode ." -> "mode."
            continue
        issue = ctx.checker.issue(w.text, clean)
        out.extend((part, w.conf, issue) for part in clean.split(" "))
    return out


def block_token_rows(block: Block, ctx: TextContext) -> list:
    """Tokens [(text, confidence, issue)] of a text block, one list per visual row (same order as
    _block_rows). A word split by a line-end hyphen is re-joined when the joined word appears
    elsewhere in the document ("require-" + "ment" -> "requirement"), otherwise the hyphen is kept
    ("track-while-" + "scan"). Paper specks read as punctuation are removed; suspicious words get an
    'issue' text."""
    raw_rows = []                                       # [[(raw text, conf)]] per row
    last = None                                         # (row index, token index) of the last token
    for row in _block_rows(block):
        current = []
        raw_rows.append(current)
        for k, w in enumerate(w for ln in row for w in ln.words):
            if k == 0 and last is not None:
                prev_text, prev_conf = raw_rows[last[0]][last[1]]
                if prev_text.endswith("-") and w.text[:1].islower():
                    head = re.search(r"([A-Za-z]+)-$", prev_text)
                    tail = re.match(r"[A-Za-z]+", w.text)
                    joined = prev_text[:-1] + w.text if head and tail and (head.group(1) + tail.group(0)).lower() in ctx.vocab \
                        else prev_text + w.text
                    raw_rows[last[0]][last[1]] = (joined, min(prev_conf, w.conf))
                    continue
            current.append((w.text, w.conf))
            last = (len(raw_rows) - 1, len(current) - 1)
    token_rows = [word_tokens((Word(t, c, (0, 0, 0, 0)) for t, c in r), ctx) for r in raw_rows]
    # the first word of each further line is marked (a 4th field), so a check can tell a full stop at
    # the end of a line from a speck in the middle of one
    return [[(r[0][0], r[0][1], r[0][2], True)] + r[1:] if k and r else r for k, r in enumerate(token_rows)]


def block_tokens(block: Block, ctx: TextContext) -> list:
    """All tokens of a text block in reading order (see block_token_rows)."""
    return [tok for row in block_token_rows(block, ctx) for tok in row]


def _text_rows(block: Block) -> list:
    """(x0, text) of each visual row of a block."""
    if not block.lines:
        return [(block.box[0], block.text)]
    return [(min(ln.box[0] for ln in row), " ".join(ln.text for ln in row).strip()) for row in _block_rows(block)]


def _on_new_line(tokens: list) -> list:
    """The tokens of a block that is joined on to the text above it: its first word starts a line
    (see block_token_rows), so a full stop just before it ends that line, not a sentence part-way."""
    return [(tokens[0][0], tokens[0][1], tokens[0][2] if len(tokens[0]) > 2 else "", True)] + list(tokens[1:]) if tokens else []


def _block_rows(block: Block) -> list:
    """Rows of a block in reading order. A paragraph joined across a page break keeps its pages
    in order (sorting all lines by height alone would put page 2 before page 1).
    A list item whose label wraps beside a value on tab stops ("Storage      -55 C to +85 C for" /
    "temperature:  up to 5 years.") is read label first, then value: its first row holds the marker,
    every line of the label and the first value line; the rest of the value follows."""
    bounds = list(block.meta.get("segments") or [0]) + [len(block.lines)]
    segments = [_rows_of(block.lines[a:z]) for a, z in zip(bounds, bounds[1:])]
    value_x, k = block.meta.get("value_x"), int(block.meta.get("label_rows") or 0)
    if value_x and k >= 2 and segments and len(segments[0]) >= k:
        # (only on the page where the item starts - more of its value may follow on the next page)
        head = segments[0][:k]
        left = [ln for row in head for ln in sorted(row, key=lambda l: l.box[0]) if ln.box[0] < value_x]
        right = [[ln for ln in row if ln.box[0] >= value_x] for row in head]
        segments[0] = [left + right[0]] + [r for r in right[1:] if r] + segments[0][k:]
    return [row for seg in segments for row in seg]


def tokens_text(tokens: list) -> tuple:
    """-> (text, [(start, end) char span of each token])"""
    spans, parts, pos = [], [], 0
    for t, *_ in tokens:
        spans.append((pos, pos + len(t)))
        parts.append(t)
        pos += len(t) + 1
    return " ".join(parts), spans


def finalize_document(doc: DocResult, settings: Settings) -> None:
    for i, b in enumerate(doc.blocks, 1):
        b.uid = i
        b.pages = b.pages or [b.page]
    ctx = build_text_context(doc.blocks, settings)
    doc._text_ctx = ctx                         # re-used by extraction and the exports
    for b in doc.blocks:
        if b.kind not in ("table", "figure") and b.lines:
            b.text = tokens_text(block_tokens(b, ctx))[0]
        if b.kind == "table" and b.table is not None:
            b.table.grid = [[ctx.checker.clean(c) if c else c for c in row] for row in b.table.grid]
    _mark_page_furniture(doc)
    _merge_across_pages(doc)
    if settings.merge_continued_tables:
        _merge_continued_tables(doc)
    for i, b in enumerate(doc.blocks, 1):      # renumber BEFORE linking captions (merges removed blocks)
        b.uid = i
    _link_captions(doc)
    # page statistics for the "Page Quality" sheet
    for p in doc.pages:
        page_lines = [ln for b in doc.blocks if b.page == p.number for ln in b.lines]
        p.n_lines = len(page_lines)
        p.mean_conf = round(float(statistics.fmean(ln.conf for ln in page_lines)), 4) if page_lines else 1.0
        p.low_conf_lines = sum(1 for ln in page_lines if ln.min_word_conf < settings.review_confidence)


def _mark_page_furniture(doc: DocResult) -> None:
    """Running headers/footers repeat on many pages near the top/bottom edge; page numbers too."""
    heights = {p.number: p.height for p in doc.pages}
    n_pages = len(doc.pages)
    candidates = []
    for b in doc.blocks:
        h = heights.get(b.page)
        if not h or b.kind in ("table", "figure") or not b.text:
            continue
        if b.box[3] < 0.085 * h:
            candidates.append((b, "page_header"))
        elif b.box[1] > 0.915 * h:
            candidates.append((b, "page_footer"))

    def norm(t):
        return re.sub(r"\d+", "#", re.sub(r"\s+", "", t.lower()))

    def notice_key(t):
        """Keyword text compared word for word, without a page number printed on the same line."""
        t = re.sub(r"^\s*(?:page\s*)?\d{1,4}(?:\s*of\s*\d{1,4})?\s+|\s+(?:page\s*)?\d{1,4}(?:\s*of\s*\d{1,4})?\s*$", " ", t, flags=re.I)
        return " ".join(t.lower().split())

    pages_by_key = defaultdict(set)             # (kind, normalized text) -> pages it appears on
    pages_by_notice = defaultdict(set)          # (kind, keyword text without page number) -> pages
    for b, kind in candidates:
        pages_by_key[(kind, norm(b.text))].add(b.page)
        pages_by_notice[(kind, notice_key(b.text))].add(b.page)
    notice_keys = list(pages_by_notice)
    keys = list(pages_by_key)
    threshold = max(1, 0.3 * (n_pages - 1))
    for b, kind in candidates:
        if RE_LIST_MARKER.match(b.text):
            continue                             # a list item is never page furniture
        if _modal_re().search(b.text):
            # Text with a requirement keyword is furniture only when the same words repeat on other pages
            # (a handling notice in every footer; OCR may garble a letter or two of a long notice). The
            # extractor still reports it once.
            key = notice_key(b.text)
            other = set(pages_by_notice[(kind, key)]) - {b.page}
            if len(other) < threshold and word_count(key) >= 6 and len(notice_keys) <= 400:
                for k2 in notice_keys:
                    if (k2[0] == kind and k2[1] != key and re.findall(r"\d+", k2[1]) == re.findall(r"\d+", key)
                            and _similar(k2[1], key, at_least=0.9) >= 0.9):
                        other |= pages_by_notice[k2] - {b.page}
            if n_pages >= 2 and len(other) >= threshold:
                b.kind = kind
                b.meta["repeated_keyword_text"] = True
                b.meta["notice_key"] = key
            continue
        key = norm(b.text)
        other_pages = set(pages_by_key[(kind, key)]) - {b.page}
        if len(other_pages) < threshold and len(key) <= 150 and len(keys) <= 400:
            # OCR noise: allow near-identical text (running headers/footers are short lines)
            for k2 in keys:
                if k2[0] == kind and k2[1] != key and _similar(k2[1], key, at_least=0.8) >= 0.8:
                    other_pages |= pages_by_key[k2] - {b.page}
        if (n_pages >= 2 and len(other_pages) >= threshold) or RE_PAGE_NUMBER.match(b.text):
            b.kind = kind


_ENDS_MID_SENTENCE = re.compile(r"(?:-|,|\b(?:and|or|the|of|to|a|an|in|for|with|by|shall|must|will|should|be|not|than|"
                                r"from|at|on|as|per|that|which))$", re.I)


def _merge_across_pages(doc: DocResult) -> None:
    """Join a paragraph that is cut by a page break ('...shall not' | next page 'exceed 30 minutes')."""
    body = [b for b in doc.blocks if b.kind not in ("page_header", "page_footer")]
    removed = set()
    for a, nxt in zip(body, body[1:]):
        if id(a) in removed or nxt.page != max(a.pages) + 1:
            continue
        if a.kind not in ("paragraph", "list_item") or nxt.kind != "paragraph" or not a.text or not nxt.text:
            continue
        para = RE_PARA_NUM.match(nxt.text)
        if (_ends_sentence(a.text) or _starts_with_delimited_id(nxt.text)
                or (para and ("." in para.group("num") or para.group("dot")) and para.group("rest")[:1].isupper()
                    and not _starts_with_unit(para.group("rest")))):
            continue                                     # the next page starts a new paragraph / requirement
        if not (nxt.text[:1].islower() or nxt.text[:1].isdigit() or _ENDS_MID_SENTENCE.search(a.text)):
            continue
        a.meta["segments"] = list(a.meta.get("segments") or [0]) + [len(a.lines)]
        a.lines = a.lines + nxt.lines
        a.pages = sorted(set(a.pages) | set(nxt.pages))
        a.meta["continued_across_pages"] = True
        joiner = "" if a.text.endswith("-") and nxt.text[:1].islower() else " "
        a.text = a.text + joiner + nxt.text
        removed.add(id(nxt))
    doc.blocks = [b for b in doc.blocks if id(b) not in removed]


def _merge_continued_tables(doc: DocResult) -> None:
    """A table that runs onto the next page usually repeats its header row (MIL-STD-961
    also repeats the title with '- Continued'). Merge those pieces into one table."""
    body = [b for b in doc.blocks if b.kind not in ("page_header", "page_footer")]
    removed = set()
    i = 0
    while i < len(body) - 1:
        a = body[i]
        if a.kind != "table" or id(a) in removed:
            i += 1
            continue
        j = i + 1
        continued_caption = None
        if j < len(body) and body[j].kind == "caption" and body[j].page == max(a.pages) + 1:
            cap = body[j]
            if cap.meta.get("continued") or (a.meta.get("caption_label") and cap.meta.get("caption_label") == a.meta.get("caption_label")):
                continued_caption = cap
                j += 1
        if j >= len(body):
            break
        b = body[j]
        if b.kind == "table" and b.page == max(a.pages) + 1 and b.table.n_cols == a.table.n_cols:
            same_header = (a.table.header_rows and b.table.header_rows and
                           _similar(" ".join(a.table.grid[0]).lower(), " ".join(b.table.grid[0]).lower(), at_least=0.85) >= 0.85)
            if same_header or continued_caption is not None:
                skip = b.table.header_rows if same_header else 0
                if a.table.cell_words and b.table.cell_words:
                    a.table.cell_words += b.table.cell_words[skip:]
                else:
                    a.table.cell_words = []
                a.table.grid += b.table.grid[skip:]
                a.table.conf += b.table.conf[skip:]
                offset = len(a.table.grid) - len(b.table.grid[skip:])
                a.table.merges += [(r0 - skip + offset, c0, r1 - skip + offset, c1) for r0, c0, r1, c1 in b.table.merges if r0 >= skip]
                a.table.cell_boxes = []
                a.table.continued_pages.append(b.page)
                a.lines = a.lines + b.lines
                a.pages = sorted(set(a.pages) | {b.page})
                removed.add(id(b))
                hidden = set()
                if continued_caption is not None:
                    continued_caption.meta["duplicate_of_table"] = True
                    hidden.add(id(continued_caption))   # keep it in the document, but out of the way here
                # stay on this table: it might continue on a third page
                body = [x for x in body if id(x) not in removed and id(x) not in hidden]
                continue
        i += 1
    doc.blocks = [b for b in doc.blocks if id(b) not in removed]


def _min_cost_assignment(cost: list) -> list:
    """Hungarian algorithm: for a cost matrix with rows <= columns, the column chosen for each row so
    that the total cost is as small as possible (each column used at most once)."""
    n, m = len(cost), len(cost[0]) if cost else 0
    inf = float("inf")
    u, v, p, way = [0.0] * (n + 1), [0.0] * (m + 1), [0] * (m + 1), [0] * (m + 1)
    for i in range(1, n + 1):
        p[0], j0 = i, 0
        minv, used = [inf] * (m + 1), [False] * (m + 1)
        while True:
            used[j0] = True
            i0, delta, j1 = p[j0], inf, 0
            for j in range(1, m + 1):
                if not used[j]:
                    cur = cost[i0 - 1][j - 1] - u[i0] - v[j]
                    if cur < minv[j]:
                        minv[j], way[j] = cur, j0
                    if minv[j] < delta:
                        delta, j1 = minv[j], j
            for j in range(m + 1):
                if used[j]:
                    u[p[j]] += delta
                    v[j] -= delta
                else:
                    minv[j] -= delta
            j0 = j1
            if p[j0] == 0:
                break
        while j0:
            j1 = way[j0]
            p[j0] = p[j1]
            j0 = j1
    result = [-1] * n
    for j in range(1, m + 1):
        if p[j]:
            result[p[j] - 1] = j - 1
    return result


def _link_captions(doc: DocResult) -> None:
    """Attach each caption to its table or figure.
    1. Known links first: a caption row taken out of a table's frame, or Docling's caption order.
    2. The side is learned: MIL-STD-961 puts table titles above tables and figure titles below figures,
       but a document is consistent with itself, so captions that have only one possible object
       decide the side used for captions sitting between two objects (the norm when unclear).
    3. On each page all remaining captions are matched to objects at once (smallest total distance),
       so one caption cannot take the table that clearly belongs to the next caption."""
    dpi_of = {p.number: p.dpi for p in doc.pages}
    taken = set()

    def link(c, target):
        c.meta["caption_of"] = target.uid
        target.meta["caption"] = c.text
        target.meta["caption_label"] = c.meta.get("caption_label", "")
        taken.add(target.uid)

    objects = defaultdict(list)                          # (page, kind) -> tables / figures on that page
    by_box = {}
    for t in doc.blocks:
        if t.kind in ("table", "figure"):
            objects[(t.page, t.kind)].append(t)
            by_box[(t.page, tuple(int(x) for x in t.box))] = t
    votes = {"table": Counter(), "figure": Counter()}
    for idx, c in enumerate(doc.blocks):
        if c.kind != "caption":
            continue
        attach = c.meta.get("attach")
        target = (doc.blocks[idx + 1] if attach == "next" and idx + 1 < len(doc.blocks) else
                  doc.blocks[idx - 1] if attach == "prev" and idx > 0 else
                  by_box.get((c.page, tuple(c.meta["caption_for_box"]))) if c.meta.get("caption_for_box") else None)
        if (target is not None and target.kind in ("table", "figure") and target.uid not in taken
                and not c.meta.get("duplicate_of_table")):    # a "(continued)" caption whose table piece was merged
            link(c, target)
            if attach in ("next", "prev") and target.kind in votes:
                votes[target.kind]["below" if attach == "next" else "above"] += 1

    def candidates(c):
        """-> {'above': (gap, block) or None, 'below': ...}: the nearest free object on each side."""
        max_gap = 1.2 * dpi_of.get(c.page, 200)          # within ~1.2 inch
        found = {"above": None, "below": None}
        for t in objects[(c.page, c.meta["caption_kind"])]:
            if t.uid in taken:
                continue
            side = "below" if t.box[1] >= c.box[1] else "above"
            gap = t.box[1] - c.box[3] if side == "below" else c.box[1] - t.box[3]
            if -0.2 * max_gap <= gap <= max_gap and (found[side] is None or gap < found[side][0]):
                found[side] = (gap, t)
        return found

    pending = [b for b in doc.blocks if b.kind == "caption" and "caption_of" not in b.meta
               and not b.meta.get("duplicate_of_table") and b.meta.get("caption_kind") in ("table", "figure")]
    for c in pending:
        found = candidates(c)
        if (found["above"] is None) != (found["below"] is None):
            votes[c.meta["caption_kind"]]["above" if found["above"] else "below"] += 1
    preferred, penalty = {}, {}
    for kind, norm in (("table", "below"), ("figure", "above")):      # side where the OBJECT sits
        total = sum(votes[kind].values())
        if total and votes[kind].most_common(1)[0][1] >= 0.75 * total:
            preferred[kind], penalty[kind] = votes[kind].most_common(1)[0][0], 2.0    # the document's own habit
        else:
            preferred[kind], penalty[kind] = norm, 1.25

    groups = defaultdict(list)
    for c in pending:
        groups[(c.page, c.meta["caption_kind"])].append(c)
    for (page, kind), caps in groups.items():
        objs = [t for t in objects[(page, kind)] if t.uid not in taken]
        if not objs:
            continue
        max_gap = 1.2 * dpi_of.get(page, 200)
        unmatched = 3.0 * max_gap                        # leaving a caption without its object is costly
        cost = []
        for c in caps:
            row = []
            for t in objs:
                side = "below" if t.box[1] >= c.box[1] else "above"
                gap = t.box[1] - c.box[3] if side == "below" else c.box[1] - t.box[3]
                if not -0.2 * max_gap <= gap <= max_gap:
                    row.append(1e9)                      # too far away: never
                else:
                    row.append(gap * (1.0 if side == preferred[kind] else penalty[kind]) if gap > 0 else gap)
            row += [unmatched] * len(caps)               # one "no object" choice per caption
            cost.append(row)
        for i, j in enumerate(_min_cost_assignment(cost)):
            if 0 <= j < len(objs) and cost[i][j] < 1e9:
                link(caps[i], objs[j])


# =============================================================================
# 8. OPTIONAL ENGINE: DOCLING  (IBM, MIT license)
#    Adds a trained layout model (Heron) and TableFormer table-structure model on
#    top of the same PP-OCRv6 OCR. Needs:  python -m pip install "docling[rapidocr]"
#    and a one-time model download from Hugging Face (or an offline copy, see README).
# =============================================================================
class EngineUnavailable(RuntimeError):
    pass


class DoclingEngine:
    def __init__(self, settings: Settings):
        models_dir = settings.docling_models_dir or os.environ.get("DOCLING_ARTIFACTS_PATH")
        if models_dir:
            # an offline model folder was given: never try to reach Hugging Face (read before docling loads)
            os.environ.setdefault("HF_HUB_OFFLINE", "1")
        try:
            from docling.datamodel.accelerator_options import AcceleratorDevice, AcceleratorOptions
            from docling.datamodel.base_models import InputFormat
            from docling.datamodel.pipeline_options import (OcrMode, PdfPipelineOptions, RapidOcrOptions,
                                                            TableFormerMode, TableStructureOptions)
            from docling.document_converter import DocumentConverter, PdfFormatOption
        except ImportError as exc:
            raise EngineUnavailable('Docling is not installed. Install it with:  python -m pip install "docling[rapidocr]"'
                                    "  - or run with the standard engine.") from exc
        self.settings = settings
        ocr_kwargs = {"mode": OcrMode.FULL_PAGE, "backend": "onnxruntime", "text_score": settings.min_text_confidence,
                      "rapidocr_params": {"Global.log_level": "error"}}
        try:  # reuse the PP-OCRv6 models that ship inside the rapidocr package (no extra download)
            import rapidocr
            models = Path(rapidocr.__file__).parent / "models"
            det, rec = models / "PP-OCRv6_det_small.onnx", models / "PP-OCRv6_rec_small.onnx"
            cls = models / "ch_ppocr_mobile_v2.0_cls_mobile.onnx"
            if det.exists() and rec.exists():
                ocr_kwargs.update({"det_model_path": str(det), "rec_model_path": str(rec)})
                if cls.exists():
                    ocr_kwargs["cls_model_path"] = str(cls)
        except ImportError:
            pass
        device = {"cpu": AcceleratorDevice.CPU, "cuda": AcceleratorDevice.CUDA, "mps": AcceleratorDevice.MPS}.get(
            settings.docling_device.lower(), AcceleratorDevice.AUTO)
        options = PdfPipelineOptions(
            do_ocr=True, do_table_structure=True, ocr_options=RapidOcrOptions(**ocr_kwargs),
            table_structure_options=TableStructureOptions(do_cell_matching=False, mode=TableFormerMode.ACCURATE),
            accelerator_options=AcceleratorOptions(device=device))
        if models_dir:
            options.artifacts_path = str(models_dir)
        self._converter = DocumentConverter(format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=options)})

    def convert(self, prepared_pages: list, work_dir: Path) -> tuple:
        """prepared_pages = [(page_number, PreparedPage)]. The cleaned page images are packed
        into a temporary PDF so Docling benefits from our deskew/rotation/cleanup too.
        Returns (blocks, {page_number: ocr_score})."""
        pdf_path = work_dir / "_docling_input.pdf"
        pdf = pdfium.PdfDocument.new()            # page by page, JPEGs are embedded without decoding (low memory)
        try:
            for _number, dpi, height, width, image_file in prepared_pages:
                w_pt, h_pt = width * 72.0 / dpi, height * 72.0 / dpi
                page = pdf.new_page(w_pt, h_pt)
                image = pdfium.PdfImage.new(pdf)
                image.load_jpeg(str(image_file), inline=False)
                image.set_matrix(pdfium.PdfMatrix().scale(w_pt, h_pt))
                page.insert_obj(image)
                page.gen_content()
                page.close()
            pdf.save(str(pdf_path))
        finally:
            pdf.close()
        try:
            result = self._converter.convert(str(pdf_path), raises_on_error=True)
        except Exception as exc:
            msg = str(exc)
            if any(k in msg.lower() for k in ("huggingface", "hf.co", "connection", "403", "offline", "artifacts")):
                raise EngineUnavailable(
                    "Docling could not load its layout/table models (" + msg[:200] + "). Download them once on a "
                    "machine with internet access (docling-tools models download), copy the folder, and pass "
                    "--docling-models FOLDER; or use the standard engine.") from exc
            raise
        scores = {}
        for idx, (page_number, *_rest) in enumerate(prepared_pages, 1):
            score = None
            for key in (idx, idx - 1):          # Docling numbers pages from 1 (older versions from 0)
                try:
                    score = float(result.confidence.pages[key].ocr_score)
                    break
                except Exception:
                    continue
            scores[page_number] = score if score is not None and score == score else 1.0   # NaN -> 1.0
        return self.blocks_from_document(result.document, prepared_pages, scores), scores

    @staticmethod
    def blocks_from_document(doc, prepared_pages: list, scores: dict) -> list:
        """Map a DoclingDocument onto our Block model (reading order is Docling's).
        prepared_pages = [(page_number, dpi, height_px, width_px, ...)]"""
        from docling_core.types.doc import ContentLayer, PictureItem, TableItem
        blocks = []
        seen_captions = set()
        layers = {ContentLayer.BODY, ContentLayer.FURNITURE}
        for item, _level in doc.iterate_items(included_content_layers=layers):
            prov = item.prov[0] if getattr(item, "prov", None) else None
            if prov is None or not (1 <= prov.page_no <= len(prepared_pages)):
                continue
            page_number, page_dpi, page_h_px = prepared_pages[prov.page_no - 1][:3]
            page_h_pt = page_h_px * 72.0 / page_dpi
            bb = prov.bbox.to_top_left_origin(page_height=page_h_pt)
            k = page_dpi / 72.0
            box = (int(bb.l * k), int(bb.t * k), int(math.ceil(bb.r * k)), int(math.ceil(bb.b * k)))
            conf = scores.get(page_number, 1.0)
            label = getattr(item.label, "value", str(item.label))
            if isinstance(item, TableItem):
                grid = [[(cell.text or "").strip() for cell in row] for row in item.data.grid]
                merges = sorted({(c.start_row_offset_idx, c.start_col_offset_idx, c.end_row_offset_idx - 1, c.end_col_offset_idx - 1)
                                 for c in item.data.table_cells if c.row_span > 1 or c.col_span > 1})
                header_rows = sum(1 for row in item.data.grid if row and all(c.column_header for c in row)) or (1 if grid and _looks_like_header(grid[0]) else 0)
                table = TableData(grid=grid, conf=[[conf] * len(r) for r in grid], merges=list(merges), header_rows=min(header_rows, 3))
                blk = Block("table", page_number, box, table=table,
                            lines=[Line(" ".join(" ".join(r) for r in grid), conf, box, _estimate_word_boxes(" ".join(" ".join(r) for r in grid), box, conf))])
                cap = item.caption_text(doc) if hasattr(item, "caption_text") else ""
                if cap:
                    blk.meta["caption"] = cap
                    for ref in item.captions:
                        seen_captions.add(ref.cref)
                blocks.append(blk)
                if cap:   # emit the caption as its own block too (Excel caption row + Word report)
                    meta = dict(caption_meta(cap, "table"), attach="next")
                    blk.meta["caption_label"] = meta["caption_label"]
                    blocks.insert(len(blocks) - 1, Block("caption", page_number, box, text=cap, meta=meta,
                                                         lines=[Line(cap, conf, box, _estimate_word_boxes(cap, box, conf))]))
                continue
            if isinstance(item, PictureItem):
                blk = Block("figure", page_number, box)
                cap = item.caption_text(doc) if hasattr(item, "caption_text") else ""
                blocks.append(blk)
                if cap:
                    for ref in item.captions:
                        seen_captions.add(ref.cref)
                    meta = dict(caption_meta(cap, "figure"), attach="prev")
                    blk.meta.update({"caption": cap, "caption_label": meta["caption_label"]})
                    blocks.append(Block("caption", page_number, box, text=cap, meta=meta,
                                        lines=[Line(cap, conf, box, _estimate_word_boxes(cap, box, conf))]))
                continue
            text = normalize_ocr_text(getattr(item, "text", "") or "")
            if not text or getattr(item, "self_ref", None) in seen_captions:
                continue
            meta = {}
            if label in ("section_header", "title"):
                kind = "heading"
                m = RE_PARA_NUM.match(text)
                meta = {"level": (m.group("num").count(".") + 1) if m else 1, "title": text}
                if m:
                    meta["number"] = m.group("num")
            elif label == "list_item":
                kind = "list_item"
                marker = (getattr(item, "marker", "") or "").strip()
                if marker and not text.startswith(marker):
                    text = f"{marker} {text}"
                meta = {"marker": marker}
            elif label == "caption":
                kind, meta = "caption", caption_meta(text, "figure")
            elif label == "page_header":
                kind = "page_header"
            elif label == "page_footer":
                kind = "page_footer"
            else:
                kind, meta = classify_text_block(text, 3)
                if kind == "heading":
                    kind, meta = "paragraph", {}
            blocks.append(Block(kind, page_number, box, text=text, meta=meta,
                                lines=[Line(text, conf, box, _estimate_word_boxes(text, box, conf))]))
        return blocks


# =============================================================================
# 9. REQUIREMENT / CAPTION / TABLE EXTRACTION
#    Deterministic rules, in this order:
#      a. text is split into sentences; a sentence with a requirement keyword is a
#         requirement (MIL-STD-961: "shall" = binding, "will" = intent, "should" = goal)
#      b. a keyword sentence ending in ":" is a STEM; each list item below it becomes
#         "<stem> <marker> <item>" (one row per item, nested lists are supported)
#      c. a table with a requirement column (by header name, or a column where most
#         cells contain a keyword) gives one requirement per row, other columns are
#         kept as attributes (e.g. "Method: T")
#      d. every table becomes one row (text grid) + its own worksheet; every caption one row
# =============================================================================
# numbers where a letter probably replaced a digit (1O0, 2l.5): flagged, never changed
RE_SUSPECT_NUMBER = re.compile(r"(?<![A-Za-z])(?:\d+[OoIil|]+\d+[A-Za-z]?|\d+[OoIil|]+(?=[.,]\d)|[OoIl|]\d+(?![A-Za-z]))(?:[.,]\d+)?")
# Words that end with a period but do NOT end a sentence. Units (ms, min, ft, lbs ...) are deliberately
# absent: "... within 100 ms. The display shall ..." is two requirements.
ABBREVIATIONS = {"e.g", "i.e", "fig", "figs", "no", "nos", "para", "paras", "ref", "refs", "rev", "sec", "vol",
                 "approx", "dept", "govt", "mr", "mrs", "dr", "vs", "cf", "al", "ch", "art", "appx", "app", "eq",
                 "eqs", "incl", "resp", "mfr", "spec", "specs", "tab", "u.s", "u.s.a", "p", "pp", "op", "cit", "viz",
                 "nom", "st", "mt", "jr", "sr", "inc", "corp", "co", "ltd"}
CONTINUATION_START = re.compile(r"^(?:This|These|That|Those|It|Its|Such|The above|Said|Where|Except|Unless|Including|"
                                r"In addition|Additionally|Also|However|Otherwise|For example|E\.g\.|I\.e\.)\b")
WEAK_REQUIREMENT_HEADERS = ("description", "text", "statement", "specification", "spec")
# verification methods (MIL-STD-961: analysis, demonstration, examination, test; NASA/SRS: inspection)
METHOD_NAMES = {"i": "Inspection", "a": "Analysis", "d": "Demonstration", "t": "Test", "e": "Examination",
                "s": "Similarity", "inspection": "Inspection", "analysis": "Analysis", "demonstration": "Demonstration",
                "demo": "Demonstration", "test": "Test", "examination": "Examination", "exam": "Examination",
                "similarity": "Similarity", "insp": "Inspection"}
VERIFICATION_HEADERS = ("method", "verification", "verif", "ver method", "vm", "v v", "qualification method", "test method")
# The sentence that defines the letters of a verification matrix, in either of its usual forms:
# "... by inspection (I), analysis (A), demonstration (D), or test (T) ..."  and  "I = Inspection".
_METHOD_WORDS = r"inspection|analysis|demonstration|examination|similarity|test"
RE_METHOD_LEGEND = re.compile(
    r"\b(" + _METHOD_WORDS + r")\b\s*[\(\[]\s*([A-Za-z])\s*[\)\]]"
    r"|[\(\[]?\b([A-Za-z])\b[\)\]]?\s*(?:=|:|-|\u2013|\u2014)\s*\b(" + _METHOD_WORDS + r")\b", re.I)
# What a verification matrix says about one paragraph, and whether the method names had to be
# deduced from the legend's column order because the matrix's own headers could not be read.
MatrixMethods = namedtuple("MatrixMethods", "methods from_legend")


def method_legend(doc: DocResult) -> list:
    """The method letters a document defines for its verification matrix, in the order written:
    "inspection (I), analysis (A), demonstration (D), or test (T)" -> [(I, Inspection), ...].

    A matrix prints its method columns in exactly that order, so this is what names the columns of
    a matrix whose own headers are printed sideways and could not be read."""
    for b in doc.blocks:
        if b.kind not in ("paragraph", "list_item", "heading", "caption"):
            continue
        order = []
        for m in RE_METHOD_LEGEND.finditer(b.text or ""):
            name, letter = (m.group(1), m.group(2)) if m.group(1) else (m.group(4), m.group(3))
            method, letter = METHOD_NAMES.get((name or "").lower()), (letter or "").upper()
            if method and letter and letter not in [l for l, _ in order]:
                order.append((letter, method))
        if len(order) >= 2:            # a real legend names at least two methods
            return order
    return []


RE_REFERENCES = re.compile(
    r"\b(?:MIL|DOD|DoD|FED|NASA|JSSG)-(?:STD|DTL|PRF|HDBK|SPEC)-\d+[A-Z]?\b"
    r"|\bMIL-[A-Z]-\d+[A-Z]?\b"
    r"|\b(?:ISO|IEC|ISO/IEC|ISO/IEC/IEEE)\s+\d+(?:-\d+)*(?::\d{4})?"
    r"|\bIEEE(?:\s+Std)?\s+\d+(?:\.\d+)?(?:-\d{4})?"
    r"|\b(?:SAE\s+)?(?:AS|ARP|AMS)\d{3,}[A-Z]?\b|\bRTCA\s+DO-\d+[A-Z]?\b|\bDO-\d{3}[A-Z]?\b"
    r"|\bSTANAG\s+\d{4}\b|\bNASM\s?\d+\b|\bASTM\s+[A-Z]\d+\b|\bANSI[/ ][\w./-]+\d"
    r"|\b(?:TABLE|Table)\s+(?:[A-Z]-)?(?:\d+(?:[-.]\d+)*|[IVXL]+\b)"
    r"|\b(?:FIGURE|Figure|FIG\.|Fig\.)\s+(?:[A-Z]-)?\d+(?:[-.]\d+)*"
    r"|\b(?:paragraph|para\.|section|Section|SECTION|appendix|Appendix|APPENDIX)\s+(?:\d+(?:\.\d+)*|[A-Z](?:\.\d+)*)\b")


def find_references(text: str) -> str:
    seen = []
    for m in RE_REFERENCES.finditer(text or ""):
        ref = re.sub(r"\s+", " ", m.group(0)).rstrip(".")
        if ref not in seen:
            seen.append(ref)
    return "; ".join(seen)


def method_name(value: str) -> str:
    """'T' -> 'Test'; 'A, T' -> 'Analysis, Test'; unknown text is kept as written."""
    parts = [p for p in re.split(r"[\s,/&+;]+", (value or "").strip().rstrip(".")) if p]
    names = [METHOD_NAMES.get(p.lower().rstrip(".")) for p in parts]
    return ", ".join(names) if parts and all(names) else (value or "").strip()
QUALITY_RULES = [
    ("Placeholder TBD/TBR/TBS", re.compile(r"\bTB[DRSC]\b|\bto be (?:determined|resolved|supplied|specified|confirmed|defined)\b", re.I)),
    ("Escape clause", re.compile(r"\b(?:as appropriate|as applicable|as required|if possible|where possible|if practical|"
                                 r"if practicable|to the extent (?:possible|practical|practicable)|as necessary|if necessary|"
                                 r"when required|if required|so far as is possible)\b", re.I)),
    ("Vague term", re.compile(r"\b(?:adequate|appropriate|sufficient|reasonable|user[- ]friendly|easy|easily|quickly|fast|"
                              r"robust|flexible|efficient|effective|minimi[sz]e|maximi[sz]e|optimi[sz]e|approximately|several|"
                              r"various|significant|state[- ]of[- ]the[- ]art|as soon as possible|timely)\b", re.I)),
    ("Open-ended list", re.compile(r"\b(?:etc\.?|and so on|including but not limited to|but not limited to)(?=\W|$)", re.I)),
    ("Ambiguous and/or", re.compile(r"\band/or\b", re.I)),
]


def split_sentences(text: str) -> list:
    """Split into sentences, respecting abbreviations, decimals, section numbers and IDs.
    Returns [(start, end)] character spans."""
    spans, start = [], 0
    for m in re.finditer(r"[.!?][\"')\]]*\s+", text):
        end = m.end()
        before = text[start:m.start() + 1]
        last_token = re.findall(r"[\w.]+\.?$", before)
        token = last_token[0].rstrip(".").lower() if last_token else ""
        nxt = text[end:end + 1]
        if not nxt or nxt.islower() or nxt.isdigit() or nxt in ",;:)":
            continue                                   # "... para. 3" / "i.e. the" -> same sentence
        if re.match(r"\([A-Za-z0-9&/-]{1,12}\)[,;:]?\s+[a-z]", text[end:]):
            continue                                   # "repair. (MTTR) shall ..." - a stray OCR period
        if token in ABBREVIATIONS or (last_token and re.fullmatch(r"[a-z]", last_token[0].rstrip("."))
                                      and not re.search(r"\d\s*[a-z]\.$", before)):
            continue                                   # "Fig. 3", "e.g. The", inline list marker "a. Search"
        spans.append((start, m.start() + 1 + len(m.group(0).rstrip()) - 1))
        start = end
    if start < len(text) and text[start:].strip():
        spans.append((start, len(text.rstrip())))
    return [(s, e) for s, e in spans if text[s:e].strip()]


class RequirementExtractor:
    def __init__(self, settings: Settings):
        self.settings = settings
        keys = sorted(settings.keywords.keys(), key=len, reverse=True)
        alternation = "|".join(r"\s+".join(re.escape(part) for part in k.split()) for k in keys)
        self.re_keyword = re.compile(r"(?<![A-Za-z])(" + alternation + r")(?![A-Za-z])", re.I)
        self.classes = {" ".join(k.lower().split()): v for k, v in settings.keywords.items()}
        self.last_number = ""            # the latest paragraph number seen, for the sequence check
        self.joined_blocks = {}          # list item block uid -> blocks holding more of its text

    # ---- helpers --------------------------------------------------------------------
    def keyword_in(self, text: str) -> Optional[str]:
        m = self.re_keyword.search(text or "")
        if not m:
            return None
        kw = " ".join(m.group(1).lower().split())
        # "May 1987" / "Will Smith": a capitalized modal right after a sentence start is fine,
        # but skip "will" used as a noun ("the will of"): rare - accepted.
        return kw

    def repair_keywords(self, text: str) -> tuple:
        """Fix OCR-garbled keywords; returns (text, [notes])."""
        notes = []
        if not self.settings.fix_ocr_keywords:
            return text, notes
        for word, pattern in OCR_KEYWORD_REPAIRS:
            if word not in self.classes and not any(word in k for k in self.classes):
                continue

            def fix(m, word=word):
                tok = m.group(0)
                if tok.lower() == word:
                    return tok
                fixed = word.capitalize() if tok[:1].isupper() else word
                notes.append(f"OCR keyword repaired: '{tok}' -> '{fixed}'")
                return fixed
            text = pattern.sub(fix, text)
        return text, notes

    def quality_notes(self, text: str) -> list:
        if not self.settings.quality_checks:
            return []
        notes = []
        for name, pattern in QUALITY_RULES:
            found = sorted({m.group(0).strip() for m in pattern.finditer(text)}, key=str.lower)
            if found:
                notes.append(f"{name}: {', '.join(found[:4])}")
        n_modal = len(self.re_keyword.findall(text))
        if n_modal > 1:
            notes.append(f"Possible compound requirement ({n_modal} keywords)")
        return notes

    def review_for_tokens(self, tokens: list, text: str, extra: list = ()) -> tuple:
        """-> (min_confidence, [review reasons]) for the OCR tokens that make up an item."""
        reasons = list(extra)
        confs = [t[1] for t in tokens]
        min_conf = min(confs) if confs else None
        # Low confidence is flagged for numbers (a misread digit changes the requirement and no
        # dictionary can catch it). For words, the dictionary/format checks are far more reliable:
        # on the test scans, low-confidence words were usually read correctly.
        low = [(t[0], t[1]) for t in tokens if token_attention((t[0], t[1], ""), self.settings)]
        if low:
            shown = ", ".join(f"'{t}' ({c:.2f})" for t, c in sorted(low, key=lambda x: x[1])[:6])
            reasons.append(f"Low OCR confidence: {shown}")
        stray = []
        for a, b, c in zip(tokens, tokens[1:], tokens[2:] + [("", 1.0, "")]):
            if (re.search(r"[A-Za-z0-9]\.$", a[0]) and not re.fullmatch(r"(?:[A-Za-z]\.){1,3}", a[0])
                    and a[0][:-1].lower() not in ABBREVIATIONS
                    and not RE_BARE_MARKER.match(b[0])          # "... limits. 1." - a lead-in and its list item
                    # a full stop that ends a LINE, with the next line starting with a number, ends a
                    # sentence ("... operating." / "2.5 C per minute ..." in a list's value column)
                    and not (len(b) > 3 and b[3] and b[0][:1].isdigit())
                    and (b[0][:1].islower() or b[0][:1].isdigit()
                         or (re.fullmatch(r"\d+\.", a[0]) and RE_UNIT.fullmatch(b[0].rstrip(".,;:")))
                         or (re.fullmatch(r"\([A-Za-z0-9&/-]{1,12}\)[,;:]?", b[0]) and c[0][:1].islower()))):
                stray.append(f"{a[0]} {b[0]}")
        if stray:
            reasons.append("Stray period? " + "; ".join(f"'{x}'" for x in stray[:5]))
        suspicious = []
        for t in tokens:
            if len(t) > 2 and t[2] and (t[0], t[2]) not in suspicious:
                suspicious.append((t[0], t[2]))
        if suspicious:
            reasons.append("Check word(s): " + "; ".join(f"'{w}' ({why})" for w, why in suspicious[:6]))
        suspects = sorted(set(RE_SUSPECT_NUMBER.findall(text)))
        if suspects:
            reasons.append("Number may contain a misread letter: " + ", ".join(suspects[:5]))
        return min_conf, reasons

    # ---- main entry ------------------------------------------------------------------
    def extract(self, doc: DocResult) -> list:
        items = []
        source = Path(doc.source).name
        vocab = document_text_context(doc, self.settings)
        self.ctx = vocab
        dpi_of = {p.number: p.dpi for p in doc.pages}
        blocks = [b for b in doc.blocks if b.kind not in ("page_header", "page_footer")]
        self.id_prefixes = self._learn_id_prefixes(blocks)
        section = ""
        self.last_number = ""            # the latest paragraph number, for the sequence check
        self.joined_blocks = {}
        i = 0
        while i < len(blocks):
            b = blocks[i]
            if b.kind == "heading" and self.keyword_in(self.repair_keywords(b.text)[0]):
                b.kind = "paragraph"             # a 'heading' that states a requirement is a paragraph
            if b.kind == "heading":
                num, title = b.meta.get("number", ""), b.meta.get("title", b.text)
                section = f"{num} {title}".strip() if num else title
                if _number_parts(num):
                    self.last_number = num
                i += 1
                continue
            if b.kind == "caption":
                if not b.meta.get("duplicate_of_table"):
                    items.append(self._caption_item(b, source, section))
                i += 1
                continue
            if b.kind == "table" and b.table is not None:
                items.extend(self._table_items(b, source, section))
                i += 1
                continue
            if b.kind == "figure":
                items.extend(self._figure_items(b, source, section))
                i += 1
                continue
            if b.kind in ("paragraph", "list_item"):
                consumed = self._text_block_items(blocks, i, source, section, vocab, dpi_of, items)
                i += consumed
                continue
            i += 1
        items.extend(self._repeated_furniture_items(doc, source))
        # verification methods and references
        matrix = self._verification_matrix(doc)
        for it in items:
            if it.item_type != "Requirement":
                continue
            it.references = find_references(it.text)
            if not it.verification and it.doc_ref:
                key = clean_para_number(it.doc_ref.split()[0].rstrip("."))
                entry = matrix.get(key)
                if entry and entry.from_legend:
                    # the columns were named from the legend's order, not from a header we read
                    it.verification = entry.methods + " (from verification matrix, columns named from the legend)"
                    it.review_reasons.append(
                        "Verification method comes from the verification matrix, whose column headers could "
                        "not be read; the methods follow the order given in the legend - check against the scan")
                elif entry:
                    it.verification = entry.methods + " (from verification matrix)"
        # item numbers
        counters = Counter()
        prefix = {"Requirement": "R", "Table": "T", "Table Caption": "TC", "Figure Caption": "FC"}
        for it in items:
            counters[it.item_type] += 1
            width = 4 if it.item_type == "Requirement" else 3
            it.item_id = f"{prefix.get(it.item_type, 'X')}-{counters[it.item_type]:0{width}d}"
        return items

    @staticmethod
    def _learn_id_prefixes(blocks: list) -> Counter:
        """Prefixes of the requirement IDs this document really uses ('SRS' in [SRS-101], 'ENV' in a
        table ID column). A bare 'MK-54 torpedo ...' is only read as an ID if MK is such a prefix."""
        prefixes = Counter()
        for b in blocks:
            if b.kind in ("paragraph", "list_item", "heading") and b.text:
                m = RE_LEADING_ID.match(b.text)
                if m and _starts_with_delimited_id(b.text):
                    prefixes[_id_prefix(m.group("id"))] += 1
                for tm in re.finditer(r"[\[(](" + RE_REQ_ID + r")[\])]", b.text):
                    if "-" in tm.group(1) or "_" in tm.group(1):
                        prefixes[_id_prefix(tm.group(1))] += 1
            elif b.kind == "table" and b.table is not None:
                for row in b.table.grid[b.table.header_rows:]:
                    if row and row[0] and re.fullmatch(r"\s*(?:" + RE_REQ_ID + r")\s*", row[0]) and re.search(r"[-_]", row[0]):
                        prefixes[_id_prefix(row[0].strip())] += 1
        return prefixes

    def _repeated_furniture_items(self, doc: DocResult, source: str) -> list:
        """Keyword text printed in the header/footer of many pages (e.g. a handling notice) is reported
        once, flagged, instead of once per page or not at all."""
        groups = []                               # [(kind, notice text, [blocks])], near-identical notices together
        for b in doc.blocks:
            if b.kind in ("page_header", "page_footer") and b.meta.get("repeated_keyword_text"):
                key = b.meta.get("notice_key") or " ".join(b.text.lower().split())
                for kind, gkey, members in groups:
                    if kind == b.kind and (gkey == key or (re.findall(r"\d+", gkey) == re.findall(r"\d+", key)
                                                          and _similar(gkey, key, at_least=0.9) >= 0.9)):
                        members.append(b)
                        break
                else:
                    groups.append((b.kind, key, [b]))
        out = []
        for kind, _key, group in groups:
            first = max(group, key=lambda x: x.conf)      # the best-read copy
            clean = re.sub(r"^\s*(?:page\s*)?\d{1,4}(?:\s*of\s*\d{1,4})?\s+|\s+(?:page\s*)?\d{1,4}(?:\s*of\s*\d{1,4})?\s*$",
                           " ", first.text, flags=re.I)
            text, notes = self.repair_keywords(" ".join(clean.split()))
            kw = self.keyword_in(text)
            if not kw:
                continue
            pages = sorted({p for b in group for p in (b.pages or [b.page])})
            where = "header" if kind == "page_header" else "footer"
            tokens = word_tokens((w for ln in first.lines for w in ln.words), self.ctx)
            if clean.strip() != first.text.strip():          # leave out the page number printed on the same line
                while tokens and re.fullmatch(r"\d{1,4}", tokens[-1][0]) and not text.rstrip().endswith(tokens[-1][0]):
                    tokens = tokens[:-1]
                while tokens and re.fullmatch(r"\d{1,4}", tokens[0][0]) and not text.lstrip().startswith(tokens[0][0]):
                    tokens = tokens[1:]
            conf, reasons = self.review_for_tokens(
                tokens, text, [f"Repeated in the page {where} on {len(pages)} pages - check whether it is a requirement"] + notes)
            out.append(ExtractedItem("Requirement", source, _pages_text(pages), "", text=text, keyword=kw,
                                     req_class=self.classes.get(kw, ""), found_in=f"Page {where} (repeated)",
                                     ocr_conf=conf, review_reasons=reasons, quality_notes=self.quality_notes(text),
                                     block_uids=[b.uid for b in group]))
        return out

    def _figure_items(self, b: Block, source: str, section: str) -> list:
        label_text = " ".join(ln.text for ln in b.lines)
        kw = self.keyword_in(self.repair_keywords(label_text)[0])
        if not kw:
            return []
        tokens = word_tokens((w for ln in b.lines for w in ln.words), self.ctx)
        conf, reasons = self.review_for_tokens(tokens, label_text, ["Keyword found inside a figure - check manually"])
        return [ExtractedItem("Requirement", source, b.page_label(), section, text=label_text, keyword=kw,
                              req_class=self.classes.get(kw, ""), found_in="Figure text", ocr_conf=conf,
                              review_reasons=reasons, block_uids=[b.uid])]

    @staticmethod
    def _verification_matrix(doc: DocResult) -> dict:
        """Read verification cross-reference matrices: a column of paragraph numbers beside two or
        more method columns (I / A / D / T ...) marked with an X.
        Returns {paragraph: MatrixMethods(methods, from_legend)}.

        The method columns are recognised by their CONTENT - a column holding nothing but marks -
        so a matrix still works when its headers are printed sideways. Each column is named from
        its own header where that could be read, and otherwise from the order of the legend
        sentence printed above the matrix. A name that came from the legend is a deduction rather
        than something read off the page, so it is reported as such and the row is flagged."""
        matrix = {}
        legend = method_legend(doc)
        for b in doc.blocks:
            t = b.table
            if b.kind != "table" or t is None or t.n_rows < 2:
                continue
            # Row 0 is treated as a header even when it did not look like one, which is what a row
            # of unreadable sideways headers looks like. With a merged banner above them
            # ("VERIFICATION METHOD"), the lowest readable line of each column is its name.
            head_rows = max(t.header_rows, 1)
            data = t.grid[head_rows:]
            if not data:
                continue
            header = [h.lower().rstrip(".") for h in column_headers(t.grid, head_rows)]
            method_cols = {c: METHOD_NAMES[h] for c, h in enumerate(header) if h in METHOD_NAMES}
            from_legend = False
            if len(method_cols) < 2:
                # Headers unreadable (the usual reason: they are printed sideways). Line the block
                # of method columns up with the legend, which lists the methods in column order.
                run = _method_column_run(data, t.n_cols)
                if run and len(run) == len(legend):
                    method_cols = {c: legend[i][1] for i, c in enumerate(run)}
                    from_legend = True
            if len(method_cols) < 2:
                continue
            para_col = None
            def para(v):                      # OCR tolerance: "3,2.1.1.." -> "3.2.1.1"
                return clean_para_number(re.sub(r"(?<=\d),(?=\d)", ".", (v or "").strip()).rstrip("."))

            for c in range(t.n_cols):
                vals = [para(r[c]) for r in data if c < len(r) and r[c].strip()]
                if c not in method_cols and vals and sum(1 for v in vals if re.fullmatch(r"(?:\d{1,2}|[A-Z])(?:\.[\dlIi|]{1,3}){1,7}\.?", v)) >= 0.6 * len(vals):
                    para_col = c
                    break
            if para_col is None:
                continue
            for r in data:
                num = para(r[para_col]) if para_col < len(r) else ""
                marks = [name for c, name in sorted(method_cols.items())
                         if c < len(r) and re.fullmatch(r"\s*(?:[xX\u00d7\u2713\u2714\u2022\u221a*+]|yes|Yes|YES)\s*", r[c] or "")]
                if num and marks:
                    matrix[num] = MatrixMethods(", ".join(dict.fromkeys(marks)), from_legend)
        return matrix

    # ---- captions and tables --------------------------------------------------------------
    def _caption_item(self, b: Block, source: str, section: str) -> ExtractedItem:
        kind = b.meta.get("caption_kind", "figure")
        tokens = word_tokens((w for ln in b.lines for w in ln.words), self.ctx)
        conf, reasons = self.review_for_tokens(tokens, b.text)
        return ExtractedItem("Table Caption" if kind == "table" else "Figure Caption", source, b.page_label(), section,
                             doc_ref=b.meta.get("caption_label", ""), text=b.text, found_in="Caption",
                             ocr_conf=conf, review_reasons=reasons, block_uids=[b.uid])

    def _table_items(self, b: Block, source: str, section: str) -> list:
        t = b.table
        caption = b.meta.get("caption", "")
        label = b.meta.get("caption_label", "")
        min_conf = min((c for row in t.conf for c in row), default=1.0)
        reasons = []
        flagged = [t.grid[r][c] for r in range(t.n_rows) for c in range(len(t.grid[r]))
                   if t.grid[r][c] and _cell_attention(t, r, c, self.ctx, self.settings)]
        if flagged:
            reasons.append(f"{len(flagged)} table cell(s) to check against the scan, e.g. " +
                           "; ".join(f"'{x[:40]}'" for x in flagged[:3]))
        suspects = sorted({m for row in t.grid for cell in row for m in RE_SUSPECT_NUMBER.findall(cell or "")})
        if suspects:
            reasons.append("Number may contain a misread letter: " + ", ".join(suspects[:5]))
        if t.continued_pages:
            reasons.append("Table continues across pages - check the joined rows")
        notes = [] if caption else ["No caption found near this table"]
        out = [ExtractedItem("Table", source, b.page_label(), section, doc_ref=label, title=caption,
                             text=t.as_text(), found_in=f"Table ({t.n_rows} rows x {t.n_cols} columns)",
                             ocr_conf=min_conf, review_reasons=reasons, quality_notes=notes, table=t, block_uids=[b.uid])]
        out.extend(self._requirement_rows_from_table(b, source, section, caption or label))
        return out

    def _requirement_rows_from_table(self, b: Block, source: str, section: str, caption: str) -> list:
        t = b.table
        s = self.settings
        if t.n_rows < 2 or t.n_cols < 1:
            return []
        header = column_headers(t.grid, t.header_rows) if t.header_rows else []
        data_start = t.header_rows
        data = t.grid[data_start:]
        if not data:
            return []
        n_cols = t.n_cols

        def norm_header(h):
            return " ".join(re.sub(r"[^a-z# ]+", " ", h.lower()).split())

        id_headers = {norm_header(x) for x in s.id_column_headers}

        def is_id_header(hn):
            """'Req ID', 'Requirement No.', 'Requirement Number', 'Para #' ... name an ID column."""
            return hn in id_headers or bool(re.search(r"(?:^|\s)(?:id|ids|no|nos|nr|num|number|#|ident|identifier|ref)$", hn))

        def is_attribute_header(hn):
            """'Requirement Title / Type / Status / Owner ...' describe the requirement - they are not its text."""
            return bool(re.search(r"(?:^|\s)(?:title|name|type|category|class|status|owner|source|priority|rationale|"
                                  r"level|phase|allocation|trace|parent|method)$", hn))

        modal_share = [sum(1 for r in data if c < len(r) and self.keyword_in(r[c])) / len(data) for c in range(n_cols)]
        mean_words = [statistics.fmean(word_count(r[c]) if c < len(r) else 0 for r in data) for c in range(n_cols)]
        text_col, header_says = None, False
        named = []           # columns whose header says "Requirement ..." (not "Requirement ID")
        for c, h in enumerate(header):
            hn = norm_header(h)
            if not hn or word_count(hn) > 3 or is_id_header(hn) or is_attribute_header(hn):
                continue
            if any(hn == k or hn.startswith(k) for k in s.requirement_table_headers):
                named.append((hn in s.requirement_table_headers or hn.endswith(("text", "statement")), modal_share[c], mean_words[c], -c))
            elif any(hn == k or hn.startswith(k) for k in WEAK_REQUIREMENT_HEADERS) and modal_share[c] >= 0.3:
                text_col = c                             # "Description" column that contains shall-statements
        if named:
            best = max(named)                            # exact "Requirement" / "... Text" first, then most shall-cells
            modal_elsewhere = max((share for c, share in enumerate(modal_share) if c != -best[3]), default=0.0)
            if (best[1] > 0 or best[2] >= 3) and not (best[1] == 0 and modal_elsewhere >= 0.5):
                text_col, header_says = -best[3], True   # every row of that column is a requirement
            elif modal_elsewhere >= 0.5:
                text_col = None                          # the shall-statements are in another column (found below)
        if text_col is None and modal_share and max(modal_share) >= 0.5:
            text_col = int(np.argmax(modal_share))
        rows_out = []
        if text_col is None:
            # not a requirement table: still catch any cell that states a requirement
            for r_off, row in enumerate(data):
                for c, cell in enumerate(row):
                    kw = self.keyword_in(cell)
                    if kw and word_count(cell) >= 4:
                        attrs = "; ".join(f"{header[k] if k < len(header) and header[k] else f'Column {k + 1}'}: {v}"
                                          for k, v in enumerate(row) if k != c and v)
                        rows_out.append(self._table_requirement(b, source, section, caption, cell, kw, "", attrs,
                                                                t.conf[data_start + r_off][c], "Table cell",
                                                                _cell_words(t, data_start + r_off, c)))
            return rows_out
        id_col = None
        for c, h in enumerate(header):
            if c != text_col and norm_header(h) and is_id_header(norm_header(h)):
                id_col = c
                break
        if id_col is None:
            for c in range(n_cols):
                if c != text_col:
                    vals = [r[c] for r in data if c < len(r) and r[c]]
                    if vals and sum(1 for v in vals if RE_ID_CELL.match(v)) >= 0.6 * len(vals):
                        id_col = c
                        break
        seen = set()
        for r_off, row in enumerate(data):
            text = (row[text_col] if text_col < len(row) else "").strip()
            if not text or (text, r_off) in seen:
                continue
            seen.add((text, r_off))
            kw = self.keyword_in(text)
            if not kw and not header_says:
                continue
            attrs = "; ".join(f"{header[c] if c < len(header) and header[c] else f'Column {c + 1}'}: {row[c]}"
                              for c in range(min(n_cols, len(row))) if c not in (text_col, id_col) and row[c])
            ref = row[id_col].strip().strip(".,;:'`[]()") if id_col is not None and id_col < len(row) else ""
            item = self._table_requirement(b, source, section, caption, text, kw or "", ref, attrs,
                                           t.conf[data_start + r_off][text_col], "Table row",
                                           _cell_words(t, data_start + r_off, text_col))
            for c, h in enumerate(header):
                if c < len(row) and row[c] and c not in (text_col, id_col) and norm_header(h).startswith(VERIFICATION_HEADERS):
                    item.verification = method_name(row[c])
            rows_out.append(item)
        return rows_out

    def _table_requirement(self, b, source, section, caption, text, kw, ref, attrs, cell_conf, found_in, cell_words=None) -> ExtractedItem:
        text, repair_notes = self.repair_keywords(text)
        kw = kw or self.keyword_in(text) or ""
        words = cell_words or [[w, cell_conf] for w in text.split()]
        tokens = word_tokens((Word(w, c, (0, 0, 0, 0)) for w, c in words), self.ctx)
        min_conf, reasons = self.review_for_tokens(tokens, text, repair_notes)
        cell_conf = min_conf if min_conf is not None else cell_conf
        return ExtractedItem("Requirement", source, b.page_label(), section, doc_ref=ref, text=text, keyword=kw,
                             req_class=self.classes.get(kw, "Listed in requirement table" if not kw else ""),
                             found_in=found_in, context=caption, attributes=attrs, ocr_conf=cell_conf,
                             review_reasons=reasons, quality_notes=self.quality_notes(text), block_uids=[b.uid])

    # ---- paragraphs and lists ----------------------------------------------------------------
    def parse_head(self, text: str) -> tuple:
        """Split '3.2.1.1 Detection range. The radar shall ...' -> ('3.2.1.1', 'Detection range', body, offset)
        and '[SRS-101] The GCS shall ...' -> ('SRS-101', '', body, offset).
        'MK-54 torpedo ...' and '3.3 V shall ...' are left alone (a designator and a quantity)."""
        ref, title, body = "", "", text
        lowered = False                  # the title's capital was read as a small letter ("3.2.7 output. The ...")
        prefixes = getattr(self, "id_prefixes", Counter())
        m = RE_LEADING_ID.match(body)
        if m and (("-" in m.group("id") or "_" in m.group("id")) and
                  (_starts_with_delimited_id(body) or prefixes.get(_id_prefix(m.group("id")), 0) >= 2)):
            ref, body = m.group("id"), m.group("rest")
        else:
            m = RE_PARA_NUM.match(body)
            if m and ("." in m.group("num") or m.group("dot")) and not _starts_with_unit(m.group("rest")):
                lowered = m.group("rest")[:1].islower() and m.group("num").count(".") >= 2 \
                    and _lowercased_title(m.group("rest"))
                if m.group("rest")[:1].isupper() or lowered:
                    ref, body = clean_para_number(m.group("num")), m.group("rest")
        if ref:
            t = re.match(r"^(?P<title>[A-Z][^.:;!?\[\]]{0,80}?)[.:]{1,3}['`·,]?(?:\s+|(?=[A-Z]))(?P<body>[A-Z\[(\"].*)$",
                         body[:1].upper() + body[1:] if lowered else body, re.S)
            if (t and word_count(t.group("title")) <= 8 and not _modal_re().search(t.group("title"))
                    and not re.search(r"\s(?:The|A|An|This|These|Each|All|Any)\s", t.group("title"))):
                title, body = t.group("title").strip(), t.group("body")
        offset = len(text) - len(body)
        return ref, title, body, offset

    def _starts_section(self, text: str) -> bool:
        """Does this paragraph begin a new SECTION, which ends a list? "3.4 Weight. ..." after 3.3 does.
        A list item's own number ("1. Input voltage:") does not, and neither does a value that happens
        to start a line ("2.5 C per minute ...", "20 Hz to 2000 Hz") - its number does not follow the
        section numbering (see plausible_next_number)."""
        m = RE_PARA_NUM.match(text or "")
        if not m:
            return False
        if _is_titled_number(text):
            return True                                   # "3.2.8 Weight. The ..." whatever its number (a skip, a misread)
        if "." not in m.group("num") and RE_NUMBERED_ITEM.match(text):
            return False                                  # "1. Input voltage: ..." - an item of the list
        return plausible_next_number(self.last_number, clean_para_number(m.group("num")))

    def _continues_item(self, nb: Block, item: Block, dpi: float) -> bool:
        """Is block `nb` more of list item `item` - its value carried on in another block, or onto the
        next page - rather than something new? It must start in one of the item's text columns (where
        its label or value starts), right under it or at the top of the next page, and must not start
        like something new: a marker, a label, a requirement ID or a titled paragraph ("3.2.4.1 Mode
        switching. The ..."). A value may start with any number ("2.5 C per minute", "3.3 V logic").

        "Right under" is measured against the item's own line spacing. After a blank line, a new
        sentence under an item that has ended its sentence is only carried on in the VALUE column (a
        second paragraph of a value on tab stops); in the item's text column it is a paragraph of its
        own, or the next item with its marker lost."""
        t = (nb.text or "").strip()
        if (not t or word_count(t) > 60 or RE_MARKER_SHAPE.match(t) or RE_LIST_MARKER.match(t)
                or RE_NUMBERED_ITEM.match(t) or RE_ITEM_LABEL.match(t) or _starts_with_delimited_id(t)
                or not item.lines):
            return False
        so_far = " ".join([item.text] + [blk.text for blk in self.joined_blocks.get(item.uid, [])])
        if _opens_sub_list(so_far):
            return False                                  # "a. ... shall include:" - what follows are its sub-items
        rows = _block_rows(item)
        label_x, value_xs = _item_columns(rows[0]) if rows else (None, [])
        at_value = any(abs(nb.box[0] - x) <= 0.1 * dpi for x in value_xs)
        if _is_titled_number(t):
            # "... as specified in" / "4.5.3 Random vibration. Levels ..." - a cross-reference that the value
            # carries on with, in its value column, and whose number does not follow the section numbering
            num = clean_para_number(RE_PARA_NUM.match(t).group("num"))
            if not (at_value and not _ends_sentence(so_far) and not plausible_next_number(self.last_number, num)):
                return False
        columns = [x for x in [label_x] + value_xs if x is not None]
        if nb.kind == "heading":
            columns = value_xs                            # a real heading never starts in a value column
        if not any(abs(nb.box[0] - x) <= 0.1 * dpi for x in columns):
            return False
        above = (self.joined_blocks.get(item.uid) or [item])[-1]     # the item's lowest part so far
        above_page = above.pages[-1] if above.pages else above.page
        if nb.page != above_page:
            return nb.page == above_page + 1                         # carried onto the next page
        text_h = statistics.median([max(ln.box[3] for ln in r) - min(ln.box[1] for ln in r) for r in rows])
        tops = [min(ln.box[1] for ln in r) for r in rows]
        bottoms = [max(ln.box[3] for ln in r) for r in rows]
        # the item's own line spacing: lines whose boxes touch or overlap (tall OCR boxes) are single-spaced, and so
        # is taken an item of one line - otherwise a blank line under it would count as "right under"
        gaps = [b - a for a, b in zip(bottoms, tops[1:])]
        row_gap = max(0.0, statistics.median(gaps)) if gaps else 0.25 * text_h
        gap = nb.box[1] - above.box[3]
        if gap <= row_gap + 0.5 * text_h:
            return True                                              # right under the item
        if gap > 2 * row_gap + 1.6 * text_h:
            return False                                             # more than one blank line
        # after one blank line: a sentence that was not finished carries on (in lower case - an item without
        # a full stop, "2. Track mode", is not unfinished); a finished one only in the value column
        return (not _ends_sentence(so_far) and (t[:1].islower() or t[:1].isdigit())) or at_value

    def _inner_paragraphs(self, text: str, body_start: int, previous: str = "") -> list:
        """[(cut, ref, title, body offset)] for numbered paragraphs that begin INSIDE one block.
        A scan often loses the blank line between "3.2.2.1 Weight. ..." and "3.2.2.2 Input power.
        ...", leaving both in the same block; without this every sentence would be filed under the
        first number.  parse_head decides what really is a heading, so a quantity ("3.3 V shall")
        and a cross-reference in running text ("... comply with 3.2.1. Testing ...") are left alone.
        A number must also FOLLOW the one before it (`previous`, then each number accepted here): in
        "3.2.4 ... per axis. 2.5 G peak ..." the 2.5 is a value, not paragraph 2.5."""
        found = []
        for m in RE_INNER_PARA_NUM.finditer(text, body_start):
            ref, title, _body, off = self.parse_head(text[m.start("num"):])
            # a paragraph with a run-in title ("3.2.8 Input power. The ...") is taken whatever its number
            # (the scan may have missed or misread the numbers in between); a bare number must follow on
            if ref and (title or plausible_next_number(previous, ref)):
                found.append((m.start(), ref, title, m.start("num") + off))
                previous = ref
        return found

    def _text_block_items(self, blocks: list, i: int, source: str, section: str, vocab: set, dpi_of: dict, items: list) -> int:
        """Handle a paragraph/list item (and, for a stem, the list below it). Returns blocks consumed."""
        b = blocks[i]
        tokens = block_tokens(b, vocab) if b.lines else [(w, 1.0, "") for w in b.text.split()]
        text, spans = tokens_text(tokens)
        marker = ""
        lead = RE_LIST_MARKER.match(text) or RE_NUMBERED_ITEM.match(text) if b.kind == "list_item" else None
        if lead:
            marker = lead.group("marker")
        body_start = len(text) - len(lead.group("rest")) if lead else 0
        ref, title, body, off = self.parse_head(text[body_start:])
        body_start += off
        own_ref = ref                            # the paragraph's own number - never a list marker
        if marker and not ref:
            ref = marker
        extra_reasons = ["Text continues across pages - check the join"] if b.meta.get("continued_across_pages") else []

        # One block can hold two numbered paragraphs when the scan lost the blank line between them
        # ("3.2.2.1 Weight. ... 3.2.2.2 Input power. ..."). Cut the sentences at every number found
        # inside the block, so each part is judged, numbered and titled on its own. The cut must come
        # BEFORE the keyword test: "... 45 kg. 3.2.2.2 The unit shall ..." is ONE sentence to
        # split_sentences (a full stop followed by a digit does not end a sentence), and cutting it
        # afterwards would lose the second requirement.
        inner = self._inner_paragraphs(text, body_start, own_ref if _number_parts(own_ref) else self.last_number)
        if _number_parts(own_ref) or inner:
            self.last_number = inner[-1][1] if inner else own_ref

        def paragraph_of(pos: int) -> int:
            """0 = the block's own paragraph, 1 = the first number found inside it, and so on."""
            return sum(1 for p in inner if p[3] <= pos)

        pieces = []
        for s0, e0 in split_sentences(text[body_start:]):
            s, e = s0 + body_start, e0 + body_start
            for cut, _ref, _title, body_at in inner:
                if s < cut < e:
                    pieces.append((s, cut))           # the end of the paragraph before the number
                    s = body_at                       # the number and its title are not requirement text
                elif cut <= s < body_at:
                    s = body_at
            if s < e:
                pieces.append((s, e))

        # sentences with requirement keywords (and the non-keyword sentences that explain them)
        units = []
        for s, e in pieces:
            sent = text[s:e].strip()
            if not sent:
                continue
            kw = self.keyword_in(self.repair_keywords(sent)[0])
            if kw:
                units.append([s, e, kw])
            elif (units and paragraph_of(units[-1][0]) == paragraph_of(s)       # never across a paragraph number
                  and CONTINUATION_START.match(sent) and not sent.endswith(":")):
                units[-1][1] = e
        if not units:
            return 1
        def head_of(u: list) -> tuple:
            """(number, title) of the paragraph a sentence actually sits in."""
            k = paragraph_of(u[0])
            return (ref, title) if k == 0 else tuple(inner[k - 1][1:3])
        heads = [head_of(u) for u in units]

        dpi = dpi_of.get(b.page, 200)
        # a stem inside the block ("... shall provide:" with the first items on the next rows of the same block)
        inline_items = []
        stem_unit = None
        # the paragraphs that follow, for a lead-in without a colon (see _plain_items)
        following = []
        for nb in blocks[i + 1:]:
            if nb.kind != "paragraph" or self._starts_section(nb.text) or _starts_with_delimited_id(nb.text):
                break
            following.append(nb)
        def plain_run_after(first_rows: list) -> int:
            """How many of the following paragraphs, with first_rows before them, are a list of plain items."""
            for n in range(len(following), -1, -1):
                if self._plain_items(first_rows + [r for blk in following[:n] for r in _text_rows(blk)], b.box[0], dpi):
                    return n
            return -1
        for u in units:
            if not _introduces_list(text[u[0]:u[1]]) or u[1] >= len(text.rstrip()) - 1:
                continue
            colon = text[u[0]:u[1]].strip().endswith(":")
            row_entries = self._row_entries(b, vocab, text, spans, start_char=u[1], whole_rows=not colon)
            if row_entries and not colon:
                # "... retested on the following day. The results ..." is a sentence, not a lead-in: without a
                # colon, the rows after it must start like items (a marker, a label on tab stops), or be a
                # run of short lines indented level with each other ("Search mode" / "Track mode")
                first = tokens_text(row_entries[0][1])[0]
                if not (RE_LIST_MARKER.match(first) or RE_NUMBERED_ITEM.match(first) or RE_MARKER_SHAPE.match(first)
                        or _is_item_label(first)
                        or plain_run_after([(x0, tokens_text(tk)[0]) for x0, tk, _ in row_entries]) >= 0):
                    row_entries = []
            if row_entries:
                stem_unit, inline_items = u, row_entries
                break
        if stem_unit is None and _introduces_list(text[units[-1][0]:units[-1][1]]):
            stem_unit = units[-1]
        last = stem_unit
        if last is not None:
            # ---- a stem: collect the list items that follow
            colon_stem = text[last[0]:last[1]].strip().endswith(":")
            children = list(inline_items)
            skipped = []
            # without a colon, unmarked items are taken only as a run of short lines indented level with
            # each other ("Own ship position and heading." / "Target range and bearing."), which may begin on
            # the lead-in's own block: one sentence after the lead-in is not a list
            inline_rows = [(x0, tokens_text(tk)[0]) for x0, tk, _ in inline_items]
            plain_run = max(0, plain_run_after(inline_rows)) if not colon_stem else 0
            j = i + 1
            while j < len(blocks):
                nb = blocks[j]
                if nb.kind in ("figure", "caption") and children and j + 1 < len(blocks):
                    after = next((x for x in blocks[j + 1:] if x.kind not in ("figure", "caption")), None)
                    prev_x0 = children[-1][0]
                    next_item = after is not None and (
                        (after.kind == "list_item" and abs(after.box[0] - prev_x0) <= 0.1 * dpi)
                        # the next item's marker was lost: it starts with its label, right of the markers
                        or (after.kind in ("paragraph", "list_item") and _is_item_label(after.text)
                            and prev_x0 - 0.1 * dpi <= after.box[0] <= prev_x0 + 1.2 * dpi))
                    if next_item:
                        skipped.append(nb)       # a stamp/signature/figure sitting between two list items
                        j += 1
                        continue
                if b.kind == "list_item" and nb.box[0] <= b.box[0] + 0.1 * dpi:
                    break                        # a sibling item, not a child ("(2) ... format:" then "(3) ...")
                if (children and children[-1][2].kind == "list_item" and nb.kind in ("paragraph", "heading")
                        and self._continues_item(nb, children[-1][2], dpi)):
                    # more of the item above - its value carried on after a larger gap or a page break,
                    # and may start with a number ("2.5 C per minute ...") - not a new item or a heading
                    x0, c_tokens, child = children[-1]
                    children[-1] = (x0, c_tokens + _on_new_line(block_tokens(nb, vocab)), child)
                    self.joined_blocks.setdefault(child.uid, []).append(nb)
                    j += 1
                    continue
                marked = bool(RE_LIST_MARKER.match(nb.text) or RE_NUMBERED_ITEM.match(nb.text)
                              or RE_MARKER_SHAPE.match(nb.text))
                if nb.kind == "list_item":
                    if children and not marked and nb.box[0] < children[0][0] - 0.1 * dpi:
                        break                    # "NOTE:  ..." back left of the items: after the list, not in it
                    children.extend(self._child_entries(nb, vocab))
                    j += 1
                    continue
                if (nb.kind == "paragraph" and nb.box[0] > b.box[0] + 0.15 * dpi and not self._starts_section(nb.text)
                        and not _starts_with_delimited_id(nb.text) and word_count(nb.text) <= 60
                        # without a colon the lead-in only takes what looks like an item
                        and (colon_stem or marked or _is_item_label(nb.text) or j <= i + plain_run)):
                    children.extend(self._child_entries(nb, vocab))   # indented items without visible bullets
                    j += 1
                    continue
                break
            # requirements stated before the stem in the same block (each under its own paragraph number)
            for u in [u for u in units if u[1] <= last[0]]:
                u_ref, u_title = head_of(u)
                items.append(self._req_item(b, source, section, u_ref, u_title, text, tokens, spans, u[0], u[1], u[2],
                                            "Paragraph", "", extra_reasons))
            stem_ref, stem_title = head_of(last)            # the list belongs to the stem's paragraph
            if not children:
                # a colon promises a list. Without one, only flag it when a short indented line follows that
                # was not taken as an item - a list that may have been missed. A sentence that just ends in
                # "... the following." followed by a new paragraph, or by a requirement of its own, is fine.
                nxt = blocks[i + 1] if i + 1 < len(blocks) else None
                missed = (nxt is not None and nxt.kind in ("paragraph", "list_item") and nxt.page == b.page
                          and nxt.box[0] > b.box[0] + 0.15 * dpi and not self._starts_section(nxt.text)
                          and all(word_count(t) <= 12 for _, t in _text_rows(nxt)) and not _modal_re().search(nxt.text))
                promised = (["Ends with ':' but no list items were found below"] if colon_stem else
                            ["Leads into a list ('the following ...') but no list items were taken from the "
                             "indented text below it - check"] if missed else [])
                items.append(self._req_item(b, source, section, stem_ref, stem_title, text, tokens, spans, last[0], last[1], last[2],
                                            "Paragraph", "", extra_reasons + promised))
                return 1
            self._list_items(b, children, source, section, stem_ref, stem_title, text, tokens, spans, last, vocab, items,
                             extra_reasons, dpi)
            for sb in skipped:                   # captions / figures that sat inside the list are exported too
                if sb.kind == "caption" and not sb.meta.get("duplicate_of_table"):
                    items.append(self._caption_item(sb, source, section))
                elif sb.kind == "figure":
                    items.extend(self._figure_items(sb, source, section))
            return j - i

        found_in = "List item" if b.kind == "list_item" else "Paragraph"
        for (s, e, kw), (u_ref, u_title) in zip(units, heads):
            items.append(self._req_item(b, source, section, u_ref, u_title, text, tokens, spans, s, e, kw, found_in, "", extra_reasons))
        return 1

    @staticmethod
    def _plain_items(rows: list, left: float, dpi: float) -> bool:
        """Do these (x0, text) lines read as a list without markers - two or more short lines (up to 12
        words) that each start with a capital or a number, state no requirement of their own, and are
        indented from the lead-in's margin `left` and level with each other?"""
        if len(rows) < 2:
            return False
        xs = [x0 for x0, _ in rows]
        return (all(0 < word_count(t) <= 12 and (t[:1].isupper() or t[:1].isdigit()) and not _modal_re().search(t)
                    for _, t in rows)
                and max(xs) - min(xs) <= 0.1 * dpi and min(xs) > left + 0.1 * dpi)

    def _row_entries(self, b: Block, ctx, text: str, spans: list, start_char: int, whole_rows: bool = False) -> list:
        """List items written on the rows after an inline stem inside the same block. With whole_rows,
        nothing is returned when the stem's own row goes on after it: those words would be lost."""
        entries, n_before = [], 0
        rows = _block_rows(b)
        for row_lines, row_tokens in zip(rows, block_token_rows(b, ctx)):
            first = n_before
            n_before += len(row_tokens)
            if not row_tokens or first >= len(spans):
                continue
            if spans[first][0] < start_char:
                if whole_rows and n_before <= len(spans) and spans[n_before - 1][1] > start_char:
                    return []                           # the lead-in's row goes on: it did not end the row
                continue
            entries.append((min(ln.box[0] for ln in row_lines), row_tokens, b))
        return entries

    def _child_entries(self, child: Block, ctx) -> list:
        """(x0, tokens, block) items of one child block. An unmarked block whose rows each start like a
        new item ('Search mode' / 'Track mode') is split into one item per row."""
        if not child.lines:
            return [(child.box[0], [(w, 1.0, "") for w in child.text.split()], child)]
        rows = _block_rows(child)
        token_rows = block_token_rows(child, ctx)
        all_tokens = [t for r in token_rows for t in r]
        text = tokens_text(all_tokens)[0]
        # ... but only when the rows start level with each other: rows that start further right hang
        # under the first row's text (a wrapped value after a wide tab), so the block is ONE item
        row_x0 = [min(ln.box[0] for ln in rl) for rl in rows]
        row_h = statistics.median([max(ln.box[3] for ln in rl) - min(ln.box[1] for ln in rl) for rl in rows])
        level = all(abs(x - row_x0[0]) <= 0.6 * row_h for x in row_x0)
        if (len(rows) > 1 and level and not (RE_LIST_MARKER.match(text) or RE_NUMBERED_ITEM.match(text))
                and all(r and word_count(" ".join(t[0] for t in r)) <= 10 and (r[0][0][:1].isupper() or r[0][0][:1].isdigit())
                        for r in token_rows)):
            return [(min(ln.box[0] for ln in rl), tr, child) for rl, tr in zip(rows, token_rows) if tr]
        return [(child.box[0], all_tokens, child)]

    @staticmethod
    def _marker_columns(children: list, tol: float) -> list:
        """x of each list item, where an item whose marker was lost ("Shock:  40 g ...", starting in
        the label column) is placed at the marker column of the neighbour whose label it lines up
        with - so it is counted as that neighbour's sibling, not as a sub-item."""
        info = []
        for x0, c_tokens, child in children:
            text = tokens_text(c_tokens)[0]
            marked = bool(RE_LIST_MARKER.match(text) or RE_NUMBERED_ITEM.match(text) or RE_MARKER_SHAPE.match(text))
            label_x = None
            if marked and child.kind == "list_item" and child.lines:
                rows = _block_rows(child)
                label_x = _item_columns(rows[0])[0] if rows else None
            info.append((marked, label_x))
        out = [x0 for x0, _, _ in children]
        texts = [tokens_text(c_tokens)[0] for _, c_tokens, _ in children]

        def skips_to_here(k: int) -> bool:
            """The marked items either side of block k, at one level, leave out exactly the markers of
            the unmarked blocks between them ("a. Search mode." / "Track mode." / "c. Ground mapping")."""
            prev_j = next((j for j in range(k - 1, -1, -1) if info[j][0]), None)
            next_j = next((j for j in range(k + 1, len(children)) if info[j][0]), None)
            if prev_j is None or next_j is None or abs(children[prev_j][0] - children[next_j][0]) > tol:
                return False
            shape = lambda t: (RE_LIST_MARKER.match(t) or RE_NUMBERED_ITEM.match(t) or RE_MARKER_SHAPE.match(t)).group("marker")
            first, last = shape(texts[prev_j]), shape(texts[next_j])
            roman = any((_marker_parts(mk) or ("", "", ""))[1].lower() in _ROMANS[1:9] for mk in (first, last))
            for _ in range(next_j - prev_j):
                first = _next_marker(first, roman)
            return bool(first) and first.lower() == last.lower()

        for k, (x0, c_tokens, _) in enumerate(children):
            text = texts[k]
            notice = text.split(":", 1)[0].strip().upper() in _NOT_LABELS
            if info[k][0] or notice or not (_is_item_label(text) or skips_to_here(k)):
                continue
            before = next((j for j in range(k - 1, -1, -1) if info[j][0]), None)
            if before is not None and _opens_sub_list(texts[before]):
                continue                        # 'a. Vibration:  as follows:' - the labels below are ITS sub-items
            m = RE_LIST_MARKER.match(texts[before]) or RE_MARKER_SHAPE.match(texts[before]) if before is not None else None
            if m and m.group("rest").strip().endswith(":") and _is_item_label(m.group("rest").strip()) \
                    and _is_item_label(text):
                continue                        # 'a. Operating:' with nothing after it - 'Temperature: ...' is ITS sub-item
            for j in (k - 1, k + 1, k - 2, k + 2):
                if 0 <= j < len(children) and info[j][0] and info[j][1] is not None and abs(x0 - info[j][1]) <= tol:
                    out[k] = children[j][0]
                    break
        return out

    @staticmethod
    def _split_lost_items(children: list, ctx, tol: float, joined: dict = None) -> list:
        """An item whose marker was lost in a single-spaced list starts exactly where a wrapped line would,
        so it arrives inside the item above:

            a.   Search mode.
                 Track mode.            <- 'b.' lost: read as a second line of 'a.'
            c.   Ground mapping mode.

        When the markers either side skip exactly one ('a.' then 'c.'), the item before the gap is split
        where ONE line in its text column starts a new sentence after a finished one; the sequence then
        gives the piece its letter (and says so). Without exactly one such line nothing is split, and the
        item is flagged instead (see _sequence_markers). A block that was joined on to the item after a
        finished sentence (see _continues_item), starting in its text column, is such a line too."""
        def marker(k):
            # (a capital the scan made of the next letter - "C." for "c." - still counts: see _lone_gap)
            t = tokens_text(children[k][1])[0]
            m = RE_LIST_MARKER.match(t) or RE_NUMBERED_ITEM.match(t) or RE_MARKER_SHAPE.match(t)
            return m.group("marker") if m and _marker_parts(m.group("marker")) else ""

        markers = [marker(k) for k in range(len(children))]
        out = []
        for k, (x0, c_tokens, child) in enumerate(children):
            out.append(children[k])
            here, after = markers[k], markers[k + 1] if k + 1 < len(children) else ""
            if not (here and after and child.lines and abs(children[k + 1][0] - x0) <= tol):
                continue
            if not RequirementExtractor._lone_gap(markers, [c[0] for c in children], k, tol):
                continue
            rows = _block_rows(child)
            token_rows = block_token_rows(child, ctx)
            own = [t[0] for row in token_rows for t in row]
            if not rows or [t[0] for t in c_tokens[:len(own)]] != own:
                continue                                 # (only the item's own text is split)
            label_x, value_xs = _item_columns(rows[0])
            in_text_column = lambda x: (label_x is not None and abs(x - label_x) <= tol
                                        and not any(abs(x - v) <= tol for v in value_xs))
            cuts = [sum(len(row) for row in token_rows[:r]) for r in range(1, len(rows))
                    if token_rows[r] and token_rows[r - 1] and in_text_column(min(ln.box[0] for ln in rows[r]))
                    and _ends_sentence(token_rows[r - 1][-1][0]) and token_rows[r][0][0][:1].isupper()]
            first_joined = ((joined or {}).get(child.uid) or [None])[0]
            if (first_joined is not None and len(c_tokens) > len(own) and own and _ends_sentence(own[-1])
                    and c_tokens[len(own)][0][:1].isupper() and in_text_column(first_joined.box[0])):
                cuts.append(len(own))                    # "Map. Ground mapping ..." joined on after "... 20 targets."
            if len(cuts) != 1:
                continue
            n = cuts[0]
            out[-1] = (x0, c_tokens[:n], child)
            out.append((x0, _on_new_line(c_tokens[n:]), child))
        return out

    @staticmethod
    def _lone_gap(markers: list, xs: list, k: int, tol: float) -> bool:
        """Do the markers of item k and the item after it skip exactly one ('a.' then 'c.'), in a list whose
        other steps at that level all count up by one? A list that counts in twos (tests 1, 3, 5) has no
        lost item, and 'a.', 'c.', 'c.', 'd.' is a misread 'b.' (see _sequence_markers), not a lost one."""
        level = [j for j in range(len(markers)) if markers[j] and abs(xs[j] - xs[k]) <= tol]
        seq = [markers[j] for j in level]
        roman = any((_marker_parts(mk) or ("", "", ""))[1].lower() in _ROMANS[1:9] for mk in seq)
        step = lambda a, n: _next_marker(a, roman) if n == 1 else _next_marker(_next_marker(a, roman), roman)
        if k not in level or level.index(k) + 1 >= len(level) or level[level.index(k) + 1] != k + 1:
            return False
        pos = level.index(k)
        if step(seq[pos], 2).lower() != seq[pos + 1].lower():
            return False
        if pos + 2 < len(seq) and seq[pos + 2].lower() == seq[pos + 1].lower():
            return False                                   # 'c.', 'c.': the first is a misread 'b.'
        others = [(a, b) for q, (a, b) in enumerate(zip(seq, seq[1:])) if q != pos]
        return all(step(a, 1).lower() == b.lower() for a, b in others)

    @staticmethod
    def _nest_by_style(children: list, columns: list, tol: float) -> list:
        """x of each list item, where a sub-list printed WITHOUT its own indent is moved one step right:

            b.  Vibration:
            (1) Random:  20 Hz to 2000 Hz ...
            (2) Sine:    5 Hz to 500 Hz ...
            c.  Shock:   40 g ...

        After an item that ends with ':' the marker style changes and a new sequence starts ((1), a., i.),
        and the parent's style comes back after it (or the list ends): those items are its sub-items. The
        parent's next letter read as a capital ('C.' after 'b.', a common misread) is the parent's style
        coming back, unless the sub-list itself is lettered in capitals (A., B., C.)."""
        def marker(k):
            t = tokens_text(children[k][1])[0]
            m = RE_LIST_MARKER.match(t) or RE_NUMBERED_ITEM.match(t) or RE_MARKER_SHAPE.match(t)
            return m.group("marker") if m and _marker_parts(m.group("marker")) else ""

        def style(mk):
            opening, core, closing = _marker_parts(mk)
            return opening, closing, "n" if core.isdigit() else "r" if len(core) > 1 else "U" if core.isupper() else "l"

        def parents_style(mk, parent, run):
            if style(mk) == style(parent):
                return True
            core = _marker_parts(mk)[1]
            capital_run = bool(run) and style(markers[run[0]])[2] == "U"
            return (len(core) == 1 and core.isupper() and style(parent)[2] == "l" and not capital_run
                    and mk.lower() == _next_marker(parent).lower())

        markers = [marker(k) for k in range(len(children))]
        out, k = list(columns), 0
        while k < len(children):
            if markers[k] and tokens_text(children[k][1])[0].rstrip().endswith(":"):
                run, j = [], k + 1
                while (j < len(children) and markers[j] and abs(columns[j] - columns[k]) <= tol
                       and not parents_style(markers[j], markers[k], run)):
                    run.append(j)
                    j += 1
                resumes = j == len(children) or (markers[j] and abs(columns[j] - columns[k]) <= tol
                                                  and parents_style(markers[j], markers[k], run))
                if run and not _previous_marker(markers[run[0]]) and resumes:
                    for r in run:
                        out[r] = columns[k] + 3 * tol
                    k = j
                    continue
            k += 1
        return out

    @staticmethod
    def _sequence_markers(children: list, tol: float) -> list:
        """[(marker, body, note, merge_into)] per list item, with each marker checked against its neighbours.

        Markers count up (a, b, c / 1, 2, 3 / i, ii, iii / (a), (b)), so an item's marker can be
        expected from the item before it. A marker that the scan misread ("C." for "c.", "e." for "c.")
        or lost is replaced by the expected one when the neighbours agree, and the change is reported
        for review. A list that really skips a letter keeps it: a valid marker is only replaced when the
        NEXT item's marker confirms the expected one. A block without a marker is never given a letter
        that a later item of the same list really carries. If it is only a piece of text ("Applied for
        1 minute per axis.") it is more of the item above at its level: merge_into is that item's index.
        A block with a requirement of its own, one that leads into a list, and a notice ('NOTE:') keep
        their own rows, unnumbered (merge_into is None)."""
        found = []
        for x0, c_tokens, _child in children:
            c_text = tokens_text(c_tokens)[0]
            m = RE_LIST_MARKER.match(c_text) or RE_NUMBERED_ITEM.match(c_text)
            shape = None if m else RE_MARKER_SHAPE.match(c_text)
            if m:
                found.append([x0, m.group("marker"), m.group("rest"), True])
            elif shape:
                found.append([x0, shape.group("marker"), shape.group("rest"), False])   # marker-like, not valid here
            else:
                found.append([x0, "", c_text, False])

        def sibling(k: int, step: int):
            """Index of the previous/next item at the same level (stops at a parent item)."""
            j = k + step
            while 0 <= j < len(found):
                if abs(found[j][0] - found[k][0]) <= tol:
                    return j
                if found[j][0] < found[k][0] - tol:
                    return None
                j += step
            return None

        cores = [(_marker_parts(f[1]) or ("", "", ""))[1].lower() for f in found]
        roman = any(c in _ROMANS[1:9] for c in cores) or (bool(cores) and cores[0] == "i")
        def capitals(x0: float) -> bool:
            """Is the list level at x0 lettered in capitals ("A.  Flow rate:", "B.  Pressure:")? Then its capitals
            are not misreads. (A sub-list below it may use small letters: each level is judged on its own.)"""
            letters = [(_marker_parts(f[1]) or ("", "", ""))[1] for f in found if abs(f[0] - x0) <= tol]
            letters = [c for c in letters if len(c) == 1 and c.isalpha()]
            return len(letters) >= 2 and all(c.isupper() for c in letters)
        def notice(text: str) -> bool:
            return (text or "").strip().split(":", 1)[0].strip().upper() in _NOT_LABELS and ":" in (text or "")

        out = []
        for k, (x0, marker, body, valid) in enumerate(found):
            note = ""
            if not marker and notice(body):
                out.append((marker, body, note, None))     # "NOTE: ..." keeps its own row, never numbered
                continue
            before, after = sibling(k, -1), sibling(k, 1)
            while before is not None and not found[before][1]:
                before = sibling(before, -1)               # expected from the last item that has a marker
            if before is not None:
                expected = _next_marker(found[before][1], roman)
                confirmed = after is not None and found[after][1].lower() == _next_marker(expected, roman).lower()
            else:
                # the first item(s): counted back from the first item after them that has a marker
                # ("Acceleration: ...", "Vibration: ...", "c. Shock: ..." -> a., b.). A single letter the scan
                # read as a capital ("C.") counts as its small letter, as it does everywhere in a list.
                j, back = after, 1
                while j is not None and not found[j][1]:
                    j, back = sibling(j, 1), back + 1
                expected = found[j][1] if j is not None else ""
                if j is not None and not found[j][3]:
                    parts = _marker_parts(expected)
                    expected = (expected if capitals(x0) else expected.lower()) \
                        if parts and len(parts[1]) == 1 and parts[1].isalpha() else ""
                for _ in range(back):
                    expected = _previous_marker(expected, roman)
                confirmed = bool(expected) and not marker
            if not marker and expected and before is not None:
                later = []
                for j in range(k + 1, len(found)):
                    if found[j][3] and abs(found[j][0] - x0) <= tol:
                        if not _previous_marker(found[j][1], roman):
                            break                       # "a." again: a second list, not more of this one
                        later.append(found[j][1])
                if expected in later:
                    # the real 'b.' comes later: this block is not 'b.'. A piece of text is more of the item
                    # above; a requirement of its own or a new lead-in keeps its own row, unnumbered.
                    own = bool(_modal_re().search(body)) or body.rstrip().endswith(":") or _introduces_list(body)
                    target = sibling(k, -1)
                    merge = None if own or target is None or tokens_text(children[target][1])[0].rstrip().endswith(":") \
                        else target
                    out.append((marker, body, note, merge))
                    continue
                else:
                    # several markers lost in a row: the next marked item, counted forward, confirms them
                    step, j = expected, after
                    while j is not None and not found[j][1]:
                        step, j = _next_marker(step, roman), sibling(j, 1)
                    confirmed = confirmed or (j is not None and found[j][1].lower() == _next_marker(step, roman).lower())
            if expected and marker != expected:
                case_only = marker.lower() == expected.lower()
                label = _is_item_label(body)
                if (not valid and marker and (case_only or confirmed)) or (valid and confirmed) \
                        or (not marker and (confirmed or (label and after is None))):
                    note = (f"List marker read as '{marker}' - taken as '{expected}' from the sequence" if marker
                            else f"List marker missing - taken as '{expected}' from the sequence")
                    marker = expected
                    found[k][1] = marker                       # the next item is checked against this one
            out.append((marker, body, note, None))
        # 'a.' followed straight by 'c.' at the same level: 'b.' may have lost its marker and been read as
        # more of 'a.' - the list is left as printed, but the item before the gap is flagged
        markers_now = [f[1] for f in found]           # (the markers as corrected above)
        for k, (x0, marker, _body, _valid) in enumerate(found):
            nxt = sibling(k, 1)
            if (marker and nxt == k + 1 and found[nxt][1] and not out[k][2] and out[k][3] is None
                    and RequirementExtractor._lone_gap(markers_now, [f[0] for f in found], k, tol)):
                out[k] = (out[k][0], out[k][1], f"The list goes from '{marker}' to '{found[nxt][1]}' - check whether "
                          f"'{_next_marker(marker, roman)}' lost its marker and was read as part of this item", None)
        return out

    def _list_items(self, stem_block, children, source, section, ref, title, text, tokens, spans, stem_unit, vocab, items, extra, dpi=200.0):
        """children = [(x0, tokens, block)] in reading order."""
        stem_s, stem_e, stem_kw = stem_unit
        tol = 0.08 * dpi                     # items within ~0.08 inch of each other are at the same level
        stem_text, stem_notes = self.repair_keywords(text[stem_s:stem_e].strip())
        stem_tokens = [tok for tok, (a, z) in zip(tokens, spans) if a < stem_e and z > stem_s]
        stack = []          # nested stems: (x0, marker, text, tokens)
        rows = []           # (block, marker_path, tokens)
        columns = self._nest_by_style(children, self._marker_columns(children, tol), tol)
        children = [(columns[k], c_tokens, child) for k, (_x0, c_tokens, child) in enumerate(children)]
        children = self._split_lost_items(children, vocab, tol, self.joined_blocks)
        # a block the sequence says is more of an item above it (its would-be letter belongs to a later
        # item) is joined to that item - the one at its own level - before the rows are built
        merged, sequenced, where = [], [], {}
        for k, (entry, seq) in enumerate(zip(children, self._sequence_markers(children, tol))):
            if seq[3] is not None and seq[3] in where:
                t = where[seq[3]]
                x0, c_tokens, child = merged[t]
                merged[t] = (x0, c_tokens + _on_new_line(entry[1]), child)
                if entry[2] is not child:
                    self.joined_blocks.setdefault(child.uid, []).append(entry[2])
                    # (with the blocks that were already joined to it: a value carried onto the next page)
                    self.joined_blocks[child.uid].extend(self.joined_blocks.pop(entry[2].uid, []))
                mk, body, note, _ = sequenced[t]
                sequenced[t] = (mk, f"{body} {tokens_text(entry[1])[0]}".strip(), note, None)
                where[k] = t
                continue
            where[k] = len(merged)
            merged.append(entry)
            sequenced.append(seq)
        children = merged
        # "a.  Search mode:   (1) Range of 100 km." with "(2) Azimuth ..." under it on the value tab stop: the
        # first sub-item sits on its parent's line. It is split off, so both sub-items are filed under 'a.'
        split, split_seq = [], []
        for k, (x0, c_tokens, child) in enumerate(children):
            nxt = children[k + 1] if k + 1 < len(children) else None
            cut = self._inline_sub_item(c_tokens, nxt, x0, tol)
            if cut:
                mk, body, note, _ = sequenced[k]
                head, sub = c_tokens[:cut], c_tokens[cut:]
                sub_text = tokens_text(sub)[0]
                sub_marker = sub_text.split()[0]
                split += [(x0, head, child), (nxt[0], sub, child)]
                split_seq += [(mk, body[:len(body) - len(sub_text)].strip(), note, None),
                              (sub_marker, sub_text[len(sub_marker):].strip(), "", None)]
                continue
            split.append((x0, c_tokens, child))
            split_seq.append(sequenced[k])
        children, sequenced = split, split_seq
        marker_notes = {}   # id(row tokens) -> notes about markers taken from the sequence (the item's and its parents')
        parents_of = {}     # id(row tokens) -> the blocks of the parent items in its path ("a. Operating:")
        for k, (x0, c_tokens, child) in enumerate(children):
            c_marker, c_body, marker_note, _ = sequenced[k]
            while stack and x0 <= stack[-1][0] + tol:
                stack.pop()
            nxt = children[k + 1] if k + 1 < len(children) else None
            if c_body.rstrip().endswith(":") and nxt is not None and nxt[0] > x0 + tol:
                stack.append((x0, c_marker, c_body.strip(), c_tokens, marker_note, child))
                continue
            path = [(entry[1], entry[2]) for entry in stack] + [(c_marker, c_body.strip())]
            path_tokens = [t for entry in stack for t in entry[3]] + c_tokens
            rows.append((child, path, path_tokens))
            parents_of[id(path_tokens)] = [entry[5] for entry in stack if entry[5] is not child]
            notes_here = [entry[4] for entry in stack if entry[4]] + ([marker_note] if marker_note else [])
            if notes_here:
                marker_notes[id(path_tokens)] = notes_here
        if not self.settings.split_list_items:
            joined = stem_text + " " + " ".join(" ".join(f"{mk} {bd}".strip() for mk, bd in path) for _, path, _ in rows)
            all_tokens = stem_tokens + [t for _, _, tk in rows for t in tk]
            all_notes = list(dict.fromkeys(n for notes in marker_notes.values() for n in notes))
            conf, reasons = self.review_for_tokens(all_tokens, joined, extra + stem_notes + all_notes)
            blocks_used = [stem_block] + [blk for r in rows for blk in parents_of.get(id(r[2]), []) + [r[0]]
                                          + self.joined_blocks.get(r[0].uid, [])]
            pages = sorted({p for blk in blocks_used for p in blk.pages})
            items.append(ExtractedItem("Requirement", source, _page_label(pages), section, doc_ref=ref, title=title,
                                       text=joined.strip(), keyword=stem_kw, req_class=self.classes.get(stem_kw, ""),
                                       found_in="List (stem + all items)", ocr_conf=conf, review_reasons=reasons,
                                       quality_notes=self.quality_notes(joined),
                                       block_uids=list(dict.fromkeys(blk.uid for blk in blocks_used))))
            return
        for child, path, path_tokens in rows:
            item_text = " ".join(f"{mk} {bd}".strip() for mk, bd in path)
            full, notes = self.repair_keywords(f"{stem_text} {item_text}")
            marker_note = marker_notes.get(id(path_tokens), [])
            conf, reasons = self.review_for_tokens(stem_tokens + path_tokens, full, extra + stem_notes + notes + marker_note)
            joined = parents_of.get(id(path_tokens), []) + self.joined_blocks.get(child.uid, [])
            pages = sorted(set(stem_block.pages) | set(child.pages) | {pg for blk in joined for pg in blk.pages})
            marker_ref = ".".join(mk.strip("().") for mk, _ in path if mk and mk not in "•·*-o◦")
            items.append(ExtractedItem("Requirement", source, _page_label(pages), section,
                                       doc_ref=(f"{ref} {marker_ref}".strip() if marker_ref else ref), title=title,
                                       text=full, keyword=stem_kw, req_class=self.classes.get(stem_kw, ""),
                                       found_in="List item under stem", context=stem_text, ocr_conf=conf,
                                       review_reasons=reasons, quality_notes=self.quality_notes(full),
                                       block_uids=list(dict.fromkeys([stem_block.uid, child.uid] + [blk.uid for blk in joined]))))

    @staticmethod
    def _inline_sub_item(c_tokens: list, nxt, x0: float, tol: float) -> int:
        """Index of the token where a first sub-item starts on its parent's line - "a. Search mode:  (1) Range
        of 100 km." - when the next item, further right, carries the next sub-marker ("(2) Azimuth ...").
        0 when there is none."""
        if nxt is None or nxt[0] <= x0 + tol:
            return 0
        nxt_m = RE_LIST_MARKER.match(tokens_text(nxt[1])[0])
        if not nxt_m or not _marker_parts(nxt_m.group("marker")):
            return 0
        wanted = _previous_marker(nxt_m.group("marker"), nxt_m.group("marker").strip("().") in _ROMANS[1:9])
        head = []
        for k, tok in enumerate(c_tokens):
            if k >= 2 and tok[0] == wanted and head[-1].endswith(":"):
                label = " ".join(head)
                m = RE_LIST_MARKER.match(label) or RE_MARKER_SHAPE.match(label)
                return k if _is_item_label((m.group("rest") if m else label).strip()) else 0
            head.append(tok[0])
        return 0

    def _req_item(self, b, source, section, ref, title, text, tokens, spans, s, e, kw, found_in, context, extra) -> ExtractedItem:
        sentence = text[s:e].strip()
        trailing = RE_TRAILING_ID.search(sentence)
        if trailing and ("-" in trailing.group("id") or "_" in trailing.group("id")) \
                and getattr(self, "id_prefixes", Counter()).get(_id_prefix(trailing.group("id")), 0) >= 2:
            trailing = trailing                       # a known requirement-ID prefix
        else:
            trailing = None                           # "(J3)" / "(CB12)" are designators: keep them in the text
        if trailing and not ref:                      # "... within 2 seconds (SRS-103)."
            ref = trailing.group("id")
            ended = sentence.endswith(".")
            sentence = sentence[:trailing.start()].rstrip()
            if ended and not sentence.endswith("."):
                sentence += "."
        fixed, notes = self.repair_keywords(sentence)
        unit_tokens = [tok for tok, (a, z) in zip(tokens, spans) if a < e and z > s]
        reasons = list(extra) + notes
        if found_in == "Paragraph" and not RE_TERMINAL.search(fixed):
            reasons.append("Sentence has no final punctuation - it may be incomplete")
        conf, reasons = self.review_for_tokens(unit_tokens, fixed, reasons)
        return ExtractedItem("Requirement", source, b.page_label(), section, doc_ref=ref, title=title, text=fixed,
                             keyword=kw, req_class=self.classes.get(kw, ""), found_in=found_in, context=context,
                             ocr_conf=conf, review_reasons=reasons, quality_notes=self.quality_notes(fixed),
                             block_uids=[b.uid])


def _cell_words(t: TableData, r: int, c: int) -> list:
    try:
        return t.cell_words[r][c]
    except (IndexError, TypeError):
        return []


def _page_label(pages: list) -> str:
    pages = sorted(set(pages))
    return str(pages[0]) if len(pages) == 1 else f"{pages[0]}-{pages[-1]}"


# =============================================================================
# 10. EXCEL EXPORT
#     Sheet "Extracted Items": one row per requirement, caption and table.
#     One sheet per table (T001, T002, ...), plus "Full Text", "Page Quality", "Run Info".
# =============================================================================
MASTER_COLUMNS = [  # (header, width, description used in the Run Info sheet)
    ("Item #", 10, "Unique ID: R- requirement, T- table, TC- table caption, FC- figure caption"),
    ("Type", 14, "Requirement / Table / Table Caption / Figure Caption"),
    ("Source File", 22, "Input document"),
    ("Page(s)", 8, "Page number(s) in the source document"),
    ("Section", 28, "Nearest heading above the item"),
    ("Doc Ref", 13, "Paragraph number, requirement ID or table/figure number printed in the document"),
    ("Title", 20, "Paragraph title (e.g. 'Detection range') or table caption"),
    ("Content", 72, "The complete requirement text, caption text, or the table as text"),
    ("Keyword", 10, "Requirement keyword found (shall, must, should, will ...)"),
    ("Class", 18, "Mandatory / Goal / Statement of intent (from the keyword)"),
    ("Found In", 20, "Paragraph, list item under a stem, table row ..."),
    ("Parent / Context", 36, "List stem or table caption the requirement belongs to"),
    ("Other Columns", 30, "Other columns of a requirement table row (e.g. parameter, priority)"),
    ("Verification", 16, "Verification method: from the requirement table, or looked up in the document's verification cross-reference matrix (I/A/D/T)"),
    ("References", 26, "Standards, tables, figures and sections cited by the requirement (e.g. MIL-STD-704F; TABLE II)"),
    ("OCR Confidence", 11, "Lowest OCR confidence of the words in this item (0-1). A high value does NOT prove the text is right"),
    ("Needs Review", 9, "Yes when something below deserves a human check"),
    ("Review Reasons", 42, "Why the item needs a check"),
    ("Quality Notes", 36, "INCOSE-style writing checks: TBD, vague terms, escape clauses, and/or ..."),
    ("Table Sheet", 11, "Link to the worksheet holding the full table"),
]
_COL = {name: i for i, (name, _w, _d) in enumerate(MASTER_COLUMNS, 1)}
_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEADER_FILL = PatternFill("solid", fgColor="1F3864")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_REVIEW_FILL = PatternFill("solid", fgColor="FCE4D6")
_LOWCONF_FILL = PatternFill("solid", fgColor="FFF2CC")
_WRAP_TOP = Alignment(wrap_text=True, vertical="top")


def _xl(value):
    """Make a value safe for Excel: strip characters Excel rejects, respect the 32,767
    character cell limit."""
    if value is None or isinstance(value, (int, float)):
        return value
    text = ILLEGAL_CHARACTERS_RE.sub("", str(value))
    if len(text) > 32000:
        text = text[:32000] + " ...[truncated - see the table sheet / Word report]"
    return text


def _put(ws, row: int, col: int, value, bold=False, fill=None, wrap=True, number_format=None):
    cell = ws.cell(row=row, column=col)
    v = _xl(value)
    cell.value = v
    if isinstance(v, str) and v.startswith("="):
        cell.data_type = "s"            # text such as "=5 V" must never become a formula
    if bold:
        cell.font = Font(bold=True)
    if fill is not None:
        cell.fill = fill
    if wrap:
        cell.alignment = _WRAP_TOP
    if number_format:
        cell.number_format = number_format
    return cell


def _safe_sheet_title(title: str, existing: set) -> str:
    title = re.sub(r"[\[\]:*?/\\']", "_", title)[:31] or "Sheet"
    base, n = title, 2
    while title.lower() in {e.lower() for e in existing}:
        suffix = f"_{n}"
        title = base[:31 - len(suffix)] + suffix
        n += 1
    return title


def write_workbook(path: Path, runs: list, settings: Settings, multi_doc: bool = False) -> Path:
    """runs = [(DocResult, [ExtractedItem])]. Returns the path actually written."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Extracted Items"
    for c, (name, width, _desc) in enumerate(MASTER_COLUMNS, 1):
        cell = _put(ws, 1, c, name, wrap=True)
        cell.font, cell.fill, cell.border = _HEADER_FONT, _HEADER_FILL, _BORDER
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.freeze_panes = "B2"
    ws.row_dimensions[1].height = 30

    sheet_names = {"Extracted Items", "Full Text", "Page Quality", "Run Info"}
    table_sheets = []           # (sheet title, item, doc)
    row = 2
    for d_index, (doc, items) in enumerate(runs, 1):
        for it in items:
            if it.item_type == "Table" and it.table is not None:
                name = _safe_sheet_title((f"D{d_index}-" if multi_doc else "") + it.item_id.replace("-", ""), sheet_names)
                sheet_names.add(name)
                it.table_sheet = name
                table_sheets.append((name, it, doc))
            values = [it.item_id if not multi_doc else f"D{d_index}-{it.item_id}", it.item_type, it.source_file, it.pages,
                      it.section, it.doc_ref, it.title, it.text, it.keyword, it.req_class, it.found_in, it.context,
                      it.attributes, it.verification, it.references, round(it.ocr_conf, 3) if it.ocr_conf is not None else None,
                      "Yes" if it.needs_review else "No", "\n".join(it.review_reasons), "\n".join(it.quality_notes),
                      it.table_sheet or None]
            for c, v in enumerate(values, 1):
                cell = _put(ws, row, c, v, number_format="0.00" if c == _COL["OCR Confidence"] else None)
                cell.border = _BORDER
            if it.table_sheet:
                link = ws.cell(row=row, column=len(MASTER_COLUMNS))
                link.hyperlink = Hyperlink(ref=link.coordinate, location=f"'{it.table_sheet}'!A1", display=it.table_sheet)
                link.style = "Hyperlink"
                link.alignment = _WRAP_TOP
            row += 1
    if row == 2:
        _put(ws, 2, 1, "-")
        _put(ws, 2, 8, "No requirements, tables or captions were found. Check the Full Text sheet and the Word report.")
        row = 3
    last = row - 1
    ref = f"A1:{get_column_letter(len(MASTER_COLUMNS))}{last}"
    xl_table = XlTable(displayName="ExtractedItems", ref=ref)
    xl_table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(xl_table)
    review_col = get_column_letter(_COL["Needs Review"])
    ws.conditional_formatting.add(f"A2:{get_column_letter(len(MASTER_COLUMNS))}{last}",
                                  FormulaRule(formula=[f'${review_col}2="Yes"'], fill=_REVIEW_FILL))

    # ---- one sheet per table ----------------------------------------------------------
    for name, it, doc in table_sheets:
        ts = wb.create_sheet(name)
        t = it.table
        ctx = document_text_context(doc, settings)
        _put(ts, 1, 1, it.title or it.doc_ref or "Table (no caption found)", bold=True, wrap=False).font = Font(bold=True, size=12)
        _put(ts, 2, 1, f"{it.item_id}  |  {it.source_file}  |  page(s) {it.pages}  |  {t.n_rows} rows x {t.n_cols} columns", wrap=False)
        back = _put(ts, 3, 1, "<< Back to Extracted Items", wrap=False)
        back.hyperlink = Hyperlink(ref=back.coordinate, location="'Extracted Items'!A1", display="Extracted Items")
        back.style = "Hyperlink"
        top = 5
        widths = [8] * t.n_cols
        merged_hidden = set()
        for r0, c0, r1, c1 in t.merges:
            for rr in range(r0, r1 + 1):
                for cc in range(c0, c1 + 1):
                    if (rr, cc) != (r0, c0):
                        merged_hidden.add((rr, cc))
        for r, grid_row in enumerate(t.grid):
            for c in range(t.n_cols):
                value = grid_row[c] if c < len(grid_row) else ""
                cell = _put(ts, top + r, c + 1, None if (r, c) in merged_hidden else (value or None))
                cell.border = _BORDER
                if r < t.header_rows:
                    cell.font, cell.fill = _HEADER_FONT, _HEADER_FILL
                elif _cell_attention(t, r, c, ctx, settings):
                    cell.fill = _LOWCONF_FILL
                widths[c] = max(widths[c], min(60, max((len(x) for x in str(value).split("\n")), default=0) + 2))
        for r0, c0, r1, c1 in t.merges:
            if r1 < t.n_rows and c1 < t.n_cols and (r1 > r0 or c1 > c0):
                ts.merge_cells(start_row=top + r0, start_column=c0 + 1, end_row=top + r1, end_column=c1 + 1)
        for c, wdt in enumerate(widths, 1):
            ts.column_dimensions[get_column_letter(c)].width = wdt
        if any(_cell_attention(t, r, c, ctx, settings) for r in range(t.header_rows, t.n_rows) for c in range(t.n_cols)):
            _put(ts, top + t.n_rows + 1, 1, "Yellow cells: check against the scan (suspicious word or low-confidence number).", wrap=False)

    # ---- Full Text -----------------------------------------------------------------------
    ft = wb.create_sheet("Full Text", 1)
    headers = [("Source File", 22), ("Page", 7), ("Block #", 8), ("Type", 12), ("Text", 90), ("Mean OCR Conf.", 10),
               ("Min OCR Conf.", 10), ("Words To Check", 36), ("Became Item(s)", 16)]
    for c, (h, wdt) in enumerate(headers, 1):
        cell = _put(ft, 1, c, h)
        cell.font, cell.fill = _HEADER_FONT, _HEADER_FILL
        ft.column_dimensions[get_column_letter(c)].width = wdt
    ft.freeze_panes = "A2"
    r = 2
    for d_index, (doc, items) in enumerate(runs, 1):
        by_block = defaultdict(list)
        for it in items:
            for uid in it.block_uids:
                by_block[uid].append(it.item_id if not multi_doc else f"D{d_index}-{it.item_id}")
        for b in doc.blocks:
            if b.kind == "table" and b.table is not None:
                text = b.table.as_text()
            elif b.kind == "figure":
                text = "[figure] " + " / ".join(ln.text for ln in b.lines)
            else:
                text = b.text
            toks = word_tokens((w for ln in b.lines for w in ln.words), document_text_context(doc, settings))
            low = sorted({tok[0] for tok in toks if token_attention(tok, settings)})
            vals = [Path(doc.source).name, b.page_label(), b.uid, b.kind.replace("_", " "), text, round(b.conf, 3),
                    round(b.min_conf, 3), ", ".join(low[:15]), ", ".join(by_block.get(b.uid, []))]
            for c, v in enumerate(vals, 1):
                _put(ft, r, c, v, fill=_LOWCONF_FILL if (c == 5 and low) else None)
            r += 1
    ft.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(r - 1, 1)}"

    # ---- Page Quality ------------------------------------------------------------------------
    pq = wb.create_sheet("Page Quality", 2)
    headers = [("Source File", 22), ("Page", 7), ("Scan dpi", 9), ("OCR dpi", 9), ("Turned (deg)", 10), ("Skew fixed (deg)", 11),
               ("Text Lines", 9), ("Mean OCR Conf.", 11), ("Lines Needing Review", 12), ("Quality", 10), ("Notes", 60)]
    for c, (h, wdt) in enumerate(headers, 1):
        cell = _put(pq, 1, c, h)
        cell.font, cell.fill = _HEADER_FONT, _HEADER_FILL
        pq.column_dimensions[get_column_letter(c)].width = wdt
    r = 2
    for doc, _items in runs:
        for p in doc.pages:
            share_low = p.low_conf_lines / max(p.n_lines, 1)
            grade = "Good" if p.mean_conf >= 0.95 and share_low <= 0.05 else ("Fair" if p.mean_conf >= 0.88 and share_low <= 0.2 else "Poor")
            failed = any(str(n).startswith("FAILED") for n in p.notes)
            if failed:
                grade = "FAILED"                  # the page could not be read: check it by hand
            elif p.n_lines == 0:
                grade = "No text"
            vals = [Path(doc.source).name, p.number, p.source_dpi or None, p.dpi or None, p.rotation, p.skew, p.n_lines,
                    None if failed else p.mean_conf, p.low_conf_lines, grade, "; ".join(p.notes)]
            for c, v in enumerate(vals, 1):
                _put(pq, r, c, v, fill=_REVIEW_FILL if (c == 10 and grade in ("Poor", "No text", "FAILED")) else None)
            r += 1

    # ---- Run Info -------------------------------------------------------------------------
    ri = wb.create_sheet("Run Info")
    ri.column_dimensions["A"].width = 30
    ri.column_dimensions["B"].width = 110
    n_req = sum(1 for _, items in runs for it in items if it.item_type == "Requirement")
    software, differences = configuration_status(docling_used=any(d.engine == "docling" for d, _ in runs))
    info = [
        ("Created", _dt.datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Tool", f"spec_extractor.py {__version__}   (build {build_id()})"),
        ("Tested configuration", "Yes - Python and package versions match constraints.txt" if not differences
         else "NO - the documented accuracy was not measured for this setup:\n" + "\n".join(differences)),
        ("Software", software.replace("; ", "\n")),
        ("Engine", runs[0][0].engine if runs else settings.engine),
        ("Documents", "\n".join(f"{Path(d.source).name}  ({len(d.pages)} pages, {d.seconds:.0f} s)" for d, _ in runs)),
        ("Requirements", n_req),
        ("Tables", sum(1 for _, items in runs for it in items if it.item_type == "Table")),
        ("Captions", sum(1 for _, items in runs for it in items if "Caption" in it.item_type)),
        ("Items needing review", sum(1 for _, items in runs for it in items if it.needs_review)),
        ("Requirement keywords", ", ".join(f"{k} = {v}" for k, v in settings.keywords.items())),
        ("List handling", "one row per list item (stem + item)" if settings.split_list_items else "one row per stem with all items"),
        ("Review threshold", f"OCR confidence below {settings.review_confidence:.2f}"),
        ("Word report look", _look_summary(settings)),
        ("Warnings", "\n".join(w for d, _ in runs for w in d.warnings) or "none"),
        ("", ""),
        ("HOW TO USE THIS WORKBOOK", "1) Filter 'Needs Review' = Yes on 'Extracted Items' and compare those rows with the scan "
                                     "(page numbers are given).\n2) Use the Word verification report to check whole pages quickly.\n"
                                     "3) 'Full Text' lists every text block the OCR found, in reading order, and which item it became.\n"
                                     "4) 'Page Quality' shows pages with poor scan quality."),
    ] + [(f"Column: {name}", desc) for name, _w, desc in MASTER_COLUMNS]
    for r, (k, v) in enumerate(info, 1):
        _put(ri, r, 1, k, bold=True)
        _put(ri, r, 2, v)
    return _save_workbook(wb, path)


def _look_summary(settings: Settings) -> str:
    if not settings.make_word_report:
        return "no Word report"
    try:
        return report_look(settings).summary
    except ReportLookError as exc:
        return "the style file or template cannot be used: " + str(exc).splitlines()[0]


def _save_workbook(wb, path: Path) -> Path:
    try:
        wb.save(path)
        return path
    except PermissionError:
        alt = path.with_name(f"{path.stem}_{_dt.datetime.now():%H%M%S}{path.suffix}")
        LOG.warning("%s is open in another program - saved as %s instead", path.name, alt.name)
        wb.save(alt)
        return alt


# =============================================================================
# 11. WORD VERIFICATION REPORT
#     Page by page: the scanned page image, then everything the tool read, with
#     low-confidence words highlighted and [R-0001] tags linking to Excel rows.
#
#     How the report LOOKS is set in word_style/report_style.ini (fonts, sizes,
#     spacing, page layout) and, optionally, by a Word template of your own
#     (.dotx/.docx) whose styles, page layout, headers and footers are copied.
#       11a  the parts of the report, their Word styles and the tool's own look
#       11b  reading the style file
#       11c  reading a Word template safely: only its styles, theme, list
#            numbering, page layout and headers/footers are used - nothing in
#            it may point to another file or to the network
#       11d  making the report document, and checking it before it is saved
#       11e  writing the report
#       11f  the starter template (make_word_template.bat)
# =============================================================================

# ---- 11a. The parts of the report --------------------------------------------------------------------------------
# Every part is written in a Word STYLE, not with formatting typed onto the text, so a template (or Word itself,
# afterwards) decides the look. The exceptions: the review marks - the highlighted words and shaded table cells that
# are the point of the report - are put directly on the words, so no template can hide them; a few labels ("Source:",
# "Text inside figure:"); the bold and italic of tags and keywords (see _styled_run); picture paragraphs; the table
# width; and (without a template) the heading-row shading. A table alignment set in the style file goes on each table;
# the tool's own centring is part of the Table Grid style.
PARAGRAPH_KEYS = ("word_style", "font", "size", "bold", "italic", "color", "alignment", "space_before",
                  "space_after", "line_spacing", "indent_left", "first_line_indent")
CHARACTER_KEYS = ("word_style", "font", "size", "bold", "italic", "color")
INHERIT = "inherit"     # (tool's own look only) remove the style's own value, so it follows the style it is based on


@dataclass(frozen=True)
class ReportPart:
    section: str            # its [section] in the style file
    kind: str               # "paragraph", "character" or "table" style
    style: str              # the Word style it uses unless the style file names another one
    based_on: str           # the part whose style a style the tool has to add is based on ("" = Normal)
    look: tuple             # the tool's own look: ((key, value), ...) - applied when there is no template,
                            # and to a style the tool has to add to a template that does not have it
    about: str              # what it is (for the starter template)
    extra_keys: tuple = ()  # settings of this section that are not style settings


REPORT_PARTS = (
    ReportPart("title", "paragraph", "Title", "", (), "the report title at the top of the first page"),
    ReportPart("heading 1", "paragraph", "Heading 1", "", (), "\"Page 12\" at the start of each scanned page"),
    ReportPart("heading 2", "paragraph", "Heading 2", "", (),
               "\"Page 12 - text read by the tool\", and the scanned document's main headings"),
    ReportPart("heading 3", "paragraph", "Heading 3", "", (), "the scanned document's second-level headings"),
    ReportPart("heading 4", "paragraph", "Heading 4", "", (), "third-level headings and below"),
    ReportPart("body", "paragraph", "Normal", "", (("font", "Calibri"), ("size", 10.0)),
               "ordinary paragraphs read from the scan"),
    ReportPart("list items", "paragraph", "List Paragraph", "",
               (("indent_left", 320040), ("first_line_indent", 0), ("contextual_spacing", False)),   # 0.35 in
               "list entries read from the scan"),
    ReportPart("requirements", "paragraph", "Requirement Text", "body", (),
               "paragraphs that became a requirement row in the Excel file", ("tag_position", "mark_keywords")),
    ReportPart("requirement tag", "character", "Requirement Tag", "",
               (("size", 8.0), ("bold", True), ("italic", False), ("color", "1F4E9A")),
               "the [R-0001] tags that link text to Excel rows"),
    ReportPart("requirement keywords", "character", "Requirement Keyword", "", (("bold", True),),
               "shall, must, should ... (only with mark_keywords = yes)"),
    ReportPart("captions", "paragraph", "Caption", "",
               (("italic", True), ("bold", False), ("color", "auto"), ("size", INHERIT), ("line_spacing", INHERIT),
                ("alignment", "center")), "table and figure titles"),
    ReportPart("tables", "table", "Table Grid", "", (("alignment", "center"),), "tables read from the scan",
               ("alignment",)),
    ReportPart("table text", "paragraph", "Table Text", "",
               (("size", 8.5), ("space_before", 0.0), ("space_after", 0.0), ("line_spacing", ("multiple", 1.0)),
                ("alignment", "left"), ("indent_left", 0), ("first_line_indent", 0)),
               "the text in table cells, and words read inside figures"),
    ReportPart("table header row", "paragraph", "Table Heading", "table text", (("bold", True),),
               "the column headings of each table", ("fill", "repeat_on_each_page")),
    ReportPart("scanned header footer", "paragraph", "Scanned Header Footer", "",
               (("size", 8.0), ("italic", True), ("color", "808080")),
               "running headers and footers printed on the scanned pages"),
    ReportPart("notes", "paragraph", "Report Note", "", (("size", 8.5), ("italic", True)),
               "the tool's own remarks: page notes, \"[figure]\""),
)
PART = {p.section: p for p in REPORT_PARTS}
# sections of the style file that are not a part of the report
OTHER_SECTIONS = {
    "document": ("template", "font"),
    "page": ("paper", "orientation", "margin_top", "margin_bottom", "margin_left", "margin_right"),
    "review marks": ("check_highlight", "likely_wrong_highlight", "likely_wrong_color", "cell_fill"),
}
# the tool's own page (used when there is no template)
BUILT_IN_PAGE = {"paper": "letter", "orientation": "portrait", "margin_top": 548640, "margin_bottom": 548640,
                 "margin_left": 640080, "margin_right": 640080}          # 0.6 in and 0.7 in, in EMU
# values used when the style file leaves a switch or review mark empty (or there is no style file)
BUILT_IN_SWITCHES = {("requirements", "tag_position"): "end", ("requirements", "mark_keywords"): False,
                     ("table header row", "repeat_on_each_page"): True,
                     ("review marks", "check_highlight"): "yellow", ("review marks", "likely_wrong_highlight"): "yellow",
                     ("review marks", "likely_wrong_color"): "C00000", ("review marks", "cell_fill"): "FFF2CC"}
BUILT_IN_HEADER_FILL = "D9E2F3"     # without a template; with one, the template's table style decides

PAPER_SIZES = {"letter": (7772400, 10058400, 1), "legal": (7772400, 12801600, 5), "tabloid": (10058400, 15544800, 3),
               "a3": (10692000, 15120000, 8), "a4": (7560000, 10692000, 9), "a5": (5328000, 7560000, 11)}
COLOR_NAMES = {"black": "000000", "gray": "808080", "grey": "808080", "red": "C00000", "dark red": "800000",
               "blue": "0070C0", "dark blue": "1F4E9A", "green": "00B050", "orange": "ED7D31", "purple": "7030A0",
               "automatic": "auto", "auto": "auto"}
HIGHLIGHT_NAMES = ("yellow", "bright green", "turquoise", "pink", "blue", "red", "dark blue", "teal", "green",
                   "violet", "dark red", "dark yellow", "gray 25", "gray 50", "black")


# ---- 11b. Reading the style file -----------------------------------------------------------------------------------
class ReportLookError(ValueError):
    """The style file or the template cannot be used. The message says why, and what to change."""


def _section_keys(section: str) -> tuple:
    if section in OTHER_SECTIONS:
        return OTHER_SECTIONS[section]
    part = PART[section]
    if part.kind == "table":
        return ("word_style",) + part.extra_keys
    return (PARAGRAPH_KEYS if part.kind == "paragraph" else CHARACTER_KEYS) + part.extra_keys


_YES, _NO = {"yes", "y", "true", "on", "1"}, {"no", "n", "false", "off", "0"}
_LENGTH_UNITS = {"cm": 360000, "mm": 36000, "in": 914400, "inch": 914400, "inches": 914400, '"': 914400, "pt": 12700}
_EXAMPLES = {"size": "size = 10.5", "space_before": "space_before = 6", "space_after": "space_after = 6",
             "color": "color = 1F4E9A", "likely_wrong_color": "likely_wrong_color = C00000",
             "line_spacing": "line_spacing = 1.15", "indent_left": "indent_left = 1 cm",
             "first_line_indent": "first_line_indent = -0.5 cm", "paper": "paper = A4"}


def _number(text: str) -> Optional[float]:
    text = text.strip().replace(",", ".")          # 10,5 is how much of Europe writes 10.5
    return float(text) if re.fullmatch(r"-?\d+(?:\.\d+)?", text) else None


def _style_value(section: str, key: str, raw: str):
    """-> (value, "") for a correct value, or (None, what to write instead)."""
    v = " ".join(raw.lower().replace(" ", " ").split())
    if key in ("word_style", "font", "template"):
        return raw.strip().strip('"').strip("'").strip(), ""
    if key == "size":
        n = _number(re.sub(r"\s*(?:pt|points?)$", "", v))
        if n is None:
            return None, f"'{raw}' is not a font size. Write a number of points, for example:  {_EXAMPLES['size']}"
        if not 1 <= n <= 200:
            return None, f"'{raw}' is outside the font sizes this tool accepts: 1 to 200 points"
        return round(n * 2) / 2, ""                 # Word stores sizes in half points
    if key in ("space_before", "space_after"):
        n = _number(re.sub(r"\s*(?:pt|points?)$", "", v))
        if n is None:
            return None, f"'{raw}' is not a spacing. Write a number of points, for example:  {_EXAMPLES[key]}"
        if not 0 <= n <= 200:
            return None, f"'{raw}' is outside the spacings this tool accepts: 0 to 200 points"
        return n, ""
    if key in ("bold", "italic", "mark_keywords", "repeat_on_each_page"):
        if v in _YES or v in _NO:
            return v in _YES, ""
        return None, f"'{raw}' should be yes or no, for example:  {key} = yes"
    if key in ("color", "likely_wrong_color"):
        hx = v.lstrip("#").upper()
        if re.fullmatch(r"[0-9A-F]{6}", hx):
            return hx, ""
        if v in COLOR_NAMES and not (key == "likely_wrong_color" and COLOR_NAMES[v] == "auto"):
            return COLOR_NAMES[v], ""
        names = ", ".join(n for n in COLOR_NAMES if n not in ("grey", "auto")
                          and not (key == "likely_wrong_color" and n == "automatic"))
        return None, (f"'{raw}' is not a colour. Write the 6-character code from Word (Font Color > More Colors > "
                      f"Custom, Hex), for example  {_EXAMPLES[key]},  or one of: {names}")
    if key in ("fill", "cell_fill"):
        hx = v.lstrip("#").upper()
        if re.fullmatch(r"[0-9A-F]{6}", hx):
            return hx, ""
        if key == "fill" and v in ("none", "no"):
            return "none", ""
        what = ("a 6-character colour code such as D9E2F3, or none" if key == "fill"
                else "a 6-character colour code such as FFF2CC (this mark cannot be switched off)")
        return None, f"'{raw}' is not a cell shading. Write {what}"
    if key in ("check_highlight", "likely_wrong_highlight"):
        v2 = " ".join(v.replace("grey", "gray").replace("-", " ").replace("%", " ").split())
        if v2 in HIGHLIGHT_NAMES:
            return v2, ""
        off = " (this mark cannot be switched off)" if v2 in ("none", "no", "off") else ""
        return None, f"'{raw}' is not one of Word's highlighter colours{off}. Choose one of: " + ", ".join(HIGHLIGHT_NAMES)
    if key == "alignment":
        names = {"left": "left", "center": "center", "centre": "center", "centered": "center",
                 "centred": "center", "right": "right", "justify": "justify", "justified": "justify"}
        if section == "tables":
            if names.get(v) in ("left", "center", "right"):
                return names[v], ""
            return None, f"'{raw}' is not a table alignment. Write left, center or right"
        if v in names:
            return names[v], ""
        return None, f"'{raw}' is not an alignment. Write left, center, right or justify"
    if key == "line_spacing":
        named = {"single": 1.0, "double": 2.0}
        if v in named:
            return ("multiple", named[v]), ""
        m = re.fullmatch(r"(at least\s+|exactly\s+)?(\d+(?:[.,]\d+)?)\s*(?:pt|points?)", v)
        if m:
            n = _number(m.group(2))
            if 1 <= n <= 200:
                return ("at_least" if (m.group(1) or "").startswith("at") else "exactly", n), ""
        n = _number(re.sub(r"\s*(?:lines?|x)$", "", v))
        if n is not None and 0.5 <= n <= 5:
            return ("multiple", n), ""
        return None, ("'" + raw + "' is not a line spacing. Write single, 1.15, 1.5 or double, or a fixed height "
                      "such as 12 pt (or at least 12 pt)")
    if key in ("indent_left", "first_line_indent") or key.startswith("margin_"):
        m = re.fullmatch(r"(-?\d+(?:[.,]\d+)?)\s*(cm|mm|in|inch|inches|\"|pt)", v)
        example = _EXAMPLES.get(key, f"{key} = 2 cm")
        if not m:
            if _number(v) is not None:
                return None, f"'{raw}' needs a unit - cm, mm or in - for example:  {example}"
            return None, f"'{raw}' is not a length. Write a number and a unit, for example:  {example}"
        emu = int(round(_number(m.group(1)) * _LENGTH_UNITS[m.group(2)]))
        if emu < 0 and key != "first_line_indent":
            return None, f"'{raw}' cannot be negative here"
        if abs(emu) > 5 * 914400:
            return None, f"'{raw}' is more than 5 inches - check the unit (cm, mm or in)"
        return emu, ""
    if key == "paper":
        if v in PAPER_SIZES:
            return v, ""
        return None, (f"'{raw}' is not a paper size this tool knows. Write Letter, A4, Legal, A3, A5 or Tabloid "
                      "(any other size can be set in your template)")
    if key == "orientation":
        if v in ("portrait", "landscape"):
            return v, ""
        return None, f"'{raw}' should be portrait or landscape"
    if key == "tag_position":
        if v in ("end", "start"):
            return v, ""
        return None, (f"'{raw}' should be end or start (the tags link the report to the Excel rows, so they always "
                      "stay)")
    return raw.strip(), ""


def _norm_name(name: str) -> str:
    return " ".join(name.replace("_", " ").lower().split())


def _line_of(lines: list, section: str, key: Optional[str] = None, occurrence: int = 1) -> int:
    """Line number of a [section] (or of a key inside it) in the style file, for the messages (0 = not found).
    `occurrence` 2 finds the second time a section appears."""
    current, seen = None, 0
    key_re = re.compile(r"\s*" + r"[\s_]+".join(map(re.escape, (key or "").split("_"))) + r"\s*[=:]", re.I)
    for number, line in enumerate(lines, 1):
        m = re.match(r"\s*\[([^\]]*)\]\s*(?:[;#].*)?$", line)
        if m:
            current = _norm_name(m.group(1))
            if current == section:
                seen += 1
                if key is None and seen == occurrence:
                    return number
            continue
        if key and current == section and key_re.match(line):
            return number
    return 0


def _at_line(number: int) -> str:
    return f"Line {number}: " if number else ""


def _style_file_message(path: Path, problems: list) -> str:
    count = "a mistake" if len(problems) == 1 else f"{len(problems)} mistakes"
    out = [f"The Word style file has {count} - no document was processed.", f"  File: {path}"]
    for line_no, where, why in sorted(problems, key=lambda p: p[0])[:12]:
        out.append((f"  Line {line_no}: {where}" if line_no else f"  {where}").rstrip())
        out.append(f"      {why}")
    if len(problems) > 12:
        out.append(f"  ... and {len(problems) - 12} more")
    out.append("Correct the file in Notepad and run again. (An empty value, or no file at all, gives the normal look.)")
    return "\n".join(out)


def _style_file_text(path: Path) -> str:
    try:
        raw = Path(path).read_bytes()
    except OSError as exc:
        raise ReportLookError(f"The Word style file {path} cannot be read ({exc.strerror or exc}).") from None
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        try:
            return raw.decode("utf-16")             # Notepad's "Unicode" (and PowerShell's Out-File)
        except UnicodeDecodeError:
            raise ReportLookError(f"The Word style file {path} is damaged. Open it in Notepad, choose File > Save As, "
                                  "set Encoding to UTF-8 and save it again.") from None
    if b"\x00" in raw:
        raise ReportLookError(f"The Word style file {path} is not a text file Notepad saved as UTF-8 or ANSI. Open it "
                              "in Notepad, choose File > Save As, set Encoding to UTF-8 and save it again.")
    try:
        return raw.decode("utf-8-sig")              # (Notepad's UTF-8 starts with a byte-order mark)
    except UnicodeDecodeError:
        return raw.decode("cp1252", errors="replace")   # Notepad's "ANSI"


def read_style_file(path: Path) -> dict:
    """{(section, key): value} of every value filled in, checked. Raises ReportLookError listing every mistake,
    with its line and what to write instead, so it can be corrected in one go before any document is read."""
    import configparser
    import difflib
    path = Path(path)
    text = _style_file_text(path)
    lines = text.splitlines()
    # no %-interpolation ("115%" must not be an error) and only ';' starts a comment inside a line, so that a
    # colour written as #1F4E9A is not taken for a comment
    parser = configparser.ConfigParser(interpolation=None, inline_comment_prefixes=(";",),
                                       comment_prefixes=("#", ";"), empty_lines_in_values=False, strict=True,
                                       default_section="\x00no default section")
    parser.optionxform = lambda k: "_".join(k.lower().split())
    try:
        parser.read_string(text, source=str(path))
    except configparser.MissingSectionHeaderError as exc:
        bracket = exc.line.strip().startswith("[")
        raise ReportLookError(_style_file_message(path, [(exc.lineno, exc.line.strip()[:60],
                              "A section name needs both brackets, such as [body]." if bracket else
                              "The settings must come after a section name in square brackets, such as [body]. Was "
                              "the first [section] line deleted?")])) from None
    except configparser.DuplicateOptionError as exc:
        raise ReportLookError(_style_file_message(path, [(exc.lineno, f"[{_norm_name(exc.section)}]",
                              f"'{exc.option}' is set twice in this section. Keep one of the two lines.")])) from None
    except configparser.DuplicateSectionError as exc:
        raise ReportLookError(_style_file_message(path, [(exc.lineno, "", f"The section [{exc.section}] appears "
                              f"twice. Put all its settings under one [{exc.section}] line.")])) from None
    except configparser.ParsingError as exc:
        raise ReportLookError(_style_file_message(path, [
            (number, lines[number - 1].strip()[:60] if 0 < number <= len(lines) else "", "This line is not a setting: write "
             "name = value, or start the line with ; to make it a comment. A section name needs both brackets, "
             "such as [body].") for number, _line in exc.errors])) from None
    except configparser.Error as exc:
        raise ReportLookError(_style_file_message(path, [(getattr(exc, "lineno", 0) or 0, "",
                                                          str(exc).splitlines()[0])])) from None
    values, problems, seen = {}, [], set()
    known = [p.section for p in REPORT_PARTS] + list(OTHER_SECTIONS)
    for section in parser.sections():
        name = _norm_name(section)
        if name not in known:
            guess = difflib.get_close_matches(name, known, 1, 0.6)
            problems.append((_line_of(lines, name), f"[{section}]", f"There is no part of the report called [{section}]."
                             + (f" Did you mean [{guess[0]}]?" if guess else "")
                             + " The sections are: " + ", ".join(f"[{s}]" for s in known)))
            continue
        if name in seen:
            problems.append((_line_of(lines, name, occurrence=2), f"[{section}]", f"The section [{name}] appears "
                             f"twice. Put all its settings under one [{name}] line."))
            continue
        seen.add(name)
        allowed = _section_keys(name)
        for key, raw in parser.items(section):
            line_no = _line_of(lines, name, key)
            where = f"[{name}]  {key} = {raw.splitlines()[0] if raw else ''}"
            if "\n" in raw:
                problems.append((line_no, where, "The line under this one is indented, so it was read as part of this "
                                 "value. Start every setting at the left edge of the line."))
                continue
            if key not in allowed:
                guess = difflib.get_close_matches(key, allowed, 1, 0.6)
                if guess:
                    hint = f" Did you mean '{guess[0]}'?"
                elif key in PARAGRAPH_KEYS and name in PART and PART[name].kind == "character":
                    hint = f" [{name}] changes words inside a line, so it has no paragraph settings such as '{key}'."
                elif key in PARAGRAPH_KEYS and name in PART and PART[name].kind == "table":
                    hint = (f" [{name}] is the table itself; the text in its cells is set in [table text] and "
                            "[table header row].")
                else:
                    hint = ""
                problems.append((line_no, where, f"'{key}' is not a setting of [{name}].{hint} The settings here "
                                 "are: " + ", ".join(allowed)))
                continue
            if key != "template":                           # (a comment straight after a value: "A4;", "Arial #x")
                raw = re.split(r";|\s#", raw, maxsplit=1)[0]
            if not raw.strip().strip('"').strip("'").strip():
                continue                                    # empty = keep the template's (or the tool's own) look
            if raw.strip().lower().startswith(("http:", "https:", "ftp:", "file:")):
                problems.append((line_no, where, "A web or network address cannot be used here: put the file in the "
                                 "word_style folder and write its name"))
                continue
            value, why = _style_value(name, key, raw)
            if why:
                problems.append((line_no, where, why))
            else:
                values[(name, key)] = value
    problems += _shared_style_conflicts(values, lines)
    problems += _unreadable_marks(values, lines)
    if problems:
        raise ReportLookError(_style_file_message(path, problems))
    return values


_HIGHLIGHT_RGB = {"yellow": "FFFF00", "bright green": "00FF00", "turquoise": "00FFFF", "pink": "FF00FF",
                  "blue": "0000FF", "red": "FF0000", "dark blue": "000080", "teal": "008080", "green": "008000",
                  "violet": "800080", "dark red": "800000", "dark yellow": "808000", "gray 25": "C0C0C0",
                  "gray 50": "808080", "black": "000000"}


def _contrast(colour_a: str, colour_b: str) -> float:
    """How well two colours (RRGGBB) stand apart, as for the contrast of text on a web page (1 = not at all)."""
    def luminance(hex_colour: str) -> float:
        channels = [int(hex_colour[i:i + 2], 16) / 255 for i in (0, 2, 4)]
        linear = [c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4 for c in channels]
        return 0.2126 * linear[0] + 0.7152 * linear[1] + 0.0722 * linear[2]
    light, dark = sorted((luminance(colour_a), luminance(colour_b)), reverse=True)
    return (light + 0.05) / (dark + 0.05)


def likely_wrong_colour(values: dict) -> str:
    """The text colour of words very likely misread: the style file's, or C00000 - or, on a highlight too dark for
    C00000 to be read (blue, red, green ...), white."""
    default = BUILT_IN_SWITCHES[("review marks", "likely_wrong_color")]
    chosen = values.get(("review marks", "likely_wrong_color"), default)
    if chosen != default:
        return chosen
    highlight = _HIGHLIGHT_RGB[values.get(("review marks", "likely_wrong_highlight"),
                                          BUILT_IN_SWITCHES[("review marks", "likely_wrong_highlight")])]
    return default if _contrast(highlight, default) >= 2.0 else "FFFFFF"


def _unreadable_marks(values: dict, lines: list) -> list:
    """The review marks cannot be switched off - so they may not be made unreadable either (black on black)."""
    def get(key):
        return values.get(("review marks", key), BUILT_IN_SWITCHES[("review marks", key)])
    highlight, colour = _HIGHLIGHT_RGB[get("likely_wrong_highlight")], likely_wrong_colour(values)
    if _contrast(highlight, colour) >= 2.0:
        return []
    key = "likely_wrong_color" if ("review marks", "likely_wrong_color") in values else "likely_wrong_highlight"
    return [(_line_of(lines, "review marks", key), f"[review marks]  {key}",
             f"Words very likely misread would be {colour} on a {get('likely_wrong_highlight')} highlight, which "
             "cannot be read. Choose a likely_wrong_color that stands out on that highlight (or another highlight).")]


def _style_name_of(section: str, values: dict) -> str:
    return values.get((section, "word_style")) or PART[section].style


def _shared_style_conflicts(values: dict, lines: list) -> list:
    """Two parts that use the same Word style cannot give it different values: say so instead of letting the
    later section win without anyone noticing."""
    by_style = defaultdict(list)
    for part in REPORT_PARTS:
        by_style[_style_name_of(part.section, values).casefold()].append(part.section)
    problems = []
    for sections in by_style.values():
        for i, first in enumerate(sections):
            for second in sections[i + 1:]:
                for key in PARAGRAPH_KEYS[1:]:
                    a, b = values.get((first, key)), values.get((second, key))
                    if a is not None and b is not None and a != b:
                        name = _style_name_of(first, values)
                        problems.append((_line_of(lines, second, key), f"[{second}]  {key}",
                                         f"[{first}] and [{second}] both use the style '{name}' but give it a "
                                         f"different {key}. Give [{second}] a style of its own (word_style), or set "
                                         f"the {key} in one place."))
    return problems


# ---- 11c. Reading a Word template safely ------------------------------------------------------------------------
# A template is never used as it is. python-docx would carry every part of it into every report, and a Word file can
# hold things that reach outside this computer when the report is opened: a link to an "attached template" on a
# network share (common in company templates), pictures linked to a file or a web address, fields that fetch a
# file, a web page or another program's data, embedded objects, macros. So every report is a NEW document, and only
# these are copied into it, each one checked: the styles, the theme (its fonts and colours), the list numbering, a
# few page-layout settings, the page size and margins, and the page headers and footers with their pictures. A link
# in a header becomes plain text, tracked changes in it are accepted, and a field other than a page number, page
# count, date or heading reference becomes the text it showed. The
# template's own text, its document properties, and everything else in the file are never used. Every finished
# report is then checked once more before it is saved (outside_references, section 11d).
_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
_O = "urn:schemas-microsoft-com:office:office"
_V = "urn:schemas-microsoft-com:vml"
_PKG_RELS = "http://schemas.openxmlformats.org/package/2006/relationships"
_CONTENT_TYPES = "http://schemas.openxmlformats.org/package/2006/content-types"
_MAIN_PART_TYPES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml": ".docx",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml": ".dotx",
    "application/vnd.ms-word.document.macroEnabled.main+xml": ".docm",
    "application/vnd.ms-word.template.macroEnabledTemplate.main+xml": ".dotm"}
_PICTURE_TYPES = {"image/png": ".png", "image/jpeg": ".jpeg", "image/gif": ".gif", "image/bmp": ".bmp",
                  "image/tiff": ".tiff", "image/x-emf": ".emf", "image/emf": ".emf", "image/x-wmf": ".wmf",
                  "image/wmf": ".wmf"}
# fields that only number or date the pages (or repeat a heading of the report): they keep working in a copied
# header or footer. Any other field (TITLE, DOCPROPERTY, INCLUDEPICTURE ...) is replaced by the text it showed, so
# what the template's header says never changes.
SAFE_FIELDS = {"PAGE", "NUMPAGES", "SECTIONPAGES", "DATE", "TIME", "PRINTDATE", "SAVEDATE", "CREATEDATE",
               "STYLEREF"}
# the only page-layout settings copied from the template's settings (none of them can point anywhere)
_LAYOUT_SETTINGS = {"view", "zoom", "mirrorMargins", "alignBordersAndEdges", "bordersDoNotSurroundHeader",
                    "bordersDoNotSurroundFooter", "gutterAtTop", "defaultTabStop", "autoHyphenation",
                    "consecutiveHyphenLimit", "hyphenationZone", "doNotHyphenateCaps", "evenAndOddHeaders",
                    "bookFoldRevPrinting", "bookFoldPrinting", "bookFoldPrintingSheets", "noPunctuationKerning",
                    "characterSpacingControl", "compat", "mathPr", "themeFontLang", "clrSchemeMapping",
                    "decimalSymbol", "listSeparator"}
# settings worth a line in the log when they are left out
_NOTED_SETTINGS = {"attachedTemplate": "its link to an attached template was left out",
                   "mailMerge": "its mail-merge data source was left out",
                   "documentProtection": "its editing protection was left out",
                   "writeProtection": "its write protection was left out",
                   "trackRevisions": "its track-changes switch was left out",
                   "docVars": "its document variables were left out",
                   "updateFields": "its 'update fields when opened' switch was left out"}
_MAX_TEMPLATE_BYTES = 50 * 1024 * 1024            # a template is a few hundred kB; this leaves room for logos
_MAX_TEMPLATE_PARTS = 2000
_MAX_UNPACKED_BYTES = 300 * 1024 * 1024           # (guards against a file that unpacks to something enormous)


@dataclass
class WordTemplate:
    """What is copied from a template into each report (all checked, see read_word_template)."""
    path: Path
    sha256: str
    styles: Optional[bytes] = None                 # (None: the template has none - Word's standard styles are used)
    numbering: Optional[bytes] = None
    settings: Optional[bytes] = None
    theme: Optional[tuple] = None                  # (xml, [pictures])
    section: Optional[bytes] = None                # the page layout: <w:sectPr> without header/footer references
    stories: list = field(default_factory=list)    # [("header" | "footer", "default" | "first" | "even", xml, [pictures])]
    notes: list = field(default_factory=list)      # what was left out or changed, for the log


def _q(ns: str, name: str) -> str:
    return f"{{{ns}}}{name}"


def _local(el) -> str:
    return el.tag.rsplit("}", 1)[-1] if isinstance(el.tag, str) else ""


def _safe_xml(data: bytes, what: str):
    """Parse XML from a Word file: no DTD, no entities, never the network. Word never writes a DTD, so a file
    that has one is refused (it could otherwise make the reader fetch or expand something)."""
    from lxml import etree
    parser = etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False, huge_tree=False)
    try:
        root = etree.fromstring(data, parser)
    except etree.XMLSyntaxError as exc:
        raise ReportLookError(f"{what} is damaged and cannot be read ({exc}).") from None
    if root.getroottree().docinfo.doctype:
        raise ReportLookError(f"{what} contains a document-type declaration, which Word never writes; the file is "
                              "not used.")
    return root


def _xml_bytes(root) -> bytes:
    from lxml import etree
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _reaches_outside(target: str, folder: str = "") -> str:
    """Why a relationship target or file reference leaves the Word file ("" when it stays inside it)."""
    t = (target or "").strip()
    low = t.lower()
    if not t:
        return ""
    if re.match(r"[a-z][a-z0-9+.-]*:", low) and not re.match(r"[a-z]:[\\/]", low):
        return f'it points to the address "{t[:100]}"'
    if t.startswith(("\\\\", "//")):
        return f'it points to the network location "{t[:100]}"'
    if re.match(r"[a-zA-Z]:[\\/]", t):
        return f'it points to the file "{t[:100]}"'
    inside = t.replace("\\", "/")
    if inside.startswith("/"):
        return ""
    resolved = posixpath.normpath(posixpath.join(folder, inside))
    if resolved == ".." or resolved.startswith(("../", "/")):
        return f'it points outside the file ("{t[:100]}")'
    return ""


def _part_folder(part_name: str) -> str:
    return posixpath.dirname(part_name.lstrip("/"))


def _rels_name(part_name: str) -> str:
    folder, name = posixpath.split(part_name.lstrip("/"))
    return posixpath.join(folder, "_rels", name + ".rels")


_Rel = namedtuple("_Rel", "kind part outside target")   # one relationship: its type, the part it leads to
                                                        # ("" when outside), why it leads outside ("" = it doesn't)


class _Package:
    """The parts of a Word file, read with the checks a file from elsewhere needs."""

    def __init__(self, data: bytes, what: str):
        import zipfile
        self.what = what
        try:
            self.zip = zipfile.ZipFile(io.BytesIO(data))
        except zipfile.BadZipFile:
            raise ReportLookError(f"{what} is not a Word .docx/.dotx file (it may be damaged).") from None
        infos = self.zip.infolist()
        if len(infos) > _MAX_TEMPLATE_PARTS or sum(i.file_size for i in infos) > _MAX_UNPACKED_BYTES:
            raise ReportLookError(f"{what} unpacks to far more than a Word template ever does; it is not used.")
        self.names = {}
        for info in infos:
            name = info.filename
            if name.startswith(("/", "\\")) or ".." in name.replace("\\", "/").split("/") or re.match(r"[A-Za-z]:", name):
                raise ReportLookError(f"{what} contains a file with an unsafe name ({name[:80]}); it is not used.")
            if info.compress_size > 4096 and info.file_size > 200 * info.compress_size:
                raise ReportLookError(f"{what} contains a part that unpacks to something enormous; it is not used.")
            self.names[name.lower()] = name
        self.types = self._content_types()

    def has(self, part_name: str) -> bool:
        return part_name.lstrip("/").lower() in self.names

    def read(self, part_name: str) -> bytes:
        return self.zip.read(self.names[part_name.lstrip("/").lower()])

    def _content_types(self) -> dict:
        if not self.has("[Content_Types].xml"):
            raise ReportLookError(f"{self.what} is not a Word .docx/.dotx file.")
        root = _safe_xml(self.read("[Content_Types].xml"), self.what)
        types = {}
        for el in root:
            if _local(el) == "Default":
                types["*." + (el.get("Extension") or "").lower()] = el.get("ContentType") or ""
            elif _local(el) == "Override":
                types[(el.get("PartName") or "").lstrip("/").lower()] = el.get("ContentType") or ""
        return types

    def content_type(self, part_name: str) -> str:
        name = part_name.lstrip("/").lower()
        return self.types.get(name) or self.types.get("*." + name.rsplit(".", 1)[-1], "")

    def rels(self, part_name: str) -> dict:
        """{rId: _Rel} of a part ('' = the package's own relationships)."""
        rels_name = "_rels/.rels" if not part_name else _rels_name(part_name)
        if not self.has(rels_name):
            return {}
        out = {}
        folder = _part_folder(part_name) if part_name else ""
        for rel in _safe_xml(self.read(rels_name), self.what):
            target = rel.get("Target") or ""
            outside = _reaches_outside(target, folder)
            if (rel.get("TargetMode") or "").lower() == "external":
                outside = outside or f'it points to "{target[:100]}" outside the file'
            name = "" if outside else (target.lstrip("/") if target.startswith("/")
                                       else posixpath.normpath(posixpath.join(folder, target.replace("\\", "/"))))
            out[rel.get("Id")] = _Rel((rel.get("Type") or "").rsplit("/", 1)[-1], name, outside, target)
        return out


def _field_name(code: str) -> str:
    m = re.match(r"\s*([A-Za-z]+|=)", code or "")
    return m.group(1).upper() if m else ""


def _field_codes(root) -> list:
    """The code of every field ("PAGE", "INCLUDEPICTURE \"http://...\"" ...). A field's code may be split over many
    runs and may hold other fields; a stack keeps each field's own code together, so a field cannot hide its name by
    putting another field in the middle of it."""
    codes = [fs.get(_q(_W, "instr"), "") for fs in root.iter(_q(_W, "fldSimple"))]
    stack = []
    for el in root.iter(_q(_W, "fldChar"), _q(_W, "instrText"), _q(_W, "delInstrText")):
        if _local(el) == "fldChar":
            kind = el.get(_q(_W, "fldCharType"))
            if kind == "begin":
                stack.append([])
            elif kind == "separate" and stack:
                stack[-1].append("\x00")           # the code ends here; text after it is the result
            elif kind == "end" and stack:
                codes.append("".join(stack.pop()).split("\x00")[0])
        elif stack and "\x00" not in stack[-1]:
            stack[-1].append(el.text or "")
        else:                                       # a field code outside any field: counted on its own
            codes.append(el.text or "")
    codes += ["".join(s).split("\x00")[0] for s in stack]
    return codes


def _unwrap(el) -> None:
    """Replace an element by its contents (a link by its text, a field by what it showed)."""
    parent = el.getparent()
    index = parent.index(el)
    for child in list(el):
        parent.insert(index, child)
        index += 1
    parent.remove(el)


def _freeze_fields(root, keep: set) -> list:
    """Turn every field whose name is not in `keep` into the text it showed; returns the names of those fields."""
    frozen = []
    for fs in list(root.iter(_q(_W, "fldSimple"))):
        name = _field_name(fs.get(_q(_W, "instr"), ""))
        if name not in keep:
            frozen.append(name or "?")
            _unwrap(fs)
    runs = list(root.iter(_q(_W, "r")))
    fields, stack, stray = [], [], []
    for i, run in enumerate(runs):
        for el in run:
            if _local(el) == "fldChar":
                kind = el.get(_q(_W, "fldCharType"))
                if kind == "begin":
                    stack.append({"begin": i, "separate": None, "code": []})
                elif kind == "separate" and stack:
                    stack[-1]["separate"] = i
                elif kind == "end" and stack:
                    fields.append(dict(stack.pop(), end=i))
            elif _local(el) == "instrText" and stack and stack[-1]["separate"] is None:
                stack[-1]["code"].append(el.text or "")
            elif _local(el) == "instrText":                        # a field code outside any field: Word ignores
                frozen.append(_field_name(el.text) or "?")          # it, but it goes all the same
                stray.append((i, el))
    fields += [dict(f, end=len(runs) - 1) for f in stack]          # (a field that never ends: to the last run)
    whole, marks = set(), set()
    for i, el in stray:
        runs[i].remove(el)
        if not any(_local(c) != "rPr" for c in runs[i]) and runs[i].getparent() is not None:
            runs[i].getparent().remove(runs[i])
    for f in fields:
        name = _field_name("".join(f["code"]))
        if name in keep:
            continue
        frozen.append(name or "?")
        code_end = f["separate"] if f["separate"] is not None else f["end"]
        whole.update(range(f["begin"] + 1, code_end))              # the runs of its code (and of fields inside it)
        marks.update({f["begin"], code_end, f["end"]})             # the field marks themselves; its result stays
    for i in sorted(marks - whole):
        for el in list(runs[i]):
            if _local(el) in ("fldChar", "instrText"):
                runs[i].remove(el)
    for i in sorted(whole | {m for m in marks if not any(_local(c) not in ("rPr", "fldChar", "instrText")
                                                          for c in runs[m])}):
        parent = runs[i].getparent()
        if parent is not None:
            parent.remove(runs[i])
    return frozen


_REVISION_DROPS = {"del", "delText", "delInstrText", "moveFrom", "moveFromRangeStart", "moveFromRangeEnd", "moveToRangeStart", "moveToRangeEnd",
                   "cellDel", "cellIns", "cellMerge", "rPrChange", "pPrChange", "sectPrChange", "tblPrChange",
                   "tblPrExChange", "trPrChange", "tcPrChange", "tblGridChange", "numberingChange",
                   "customXmlDelRangeStart", "customXmlDelRangeEnd", "customXmlInsRangeStart", "customXmlInsRangeEnd",
                   "customXmlMoveFromRangeStart", "customXmlMoveFromRangeEnd", "customXmlMoveToRangeStart",
                   "customXmlMoveToRangeEnd"}


def _accept_revisions(root) -> int:
    """Accept every tracked change, as Word's 'Accept All Changes' does: deleted text goes, inserted text stays, the
    reviser's names and dates go. (Otherwise a marking changed from SECRET to UNCLASSIFIED but never accepted would
    carry the struck-out SECRET into every report.) Returns the number of changes."""
    changes = [el for el in root.iter() if isinstance(el.tag, str) and el.tag.startswith(f"{{{_W}}}")
               and _local(el) in _REVISION_DROPS | {"ins", "moveTo"}]
    deleted = (_q(_W, "del"), _q(_W, "moveFrom"))
    for row in list(root.iter(_q(_W, "tr"))):                     # a deleted table row goes (and an emptied table)
        props = row.find(_q(_W, "trPr"))
        if props is not None and any(props.find(tag) is not None for tag in deleted) and row.getparent() is not None:
            table = row.getparent()
            table.remove(row)
            if table.find(_q(_W, "tr")) is None and table.getparent() is not None:
                table.getparent().remove(table)
    for p in list(root.iter(_q(_W, "p"))):                        # a deleted paragraph mark joins the paragraph to
        mark = p.find(f"{_q(_W, 'pPr')}/{_q(_W, 'rPr')}")         # the next one, as Word's Accept All does
        if mark is None or not any(mark.find(tag) is not None for tag in deleted) or p.getparent() is None:
            continue
        content = [c for c in p if _local(c) != "pPr" and not (_local(c) in ("del", "moveFrom"))]
        following = p.getnext()
        if following is not None and _local(following) == "p":
            props = following.find(_q(_W, "pPr"))
            at = 1 if props is not None else 0
            for c in content:
                following.insert(at, c)
                at += 1
            p.getparent().remove(p)
        elif not content:
            p.getparent().remove(p)
    for el in changes:
        if el.getparent() is None:
            continue
        if _local(el) in ("ins", "moveTo"):
            _unwrap(el)
        else:
            el.getparent().remove(el)
    return len(changes)


def _clean_story(root, package: "_Package", part_name: str, where: str, notes: list) -> tuple:
    """A header or footer of the template, made safe to copy: (xml, [(rId, extension, content type, data)]).
    Raises ReportLookError for something that cannot be made safe without losing what the header shows."""
    rels = package.rels(part_name)
    what = f"The {where} of the template"
    if _accept_revisions(root):
        notes.append(f"tracked changes in the {where} were accepted")
    # links become plain text (the text stays, the address goes)
    links = 0
    for link in list(root.iter(_q(_W, "hyperlink"))):
        if link.get(_q(_R, "id")):
            _unwrap(link)
            links += 1
    for el in list(root.iter()):
        if _local(el) in ("hlinkClick", "hlinkHover", "hlinkMouseOver") and el.getparent() is not None:
            el.getparent().remove(el)
            links += 1
        elif _local(el) == "dataBinding":                   # content control tied to the template's hidden data
            el.getparent().remove(el)
        elif _local(el) == "ext" and el.getparent() is not None and _local(el.getparent()) == "extLst" \
                and any(k.startswith(f"{{{_R}}}") for d in el.iter() for k in d.attrib):
            el.getparent().remove(el)                       # e.g. the SVG copy of a picture (its PNG stays)
    if links:
        notes.append(f"{links} link(s) in the {where} became plain text")
    frozen = _freeze_fields(root, SAFE_FIELDS)
    if frozen:
        notes.append(f"field(s) {', '.join(sorted(set(frozen)))} in the {where} became the text they showed")
    for el in root.iter():
        name = _local(el)
        if name in ("object", "control", "altChunk", "subDoc", "OLEObject"):
            raise ReportLookError(f"{what} holds an embedded object or control (<{name}>), which cannot be copied "
                                  "safely. In Word, delete it from the header/footer (a picture of it is fine), "
                                  "save, and run again.")
        for attr in ("href", _q(_O, "href"), "src"):
            if el.get(attr) is not None and name in ("imagedata", "shape", "rect", "roundrect", "oval", "image"):
                del el.attrib[attr]                          # VML: the original file name or link of a picture
    pictures, seen = [], {}
    for el in root.iter():
        for attr, value in list(el.attrib.items()):
            if not (attr.startswith(f"{{{_R}}}") or attr == _q(_O, "relid")):
                continue
            rel = rels.get(value)
            if attr == _q(_R, "link") or (rel and rel.outside):
                reason = rel.outside if rel and rel.outside else "it is linked to a file instead of being stored in the template"
                raise ReportLookError(f"{what} holds a picture that is not stored in the template - {reason}. In Word, "
                                      "delete it and insert it again with Insert > Pictures > This Device (do not "
                                      "choose 'Link to File'), save, and run again.")
            if rel is None:
                raise ReportLookError(f"{what} refers to a picture or object that is missing from the file. In Word, "
                                      "insert it again (or remove it), save, and run again.")
            kind, target = rel.kind, rel.part
            ctype = package.content_type(target) if target else ""
            if kind != "image" or ctype not in _PICTURE_TYPES or not package.has(target):
                raise ReportLookError(f"{what} holds something other than text and pictures ({kind or 'unknown'}"
                                      f"{', ' + ctype if ctype else ''}), which cannot be copied safely. In Word, "
                                      "remove it from the header/footer (or replace it by a picture), save, and run "
                                      "again.")
            if value not in seen:
                data = package.read(target)
                if len(data) > 20 * 1024 * 1024:
                    raise ReportLookError(f"{what} holds a picture of {len(data) // 1048576} MB; make it smaller "
                                          "(Picture Format > Compress Pictures) and run again.")
                seen[value] = True
                pictures.append((value, _PICTURE_TYPES[ctype], ctype, data))
    bad = [c for c in _field_codes(root) if _field_name(c) not in SAFE_FIELDS]
    if bad:                                                   # (cannot happen after _freeze_fields: a last guard)
        raise ReportLookError(f"{what} holds a field that could not be made safe ({bad[0].strip()[:60]}).")
    return _xml_bytes(root), pictures


def read_word_template(path: Path) -> WordTemplate:
    """Read and check a Word template (.dotx or .docx), keeping only what the report may copy (see above).
    Raises ReportLookError, saying what to do in Word, for a file that cannot be used."""
    try:
        return _read_word_template(Path(path))
    except ReportLookError:
        raise
    except Exception as exc:                  # a damaged or unusual file: a clear message, never a crash
        raise ReportLookError(f"The template '{Path(path).name}' could not be read ({type(exc).__name__}: {exc}). "
                              "Open it in Word, save it again as 'Word Template (*.dotx)' and use the new file.") from None


def _read_word_template(path: Path) -> WordTemplate:
    name = path.name
    suffix = path.suffix.lower()
    if suffix in (".docm", ".dotm"):
        raise ReportLookError(f"'{name}' can hold macros, so it is not used. In Word, open it and choose File > Save "
                              "As > 'Word Template (*.dotx)', then write the new name in the style file.")
    if suffix in (".doc", ".dot", ".rtf", ".odt", ".ott"):
        raise ReportLookError(f"'{name}' is not in Word's current format. In Word, open it and choose File > Save As "
                              "> 'Word Template (*.dotx)', then write the new name in the style file.")
    if suffix not in (".docx", ".dotx"):
        raise ReportLookError(f"'{name}' is not a Word template: use a .dotx or .docx file.")
    try:
        size = path.stat().st_size
        if size > _MAX_TEMPLATE_BYTES:
            raise ReportLookError(f"'{name}' is {size // 1048576} MB - far larger than a template. Use a template "
                                  "without document content.")
        data = path.read_bytes()
    except OSError as exc:
        raise ReportLookError(f"The Word template '{path}' cannot be read ({exc.strerror or exc}).") from None
    what = f"The template '{name}'"
    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        raise ReportLookError(f"{what} is protected with a password (or is an old-format Word file) and cannot be "
                              "read. Save an unprotected copy as 'Word Template (*.dotx)' and use that.")
    package = _Package(data, what)
    main = next((rel.part for rel in package.rels("").values() if rel.kind == "officeDocument" and rel.part), "")
    if not main or not package.has(main):
        raise ReportLookError(f"{what} has no document in it; it is not a Word .docx/.dotx file.")
    main_type = package.content_type(main)
    if _MAIN_PART_TYPES.get(main_type) in (".docm", ".dotm") or any(
            "macroenabled" in t.lower() or "vbaproject" in t.lower() or "activex" in t.lower()
            for t in package.types.values()):
        raise ReportLookError(f"{what} contains macros or ActiveX controls, so it is not used. In Word, choose File > "
                              "Save As > 'Word Template (*.dotx)' (which removes them) and use the new file.")
    if main_type not in _MAIN_PART_TYPES:
        raise ReportLookError(f"{what} is not a Word document or template (it says it holds '{main_type or '?'}'). "
                              "If it was saved as 'Strict Open XML', save it again as 'Word Template (*.dotx)'.")
    rels = package.rels(main)
    t = WordTemplate(path=path, sha256=hashlib.sha256(data).hexdigest())
    notes = t.notes

    def first(kind: str) -> str:
        found = [rel.part for rel in rels.values() if rel.kind == kind and rel.part and package.has(rel.part)]
        return found[0] if found else ""

    # styles: the heart of the template
    styles_part = first("styles")
    if styles_part:
        styles = _safe_xml(package.read(styles_part), what)
        if _local(styles) != "styles":
            raise ReportLookError(f"{what} is damaged: its style list cannot be read.")
        for el in styles.iter():                   # (styles never point anywhere; drop anything that tries to)
            for attr in [a for a in el.attrib if a.startswith(f"{{{_R}}}")]:
                del el.attrib[attr]
        if _accept_revisions(styles):
            notes.append("tracked changes to its styles were accepted")
        t.styles = _xml_bytes(styles)
    else:
        notes.append("it has no styles of its own, so Word's standard styles are used")
    # list numbering (heading styles may use it): picture bullets are left out
    numbering_part = first("numbering")
    if numbering_part:
        numbering = _safe_xml(package.read(numbering_part), what)
        for el in list(numbering.iter(_q(_W, "numPicBullet"), _q(_W, "lvlPicBulletId"))):
            el.getparent().remove(el)
        _accept_revisions(numbering)
        if any(a.startswith(f"{{{_R}}}") for el in numbering.iter() for a in el.attrib):
            notes.append("its list numbering was left out (it refers to other parts of the file)")
        else:
            t.numbering = _xml_bytes(numbering)
    # page-layout settings: only the known harmless ones
    settings_part = first("settings")
    if settings_part:
        settings = _safe_xml(package.read(settings_part), what)
        settings_rels = package.rels(settings_part)
        for el in list(settings):
            local = _local(el)
            if local in _LAYOUT_SETTINGS:
                continue
            if local in _NOTED_SETTINGS:
                link = next((rel.target for rel in settings_rels.values() if rel.kind == "attachedTemplate"), "") \
                    if local == "attachedTemplate" else ""
                notes.append(_NOTED_SETTINGS[local] + (f" ({link[:100]})" if link else ""))
            settings.remove(el)
        t.settings = _xml_bytes(settings)
    # theme: fonts and colours (theme pictures are rare; a theme that needs more than its XML is left out)
    theme_part = first("theme")
    if theme_part:
        theme = _safe_xml(package.read(theme_part), what)
        if any(a.startswith(f"{{{_R}}}") for el in theme.iter() for a in el.attrib):
            notes.append("its theme was left out (it holds pictures); the standard theme fonts and colours are used")
        else:
            t.theme = (_xml_bytes(theme), [])
    # page layout and headers/footers: those of the last section, with the ones it takes over from the sections
    # before it ("Link to Previous")
    document = _safe_xml(package.read(main), what)
    body = document.find(_q(_W, "body"))
    sections = [s for s in body.iter(_q(_W, "sectPr")) if _local(s.getparent()) in ("body", "pPr")] \
        if body is not None else []                      # (not the old layout kept by a tracked change)
    references = {}
    for sect in sections:
        for ref in sect:
            if _local(ref) in ("headerReference", "footerReference"):
                references[(_local(ref)[:6], ref.get(_q(_W, "type")) or "default")] = ref.get(_q(_R, "id"))
    if sections:
        last = copy.deepcopy(sections[-1])
        for el in list(last):
            if _local(el) in ("headerReference", "footerReference", "printerSettings", "sectPrChange"):
                last.remove(el)
            elif _local(el) == "cols" and (el.get(_q(_W, "num")) or "1").strip() != "1":
                last.remove(el)
                notes.append("its text columns were left out (the report is set in one column)")
        for attr in [a for a in last.attrib if a.startswith(f"{{{_W}}}rsid")]:
            del last.attrib[attr]
        if any(a.startswith(f"{{{_R}}}") for el in last.iter() for a in el.attrib):
            notes.append("its page layout was left out (it refers to other parts of the file)")
        else:
            t.section = _xml_bytes(last)
        if len(sections) > 1:
            notes.append(f"its {len(sections)} sections were reduced to the page layout of the last one")
    for (kind, hf_type), rid in sorted(references.items()):
        found = rels.get(rid)
        where = {"default": "", "first": "first-page ", "even": "even-page "}.get(hf_type, "") + kind
        if not found or not found.part or not package.has(found.part):
            notes.append(f"its {where} was left out (it is missing from the file)")
            continue
        root = _safe_xml(package.read(found.part), what)
        xml, pictures = _clean_story(root, package, found.part, where, notes)
        t.stories.append((kind, hf_type, xml, pictures))
    if (body is not None and any(_local(el) in ("p", "tbl", "sdt") for el in body)
            and any((txt.text or "").strip() for txt in body.iter(_q(_W, "t")))):
        notes.append("its own text was left out (only its styles, page layout, headers and footers are used)")
    return t


# ---- 11d. Making the report document ---------------------------------------------------------------------------
@dataclass
class ReportLook:
    """Everything that decides how a report looks: the style file's values and the template (both checked)."""
    values: dict = field(default_factory=dict)       # {(section, key): value} filled in the style file
    style_file: Optional[Path] = None
    template: Optional[WordTemplate] = None
    lines: list = field(default_factory=list)        # the style file's lines (for the line numbers in messages)
    template_signature: tuple = ()                   # (to notice when the template file is changed)

    def get(self, section: str, key: str):
        value = self.values.get((section, key))
        return BUILT_IN_SWITCHES.get((section, key)) if value is None else value

    @property
    def summary(self) -> str:
        """One line for the log, the Run Info sheet and the report itself."""
        parts = []
        if self.template is not None:
            parts.append(f"template {self.template.path.name} (sha256 {self.template.sha256[:12]})")
        if self.style_file is not None:
            n = sum(1 for (s, k), v in self.values.items()
                    if not (s, k) == ("document", "template") and BUILT_IN_SWITCHES.get((s, k)) != v)
            parts.append(f"{self.style_file.name} ({n} value{'s' if n != 1 else ''} changed)" if n
                         else f"{self.style_file.name} (nothing changed)")
        return ", ".join(parts) if parts else "the tool's own look"


_LOOK_CACHE = {}


def _file_signature(path: Optional[Path]) -> tuple:
    try:
        st = Path(path).stat()
        return (str(Path(path).resolve()), st.st_mtime_ns, st.st_size)
    except (OSError, TypeError):
        return (str(path), None, None)


def report_look(settings: Settings) -> ReportLook:
    """The look of the Word report: the style file (Settings.word_style_file) and the template it names, read and
    checked once and remembered while neither file changes. Raises ReportLookError with the reason.
    A run checks the look at start-up and keeps it in Settings.word_look, so editing the style file during a long
    batch changes nothing until the next run."""
    if settings.word_look is not None:
        return settings.word_look
    path = settings.word_style_file
    if path is not None and not Path(path).exists():
        if Path(path) == DEFAULT_WORD_STYLE_FILE:
            path = None                                  # (only spec_extractor.py was copied: the built-in look)
        else:
            raise ReportLookError(f"The Word style file {path} does not exist.")
    key = _file_signature(path) if path is not None else ("built-in",)
    cached = _LOOK_CACHE.get(key)
    if cached is not None and (cached.template is None
                               or _file_signature(cached.template.path) == cached.template_signature):
        return cached
    look = ReportLook()
    if path is not None:
        path = Path(path)
        look.style_file = path
        look.values = read_style_file(path)
        look.lines = _style_file_text(path).splitlines()
        template = look.values.get(("document", "template"))
        if template:
            template_path = Path(template).expanduser()
            if not template_path.is_absolute():          # a name is looked for next to the style file
                template_path = path.parent / template_path
            if not template_path.exists():
                here = sorted(p.name for p in template_path.parent.glob("*.do[ct]x")) if template_path.parent.is_dir() else []
                raise ReportLookError(
                    f"The Word template named in {path.name} was not found.\n"
                    f"  {_at_line(_line_of(look.lines, 'document', 'template'))}[document]  template = {template}\n"
                    f"      Looked for {template_path}."
                    + (f" Word files in that folder: {', '.join(here)}" if here else ""))
            look.template = read_word_template(template_path)
    look.template_signature = _file_signature(look.template.path) if look.template else ()
    # build one document now, so that a style name that does not exist, or anything in the template that cannot be
    # used, is reported before any page is read - not after an hour of OCR
    try:
        probe, _styles = _report_document(look)
        section = probe.sections[0]
        width = (section.page_width - section.left_margin - section.right_margin - (section.gutter or 0)) / 914400
        height = (section.page_height - abs(section.top_margin) - abs(section.bottom_margin)) / 914400
        _usable_page(probe)                   # (reads the template's header pictures: a damaged one is caught here)
    except ReportLookError:
        raise
    except Exception as exc:                  # something in a template python-docx cannot use: a message, not a crash
        if look.template is None:
            raise
        raise ReportLookError(f"The template '{look.template.path.name}' could not be used ({type(exc).__name__}: "
                              f"{exc}). Open it in Word, save it again as 'Word Template (*.dotx)' and use the new "
                              "file.") from None
    if width < 1.5 or height < 1.5:
        lines = [f"  {_at_line(_line_of(look.lines, 'page', k))}[page]  {k}" for k in OTHER_SECTIONS["page"]
                 if ("page", k) in look.values]
        raise ReportLookError(
            f"The page of the report would leave only {max(width, 0):.1f} x {max(height, 0):.1f} in for the text: the "
            "margins are too wide for the paper.\n"
            + ("\n".join(lines) + "\n" if lines else "  (the margins come from the template)\n")
            + "      Make the margins smaller, or the paper larger, in "
            + (look.style_file.name if look.style_file else "the style file") + ".")
    buffer = io.BytesIO()
    probe.save(buffer)
    problems = outside_references(buffer.getvalue())
    if problems:
        raise ReportLookError("The report made with this look would point outside this computer, so it is not "
                              "used:\n  " + "\n  ".join(problems[:5]))
    _LOOK_CACHE[key] = look
    return look


_STANDARD_STYLES = None


def _standard_styles():
    """Word's standard styles as python-docx's own template defines them (Title, Heading 1-9, Caption, List
    Paragraph, Table Grid, Body Text ...), used for any of them a template does not have."""
    global _STANDARD_STYLES
    if _STANDARD_STYLES is None:
        _STANDARD_STYLES = docx.Document().styles.element
    return _STANDARD_STYLES


def _style_names(el) -> set:
    from docx.styles import BabelFish
    name_el = el.find(_q(_W, "name"))
    name = name_el.get(_q(_W, "val")) if name_el is not None else ""
    return {name.casefold(), BabelFish.internal2ui(name).casefold(), (el.get(_q(_W, "styleId")) or "").casefold()}


def _find_style_element(root, name: str):
    """A style by its name in Word ('Heading 1'), its stored English name ('heading 1') or its id, ignoring case."""
    from docx.styles import BabelFish
    wanted = {name.casefold(), BabelFish.ui2internal(name).casefold()}
    return next((el for el in root.findall(_q(_W, "style")) if _style_names(el) & wanted), None)


def _free_style_id(root, wanted: str, ignore=None) -> str:
    """A style id no other style uses yet ('TableText', else 'TableText2' ...)."""
    used = {el.get(_q(_W, "styleId")) for el in root.findall(_q(_W, "style")) if el is not ignore}
    new, n = wanted, 1
    while new in used:
        n += 1
        new = f"{wanted}{n}"
    return new


def _copy_standard_style(d, name: str, added: set, copied: dict):
    """Copy a standard style from python-docx's template into the report; None when Word has no standard style of
    that name. The styles it refers to (based on, linked to, next) are matched BY NAME in the report first - a
    template from a non-English Word stores Normal under the id 'Standard' or 'Normale' - and the report's default
    style of a kind stands in for python-docx's (a template's default paragraph style need not be called Normal).
    Only a style the report has in no form is copied as well. The ids of the copies are recorded in `added`."""
    standard = _standard_styles()
    source = _find_style_element(standard, name)
    if source is None:
        return None
    root = d.styles.element                              # (copied: python-docx's style id -> the report's)

    def names_of(el) -> set:
        from docx.styles import BabelFish
        name_el = el.find(_q(_W, "name"))
        stored = name_el.get(_q(_W, "val")) if name_el is not None else ""
        return {stored.casefold(), BabelFish.internal2ui(stored).casefold()} - {""}

    def in_report(src):
        kind, names = src.get(_q(_W, "type")), names_of(src)
        same_kind = [el for el in root.findall(_q(_W, "style")) if el.get(_q(_W, "type")) == kind]
        found = next((el for el in same_kind if names_of(el) & names), None)
        if found is None:                                # (no style of that name: one with python-docx's id)
            found = next((el for el in same_kind if el.get(_q(_W, "styleId")) == src.get(_q(_W, "styleId"))), None)
        if found is None and src.get(_q(_W, "default")) == "1":
            found = next((el for el in root.findall(_q(_W, "style")) if el.get(_q(_W, "default")) == "1"
                          and el.get(_q(_W, "type")) == src.get(_q(_W, "type"))), None)
        return found

    def copy_style(src) -> str:
        src_id = src.get(_q(_W, "styleId"))
        if src_id in copied:
            return copied[src_id]
        found = in_report(src)
        if found is not None:
            copied[src_id] = found.get(_q(_W, "styleId"))
            return copied[src_id]
        new = copy.deepcopy(src)
        new.attrib.pop(_q(_W, "default"), None)
        for child in list(new):
            if _local(child) == "rsid":
                new.remove(child)
        for num in new.findall(f".//{_q(_W, 'numPr')}"):            # (list styles point at python-docx's lists)
            num.getparent().remove(num)
        new.set(_q(_W, "styleId"), _free_style_id(root, src_id))
        copied[src_id] = new.get(_q(_W, "styleId"))
        added.add(copied[src_id])
        root.append(new)
        for tag in ("basedOn", "link", "next"):
            ref = new.find(_q(_W, tag))
            if ref is None:
                continue
            dep = next((s for s in standard.findall(_q(_W, "style")) if s.get(_q(_W, "styleId")) == ref.get(_q(_W, "val"))),
                       None)
            if dep is None:
                new.remove(ref)
            else:
                ref.set(_q(_W, "val"), copy_style(dep))
        return copied[src_id]

    style_id = copy_style(source)
    return next(el for el in root.findall(_q(_W, "style")) if el.get(_q(_W, "styleId")) == style_id)


_STYLE_TYPES = {"paragraph": WD_STYLE_TYPE.PARAGRAPH, "character": WD_STYLE_TYPE.CHARACTER,
                "table": WD_STYLE_TYPE.TABLE}


def _report_style(d, part: ReportPart, look: ReportLook, styles: dict, added: set, copied: dict):
    """-> (style, added): the style a part uses, added to the document when it does not have it yet (the ids of
    styles added to this document are in `added`)."""
    from docx.styles.style import StyleFactory
    name = _style_name_of(part.section, look.values)
    named = (part.section, "word_style") in look.values
    el = _find_style_element(d.styles.element, name)
    if el is None and name.casefold() == "normal":      # (a template whose default paragraph style has another name)
        el = d.styles.default(WD_STYLE_TYPE.PARAGRAPH).element if d.styles.default(WD_STYLE_TYPE.PARAGRAPH) else None
    if el is None:
        el = _copy_standard_style(d, name, added, copied)
    if el is None and named:
        known = sorted({StyleFactory(s).name for s in d.styles.element.findall(_q(_W, "style"))
                        if (s.get(_q(_W, "type")) or "paragraph") == part.kind})
        import difflib
        similar = difflib.get_close_matches(name, known, 3, 0.6)
        where = f"the template {look.template.path.name}" if look.template else "Word's standard styles"
        raise ReportLookError(
            f"A style named in {look.style_file.name if look.style_file else 'the style file'} does not exist.\n"
            f"  {_at_line(_line_of(look.lines, part.section, 'word_style'))}[{part.section}]  word_style = {name}\n"
            f"      {where} has no {part.kind} style called '{name}'."
            + (f" Similar: {', '.join(repr(s) for s in similar)}." if similar else "")
            + " Write the name as Word's Styles pane shows it (for built-in styles, the English name), or leave "
              "word_style empty.")
    if el is None:                                  # one of the report's own styles: add it
        style = d.styles.add_style(name, _STYLE_TYPES[part.kind])
        style.element.set(_q(_W, "styleId"), _free_style_id(d.styles.element, style.style_id, ignore=style.element))
        added.add(style.style_id)                   # (a template may already use the id, e.g. 'TableText')
        base = styles.get(part.based_on) if part.based_on else None
        if part.kind == "paragraph":
            style.base_style = base or d.styles.default(WD_STYLE_TYPE.PARAGRAPH)
        style.quick_style = True                    # shown in Word's Styles gallery, so it is easy to restyle
        style.priority = 1
        style.hidden = False
        style.unhide_when_used = False
        return style, True
    style = StyleFactory(el)
    if style.type != _STYLE_TYPES[part.kind]:
        kinds = {WD_STYLE_TYPE.PARAGRAPH: "paragraph", WD_STYLE_TYPE.CHARACTER: "character",
                 WD_STYLE_TYPE.TABLE: "table", WD_STYLE_TYPE.LIST: "list"}
        need = {"paragraph": "a paragraph style", "character": "a character style (one that changes words inside "
                "a line, such as Strong)", "table": "a table style (such as Table Grid)"}[part.kind]
        raise ReportLookError(
            f"A style named in the style file is of the wrong kind.\n"
            f"  {_at_line(_line_of(look.lines, part.section, 'word_style'))}[{part.section}]  word_style = {name}\n"
            f"      '{name}' is a {kinds.get(style.type, '?')} style; [{part.section}] needs {need}.")
    return style, style.style_id in added


def _set_font_name(rPr, name: str) -> None:
    """A font for a style. Word's theme-font settings win over a font name, so they are removed."""
    fonts = rPr.get_or_add_rFonts()
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme", "hint"):
        fonts.attrib.pop(_q(_W, attr), None)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(_q(_W, attr), name)


def _apply_look(style, look: dict, styles_root=None) -> None:
    """Put formatting values (from the tool's own look or the style file) into a style's definition."""
    from docx.oxml import OxmlElement
    rPr = style.element.get_or_add_rPr()
    run_keys = ("font", "size", "bold", "italic", "color")
    for key, value in look.items():
        if key == "font":
            _set_font_name(rPr, value)
        elif key == "size":
            for tag in ("sz", "szCs"):
                for old in rPr.findall(_q(_W, tag)):
                    rPr.remove(old)
            if value != INHERIT:
                style.font.size = Pt(value)
                size_cs = OxmlElement("w:szCs")
                size_cs.set(_q(_W, "val"), str(int(round(value * 2))))
                rPr.find(_q(_W, "sz")).addnext(size_cs)
        elif key == "bold":
            style.font.bold = value
            style.font.cs_bold = value
        elif key == "italic":
            style.font.italic = value
            style.font.cs_italic = value
        elif key == "color":
            rPr._remove_color()
            rPr.get_or_add_color().set(_q(_W, "val"), value)
    if style.type == WD_STYLE_TYPE.TABLE and "alignment" in look:      # (a table style: where its tables sit)
        table_props = style.element.find(_q(_W, "tblPr"))
        if table_props is None:
            from docx.oxml import parse_xml
            table_props = parse_xml(f'<w:tblPr xmlns:w="{_W}"/>')
            later = next((c for c in style.element if _local(c) in ("trPr", "tcPr", "tblStylePr")), None)
            later.addprevious(table_props) if later is not None else style.element.append(table_props)
        table_props.alignment = {"left": WD_TABLE_ALIGNMENT.LEFT, "center": WD_TABLE_ALIGNMENT.CENTER,
                                 "right": WD_TABLE_ALIGNMENT.RIGHT}[look["alignment"]]
    if style.type == WD_STYLE_TYPE.PARAGRAPH:
        pf = style.paragraph_format
        for key, value in look.items():
            if key == "alignment":
                pf.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                                "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}[value]
            elif key == "space_before":
                pf.space_before = Pt(value)
            elif key == "space_after":
                pf.space_after = Pt(value)
            elif key == "line_spacing":
                if value == INHERIT:
                    spacing = style.element.get_or_add_pPr().find(_q(_W, "spacing"))
                    for attr in ("line", "lineRule"):
                        if spacing is not None:
                            spacing.attrib.pop(_q(_W, attr), None)
                else:
                    kind, amount = value
                    if kind == "multiple":
                        pf.line_spacing = float(amount)
                    else:
                        pf.line_spacing = Pt(amount)
                        pf.line_spacing_rule = WD_LINE_SPACING.AT_LEAST if kind == "at_least" else WD_LINE_SPACING.EXACTLY
            elif key == "indent_left":
                pf.left_indent = Emu(value)
            elif key == "first_line_indent":
                pf.first_line_indent = Emu(value)          # (negative = a hanging indent)
            elif key == "contextual_spacing":
                for el in style.element.get_or_add_pPr().findall(_q(_W, "contextualSpacing")):
                    el.set(_q(_W, "val"), "1" if value else "0")
    link = style.element.find(_q(_W, "link"))
    run_look = {k: v for k, v in look.items() if k in run_keys}
    if styles_root is not None and link is not None and run_look:      # e.g. Heading 1 -> Heading 1 Char:
        from docx.styles.style import StyleFactory                    # Word keeps the two in step
        linked = next((s for s in styles_root.findall(_q(_W, "style"))
                       if s.get(_q(_W, "styleId")) == link.get(_q(_W, "val"))), None)
        if linked is not None and (linked.get(_q(_W, "type")) or "") == "character":
            _apply_look(StyleFactory(linked), run_look)


def _without_numbering(style, styles_root) -> None:
    """Template heading and list styles often number their paragraphs ("1.1"). The report's paragraphs carry the
    scan's own numbers, so the numbering is switched off in the report's copy of the style."""
    el, seen = style.element, set()
    while el is not None and el.get(_q(_W, "styleId")) not in seen:
        seen.add(el.get(_q(_W, "styleId")))
        pPr = el.find(_q(_W, "pPr"))
        num = pPr.find(_q(_W, "numPr")) if pPr is not None else None
        if num is not None:
            num_id = num.find(_q(_W, "numId"))
            if num_id is not None and num_id.get(_q(_W, "val")) == "0":
                return
            own = style.element.get_or_add_pPr().get_or_add_numPr()
            for child in list(own):
                own.remove(child)
            own.get_or_add_numId().val = 0
            return
        based = el.find(_q(_W, "basedOn"))
        el = next((s for s in styles_root.findall(_q(_W, "style"))
                   if based is not None and s.get(_q(_W, "styleId")) == based.get(_q(_W, "val"))), None)


def _attach_pictures(part, xml: bytes, pictures: list, store: dict) -> bytes:
    """Give the pictures of a copied header/footer (or theme) to its new part, and point the XML at them."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part
    if not pictures:
        return xml
    new_ids = {}
    for old_id, extension, content_type, data in pictures:
        digest = hashlib.sha1(data).hexdigest()
        if digest not in store:
            store[digest] = Part(PackURI(f"/word/media/template_picture{len(store) + 1}{extension}"),
                                 content_type, data, part.package)
        new_ids[old_id] = part.relate_to(store[digest], RT.IMAGE)
    root = _safe_xml(xml, "a header of the template")
    for el in root.iter():
        for attr, value in list(el.attrib.items()):
            if (attr.startswith(f"{{{_R}}}") or attr == _q(_O, "relid")) and value in new_ids:
                el.set(attr, new_ids[value])
    return _xml_bytes(root)


def _install_template(d, t: WordTemplate) -> None:
    """Copy the checked parts of a template into a new report document."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import parse_xml
    if t.styles:
        d.part._styles_part._element = parse_xml(t.styles)
    d.part.numbering_part._element = parse_xml(t.numbering or (
        b'<w:numbering xmlns:w="' + _W.encode() + b'"/>'))       # (python-docx's own lists must not number its styles)
    if t.settings:
        settings_part = d.part._settings_part
        settings_part._element = settings_part._settings = parse_xml(t.settings)   # (python-docx keeps both)
    for rid, rel in list(d.part.rels.items()):           # Word 2010's copy of python-docx's own styles: not wanted
        if rel.reltype.endswith("/stylesWithEffects"):
            del d.part.rels[rid]
    pictures = {}
    if t.theme:
        theme_part = d.part.part_related_by(RT.THEME)
        theme_part._blob = _attach_pictures(theme_part, t.theme[0], t.theme[1], pictures)
    section = d.sections[0]
    drawing_ids = 100000
    if t.section:
        new = parse_xml(t.section)
        sect = section._sectPr
        for child in list(sect):
            sect.remove(child)
        for child in list(new):
            sect.append(child)
    for kind, hf_type, xml, story_pictures in t.stories:
        hf = {("header", "default"): section.header, ("header", "first"): section.first_page_header,
              ("header", "even"): section.even_page_header, ("footer", "default"): section.footer,
              ("footer", "first"): section.first_page_footer, ("footer", "even"): section.even_page_footer
              }.get((kind, hf_type))
        if hf is None:
            continue
        hf.is_linked_to_previous = False
        part = hf.part
        story = _safe_xml(_attach_pictures(part, xml, story_pictures, pictures), "a header of the template")
        for drawing in story.iter("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr"):
            drawing_ids += 1                         # (picture ids must not repeat those of the report's own
            drawing.set("id", str(drawing_ids))      # pictures, which python-docx numbers from 1)
        part._element = parse_xml(_xml_bytes(story))


def _setup_page(d, look: ReportLook) -> None:
    """Page size, orientation and margins: the template's (or the tool's own), then the style file's."""
    values = {k: look.values[("page", k)] for k in OTHER_SECTIONS["page"] if ("page", k) in look.values}
    if look.template is None:
        values = {**BUILT_IN_PAGE, **values}
    for section in d.sections:
        if section.page_width is None or section.page_height is None:     # (a template without a page size)
            section.page_width, section.page_height = Emu(PAPER_SIZES["letter"][0]), Emu(PAPER_SIZES["letter"][1])
        for side, default in (("top", 914400), ("bottom", 914400), ("left", 914400), ("right", 914400),
                              ("header_distance", 457200), ("footer_distance", 457200), ("gutter", 0)):
            name = side if side in ("header_distance", "footer_distance", "gutter") else f"{side}_margin"
            if getattr(section, name) is None:                           # (... or without margins: Word's own)
                setattr(section, name, Emu(default))
        landscape = values.get("orientation", "landscape" if section.page_width > section.page_height else "portrait") \
            == "landscape"
        if "paper" in values:
            width, height, code = PAPER_SIZES[values["paper"]]
            section.page_width, section.page_height = Emu(width), Emu(height)
            section._sectPr.get_or_add_pgSz().set(_q(_W, "code"), str(code))
        if (section.page_width > section.page_height) != landscape:
            section.page_width, section.page_height = section.page_height, section.page_width
        section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        for key in ("margin_top", "margin_bottom", "margin_left", "margin_right"):
            if key in values:
                setattr(section, key.replace("margin_", "") + "_margin", Emu(values[key]))


def _same_style_conflicts(look: ReportLook, styles: dict) -> None:
    """Like _shared_style_conflicts, for parts that turn out to use the same style although the style file names it
    differently ('Heading1' is the id of 'Heading 1'; 'Normal' may be the template's 'Body')."""
    by_id = defaultdict(list)
    for part in REPORT_PARTS:
        by_id[styles[part.section].style_id].append(part.section)
    problems = []
    for sections in by_id.values():
        for i, first in enumerate(sections):
            for second in sections[i + 1:]:
                for key in PARAGRAPH_KEYS[1:]:
                    a, b = look.values.get((first, key)), look.values.get((second, key))
                    if a is not None and b is not None and a != b:
                        problems.append((_line_of(look.lines, second, key), f"[{second}]  {key}",
                                         f"[{first}] and [{second}] both use the style '{styles[first].name}' but give "
                                         f"it a different {key}. Give [{second}] a style of its own (word_style), or "
                                         f"set the {key} in one place."))
    if problems:
        raise ReportLookError(_style_file_message(look.style_file or Path("report_style.ini"), problems))


def _report_document(look: ReportLook) -> tuple:
    """A new, empty report document with every style it uses, set up as the look says: (document, {part: style}).
    Order of precedence: the template (or python-docx's standard document), then the tool's own look - applied
    when there is no template, and to styles the tool has to add to one - then the [document] font, then every
    value filled in the style file."""
    d = docx.Document()
    if look.template is not None:
        _install_template(d, look.template)
    styles, added_ids, copied = {}, set(), {}
    for part in REPORT_PARTS:
        style, added = _report_style(d, part, look, styles, added_ids, copied)
        styles[part.section] = style
        own = _style_name_of(part.section, look.values).casefold() == part.style.casefold()
        if own and (look.template is None or added):
            _apply_look(style, dict(part.look), d.styles.element)
    _same_style_conflicts(look, styles)
    font = look.values.get(("document", "font"))
    if font:
        defaults = d.styles.element.find(_q(_W, "docDefaults"))
        rpr_default = defaults.find(_q(_W, "rPrDefault")) if defaults is not None else None
        if rpr_default is not None and rpr_default.find(_q(_W, "rPr")) is not None:
            _set_font_name(rpr_default.find(_q(_W, "rPr")), font)
        for style in {id(s): s for s in styles.values()}.values():
            if style.type != WD_STYLE_TYPE.TABLE:
                _apply_look(style, {"font": font}, d.styles.element)
    for part in REPORT_PARTS:
        changes = {k: look.values[(part.section, k)] for k in PARAGRAPH_KEYS[1:] if (part.section, k) in look.values}
        if changes and part.kind != "table":
            _apply_look(styles[part.section], changes, d.styles.element)
        if part.kind == "paragraph":
            _without_numbering(styles[part.section], d.styles.element)
    _setup_page(d, look)
    zoom = d.settings.element.find(_q(_W, "zoom"))
    if zoom is not None and zoom.get(_q(_W, "percent")) is None:    # (python-docx's own settings leave out this
        zoom.set(_q(_W, "percent"), "100")                          # required value)
    properties = d.core_properties                      # nothing of python-docx's or the template's own
    for child in list(properties._element):
        properties._element.remove(child)
    now = _dt.datetime.now(_dt.timezone.utc).replace(microsecond=0)
    properties.title = "Extraction Verification Report"
    properties.revision = 1
    properties.created = properties.modified = now
    return d, styles


# The finished report is checked before it is saved: whatever happened above, a report that could make Word reach
# outside this computer is never written. Only the kinds of relationship a report needs are allowed.
_REPORT_RELATIONSHIPS = {"officeDocument", "core-properties", "extended-properties", "thumbnail", "styles",
                         "stylesWithEffects", "settings", "webSettings", "fontTable", "theme", "numbering", "customXml",
                         "customXmlProps", "header", "footer", "image", "footnotes", "endnotes"}
_FORBIDDEN_ELEMENTS = {"altChunk", "subDoc", "object", "control", "attachedTemplate", "mailMerge", "frameset",
                       "OLEObject", "hyperlink", "hlinkClick", "hlinkHover", "dataBinding"}


def outside_references(data: bytes) -> list:
    """Everything in a Word file (as bytes) that could make Word reach outside this computer, or run something,
    when the file is opened, printed or its fields are updated - as sentences. [] means none."""
    import zipfile
    problems = []
    try:
        z = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        return ["it is not a Word file"]
    names = z.namelist()
    if "[Content_Types].xml" in names:
        types = z.read("[Content_Types].xml").lower()
        for bad in (b"macroenabled", b"vbaproject", b"activex", b"oleobject"):
            if bad in types:
                problems.append(f"it declares {bad.decode()} content")
    for name in names:
        low = name.lower()
        if low.endswith(".rels"):
            folder = _part_folder(name.replace("_rels/", "").rsplit(".rels", 1)[0])
            for rel in _safe_xml(z.read(name), "the report"):
                kind = (rel.get("Type") or "").rsplit("/", 1)[-1]
                outside = _reaches_outside(rel.get("Target") or "", folder)
                if (rel.get("TargetMode") or "").lower() == "external" or outside:
                    problems.append(f"{name}: a {kind} relationship points outside the file ({(rel.get('Target') or '')[:80]})")
                elif kind not in _REPORT_RELATIONSHIPS:
                    problems.append(f"{name}: a {kind} relationship, which a report never needs")
        elif low.endswith(".xml"):
            content = z.read(name)
            if not any(k in content for k in (b"fldChar", b"fldSimple", b"instrText", b"InstrText", b"altChunk",
                                              b"subDoc",
                                              b"object", b"control", b"attachedTemplate", b"mailMerge", b"frameset",
                                              b"OLEObject", b"href", b"hlink", b"dataBinding", b"HyperlinkBase",
                                              b"src=")):
                continue                                   # (most parts: nothing to look at closely)
            root = _safe_xml(content, "the report")
            for code in _field_codes(root):
                if _field_name(code) not in SAFE_FIELDS:
                    problems.append(f"{name}: a {_field_name(code) or '?'} field ({code.strip()[:60]})")
            for el in root.iter():
                local = _local(el)
                if local in _FORBIDDEN_ELEMENTS and (el.tag.startswith(f"{{{_W}}}") or el.tag.startswith(f"{{{_O}}}")
                                                     or local in ("hlinkClick", "hlinkHover")):
                    problems.append(f"{name}: a <{local}> element")
                if local == "HyperlinkBase" and (el.text or "").strip():
                    problems.append(f"{name}: a hyperlink base ({el.text.strip()[:60]})")
                for attr in ("href", _q(_O, "href"), "src"):
                    if el.get(attr) and el.tag.startswith(f"{{{_V}}}"):
                        problems.append(f"{name}: a picture linked to {el.get(attr)[:80]}")
    return problems


# ---- 11e. Writing the report -------------------------------------------------------------------------------------
def _docx_text(text: str) -> str:
    return ILLEGAL_CHARACTERS_RE.sub("", text or "")


def _shade_cell(cell, hex_fill: str):
    """Cell shading, put where Word expects it among the cell's properties."""
    from docx.oxml import OxmlElement
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(_q(_W, "shd")):
        tc_pr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(_q(_W, "val"), "clear")
    shd.set(_q(_W, "color"), "auto")
    shd.set(_q(_W, "fill"), hex_fill)
    before = [tc_pr.find(_q(_W, t)) for t in ("tcBorders", "vMerge", "hMerge", "gridSpan", "tcW", "cnfStyle")]
    anchor = next((el for el in before if el is not None), None)
    if anchor is not None:
        anchor.addnext(shd)
    else:
        tc_pr.insert(0, shd)


def _highlight(name: str):
    return getattr(WD_COLOR_INDEX, name.upper().replace(" ", "_"))


def _mark_review(run, level: str, look: ReportLook) -> None:
    """The review marks (always put on the words themselves, so no template can hide them)."""
    if level == "check":
        run.font.highlight_color = _highlight(look.get("review marks", "check_highlight"))
    elif level == "likely_wrong":
        run.font.highlight_color = _highlight(look.get("review marks", "likely_wrong_highlight"))
        run.font.color.rgb = RGBColor.from_string(likely_wrong_colour(look.values))
        run.bold = True


def _styled_run(paragraph, text: str, style):
    """A run in a character style (a tag, a keyword) that looks as that style says in any paragraph. Bold and
    italic are 'toggles' in Word: a bold character style inside a bold paragraph style would turn the words NOT
    bold, and 'not italic' cannot undo an italic paragraph style. So the style's own bold and italic are also put on
    the run itself."""
    run = paragraph.add_run(text, style=style)
    if style.font.bold is not None:
        run.bold = style.font.bold
    if style.font.italic is not None:
        run.italic = style.font.italic
    return run


def _keyword_tokens(tokens: list, keyword_re) -> set:
    """Positions of the tokens that are (part of) a requirement keyword."""
    if keyword_re is None or not tokens:
        return set()
    starts, text = [], ""
    for token in tokens:
        starts.append(len(text))
        text += token[0] + " "
    marked = set()
    for m in keyword_re.finditer(text):
        marked.update(i for i, s in enumerate(starts) if s < m.end() and s + len(tokens[i][0]) > m.start())
    return marked


def _add_confidence_runs(paragraph, tokens: list, settings: Settings, look: ReportLook,
                         keyword_re=None, keyword_style=None):
    """Write words, highlighting the ones the OCR was not sure about (and, if asked, the requirement keywords)."""
    keywords = _keyword_tokens(tokens, keyword_re)
    group, group_key = [], None

    def emit():
        if group:
            text = _docx_text(" ".join(group)) + " "
            run = _styled_run(paragraph, text, keyword_style) if group_key[1] else paragraph.add_run(text)
            _mark_review(run, group_key[0], look)

    for i, token in enumerate(tokens):
        key = (token_attention(token, settings), i in keywords)
        if key != group_key:
            emit()
            group, group_key = [], key
        group.append(token[0])
    emit()


def _cell_attention(t: "TableData", r: int, c: int, ctx: "TextContext", settings: Settings) -> bool:
    words = _cell_words(t, r, c)
    if not words:
        text = t.grid[r][c] if r < len(t.grid) and c < len(t.grid[r]) else ""
        conf = t.conf[r][c] if r < len(t.conf) and c < len(t.conf[r]) else 1.0
        words = [[w, conf] for w in text.split()]
    tokens = word_tokens((Word(w, cf, (0, 0, 0, 0)) for w, cf in words), ctx)
    return any(token_attention(tok, settings) for tok in tokens)


_WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
_CSS_LENGTH = {"pt": 12700, "in": 914400, "cm": 360000, "mm": 36000, "px": 9525, "": 12700}


def _length_emu(value, unit: str = "") -> int:
    """A size from a picture's settings, in EMU; 0 when it cannot be read (a damaged file must not stop a run)."""
    try:
        return int(float(value) * _CSS_LENGTH.get(unit, 0)) if unit != "emu" else int(float(value))
    except (TypeError, ValueError):
        return 0


def _block_height(el) -> int:
    """Rough height (EMU) of a paragraph or table of a header or footer: a line of text, or the tallest picture
    placed in the line (DrawingML or the older VML) when that is taller; a table row is as tall as its tallest cell.
    Floating pictures and text boxes take no room, and Word's second copy of a drawing (mc:Fallback) is skipped."""
    if _local(el) == "tbl":
        return sum(max([sum(_block_height(b) for b in cell if _local(b) in ("p", "tbl"))
                        for cell in row.iter(_q(_W, "tc"))] or [0])
                   for row in el if _local(row) == "tr")
    tallest = 228600                                          # a line of text: about 0.25 in
    stack = list(el)
    while stack:
        child = stack.pop()
        if not isinstance(child.tag, str) or _local(child) in ("anchor", "Fallback", "txbxContent", "textbox"):
            continue
        if child.tag == f"{{{_WP}}}inline":
            extent = child.find(f"{{{_WP}}}extent")
            tallest = max(tallest, _length_emu(extent.get("cy") if extent is not None else 0, "emu"))
            continue
        style = child.get("style") if child.tag.startswith(f"{{{_V}}}") else None
        if style and "position:absolute" not in style:
            m = re.search(r"(?:^|;)\s*height:\s*(\d+(?:\.\d+)?|\.\d+)\s*(pt|in|cm|mm|px|)\s*(?:;|$)", style)
            if m:
                tallest = max(tallest, _length_emu(m.group(1), m.group(2)))
            continue
        stack.extend(child)
    return tallest


def _story_height(story) -> int:
    """Rough height (EMU) of a header or footer: its paragraphs and tables one under the other."""
    return sum(_block_height(el) for el in story if _local(el) in ("p", "tbl"))


def _on(el) -> bool:
    """A Word on/off setting: on when present, unless its value says off."""
    return el is not None and (el.get(_q(_W, "val")) or "true").lower() in ("1", "true", "on")


def _usable_page(d) -> tuple:
    """Width and height (inches) of the text area on a page of the report: the page less its margins and gutter,
    and less any part of a header or footer that reaches into the text area (a tall logo pushes the text down)."""
    section = d.sections[0]
    settings = d.settings.element
    gutter_at_top = _on(settings.find(_q(_W, "gutterAtTop")))
    gutter = section.gutter or 0
    width = section.page_width - section.left_margin - section.right_margin - (0 if gutter_at_top else gutter)
    height = section.page_height - abs(section.top_margin) - abs(section.bottom_margin) - (gutter if gutter_at_top else 0)
    even = _on(settings.find(_q(_W, "evenAndOddHeaders")))
    for stories, distance, margin in (((section.header, section.even_page_header), section.header_distance,
                                       section.top_margin),
                                      ((section.footer, section.even_page_footer), section.footer_distance,
                                       section.bottom_margin)):
        reach = 0
        for k, hf in enumerate(stories):
            if (k == 1 and not even) or hf.is_linked_to_previous:
                continue
            reach = max(reach, (distance or 0) + _story_height(hf._element) - abs(margin or 0))
        height -= max(0, reach)
    return width / 914400, height / 914400


def write_word_report(path: Path, doc: DocResult, items: list, settings: Settings) -> Path:
    look = report_look(settings)
    d, styles = _report_document(look)
    usable_w, usable_h = _usable_page(d)
    table_alignment = look.values.get(("tables", "alignment"))     # (else the table style decides)
    tag_start = look.get("requirements", "tag_position") == "start"
    keyword_re = None
    if look.get("requirements", "mark_keywords"):
        words = sorted(settings.keywords, key=len, reverse=True)
        keyword_re = re.compile(r"(?<![A-Za-z])(?:" + "|".join(r"\s+".join(map(re.escape, w.split())) for w in words)
                                + r")(?![A-Za-z])", re.I)
    header_fill = look.values.get(("table header row", "fill")) or (BUILT_IN_HEADER_FILL if look.template is None
                                                                     else "none")
    style_ids = {section: style.style_id for section, style in styles.items()}

    tags = defaultdict(list)
    requirement_blocks = set()
    for it in items:
        for uid in it.block_uids:
            tags[uid].append(it.item_id)
            if it.item_type == "Requirement":
                requirement_blocks.add(uid)
    vocab = document_text_context(doc, settings)

    def paragraph(section: str, text: str = "", page_break: bool = False):
        p = d.add_paragraph()
        p._p.style = style_ids[section]                     # (faster than by style object, for long reports)
        if text:
            p.add_run(_docx_text(text))
        if page_break:
            p.paragraph_format.page_break_before = True     # (never an extra empty paragraph: no blank pages)
        return p

    def tag(p, ids: list, where: str, alone: bool = False):
        if ids and (where == "start") == tag_start:
            text = "[" + ", ".join(ids) + "]"
            if not alone:                                   # (a tag on a line of its own has no spaces around it)
                text = text + "  " if tag_start else "  " + text
            _styled_run(p, text, styles["requirement tag"])

    def picture(image: Path, width: float):
        p = d.add_paragraph()                               # a picture: centred, no indent, single spacing (a
        pf = p.paragraph_format                             # template's fixed line height would cut it off)
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.left_indent = pf.right_indent = pf.first_line_indent = Emu(0)
        pf.line_spacing = 1.0
        p.add_run().add_picture(str(image), width=Inches(width))

    paragraph("title", "Extraction Verification Report")
    p = paragraph("body")
    p.add_run("Source: ").bold = True
    p.add_run(_docx_text(Path(doc.source).name))
    p.add_run(f"\nPages: {len(doc.pages)}    Engine: {doc.engine}    Created: {_dt.datetime.now():%Y-%m-%d %H:%M}")
    _software, differences = configuration_status(docling_used=doc.engine == "docling")
    p.add_run(f"\nTool: spec_extractor.py {__version__} (build {build_id()})    Tested configuration: "
              + ("yes" if not differences else "NO - see 'Run Info' in the Excel file or spec_extractor_log.txt"))
    n_req = sum(1 for it in items if it.item_type == "Requirement")
    p.add_run(f"\nRequirements: {n_req}    Tables: {sum(1 for it in items if it.item_type == 'Table')}    "
              f"Captions: {sum(1 for it in items if 'Caption' in it.item_type)}    "
              f"Items needing review: {sum(1 for it in items if it.needs_review)}")
    p.add_run(f"\nReport look: {_docx_text(look.summary)}")
    paragraph("heading 2", "How to read this report")
    legend = paragraph("body")
    _mark_review(legend.add_run("words marked like this"), "check", look)
    legend.add_run(" = check against the scan: the word failed a check (not in the dictionary, accent marks, "
                   f"stray punctuation, a near-duplicate ID) or is a number read with OCR confidence below "
                   f"{settings.review_confidence:.2f};  ")
    _mark_review(legend.add_run("words marked like this"), "likely_wrong", look)
    legend.add_run(" = very low OCR confidence (likely misread);  ")
    _styled_run(legend, "[R-0001]", styles["requirement tag"])
    legend.add_run(" = the row ID in the Excel workbook that this text became;  ")
    if keyword_re is not None:
        _styled_run(legend, "shall", styles["requirement keywords"])
        legend.add_run(" = a requirement keyword;  ")
    legend.add_run("[header] / [footer] = page header/footer text read from the scan.")
    paragraph("body", "Suggested check: compare each page image with the text below it. Every sentence with a "
                      "requirement keyword should carry an [R-...] tag; tables should match cell by cell.")

    blocks_by_page = defaultdict(list)
    for b in doc.blocks:
        blocks_by_page[b.page].append(b)
    furniture_rank = {"page_header": 0, "page_footer": 2}      # running header first, footer last on each page
    for page_blocks in blocks_by_page.values():
        page_blocks.sort(key=lambda blk: furniture_rank.get(blk.kind, 1))     # (stable: reading order is kept)
    cache_dir = Path(doc.cache_dir) if doc.cache_dir else None

    for page in doc.pages:
        paragraph("heading 1", f"Page {page.number}", page_break=True)
        if page.notes:
            paragraph("notes", "Notes: " + "; ".join(page.notes))
        if page.notes and any(n.startswith("FAILED") for n in page.notes):
            paragraph("notes").add_run("This page could not be read - check the scan manually.").bold = True
        if settings.word_page_images and page.image_file and cache_dir and (cache_dir / page.image_file).exists():
            aspect = page.height / max(page.width, 1)
            # (room is kept for the "Page N" heading and a line of notes above the picture)
            picture(cache_dir / page.image_file, max(1.0, min(usable_w, (usable_h - 1.2) / aspect)))
            paragraph("heading 2", f"Page {page.number} - text read by the tool", page_break=True)
        for b in blocks_by_page.get(page.number, []):
            ids = tags.get(b.uid, [])
            if b.kind == "heading":
                level = min(4, 1 + int(b.meta.get("level", 1)))
                paragraph(f"heading {level}", b.text)
            elif b.kind in ("page_header", "page_footer"):
                para = paragraph("scanned header footer")
                tag(para, ids, "start")
                para.add_run(("[header] " if b.kind == "page_header" else "[footer] ") + _docx_text(b.text))
                tag(para, ids, "end")                       # (a repeated notice that was reported as a requirement)
            elif b.kind == "table" and b.table is not None:
                t = b.table
                if b.meta.get("caption") and not any(x.kind == "caption" and x.meta.get("caption_of") == b.uid
                                                     for x in blocks_by_page[page.number]):
                    paragraph("captions", b.meta["caption"])
                if ids and tag_start:
                    tag(paragraph("body"), ids, "start", alone=True)
                tbl = d._body.add_table(t.n_rows, t.n_cols, Emu(int(usable_w * 914400)))   # (the text width,
                tbl.style = styles["tables"]                                                # less any gutter)
                if table_alignment:
                    tbl.alignment = {"left": WD_TABLE_ALIGNMENT.LEFT, "center": WD_TABLE_ALIGNMENT.CENTER,
                                     "right": WD_TABLE_ALIGNMENT.RIGHT}[table_alignment]
                table_look = tbl._tbl.tblPr.find(_q(_W, "tblLook"))
                if table_look is not None:                  # a heading row; no first-column or banded-row look
                    table_look.attrib.pop(_q(_W, "val"), None)   # from a template's table style
                    for attr, value in (("firstRow", "1"), ("lastRow", "0"), ("firstColumn", "0"),
                                        ("lastColumn", "0"), ("noHBand", "1"), ("noVBand", "1")):
                        table_look.set(_q(_W, attr), value)
                hidden = set()
                table_rows = list(tbl.rows)       # (Table.cell() re-reads the whole table on every call: slow)
                for r0, c0, r1, c1 in t.merges:
                    if r1 < t.n_rows and c1 < t.n_cols and (r1 > r0 or c1 > c0):
                        try:
                            table_rows[r0].cells[c0].merge(table_rows[r1].cells[c1])
                        except Exception:
                            continue
                        hidden |= {(rr, cc) for rr in range(r0, r1 + 1) for cc in range(c0, c1 + 1)} - {(r0, c0)}
                if look.get("table header row", "repeat_on_each_page"):
                    from docx.oxml import OxmlElement
                    for r_i in range(min(t.header_rows, t.n_rows)):     # Word's "Repeat as header row"
                        tr_pr = table_rows[r_i]._tr.get_or_add_trPr()
                        if tr_pr.find(_q(_W, "tblHeader")) is None:
                            tr_pr.append(OxmlElement("w:tblHeader"))
                grid_cells = tbl._cells           # one cell object per grid position, read once
                for r_i, row in enumerate(t.grid):
                    heading_row = r_i < t.header_rows
                    for c_i in range(t.n_cols):
                        if (r_i, c_i) in hidden:
                            continue
                        cell = grid_cells[r_i * t.n_cols + c_i]
                        cell.text = ""
                        cell_p = cell.paragraphs[0]
                        cell_p._p.style = style_ids["table header row" if heading_row else "table text"]
                        cell_p.add_run(_docx_text(row[c_i] if c_i < len(row) else ""))
                        if heading_row:
                            if header_fill != "none":
                                _shade_cell(cell, header_fill)
                        elif _cell_attention(t, r_i, c_i, vocab, settings):
                            _shade_cell(cell, look.get("review marks", "cell_fill"))
                if ids and not tag_start:
                    tag(paragraph("body"), ids, "end", alone=True)
            elif b.kind == "figure":
                img_file = b.meta.get("image_file")
                if img_file and cache_dir and (cache_dir / img_file).exists():
                    fw = max(1, b.box[2] - b.box[0]) / max(page.dpi, 1)
                    picture(cache_dir / img_file, max(0.5, min(usable_w, max(1.5, fw))))
                else:
                    paragraph("notes", "[figure]")
                if b.lines:
                    fp = paragraph("table text")
                    fp.add_run("Text inside figure: ").italic = True
                    _add_confidence_runs(fp, word_tokens((w for ln in b.lines for w in ln.words), vocab), settings, look)
            else:
                requirement = b.uid in requirement_blocks
                section = ("list items" if b.kind == "list_item" else "captions" if b.kind == "caption"
                           else "requirements" if requirement else "body")
                para = paragraph(section)
                tag(para, ids, "start")
                tokens = block_tokens(b, vocab) if b.lines else [(w, 1.0, "") for w in b.text.split()]
                _add_confidence_runs(para, tokens, settings, look, keyword_re if requirement else None,
                                     styles["requirement keywords"])
                if b.meta.get("continued_across_pages"):
                    label = para.add_run(" [continues across pages]")
                    label.italic = True
                    label.font.size = styles["notes"].font.size or Pt(8)
                tag(para, ids, "end")
    return _save_word_file(d, path)


def _save_word_file(d, path: Path) -> Path:
    """Save a report - only after checking that nothing in it points outside this computer."""
    buffer = io.BytesIO()
    d.save(buffer)
    data = buffer.getvalue()
    problems = outside_references(data)
    if problems:
        raise RuntimeError("the Word report was not saved, because it would point outside this computer: "
                           + "; ".join(problems[:3]))
    try:
        Path(path).write_bytes(data)
        return path
    except PermissionError:
        alt = path.with_name(f"{path.stem}_{_dt.datetime.now():%H%M%S}{path.suffix}")
        LOG.warning("%s is open in another program - saved as %s instead", path.name, alt.name)
        alt.write_bytes(data)
        return alt


# ---- 11f. The starter template -----------------------------------------------------------------------------------
def write_starter_template(target: Path, look: ReportLook) -> Path:
    """A Word file with every style of the report and a sample of each part, made with the current look (the
    template and style file, if any). Restyle it in Word and name it as the template in report_style.ini.
    `target` is a .docx file name or a folder (then report_template.docx, or report_template (2).docx ...).
    An existing file is never overwritten."""
    target = Path(target)
    if not target.exists() and not target.suffix:          # a folder that is not there yet (e.g. word_style)
        target.mkdir(parents=True, exist_ok=True)
    if target.is_dir():
        path, n = target / "report_template.docx", 2
        while path.exists():
            path, n = target / f"report_template ({n}).docx", n + 1
    else:
        path = target
        if path.suffix.lower() != ".docx":
            raise ReportLookError(f"Write the starter template as a .docx file (not '{path.name}'): Word opens a "
                                  ".docx for editing, while double-clicking a .dotx makes a new document instead.")
        if path.exists():
            raise ReportLookError(f"{path} already exists and is not overwritten. Choose another name, or give a "
                                  "folder to get the next free name.")
        path.parent.mkdir(parents=True, exist_ok=True)
    d, styles = _report_document(look)
    style_file = look.style_file.name if look.style_file else "report_style.ini"

    def paragraph(section: str, text: str = ""):
        p = d.add_paragraph(style=styles[section])
        if text:
            p.add_run(text)
        return p

    def label(part: ReportPart):
        paragraph("notes", f"[{part.section}] - style '{styles[part.section].name}': {part.about}")

    paragraph("title", "Word template for the verification report")
    for text in (
            "This file holds the look of the Word reports. Change it here in Word and every report follows:",
            "1. Styles: open the Styles pane (Home tab: the small arrow at the bottom right of the Styles group, or "
            "Alt+Ctrl+Shift+S), right-click a style and choose Modify: font, size, colour, spacing, indents.",
            "2. A logo, a classification marking or page numbers: double-click the top or bottom of the page and edit "
            "the header or footer. Never type them in the text: the text of this file is not copied into reports.",
            "3. Page size and margins: the Layout tab.",
            f"4. Save it (keep the .docx format) in the word_style folder and write its name in {style_file}:   "
            f"template = {path.name}",
            f"Values filled in {style_file} still win over this file: empty them to let this file decide."):
        paragraph("body", text)
    by_section = {p.section: p for p in REPORT_PARTS}
    label(by_section["title"])
    paragraph("title", "Extraction Verification Report")
    for section, text in (("heading 1", "Page 12"), ("heading 2", "Page 12 - text read by the tool"),
                          ("heading 3", "3.2 Electrical requirements"), ("heading 4", "3.2.1 Power input")):
        label(by_section[section])
        paragraph(section, text)
    label(by_section["body"])
    paragraph("body", "This is ordinary text read from the scan. It shows how a paragraph of the scanned document "
                      "looks in the report, with the spacing between paragraphs.")
    label(by_section["list items"])
    for text in ("a. The first entry of a list.", "b. The second entry of a list."):
        paragraph("list items", text)
    label(by_section["requirements"])
    p = paragraph("requirements")
    # (the samples carry the styles only, so that changing a style in Word shows on them at once)
    if look.get("requirements", "tag_position") == "start":
        p.add_run("[R-0001]  ", style=styles["requirement tag"])
    p.add_run("3.2.1.1 Output voltage. The power supply ")
    p.add_run("shall", style=styles["requirement keywords"])
    p.add_run(" deliver 28 V DC within 0.5 V.")
    if look.get("requirements", "tag_position") != "start":
        p.add_run("  [R-0001]", style=styles["requirement tag"])
    for section in ("requirement tag", "requirement keywords"):
        label(by_section[section])
    label(by_section["captions"])
    paragraph("captions", "TABLE I. Power characteristics")
    label(by_section["tables"])
    label(by_section["table header row"])
    label(by_section["table text"])
    table = d._body.add_table(3, 3, Emu(int(_usable_page(d)[0] * 914400)))
    table.style = styles["tables"]
    alignment = look.values.get(("tables", "alignment"))
    if alignment:
        table.alignment = {"left": WD_TABLE_ALIGNMENT.LEFT, "center": WD_TABLE_ALIGNMENT.CENTER,
                           "right": WD_TABLE_ALIGNMENT.RIGHT}[alignment]
    # The tool's own heading-row shading goes into the table style here (a report puts it on the cells), so that it
    # can be changed in Word, and a report made with this file as its template shades its heading rows the same way.
    # A shading set in the style file is put on the cells, as in the reports.
    explicit_fill = look.values.get(("table header row", "fill"))
    fill = BUILT_IN_HEADER_FILL if look.template is None and explicit_fill is None else "none"
    table_style = styles["tables"].element
    if fill != "none" and not any(el.get(_q(_W, "type")) == "firstRow" for el in table_style.findall(_q(_W, "tblStylePr"))):
        from docx.oxml import parse_xml
        table_style.append(parse_xml(f'<w:tblStylePr xmlns:w="{_W}" w:type="firstRow"><w:tcPr><w:shd w:val="clear" '
                                     f'w:color="auto" w:fill="{fill}"/></w:tcPr></w:tblStylePr>'))
    for r, row in enumerate((("Parameter", "Value", "Method"), ("Voltage", "28 V", "Test"),
                             ("Ripple", "0.5 V", "Analysis"))):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.paragraphs[0].style = styles["table header row" if r == 0 else "table text"]
            cell.paragraphs[0].add_run(text)
            if r == 0 and explicit_fill not in (None, "none"):
                _shade_cell(cell, explicit_fill)
            elif r == 2 and c == 1:
                _shade_cell(cell, look.get("review marks", "cell_fill"))
    label(by_section["scanned header footer"])
    paragraph("scanned header footer", "[header] SPEC-1234 Rev. A")
    paragraph("scanned header footer", "[footer] Page 3 of 40")
    label(by_section["notes"])
    paragraph("notes", "Notes: page turned 180 degrees")
    paragraph("notes", "[review marks] - put on the words themselves, set in the style file:")
    p = paragraph("body", "Words to check are ")
    _mark_review(p.add_run("marked like this"), "check", look)
    p.add_run(", and words that were very likely misread ")
    _mark_review(p.add_run("like this"), "likely_wrong", look)
    p.add_run(". Table cells holding a word to check are shaded as in the table above.")
    d.core_properties.title = "Word template for the verification report"
    return _save_word_file(d, path)


# =============================================================================
# 12. BATCH RUNNER, CACHING AND COMMAND LINE
# =============================================================================
def _file_sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _imwrite(path: Path, img: "np.ndarray", params: Optional[list] = None) -> bool:
    """cv2.imwrite cannot write to non-ASCII paths on Windows; encode in memory, then write the bytes.
    Returns False (with a warning) instead of stopping when the file cannot be written."""
    try:
        ok, buf = cv2.imencode(Path(path).suffix or ".png", img, params or [])
        if ok:
            buf.tofile(str(path))
        return bool(ok)
    except (OSError, cv2.error) as exc:
        LOG.warning("  could not write %s (%s)", Path(path).name, exc)
        return False


def _imread_gray(path: Path) -> "np.ndarray":
    return cv2.imdecode(np.fromfile(str(path), dtype=np.uint8), cv2.IMREAD_GRAYSCALE)


def _save_jpeg(gray: "np.ndarray", path: Path, scale: float, quality: int = 70) -> bool:
    img = cv2.resize(gray, None, fx=scale, fy=scale, interpolation=cv2.INTER_AREA) if scale < 1 else gray
    return _imwrite(path, img, [cv2.IMWRITE_JPEG_QUALITY, quality])


class SpecExtractor:
    """Runs the whole pipeline for one file at a time; OCR models are loaded once."""

    def __init__(self, settings: Settings):
        self.settings = settings
        configure_keywords(settings.keywords)
        self._ocr = None
        self._orientation = None
        self._docling = None

    @property
    def ocr(self) -> RapidOcrEngine:
        if self._ocr is None:
            LOG.info("Loading OCR model (PP-OCRv6 via RapidOCR)...")
            self._ocr = RapidOcrEngine(self.settings)
        return self._ocr

    @property
    def orientation(self) -> OrientationDetector:
        if self._orientation is None:
            self._orientation = OrientationDetector()
        return self._orientation

    def process_file(self, path: Path, out_dir: Path, output_stem: str) -> tuple:
        s = self.settings
        start = time.time()
        page_key = hashlib.sha1(s.page_selection.encode()).hexdigest()[:6] if s.page_selection else "all"
        safe_stem = re.sub(r"[^A-Za-z0-9._-]+", "_", path.stem)[:40]
        cache_dir = out_dir / "_cache" / f"{safe_stem}_{_file_sha256(path)[:12]}_{s.ocr_fingerprint()}_{page_key}"
        result_file = cache_dir / "result.json"
        doc = None
        if s.use_cache and result_file.exists():
            try:
                doc = DocResult.from_json(json.loads(result_file.read_text(encoding="utf-8")))
                doc.cache_dir = str(cache_dir)
                doc.source = str(path)            # the same file may have been moved, renamed or copied
                LOG.info("  using cached OCR results (delete the _cache folder or use --no-cache to re-read)")
            except Exception as exc:
                LOG.warning("  cache unreadable (%s) - reading the document again", exc)
                doc = None
        if doc is None:
            cache_dir.mkdir(parents=True, exist_ok=True)
            doc = self._read_document(path, cache_dir)
            failed = [p.number for p in doc.pages if any(str(n).startswith("FAILED") for n in p.notes)
                      or (p.width and not p.image_file)]            # a page or its report image could not be saved
            if failed:
                LOG.info("  results not cached because page(s) %s were not completely read or saved - the next run "
                         "tries again", _pages_text(failed))
            elif doc.engine == s.engine or not doc.engine.startswith("standard ("):
                try:                                 # (a Docling fallback run is not cached as a Docling result)
                    tmp = result_file.with_suffix(".tmp")
                    tmp.write_text(json.dumps(doc.to_json(), ensure_ascii=False), encoding="utf-8")
                    os.replace(tmp, result_file)
                except OSError as exc:
                    LOG.warning("  could not save the OCR cache (%s) - the results are still written", exc)
        finalize_document(doc, s)
        items = RequirementExtractor(s).extract(doc)
        outputs = {}
        if s.make_excel:
            outputs["excel"] = write_workbook(out_dir / f"{output_stem}_extracted.xlsx", [(doc, items)], s)
        if s.make_word_report:
            outputs["word"] = write_word_report(out_dir / f"{output_stem}_verification.docx", doc, items, s)
        failed_pages = [p.number for p in doc.pages if any(str(n).startswith("FAILED") for n in p.notes)]
        LOG.info("  done in %.0f s: %d requirements, %d tables, %d captions, %d items flagged for review%s",
                 time.time() - start, sum(1 for i in items if i.item_type == "Requirement"),
                 sum(1 for i in items if i.item_type == "Table"), sum(1 for i in items if "Caption" in i.item_type),
                 sum(1 for i in items if i.needs_review),
                 f" - {len(failed_pages)} page(s) could NOT be read: {_pages_text(failed_pages)} (see Page Quality)"
                 if failed_pages else "")
        return doc, items, outputs

    def _read_document(self, path: Path, cache_dir: Path) -> DocResult:
        s = self.settings
        doc = DocResult(source=str(path), engine=s.engine, cache_dir=str(cache_dir))
        t0 = time.time()
        total = len(parse_page_selection(s.page_selection, count_pages(path)))
        if total == 0:
            doc.warnings.append(f"{path.name}: no pages selected")
            LOG.warning("  no pages selected (--pages %s, the file has %d page(s))", s.page_selection, count_pages(path))
            return doc
        docling_pages = []     # (page number, dpi, height, width, cleaned image file) - images stay on disk
        for n, raw in enumerate(iter_raw_pages(path, s), 1):
            page_start = time.time()
            if raw.error:
                self._page_failed(doc, raw.number, raw.error)
                continue
            try:
                prep = prepare_page(raw, s, self.orientation, self.ocr)
                h, w = prep.ocr_image.shape[:2]
                info = PageInfo(number=raw.number, width=int(w), height=int(h), source_dpi=round(float(raw.dpi), 1),
                                dpi=round(float(prep.dpi), 1), rotation=prep.rotation, skew=prep.skew, notes=list(prep.notes))
                # page image for the Word report (always cached, so later runs can include it)
                info.image_file = f"page_{raw.number:04d}.jpg"
                if not _save_jpeg(prep.view_image, cache_dir / info.image_file, min(1.0, s.word_image_dpi / max(prep.dpi, 1))):
                    info.image_file = ""
                if s.save_debug_images:
                    _imwrite(cache_dir / f"page_{raw.number:04d}_cleaned.png", prep.ocr_image)
                if s.engine == "docling":
                    ocr_file = cache_dir / f"page_{raw.number:04d}_docling.jpg"
                    if not _imwrite(ocr_file, prep.ocr_image, [cv2.IMWRITE_JPEG_QUALITY, 92]):
                        raise OSError(f"could not write the temporary page image {ocr_file.name}")
                    docling_pages.append((raw.number, float(prep.dpi), int(h), int(w), ocr_file))
                    doc.pages.append(info)
                    continue
                lines = self.ocr.read(prep.ocr_image)
                blocks = analyze_page(raw.number, prep, lines, s, ocr=self.ocr)
                self._save_figure_images(blocks, prep.view_image, prep.dpi, cache_dir)
            except Exception as exc:
                LOG.debug(traceback.format_exc())
                self._page_failed(doc, raw.number, f"{type(exc).__name__}: {exc}")
                continue
            doc.pages.append(info)
            doc.blocks.extend(blocks)
            eta = (time.time() - t0) / n * (total - n)
            LOG.info("  page %d (%d of %d): %d text lines, %d tables, %d figures  [%.1f s; about %s left]",
                     raw.number, n, total, len(lines), sum(1 for b in blocks if b.kind == "table"),
                     sum(1 for b in blocks if b.kind == "figure"), time.time() - page_start, _human_time(eta))
        if s.engine == "docling" and docling_pages:
            try:
                if self._docling is None:
                    LOG.info("Loading Docling models (first use downloads them - this can take a while)...")
                    self._docling = DoclingEngine(s)
                blocks, _scores = self._docling.convert(docling_pages, cache_dir)
                for number, dpi, _h, _w, image_file in docling_pages:
                    img = _imread_gray(image_file)
                    self._save_figure_images([b for b in blocks if b.page == number], img, dpi, cache_dir)
                doc.blocks = blocks
            except Exception as exc:              # missing models, no network, or any Docling error
                LOG.debug(traceback.format_exc())
                msg = f"Docling engine could not run - used the standard engine instead. {exc}"
                LOG.warning("  %s", msg)
                doc.warnings.append(msg)
                doc.engine = "standard (Docling unavailable)"
                s.engine = "standard"           # do not retry Docling for the remaining files
                for number, dpi, _h, _w, image_file in docling_pages:
                    try:
                        img = _imread_gray(image_file)
                        prep = PreparedPage(img, img, dpi)
                        lines = self.ocr.read(img)
                        blocks = analyze_page(number, prep, lines, s, ocr=self.ocr)
                        self._save_figure_images(blocks, img, dpi, cache_dir)
                        doc.blocks.extend(blocks)
                        LOG.info("  page %d: %d text lines (standard engine)", number, len(lines))
                    except Exception as page_exc:
                        self._page_failed(doc, number, f"{type(page_exc).__name__}: {page_exc}", keep_info=True)
        doc.seconds = round(time.time() - t0, 1)
        return doc

    @staticmethod
    def _page_failed(doc: DocResult, number: int, reason: str, keep_info: bool = False):
        msg = f"page {number} FAILED and was skipped: {reason}"
        LOG.error("  %s", msg)
        doc.warnings.append(msg)
        existing = next((p for p in doc.pages if p.number == number), None) if keep_info else None
        if existing is not None:
            existing.notes.append(f"FAILED: {reason}")
        else:
            doc.pages.append(PageInfo(number=number, width=0, height=0, source_dpi=0.0, dpi=0.0,
                                      notes=[f"FAILED: {reason}"]))

    @staticmethod
    def _save_figure_images(blocks: list, image: "np.ndarray", dpi: float, cache_dir: Path):
        for k, b in enumerate(b for b in blocks if b.kind == "figure"):
            x0, y0, x1, y1 = b.box
            pad = int(0.05 * dpi)
            h, w = image.shape[:2]
            crop = image[max(0, y0 - pad):min(h, y1 + pad), max(0, x0 - pad):min(w, x1 + pad)]
            if crop.size:
                name = f"figure_p{b.page:04d}_{k + 1}.png"
                if _imwrite(cache_dir / name, crop):
                    b.meta["image_file"] = name


def _human_time(seconds: float) -> str:
    seconds = max(0, int(seconds))
    if seconds < 60:
        return f"{seconds} s"
    if seconds < 3600:
        return f"{seconds // 60} min"
    return f"{seconds // 3600} h {(seconds % 3600) // 60} min"


def collect_input_files(inputs: list, recursive: bool) -> list:
    files = []
    for item in inputs:
        p = Path(item).expanduser()
        if p.is_dir():
            pattern = "**/*" if recursive else "*"
            files += sorted(f for f in p.glob(pattern) if f.is_file() and f.suffix.lower() in SUPPORTED_EXTENSIONS
                            and "_cache" not in f.parts and not f.name.startswith("~$"))
        elif p.is_file() and p.suffix.lower() in SUPPORTED_EXTENSIONS:
            files.append(p)
        elif p.exists():
            LOG.warning("Skipping %s (supported types: %s)", p.name, ", ".join(sorted(SUPPORTED_EXTENSIONS)))
        else:
            LOG.warning("Not found: %s", p)
    unique, seen = [], set()
    for f in files:
        key = str(f.resolve()).lower()
        if key not in seen:
            seen.add(key)
            unique.append(f)
    return unique


def unique_output_stems(files: list) -> dict:
    """File -> name used for its outputs (<name>_extracted.xlsx). Names never collide, so one run cannot
    overwrite its own results: 'spec.pdf' + 'spec.tif' -> spec_pdf / spec_tif, and the same name in two
    folders gets the folder name in front (B_spec_pdf), then a counter if that is still not unique."""
    stems = Counter(f.stem.lower() for f in files)
    output_stems, used = {}, set()
    for f in files:
        stem = f.stem if stems[f.stem.lower()] == 1 else f"{f.stem}_{f.suffix.lstrip('.').lower()}"
        candidate = stem if stem.lower() not in used else f"{f.parent.name}_{stem}"
        k = 2
        while candidate.lower() in used:
            candidate = f"{stem}_{k}"
            k += 1
        used.add(candidate.lower())
        output_stems[f] = candidate
    return output_stems


def _pick_inputs_with_dialog() -> list:
    """File picker for people who double-click the script instead of using a terminal."""
    try:
        import tkinter as tk
        from tkinter import filedialog
        root = tk.Tk()
        root.withdraw()
        root.attributes("-topmost", True)
        files = filedialog.askopenfilenames(
            title="Choose scanned PDF/TIFF files  (press Cancel to choose a whole folder instead)",
            filetypes=[("Scanned documents", "*.pdf *.tif *.tiff *.png *.jpg *.jpeg *.bmp"), ("All files", "*.*")])
        if files:
            root.destroy()
            return [Path(f) for f in files]
        folder = filedialog.askdirectory(title="Choose a folder of scanned documents")
        root.destroy()
        return [Path(folder)] if folder else []
    except Exception:
        return []


def build_arg_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        prog="spec_extractor.py",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        description="Extract requirements, tables and captions from scanned PDF/TIFF documents into Excel, "
                    "with a Word report for checking the OCR.",
        epilog="examples:\n"
               "  python spec_extractor.py C:\\Scans\\spec.pdf\n"
               "  python spec_extractor.py C:\\Scans -r -o C:\\Results\n"
               "  python spec_extractor.py spec.tif --pages 1-10 --keywords shall,must\n"
               "  python spec_extractor.py spec.pdf --engine docling --docling-models D:\\docling_models\n")
    p.add_argument("inputs", nargs="*", help="PDF/TIFF/PNG/JPG files and/or folders (omit to open a file picker)")
    g = p.add_argument_group("input / output")
    g.add_argument("-o", "--output", help="output folder (default: 'extraction_output' next to the input)")
    g.add_argument("-r", "--recursive", action="store_true", help="also process files in sub-folders")
    g.add_argument("--pages", default="", help="pages to process, e.g. 1-5,9 (default: all)")
    g.add_argument("--no-word", action="store_true", help="do not write the Word verification report")
    g.add_argument("--no-page-images", action="store_true", help="leave page images out of the Word report (smaller file)")
    g.add_argument("--word-style", metavar="FILE",
                   help="style file for the Word report's look and template (default: word_style/report_style.ini)")
    g.add_argument("--write-word-template", metavar="FILE_OR_FOLDER",
                   help="write a Word file with every style of the report and a sample of each, to restyle in Word "
                        "and use as the template (a folder gets report_template.docx), then stop")
    g.add_argument("--no-combined", action="store_true", help="no ALL_DOCUMENTS workbook for multi-file runs")
    g.add_argument("--no-cache", action="store_true", help="always re-run OCR (ignore cached results)")
    g.add_argument("--debug-images", action="store_true", help="save the cleaned page images in the _cache folder")
    g = p.add_argument_group("requirements")
    g.add_argument("--keywords", help="comma-separated requirement keywords (default: shall,must,is required to,"
                                      "are required to,should,will)")
    g.add_argument("--no-will", action="store_true", help='do not report "will" statements')
    g.add_argument("--no-should", action="store_true", help='do not report "should" statements')
    g.add_argument("--keep-list-stems", action="store_true", help="one row per stem with all its list items")
    g.add_argument("--no-keyword-repair", action="store_true", help="do not repair OCR-garbled keywords such as 'sha11'")
    g.add_argument("--no-quality-notes", action="store_true", help="skip INCOSE-style writing checks")
    g = p.add_argument_group("OCR and image clean-up")
    g.add_argument("--engine", choices=["standard", "docling"], default="standard",
                   help="standard = pip-only (default); docling = adds IBM Docling layout/table AI models")
    g.add_argument("--dpi", type=int, help="working resolution for OCR (default 220)")
    g.add_argument("--review-threshold", type=float, help="flag numbers read below this OCR confidence (default 0.70)")
    g.add_argument("--min-confidence", type=float, help="drop OCR text below this confidence as noise (default 0.30)")
    g.add_argument("--no-rotate", action="store_true", help="do not auto-rotate sideways/upside-down pages")
    g.add_argument("--no-deskew", action="store_true", help="do not straighten skewed scans")
    g.add_argument("--no-clean", action="store_true", help="no background flattening / denoising")
    g.add_argument("--remove-color-marks", action="store_true", help="fade colored stamps, pen marks and highlighter")
    g.add_argument("--no-borderless-tables", action="store_true", help="only detect tables that have ruling lines")
    g.add_argument("--no-inferred-rows", action="store_true",
                   help="do not work out a table's missing row separators from the text (tables whose "
                        "columns are ruled but whose rows are not will then be missed)")
    g.add_argument("--no-sideways-headers", action="store_true",
                   help="do not turn a table's empty header band upright to read column headers "
                        "that were printed sideways")
    g.add_argument("--threads", type=int, default=0, help="CPU threads for OCR (default: automatic)")
    g.add_argument("--docling-models", help="folder with pre-downloaded Docling models (offline use)")
    g.add_argument("--device", choices=["auto", "cpu", "cuda", "mps"], default="auto", help="device for the Docling engine")
    p.add_argument("-v", "--verbose", action="store_true", help="more detailed log")
    p.add_argument("--version", action="version", version=f"%(prog)s {__version__} (build {build_id()})")
    return p


KEYWORD_CLASSES = {"shall": "Mandatory", "must": "Mandatory", "is required to": "Mandatory", "are required to": "Mandatory",
                   "required": "Mandatory", "should": "Goal / recommended", "will": "Statement of intent",
                   "may": "Permissive / optional", "can": "Permissive / optional"}


def settings_from_args(args) -> Settings:
    s = Settings()
    if args.output:
        s.output_dir = Path(args.output)
    s.recursive = args.recursive
    s.page_selection = args.pages or ""
    s.make_word_report = not args.no_word
    s.word_page_images = not args.no_page_images
    if args.word_style:
        s.word_style_file = Path(args.word_style).expanduser()
    s.combined_workbook = not args.no_combined
    s.use_cache = not args.no_cache
    s.save_debug_images = args.debug_images
    if args.keywords:
        words = [" ".join(k.strip().lower().split()) for k in re.split(r"[,;|]", args.keywords) if k.strip()]
        s.keywords = {k: KEYWORD_CLASSES.get(k, "Custom keyword") for k in words}
    if args.no_will:
        s.keywords.pop("will", None)
    if args.no_should:
        s.keywords.pop("should", None)
    if not s.keywords:
        raise SystemExit("No requirement keywords left - check --keywords / --no-will / --no-should.")
    s.split_list_items = not args.keep_list_stems
    s.fix_ocr_keywords = not args.no_keyword_repair
    s.quality_checks = not args.no_quality_notes
    s.engine = args.engine
    if args.dpi:
        s.target_dpi = int(np.clip(args.dpi, 120, 400))
    if args.review_threshold is not None:
        s.review_confidence = float(np.clip(args.review_threshold, 0.0, 1.0))
    if args.min_confidence is not None:
        s.min_text_confidence = float(np.clip(args.min_confidence, 0.0, 0.95))
    s.auto_rotate = not args.no_rotate
    s.deskew = not args.no_deskew
    s.clean_background = s.denoise = not args.no_clean
    s.remove_color_marks = args.remove_color_marks
    s.borderless_tables = not args.no_borderless_tables
    s.infer_missing_rules = not args.no_inferred_rows
    s.sideways_headers = not args.no_sideways_headers
    s.ocr_threads = max(0, args.threads or 0)
    if args.docling_models:
        s.docling_models_dir = Path(args.docling_models)
    s.docling_device = args.device
    return s


def _setup_logging(out_dir: Optional[Path], verbose: bool):
    # Windows consoles/pipes may not be able to print every character of a file name: never crash on it
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(errors="replace")
        except (AttributeError, ValueError):
            pass
    root = logging.getLogger()
    root.handlers.clear()
    root.setLevel(logging.DEBUG)
    console = logging.StreamHandler(sys.stdout)
    console.setLevel(logging.DEBUG if verbose else logging.INFO)
    console.setFormatter(logging.Formatter("%(message)s"))
    root.addHandler(console)
    if out_dir is not None:
        fh = logging.FileHandler(out_dir / "spec_extractor_log.txt", encoding="utf-8")
        fh.setLevel(logging.DEBUG)
        fh.setFormatter(logging.Formatter("%(asctime)s %(levelname)s %(message)s"))
        root.addHandler(fh)
    _quiet_third_party_loggers()


def report_configuration(docling_used: bool) -> None:
    """Log whether this run uses a tested Python and the tested package versions (details in the log file)."""
    software, differences = configuration_status(docling_used)
    LOG.debug("Software: %s", software)
    if not differences:
        LOG.info("Tested configuration: yes (%s with the package versions in %s)",
                 software.split(";")[0], CONSTRAINTS_FILE.name)
        return
    LOG.warning("NOT the tested configuration - the documented accuracy was not measured for:")
    for difference in differences:
        LOG.warning("  - %s", difference)
    if any(d != DOCLING_NOTE for d in differences):
        LOG.warning("  To use the tested setup: delete the .venv folder next to spec_extractor.py and run "
                    "setup_windows.bat again")
        LOG.warning("  (other systems: python -m pip install -r requirements.txt -c constraints.txt)")


def _write_starter_template_command(args) -> int:
    """--write-word-template: write the starter template, then stop (no documents are read)."""
    _setup_logging(None, args.verbose)
    settings = settings_from_args(args)
    try:
        look = report_look(settings)
        path = write_starter_template(Path(args.write_word_template).expanduser(), look)
    except ReportLookError as exc:
        LOG.error("%s", exc)
        return 2
    except OSError as exc:
        LOG.error("The Word template could not be written to %s (%s). Choose a folder you can write to.",
                  args.write_word_template, exc.strerror or exc)
        return 2
    style_file = Path(settings.word_style_file or DEFAULT_WORD_STYLE_FILE)
    LOG.info("%s", version_line())
    LOG.info("Word template written: %s", path)
    LOG.info("Restyle it in Word (Home > Styles: right-click a style > Modify; double-click the top or bottom of a page "
             "for the header and footer), save it, then write its name in the style file:")
    # (a name is looked for next to the style file, so anything else is given in full)
    LOG.info("    %s:   template = %s", style_file, path.name if path.resolve().parent == style_file.resolve().parent
             else path.resolve())
    if os.name == "nt" and os.environ.get("SPEC_EXTRACTOR_LAUNCHER"):     # (started from make_word_template.bat)
        try:
            os.startfile(str(path))                                      # open it in Word on this computer
        except OSError:
            pass
    return 0


def _pause_if_dialog(used_dialog: bool, code: int) -> int:
    """When the script was double-clicked (files chosen in the file window), keep its window open until Enter is
    pressed, so the messages can be read. (The .bat launchers pause by themselves.)"""
    if used_dialog and sys.stdin and sys.stdin.isatty() and not os.environ.get("SPEC_EXTRACTOR_LAUNCHER"):
        try:
            input("\nPress Enter to close this window...")
        except EOFError:
            pass
    return code


def main(argv: Optional[list] = None) -> int:
    args = build_arg_parser().parse_args(argv)
    if args.write_word_template:
        return _write_starter_template_command(args)
    used_dialog = False
    inputs = args.inputs
    if not inputs:
        inputs = _pick_inputs_with_dialog()
        used_dialog = bool(inputs)
        if not inputs:
            build_arg_parser().print_help()
            return 1
    settings = settings_from_args(args)
    _setup_logging(None, args.verbose)
    try:
        parse_page_selection(settings.page_selection, 10 ** 6)
    except ValueError as exc:
        LOG.error("--pages: %s", exc)
        return _pause_if_dialog(used_dialog, 2)
    configure_keywords(settings.keywords)
    look = None
    if settings.make_word_report:          # a mistake in the style file or template is reported now, not after the OCR
        try:
            look = report_look(settings)
        except ReportLookError as exc:
            LOG.error("%s", exc)
            return _pause_if_dialog(used_dialog, 2)
        settings.word_look = look          # (the look checked now is the one every report of this run uses)
    files = collect_input_files(inputs, settings.recursive)
    if not files:
        LOG.error("No PDF/TIFF/image files found in: %s", ", ".join(map(str, inputs)))
        return _pause_if_dialog(used_dialog, 1)
    first = Path(inputs[0]).expanduser()
    out_dir = settings.output_dir or ((first if first.is_dir() else first.parent) / "extraction_output")
    out_dir.mkdir(parents=True, exist_ok=True)
    _setup_logging(out_dir, args.verbose)
    LOG.info("%s - %d file(s) - engine: %s - output: %s", version_line(), len(files), settings.engine, out_dir)
    report_configuration(docling_used=settings.engine == "docling")
    if look is not None:
        LOG.info("Word report look: %s", look.summary)
        for note in look.template.notes if look.template else []:
            LOG.info("  from the template, %s", note)

    output_stems = unique_output_stems(files)
    extractor = SpecExtractor(settings)
    runs, failures = [], []
    for n, f in enumerate(files, 1):
        stem = output_stems[f]
        LOG.info("")
        LOG.info("[%d/%d] %s", n, len(files), f.name if stem == f.stem else f"{f.name}   (from {f.parent} -> {stem}_...)")
        try:
            doc, items, outputs = extractor.process_file(f, out_dir, stem)
            runs.append((doc, items))
            for kind, pth in outputs.items():
                LOG.info("  %-5s -> %s", kind, pth)
        except EngineUnavailable as exc:
            LOG.error("  %s", exc)
            return 3
        except KeyboardInterrupt:
            LOG.error("Stopped by user.")
            return 130
        except Exception as exc:
            reason = str(exc)
            if "password" in reason.lower():
                reason = "the PDF is password-protected - save an unprotected copy and run again"
            elif "format error" in reason.lower() or "cannot identify image file" in reason.lower():
                reason = "the file is damaged or is not really a PDF/TIFF/image"
            LOG.error("  FAILED: %s (details in spec_extractor_log.txt)", reason)
            LOG.debug(traceback.format_exc())
            failures.append((f.name, reason))

    if settings.combined_workbook and len(runs) > 1 and settings.make_excel:
        combined = write_workbook(out_dir / f"ALL_DOCUMENTS_{_dt.datetime.now():%Y%m%d_%H%M}.xlsx", runs, settings, multi_doc=True)
        LOG.info("")
        LOG.info("Combined workbook -> %s", combined)
    LOG.info("")
    LOG.info("Finished: %d document(s) processed, %d failed.", len(runs), len(failures))
    for name, err in failures:
        LOG.info("  failed: %s - %s", name, err)
    return _pause_if_dialog(used_dialog, 0 if not failures else 1)


if __name__ == "__main__":
    sys.exit(main())
